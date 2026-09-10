#!/usr/bin/env python
"""Apply professional rules and create a clean species registration workbook."""

from __future__ import annotations

import argparse
import re
from collections import Counter
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from validar_cadastro_especies import (
    CADASTRO_COLUMNS,
    _normalized,
    query_ctfb,
    query_gbif,
    query_worms,
)


TAXONOMY_FIELDS = {
    "Reino": "kingdom",
    "Filo": "phylum",
    "Classe": "class",
    "Ordem": "order",
    "Familia": "family",
    "Genero": "genus",
    "Autor_e_Ano": "authorship",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    return parser.parse_args()


def text(value: Any) -> str:
    if pd.isna(value):
        return "N.A."
    value = str(value).strip()
    return value if value else "N.A."


def original_value(row: pd.Series, field: str) -> str:
    return text(row.get(f"input_{field}", row.get(field)))


def majority(results: list[dict[str, Any]], field: str) -> str:
    values = [str(result.get(field) or "").strip() for result in results if result.get(field)]
    normalized = [_normalized(value) for value in values]
    counts = Counter(normalized)
    if not counts:
        return ""
    winner, count = counts.most_common(1)[0]
    if count < 2:
        return ""
    return next(value for value in values if _normalized(value) == winner)


def query_candidate(name: str, rank: str, kingdom: str = "") -> list[dict[str, Any]]:
    results = [
        query_worms(name, expected_kingdom=kingdom, expected_rank=rank),
        query_ctfb(name),
        query_gbif(name, kingdom, 85),
    ]
    for result in results:
        if _normalized(result.get("rank")) != _normalized(rank):
            result["accepted"] = False
    return results


def accepted_count(results: list[dict[str, Any]], name: str) -> int:
    return sum(
        bool(result.get("accepted"))
        and _normalized(result.get("accepted_name")) == _normalized(name)
        for result in results
    )


def apply_taxonomy(final: dict[str, str], results: list[dict[str, Any]]) -> list[str]:
    changed: list[str] = []
    for output_field, source_field in TAXONOMY_FIELDS.items():
        value = majority(results, source_field)
        if value and _normalized(final[output_field]) != _normalized(value):
            final[output_field] = value
            changed.append(output_field)
    return changed


def classify_sp(row: pd.Series, final: dict[str, str], original_name: str) -> tuple[str, str]:
    match = re.fullmatch(r"\s*([A-Za-z][A-Za-z-]+)\s+sp\.\s*", original_name, flags=re.IGNORECASE)
    if not match:
        return "", ""
    supplied_genus = match.group(1)
    seed = query_gbif(supplied_genus, "Animalia", 85)
    candidate = supplied_genus
    if (
        _normalized(seed.get("rank")) == "genus"
        and seed.get("accepted_name")
        and SequenceMatcher(None, _normalized(supplied_genus), _normalized(seed["accepted_name"])).ratio() >= 0.85
    ):
        candidate = str(seed["accepted_name"])
    results = query_candidate(candidate, "genus", "Animalia")
    if accepted_count(results, candidate) < 2:
        return "SP_MANTIDO", "nome com sp. mantido; classificacao sem confirmacao em duas bases"
    final["Nome_Cientifico"] = f"{candidate} sp."
    changed = apply_taxonomy(final, results)
    correction = _normalized(candidate) != _normalized(supplied_genus)
    action = "SP_CORRIGIDO_E_CLASSIFICADO" if correction else "SP_MANTIDO_E_CLASSIFICADO"
    detail = "classificacao confirmada; campos atualizados: " + ", ".join(changed or ["nenhum"])
    return action, detail


def correct_spelling(row: pd.Series, final: dict[str, str], original_name: str) -> tuple[str, str]:
    candidate = text(row.get("opyta_gbif_accepted_name"))
    rank = text(row.get("opyta_gbif_rank"))
    try:
        confidence = int(float(row.get("opyta_gbif_confidence") or 0))
    except (TypeError, ValueError):
        confidence = 0
    if candidate == "N.A." or _normalized(candidate) == _normalized(original_name):
        return "", ""
    if confidence < 85 or _normalized(rank) != "species":
        return "", ""
    if len(_normalized(original_name).split()) != len(_normalized(candidate).split()):
        return "", ""
    similarity = SequenceMatcher(None, _normalized(original_name), _normalized(candidate)).ratio()
    if similarity < 0.85:
        return "", ""
    results = query_candidate(candidate, "species")
    if accepted_count(results, candidate) < 2:
        return "", ""
    final["Nome_Cientifico"] = candidate
    changed = apply_taxonomy(final, results)
    return (
        "ORTOGRAFIA_CORRIGIDA",
        f"{original_name} -> {candidate}; campos atualizados: {', '.join(changed or ['nenhum'])}",
    )


def style_sheet(worksheet, color: str) -> None:
    fill = PatternFill("solid", fgColor=color)
    for cell in worksheet[1]:
        cell.fill = fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 32
    for column in worksheet.columns:
        letter = column[0].column_letter
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
        worksheet.column_dimensions[letter].width = max(width, 14)
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def main() -> int:
    args = parse_args()
    source = pd.read_excel(args.input, sheet_name="Cadastro_Especies", keep_default_na=False)
    final_rows: list[dict[str, str]] = []
    changes: list[list[str]] = []

    for _, row in source.iterrows():
        original_name = original_value(row, "Nome_Cientifico")
        final = {field: original_value(row, field) for field in CADASTRO_COLUMNS}
        action, detail = classify_sp(row, final, original_name)
        if not action:
            action, detail = correct_spelling(row, final, original_name)
        if not action:
            status = text(row.get("opyta_validation_status"))
            if status == "approved_consensus":
                final = {field: text(row.get(field)) for field in CADASTRO_COLUMNS}
                action = "CONSENSO_3_BASES"
                detail = "classificacao aceita por WoRMS, CTFB e GBIF"
            else:
                action = "MANTIDO_PROFISSIONAL"
                detail = "registro parcial, conflitante ou nao encontrado; valor profissional mantido"
        final_rows.append(final)
        if action != "MANTIDO_PROFISSIONAL":
            changes.append([original_name, final["Nome_Cientifico"], action, detail])

    workbook = Workbook()
    cadastro = workbook.active
    cadastro.title = "Cadastro_Especies"
    cadastro.append(CADASTRO_COLUMNS)
    for row in final_rows:
        cadastro.append([row[field] for field in CADASTRO_COLUMNS])
    style_sheet(cadastro, "245B78")

    altered = workbook.create_sheet("O_QUE_MUDOU")
    altered.append(["Nome informado", "Nome final", "Regra aplicada", "Detalhe"])
    for row in changes:
        altered.append(row)
    style_sheet(altered, "2F6B4F")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(f"[professional-rules] rows={len(final_rows)} changes={len(changes)} output={args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
