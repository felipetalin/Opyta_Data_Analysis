#!/usr/bin/env python
"""Create a compact workbook for human approval of taxonomic validation."""

from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    return parser.parse_args()


def text(value: object) -> str:
    if pd.isna(value):
        return "N.A."
    value = str(value).strip()
    return value if value else "N.A."


def prepare_sheet(worksheet, headers: list[str], rows: list[list[str]], color: str) -> None:
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)

    header_fill = PatternFill("solid", fgColor=color)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 34
    widths = [22, 30, 30, 18, 20, 20, 20, 22, 22, 25, 55]
    for index, width in enumerate(widths[: len(headers)], start=1):
        worksheet.column_dimensions[worksheet.cell(1, index).column_letter].width = width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def main() -> int:
    args = parse_args()
    dataframe = pd.read_excel(args.input, sheet_name="Cadastro_Especies", keep_default_na=False)
    approved = dataframe[dataframe["opyta_validation_status"] == "approved_consensus"]
    pending = dataframe[dataframe["opyta_validation_status"] != "approved_consensus"]

    workbook = Workbook()
    approve_sheet = workbook.active
    approve_sheet.title = "APROVAR_15"
    approve_headers = [
        "APROVAR? (SIM/NAO)",
        "Nome informado",
        "Nome validado",
        "Reino",
        "Filo",
        "Classe",
        "Ordem",
        "Familia",
        "Genero",
        "Autor e ano",
        "Comprovacao",
    ]
    approve_rows = [
        [
            "",
            text(row["input_Nome_Cientifico"]),
            text(row["Nome_Cientifico"]),
            text(row["Reino"]),
            text(row["Filo"]),
            text(row["Classe"]),
            text(row["Ordem"]),
            text(row["Familia"]),
            text(row["Genero"]),
            text(row["Autor_e_Ano"]),
            "WoRMS + CTFB + GBIF concordam",
        ]
        for _, row in approved.iterrows()
    ]
    prepare_sheet(approve_sheet, approve_headers, approve_rows, "2F6B4F")
    approval_validation = DataValidation(type="list", formula1='"SIM,NAO"', allow_blank=True)
    approve_sheet.add_data_validation(approval_validation)
    approval_validation.add(f"A2:A{len(approve_rows) + 1}")
    approve_sheet.conditional_formatting.add(
        f"A2:A{len(approve_rows) + 1}",
        FormulaRule(formula=["A2=\"SIM\""], fill=PatternFill("solid", fgColor="C6EFCE")),
    )
    approve_sheet.conditional_formatting.add(
        f"A2:A{len(approve_rows) + 1}",
        FormulaRule(formula=["A2=\"NAO\""], fill=PatternFill("solid", fgColor="FFC7CE")),
    )

    pending_sheet = workbook.create_sheet("PENDENTES_43")
    pending_headers = [
        "DECISAO DO PROFISSIONAL",
        "Nome informado",
        "Situacao",
        "Motivo",
        "Nome no WoRMS",
        "Nome no CTFB",
        "Nome no GBIF",
        "Nome mantido",
    ]
    labels = {
        "review_partial": "Encontrado parcialmente",
        "review_conflict": "Conflito entre bases",
        "not_found": "Nao encontrado",
    }
    pending_rows = [
        [
            "",
            text(row["input_Nome_Cientifico"]),
            labels.get(text(row["opyta_validation_status"]), text(row["opyta_validation_status"])),
            text(row["opyta_validation_notes"]),
            text(row.get("opyta_worms_accepted_name")),
            text(row.get("opyta_ctfb_accepted_name")),
            text(row.get("opyta_gbif_accepted_name")),
            text(row["Nome_Cientifico"]),
        ]
        for _, row in pending.iterrows()
    ]
    prepare_sheet(pending_sheet, pending_headers, pending_rows, "9C6500")
    decision_validation = DataValidation(type="list", formula1='"MANTER,CORRIGIR"', allow_blank=True)
    pending_sheet.add_data_validation(decision_validation)
    decision_validation.add(f"A2:A{len(pending_rows) + 1}")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(f"[simple-review] approved={len(approve_rows)} pending={len(pending_rows)} output={args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
