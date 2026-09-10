#!/usr/bin/env python
"""Apply fitoplankton consensus rules and preserve a per-source audit trail."""

from __future__ import annotations

import argparse
import re
from collections import Counter
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from validar_cadastro_especies import CADASTRO_COLUMNS, _normalized


SOURCES = ("gbif", "col", "diatombase")
PRIMARY_SOURCE_ORDER = ("diatombase", "col", "gbif")
SOURCE_LABELS = {
    "gbif": "GBIF",
    "col": "Catalogue of Life",
    "diatombase": "DiatomBase",
}
TAXONOMY_FIELDS = {
    "Reino": "kingdom",
    "Filo": "phylum",
    "Classe": "class",
    "Ordem": "order",
    "Familia": "family",
    "Genero": "genus",
    "Autor_e_Ano": "authorship",
}
CLASSIFICATION_FIELDS = ("Reino", "Filo", "Classe", "Ordem", "Familia", "Genero")
PARTIAL_SOURCE_FIELDS = {
    "Reino": "kingdom",
    "Filo": "phylum",
    "Classe": "class",
    "Ordem": "order",
    "Familia": "family",
    "Genero": "genus",
}
TRACE_FIELDS = (
    "found",
    "accepted",
    "accepted_name",
    "matched_name",
    "status",
    "rank",
    "kingdom",
    "phylum",
    "class",
    "order",
    "family",
    "genus",
    "authorship",
    "url",
    "notes",
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    return parser.parse_args()


def raw_text(value: Any) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip()


def text(value: Any) -> str:
    value = raw_text(value)
    return value if value and _normalized(value) not in {"n.a.", "na"} else "N.A."


def is_available(value: Any) -> bool:
    return text(value) != "N.A."


def is_true(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return _normalized(value) in {"true", "1", "yes", "sim"}


def original_value(row: pd.Series, field: str) -> str:
    return text(row.get(f"input_{field}", row.get(field)))


def scientific_name_format(value: Any) -> str:
    value = "" if pd.isna(value) else str(value)
    value = " ".join(value.split())
    value = re.sub(r"\bsp\.?(?=\s|$)", "sp.", value, flags=re.IGNORECASE)
    value = re.sub(r"\bn\.?\s*i\.?(?=\s|$)", "n.i.", value, flags=re.IGNORECASE)
    return value or "N.A."


def consensus(
    row: pd.Series,
    source_field: str,
    minimum_sources: int = 2,
) -> tuple[str, list[str], bool]:
    values: list[tuple[str, str]] = []
    for source in SOURCES:
        if not is_true(row.get(f"opyta_{source}_accepted")):
            continue
        value = raw_text(row.get(f"opyta_{source}_{source_field}"))
        if value:
            values.append((source, value))
    if len(values) < minimum_sources:
        return "", [SOURCE_LABELS[source] for source, _ in values], False
    distinct = {_normalized(value) for _, value in values}
    if len(distinct) != 1:
        return "", [SOURCE_LABELS[source] for source, _ in values], True
    return values[0][1], [SOURCE_LABELS[source] for source, _ in values], False


def source_is_usable(
    row: pd.Series,
    source: str,
    partial_taxon: str = "",
    partial_rank: str = "",
) -> bool:
    if not is_true(row.get(f"opyta_{source}_accepted")):
        return False
    if not partial_taxon or partial_rank not in PARTIAL_SOURCE_FIELDS:
        return True
    source_rank = raw_text(row.get(f"opyta_{source}_rank"))
    accepted_name = raw_text(row.get(f"opyta_{source}_accepted_name"))
    rank_value = raw_text(row.get(f"opyta_{source}_{PARTIAL_SOURCE_FIELDS[partial_rank]}"))
    return (
        _normalized(source_rank) == _normalized(PARTIAL_SOURCE_FIELDS[partial_rank])
        and _normalized(accepted_name) == _normalized(partial_taxon)
        and _normalized(rank_value) == _normalized(partial_taxon)
    )


def resolve_taxonomy(
    row: pd.Series,
    source_field: str,
    partial_taxon: str = "",
    partial_rank: str = "",
) -> tuple[str, list[str], str]:
    values: list[tuple[str, str]] = []
    for source in SOURCES:
        if not source_is_usable(row, source, partial_taxon, partial_rank):
            continue
        value = raw_text(row.get(f"opyta_{source}_{source_field}"))
        if value:
            values.append((source, value))
    counts = Counter(_normalized(value) for _, value in values)
    if counts:
        winner, count = counts.most_common(1)[0]
        if count >= 2:
            agreeing = [(source, value) for source, value in values if _normalized(value) == winner]
            return (
                agreeing[0][1],
                [SOURCE_LABELS[source] for source, _ in agreeing],
                "CONSENSO_2_BASES",
            )
    for source in PRIMARY_SOURCE_ORDER:
        for found_source, value in values:
            if found_source == source:
                return value, [SOURCE_LABELS[source]], f"BASE_PRINCIPAL_{SOURCE_LABELS[source]}"
    return "", [], "SEM_RETORNO"


def partial_identification_rank(row: pd.Series, name: str) -> str:
    if re.fullmatch(r"[A-Za-z][A-Za-z-]+\s+sp\.", name, flags=re.IGNORECASE):
        return "Genero"
    if not re.search(r"\bn\.?\s*i\.?(?=\s|$)", name, flags=re.IGNORECASE):
        return ""
    taxon = name.split()[0]
    suffix_ranks = (
        ("aceae", "Familia"),
        ("ales", "Ordem"),
        ("phyceae", "Classe"),
        ("phyta", "Filo"),
    )
    for suffix, rank in suffix_ranks:
        if taxon.casefold().endswith(suffix):
            return rank
    for field in reversed(CLASSIFICATION_FIELDS):
        if is_available(original_value(row, field)):
            return field
    return ""


def style_sheet(worksheet: Any, color: str) -> None:
    fill = PatternFill("solid", fgColor=color)
    for cell in worksheet[1]:
        cell.fill = fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 32
    for column in worksheet.columns:
        letter = column[0].column_letter
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 55)
        worksheet.column_dimensions[letter].width = max(width, 14)
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def main() -> int:
    args = parse_args()
    source = pd.read_excel(args.input, sheet_name="Cadastro_Especies", keep_default_na=False)
    final_rows: list[dict[str, str]] = []
    changes: list[list[str]] = []
    pending: list[list[str]] = []
    pending_keys: set[str] = set()
    source_rows: list[list[Any]] = []
    traced: set[tuple[str, str]] = set()

    for _, row in source.iterrows():
        original_name_raw = row.get("input_Nome_Cientifico", row.get("Nome_Cientifico"))
        original_name = "" if pd.isna(original_name_raw) else str(original_name_raw)
        final = {field: original_value(row, field) for field in CADASTRO_COLUMNS}
        changed_fields: list[str] = []
        field_sources: set[str] = set()
        resolution_methods: set[str] = set()
        consensus_name_changed = False

        formatted_name = scientific_name_format(original_name_raw)
        if original_name != formatted_name:
            final["Nome_Cientifico"] = formatted_name
            changed_fields.append("Nome_Cientifico")

        status = raw_text(row.get("opyta_validation_status"))
        accepted_name, name_sources, name_conflict = consensus(row, "accepted_name", 3)
        if status == "approved_consensus" and accepted_name:
            accepted_name = scientific_name_format(accepted_name)
            if _normalized(final["Nome_Cientifico"]) != _normalized(accepted_name):
                final["Nome_Cientifico"] = accepted_name
                if "Nome_Cientifico" not in changed_fields:
                    changed_fields.append("Nome_Cientifico")
                field_sources.update(name_sources)
                consensus_name_changed = True

        partial_rank = partial_identification_rank(row, formatted_name)
        partial_taxon = formatted_name.split()[0] if partial_rank else ""
        partial_rank_index = (
            CLASSIFICATION_FIELDS.index(partial_rank)
            if partial_rank in CLASSIFICATION_FIELDS else
            None
        )
        resolved_fields: set[str] = set()
        for output_field, source_field in TAXONOMY_FIELDS.items():
            if partial_rank:
                if output_field == "Autor_e_Ano":
                    continue
                field_index = CLASSIFICATION_FIELDS.index(output_field)
                if partial_rank_index is not None and field_index > partial_rank_index:
                    if is_available(final[output_field]):
                        final[output_field] = "N.A."
                        changed_fields.append(output_field)
                    continue
            value, sources, method = resolve_taxonomy(
                row,
                source_field,
                partial_taxon,
                partial_rank,
            )
            if not value:
                continue
            resolved_fields.add(output_field)
            if value and _normalized(final[output_field]) != _normalized(value):
                final[output_field] = value
                changed_fields.append(output_field)
                field_sources.update(sources)
                resolution_methods.add(method)

        final_rows.append(final)
        if changed_fields:
            action = (
                "CONSENSO_3_BASES" if consensus_name_changed else
                "NORMALIZACAO_FORMATO" if changed_fields == ["Nome_Cientifico"] else
                "CLASSIFICACAO_PARCIAL_MENOR_TAXON" if partial_rank else
                "CLASSIFICACAO_2_BASES_OU_PRINCIPAL"
            )
            source_detail = (
                f"fontes concordantes: {', '.join(sorted(field_sources))}"
                if field_sources else
                "padronizacao de espacos e marcadores taxonomicos"
            )
            method_detail = (
                f"; criterio: {', '.join(sorted(resolution_methods))}"
                if resolution_methods else
                ""
            )
            detail = f"campos atualizados: {', '.join(changed_fields)}; {source_detail}{method_detail}"
            changes.append([original_name, final["Nome_Cientifico"], action, detail])

        accepted_names = {
            _normalized(row.get(f"opyta_{source_name}_accepted_name"))
            for source_name in SOURCES
            if is_true(row.get(f"opyta_{source_name}_accepted"))
            and raw_text(row.get(f"opyta_{source_name}_accepted_name"))
        }
        name_conflict = not partial_rank and len(accepted_names) > 1
        no_classification = not (resolved_fields & set(CLASSIFICATION_FIELDS))
        if no_classification or (partial_rank and partial_rank_index is None):
            pending_key = _normalized(formatted_name)
            if pending_key not in pending_keys:
                pending_keys.add(pending_key)
                pending_fields = []
                if no_classification:
                    pending_fields.append("Classificacao")
                pending.append(
                    [
                        formatted_name,
                        "SEM_RETORNO",
                        ", ".join(pending_fields) or "Menor nivel taxonomico nao identificado",
                        raw_text(row.get("opyta_validation_notes")) or "bases sem classificacao utilizavel",
                    ]
                )

        for source_name in SOURCES:
            trace_key = (original_name, source_name)
            if trace_key in traced:
                continue
            traced.add(trace_key)
            source_rows.append(
                [
                    original_name,
                    SOURCE_LABELS[source_name],
                    *[text(row.get(f"opyta_{source_name}_{field}")) for field in TRACE_FIELDS],
                ]
            )

    workbook = Workbook()
    cadastro = workbook.active
    cadastro.title = "Cadastro_Especies"
    cadastro.append(CADASTRO_COLUMNS)
    for row in final_rows:
        cadastro.append([row[field] for field in CADASTRO_COLUMNS])
    style_sheet(cadastro, "245B78")

    altered = workbook.create_sheet("O_QUE_MUDOU")
    altered.append(["Nome informado", "Nome final", "Regra aplicada", "Detalhe"])
    if changes:
        for row in changes:
            altered.append(row)
    else:
        altered.append(["-", "-", "SEM_ALTERACOES", "Nenhum campo atingiu consenso suficiente."])
    style_sheet(altered, "2F6B4F")

    sources = workbook.create_sheet("FONTES_CONSULTADAS")
    sources.append(
        [
            "Nome informado",
            "Base",
            *[
                {
                    "found": "Encontrado",
                    "accepted": "Aceito",
                    "accepted_name": "Nome aceito",
                    "matched_name": "Nome localizado",
                    "status": "Status",
                    "rank": "Categoria",
                    "kingdom": "Reino",
                    "phylum": "Filo",
                    "class": "Classe",
                    "order": "Ordem",
                    "family": "Familia",
                    "genus": "Genero",
                    "authorship": "Autor e ano",
                    "url": "URL",
                    "notes": "Observacao",
                }[field]
                for field in TRACE_FIELDS
            ],
        ]
    )
    for row in source_rows:
        sources.append(row)
    style_sheet(sources, "5B4A70")

    review = workbook.create_sheet("PENDENCIAS_TAXONOMICAS")
    review.append(["Nome informado", "Status", "Campos em conflito", "Motivo"])
    if pending:
        for row in pending:
            review.append(row)
    else:
        review.append(["-", "APROVADO", "N.A.", "Nenhuma pendencia taxonomica."])
    style_sheet(review, "8A5A2B")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(
        f"[fitoplankton-rules] rows={len(final_rows)} changes={len(changes)} "
        f"pending={len(pending)} output={args.output}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
