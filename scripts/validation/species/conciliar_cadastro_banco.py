#!/usr/bin/env python
"""Compare a validated species workbook with the master database without writing to it."""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import text as sql_text

from validar_cadastro_especies import _normalized


SOURCE_FIELDS = {
    "Nome_Cientifico": "Nome aceito",
    "Reino": "Reino",
    "Filo": "Filo",
    "Classe": "Classe",
    "Ordem": "Ordem",
    "Familia": "Familia",
    "Genero": "Genero",
    "Autor_e_Ano": "Autor e ano",
}
DB_FIELDS = {
    "Nome_Cientifico": "nome_cientifico",
    "Reino": "reino",
    "Filo": "filo",
    "Classe": "classe",
    "Ordem": "ordem",
    "Familia": "familia",
    "Genero": "genero",
    "Autor_e_Ano": "autor_e_ano",
}
OUTPUT_COLUMNS = [
    "Nome_planilha",
    "Nome_aceito_consenso",
    "Acao_proposta",
    "Pronto_para_migrar",
    "id_especie_banco",
    "Nome_no_banco",
    "Correspondencia",
    "Grupo_no_banco",
    "Campos_para_atualizar",
    "Conflitos",
    "Aliases_consultados",
    "Status_taxonomico",
    "Observacao",
]
PRIMARY_BASE_ORDER = ("DiatomBase", "Catalogue of Life", "GBIF")
PARTIAL_RANK_COLUMNS = {
    "Reino": ("kingdom", "Reino"),
    "Filo": ("phylum", "Filo"),
    "Classe": ("class", "Classe"),
    "Ordem": ("order", "Ordem"),
    "Familia": ("family", "Familia"),
    "Genero": ("genus", "Genero"),
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output-json", required=True, type=Path)
    parser.add_argument("--opyta-data-root", required=True, type=Path)
    return parser.parse_args()


def raw_text(value: Any) -> str:
    if pd.isna(value):
        return ""
    value = str(value).strip()
    return "" if _normalized(value) in {"n.a.", "na", "none", "nan"} else value


def is_true(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return _normalized(value) in {"true", "1", "yes", "sim"}


def consensus(source_rows: pd.DataFrame, column: str, minimum: int) -> tuple[str, bool, list[str]]:
    accepted = source_rows[source_rows["Aceito"].map(is_true)]
    pairs = [
        (raw_text(row["Base"]), raw_text(row[column]))
        for _, row in accepted.iterrows()
        if raw_text(row[column])
    ]
    if len(pairs) < minimum:
        return "", False, [source for source, _ in pairs]
    distinct = {_normalized(value) for _, value in pairs}
    if len(distinct) != 1:
        return "", True, [source for source, _ in pairs]
    return pairs[0][1], False, [source for source, _ in pairs]


def partial_context(name: str) -> tuple[str, str]:
    name = " ".join(name.split())
    taxon = name.split()[0] if name else ""
    if re.fullmatch(r"[A-Za-z][A-Za-z-]+\s+sp\.", name, flags=re.IGNORECASE):
        return taxon, "Genero"
    if not re.search(r"\bn\.?\s*i\.?(?=\s|$)", name, flags=re.IGNORECASE):
        return "", ""
    for suffix, rank in (("aceae", "Familia"), ("ales", "Ordem"), ("phyceae", "Classe"), ("phyta", "Filo")):
        if taxon.casefold().endswith(suffix):
            return taxon, rank
    return taxon, ""


def source_row_usable(row: pd.Series, partial_taxon: str, partial_rank: str) -> bool:
    if not is_true(row["Aceito"]):
        return False
    if not partial_taxon or partial_rank not in PARTIAL_RANK_COLUMNS:
        return True
    expected_rank, rank_column = PARTIAL_RANK_COLUMNS[partial_rank]
    return (
        _normalized(row["Categoria"]) == expected_rank
        and _normalized(row["Nome aceito"]) == _normalized(partial_taxon)
        and _normalized(row[rank_column]) == _normalized(partial_taxon)
    )


def resolve_classification(
    source_rows: pd.DataFrame,
    column: str,
    partial_taxon: str = "",
    partial_rank: str = "",
) -> tuple[str, list[str], str]:
    accepted = source_rows[
        source_rows.apply(
            lambda row: source_row_usable(row, partial_taxon, partial_rank),
            axis=1,
        )
    ]
    pairs = [
        (raw_text(row["Base"]), raw_text(row[column]))
        for _, row in accepted.iterrows()
        if raw_text(row[column])
    ]
    counts = Counter(_normalized(value) for _, value in pairs)
    if counts:
        winner, count = counts.most_common(1)[0]
        if count >= 2:
            agreeing = [(source, value) for source, value in pairs if _normalized(value) == winner]
            return agreeing[0][1], [source for source, _ in agreeing], "CONSENSO_2_BASES"
    for primary in PRIMARY_BASE_ORDER:
        for source, value in pairs:
            if source == primary:
                return value, [source], f"BASE_PRINCIPAL_{primary}"
    return "", [], "SEM_RETORNO"


def style_sheet(worksheet: Any) -> None:
    fill = PatternFill("solid", fgColor="355F7A")
    for cell in worksheet[1]:
        cell.fill = fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 32
    for column in worksheet.columns:
        letter = column[0].column_letter
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 55)
        worksheet.column_dimensions[letter].width = max(width, 14)
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def main() -> int:
    args = parse_args()
    if str(args.opyta_data_root) not in sys.path:
        sys.path.insert(0, str(args.opyta_data_root))
    from core.engine import get_engine

    catalog = pd.read_excel(args.input, sheet_name="Cadastro_Especies", keep_default_na=False)
    sources = pd.read_excel(args.input, sheet_name="FONTES_CONSULTADAS", keep_default_na=False)
    changes = pd.read_excel(args.input, sheet_name="O_QUE_MUDOU", keep_default_na=False)
    pending = pd.read_excel(args.input, sheet_name="PENDENCIAS_TAXONOMICAS", keep_default_na=False)

    name_map = {
        _normalized(row["Nome informado"]): raw_text(row["Nome final"])
        for _, row in changes.iterrows()
        if raw_text(row.get("Nome informado")) not in {"", "-"}
        and raw_text(row.get("Nome final")) not in {"", "-"}
    }
    catalog_index = {
        _normalized(row["Nome_Cientifico"]): row
        for _, row in catalog.iterrows()
        if raw_text(row.get("Nome_Cientifico"))
    }
    pending_status = {
        _normalized(row["Nome informado"]): raw_text(row["Status"])
        for _, row in pending.iterrows()
        if raw_text(row.get("Nome informado")) not in {"", "-"}
    }
    professional_overrides = {
        _normalized(row["Nome final"])
        for _, row in changes.iterrows()
        if raw_text(row.get("Regra aplicada")) == "INFORMACAO_PROFISSIONAL_FORNECIDA"
        and raw_text(row.get("Nome final"))
    }

    engine = get_engine()
    with engine.connect() as connection:
        db_rows = connection.execute(
            sql_text(
                "SELECT id_especie, nome_cientifico, grupo_biologico, reino, filo, classe, "
                "ordem, familia, genero, autor_e_ano FROM especies"
            )
        ).mappings().all()
    db_index: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in db_rows:
        if raw_text(row.get("nome_cientifico")):
            db_index[_normalized(row["nome_cientifico"])].append(dict(row))

    reconciled: list[dict[str, Any]] = []
    processed_names: set[str] = set()
    for input_name in sources["Nome informado"].drop_duplicates().tolist():
        input_name = raw_text(input_name)
        input_key = _normalized(input_name)
        if not input_name or input_key in processed_names:
            continue
        processed_names.add(input_key)
        source_rows = sources[sources["Nome informado"].map(_normalized) == input_key]
        final_name = name_map.get(input_key, input_name)
        professional_override = _normalized(final_name) in professional_overrides
        partial_taxon, partial_rank = partial_context(final_name)
        final_row = catalog_index.get(_normalized(final_name), {})
        accepted_name, name_conflict, _ = consensus(source_rows, "Nome aceito", 3)

        aliases = {input_name, final_name}
        for column in ("Nome aceito", "Nome localizado"):
            aliases.update(raw_text(value) for value in source_rows[column].tolist() if raw_text(value))
        matched_by_id: dict[Any, dict[str, Any]] = {}
        matched_aliases: list[str] = []
        for alias in aliases:
            for db_row in db_index.get(_normalized(alias), []):
                matched_by_id[db_row["id_especie"]] = db_row
                matched_aliases.append(alias)

        supported_values: dict[str, str] = {}
        for output_field, source_column in SOURCE_FIELDS.items():
            if output_field == "Nome_Cientifico":
                value, _, _ = consensus(source_rows, source_column, 3)
            else:
                value, _, _ = resolve_classification(
                    source_rows,
                    source_column,
                    partial_taxon,
                    partial_rank,
                )
            if value:
                supported_values[output_field] = value
        accepted_name_values = {
            _normalized(row["Nome aceito"])
            for _, row in source_rows[source_rows["Aceito"].map(is_true)].iterrows()
            if raw_text(row["Nome aceito"])
        }
        external_conflicts = ["Nome_Cientifico"] if len(accepted_name_values) > 1 else []

        db_row = next(iter(matched_by_id.values())) if len(matched_by_id) == 1 else None
        updates: list[str] = []
        local_conflicts: list[str] = []
        correspondence = "N.A."
        if db_row:
            correspondence = (
                "NOME_DIRETO"
                if _normalized(db_row["nome_cientifico"]) == _normalized(final_name) else
                "ALIAS_OU_SINONIMO"
            )
            final_group = raw_text(final_row.get("Grupo_Biologico"))
            db_group = raw_text(db_row.get("grupo_biologico"))
            if final_group and db_group and _normalized(final_group) != _normalized(db_group):
                local_conflicts.append("Grupo_Biologico")

            for output_field, db_field in DB_FIELDS.items():
                proposed = (
                    accepted_name
                    if output_field == "Nome_Cientifico" and accepted_name else
                    raw_text(final_row.get(output_field))
                )
                current = raw_text(db_row.get(db_field))
                if not proposed or _normalized(proposed) == _normalized(current):
                    continue
                supported = supported_values.get(output_field, "")
                if professional_override and output_field != "Nome_Cientifico":
                    updates.append(output_field)
                elif supported and _normalized(supported) == _normalized(proposed):
                    updates.append(output_field)

        if len(matched_by_id) > 1:
            action = "REVISAR_CONFLITO"
            local_conflicts.append("MULTIPLOS_REGISTROS_NO_BANCO")
        elif local_conflicts:
            action = "REVISAR_CONFLITO"
        elif db_row and updates:
            action = "ATUALIZAR_TAXONOMIA"
        elif db_row:
            action = "REUTILIZAR_REGISTRO"
        else:
            action = "INSERIR_NOVA"

        status = pending_status.get(_normalized(input_name), "APROVADO")
        ready = "NAO" if action == "REVISAR_CONFLITO" or status == "SEM_RETORNO" else "SIM"
        conflicts = sorted(set(external_conflicts + local_conflicts))
        observation = (
            "Banco consultado somente para conciliacao; fontes externas definem a proposta."
        )
        if professional_override:
            observation += " Classificacao confirmada diretamente pelo profissional."
        if external_conflicts:
            observation += " Conflito de nome preserva o valor profissional."
        reconciled.append(
            {
                "Nome_planilha": final_name,
                "Nome_aceito_consenso": accepted_name or "N.A.",
                "Acao_proposta": action,
                "Pronto_para_migrar": ready,
                "id_especie_banco": db_row["id_especie"] if db_row else "N.A.",
                "Nome_no_banco": raw_text(db_row.get("nome_cientifico")) if db_row else "N.A.",
                "Correspondencia": correspondence,
                "Grupo_no_banco": raw_text(db_row.get("grupo_biologico")) if db_row else "N.A.",
                "Campos_para_atualizar": ", ".join(updates) or "N.A.",
                "Conflitos": ", ".join(conflicts) or "N.A.",
                "Aliases_consultados": " | ".join(sorted(aliases)),
                "Status_taxonomico": status or "N.A.",
                "Observacao": observation,
            }
        )

    workbook = load_workbook(args.input)
    if "CONCILIACAO_BANCO" in workbook.sheetnames:
        del workbook["CONCILIACAO_BANCO"]
    worksheet = workbook.create_sheet("CONCILIACAO_BANCO")
    worksheet.append(OUTPUT_COLUMNS)
    for row in reconciled:
        worksheet.append([row[column] for column in OUTPUT_COLUMNS])
    style_sheet(worksheet)
    workbook.save(args.input)

    counts: dict[str, int] = defaultdict(int)
    for row in reconciled:
        counts[row["Acao_proposta"]] += 1
    payload = {
        "created_at": datetime.now(timezone.utc).isoformat(),
        "database_mode": "read_only",
        "counts": dict(sorted(counts.items())),
        "not_ready": sum(row["Pronto_para_migrar"] == "NAO" for row in reconciled),
        "rows": reconciled,
    }
    args.output_json.parent.mkdir(parents=True, exist_ok=True)
    args.output_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"counts": payload["counts"], "not_ready": payload["not_ready"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
