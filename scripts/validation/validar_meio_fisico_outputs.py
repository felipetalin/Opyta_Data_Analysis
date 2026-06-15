#!/usr/bin/env python
"""Post-run audit for Meio Fisico XLSX outputs.

Creates a compact JSON manifest with input hashes, git context and checks that
B2, B4 and B11 agree on violation counts.
"""

from __future__ import annotations

import hashlib
import json
import os
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

REPO_ROOT = next(
    parent for parent in Path(__file__).resolve().parents
    if (parent / "src" / "opyta_analysis").exists()
)
SRC_ROOT = REPO_ROOT / "src"
if str(SRC_ROOT) not in sys.path:
    sys.path.insert(0, str(SRC_ROOT))

from opyta_analysis.meio_fisico.rules import latest_generated  # noqa: E402


CLIENT_ROOT = Path(os.environ.get(
    "OPYTA_MF_CLIENT_ROOT",
    r"G:/Meu Drive/Opyta/Clientes/Clientes/Clientes/Ferreira Rocha/SAM Metais/Produtos",
))
SRC_RESULTADOS = CLIENT_ROOT / "Migra\u00e7\u00e3o" / "F\u00edsico" / "Resultados_Meio_Fisico.xlsx"
SRC_CADASTRO = CLIENT_ROOT / "Migra\u00e7\u00e3o" / "F\u00edsico" / "cadastro_parametros_opyta.xlsx"
OUT_ROOT = CLIENT_ROOT / "Resultados" / "Meio_f\u00edsico"

RULESET_VERSION = "meio_fisico_gold_v1.1"

MATRIZES = {
    "\u00c1gua Superficial": {
        "sub": "Superficial",
        "b2": "01_Conformidade_Agua_Superficial.xlsx",
    },
    "\u00c1gua Subterr\u00e2nea": {
        "sub": "Subterr\u00e2nea",
        "b2": "01_Conformidade_Agua_Subterranea.xlsx",
    },
    "Sedimento": {
        "sub": "Sedimentos",
        "b2": "01_Conformidade_Sedimento.xlsx",
    },
}


def _run_git(args: list[str]) -> str | None:
    try:
        result = subprocess.run(
            ["git", *args],
            cwd=REPO_ROOT,
            text=True,
            capture_output=True,
            check=False,
        )
    except OSError:
        return None
    if result.returncode != 0:
        return None
    return result.stdout.strip()


def _sha256(path: Path) -> str | None:
    if not path.exists():
        return None
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _read_pct(path: Path, sheet_name: str | int = 0) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name=sheet_name)
    if "N_violacoes" not in df.columns and "N_Violacoes" in df.columns:
        df = df.rename(columns={"N_Violacoes": "N_violacoes"})
    if "N_amostras" not in df.columns and "N_Amostras" in df.columns:
        df = df.rename(columns={"N_Amostras": "N_amostras"})
    return df


def _pct_records(df: pd.DataFrame) -> dict[str, dict[str, float]]:
    out: dict[str, dict[str, float]] = {}
    for _, row in df.iterrows():
        param = str(row.get("Parametro", row.get("Par\u00e2metro", ""))).strip()
        if not param:
            continue
        out[param] = {
            "n_amostras": float(row.get("N_amostras", 0) or 0),
            "n_violacoes": float(row.get("N_violacoes", 0) or 0),
            "pct_violacao": float(row.get("Pct_Violacao", 0) or 0),
        }
    return out


def _count_pink_cells(path: Path) -> int:
    wb = load_workbook(path, data_only=True)
    ws = wb["Conformidade"] if "Conformidade" in wb.sheetnames else wb[wb.sheetnames[0]]
    total = 0
    for row in ws.iter_rows():
        for cell in row:
            fill = cell.fill
            if not fill or not fill.fgColor:
                continue
            rgb = fill.fgColor.rgb
            if rgb in ("FFFFC7CE", "00FFC7CE"):
                total += 1
    return total


def _compare_pct(b4: pd.DataFrame, b11: pd.DataFrame) -> list[str]:
    errors: list[str] = []
    b4r = _pct_records(b4)
    b11r = _pct_records(b11)
    if set(b4r) != set(b11r):
        missing_b11 = sorted(set(b4r) - set(b11r))
        extra_b11 = sorted(set(b11r) - set(b4r))
        if missing_b11:
            errors.append(f"missing_in_b11={missing_b11[:10]}")
        if extra_b11:
            errors.append(f"extra_in_b11={extra_b11[:10]}")
    for param in sorted(set(b4r) & set(b11r)):
        a = b4r[param]
        b = b11r[param]
        for key in ("n_amostras", "n_violacoes"):
            if int(round(a[key])) != int(round(b[key])):
                errors.append(f"{param}: {key} B4={a[key]} B11={b[key]}")
        if abs(a["pct_violacao"] - b["pct_violacao"]) > 1e-6:
            errors.append(f"{param}: pct B4={a['pct_violacao']} B11={b['pct_violacao']}")
    return errors


def audit_matrix(matriz: str, cfg: dict[str, str]) -> dict[str, Any]:
    subdir = OUT_ROOT / cfg["sub"]
    b2_path = latest_generated(subdir / cfg["b2"])
    b4_path = latest_generated(subdir / "04_Pct_Violacao.xlsx")
    b11_path = latest_generated(subdir / "11_Sintese_Executiva.xlsx")

    result: dict[str, Any] = {
        "matriz": matriz,
        "subdir": str(subdir),
        "files": {
            "b2": str(b2_path),
            "b4": str(b4_path),
            "b11": str(b11_path),
        },
        "checks": [],
        "errors": [],
    }

    missing = [str(p) for p in (b2_path, b4_path, b11_path) if not p.exists()]
    if missing:
        result["errors"].append(f"missing_files={missing}")
        return result

    b4 = _read_pct(b4_path)
    b11 = _read_pct(b11_path, sheet_name="Pct_Violacao")
    b4_viol_total = int(round(float(b4["N_violacoes"].fillna(0).sum())))
    b11_viol_total = int(round(float(b11["N_violacoes"].fillna(0).sum())))
    b2_pink = _count_pink_cells(b2_path)

    result["summary"] = {
        "b2_violation_cells": b2_pink,
        "b4_violation_total": b4_viol_total,
        "b11_violation_total": b11_viol_total,
        "b4_violated_params": int((b4["N_violacoes"].fillna(0) > 0).sum()),
    }

    if b2_pink != b4_viol_total:
        result["errors"].append(f"B2_vs_B4_total: B2={b2_pink} B4={b4_viol_total}")
    if b4_viol_total != b11_viol_total:
        result["errors"].append(f"B4_vs_B11_total: B4={b4_viol_total} B11={b11_viol_total}")
    result["errors"].extend(_compare_pct(b4, b11))

    if not result["errors"]:
        result["checks"].append("B2_B4_B11_violation_counts_match")
    return result


def main() -> int:
    OUT_ROOT.mkdir(parents=True, exist_ok=True)
    matrices = [audit_matrix(matriz, cfg) for matriz, cfg in MATRIZES.items()]
    errors = [err for item in matrices for err in item.get("errors", [])]
    status = "OK" if not errors else "ERROR"
    manifest = {
        "generated_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "status": status,
        "ruleset_version": RULESET_VERSION,
        "client_root": str(CLIENT_ROOT),
        "output_root": str(OUT_ROOT),
        "git": {
            "branch": _run_git(["branch", "--show-current"]),
            "commit": _run_git(["rev-parse", "--short", "HEAD"]),
            "dirty": bool(_run_git(["status", "--short"])),
        },
        "inputs": {
            "resultados": {"path": str(SRC_RESULTADOS), "sha256": _sha256(SRC_RESULTADOS)},
            "cadastro": {"path": str(SRC_CADASTRO), "sha256": _sha256(SRC_CADASTRO)},
        },
        "matrices": matrices,
    }

    out = OUT_ROOT / "12_Auditoria_Execucao.json"
    out.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"[audit] {status} -> {out}")
    if errors:
        for err in errors[:20]:
            print(f"  [ERRO] {err}")
    return 0 if status == "OK" else 1


if __name__ == "__main__":
    raise SystemExit(main())
