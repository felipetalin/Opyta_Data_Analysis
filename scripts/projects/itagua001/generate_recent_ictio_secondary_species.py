"""Gera a lista secundária recente de espécies de ictiofauna do ITAGUA001.

O recorte é calculado automaticamente a partir dos dois anos mais recentes
presentes em ``biota_analise_consolidada``. A taxonomia e os status são
obtidos da tabela ``especies`` do Supabase do Opyta.
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


REPO_ROOT = Path(__file__).resolve().parents[3]
if str(REPO_ROOT / "src") not in sys.path:
    sys.path.insert(0, str(REPO_ROOT / "src"))

from opyta_analysis.supabase_client import get_client, paginate  # noqa: E402


PROJECT_ID = 165
PROJECT_CODE = "ITAGUA001"
GROUP = "Ictiofauna"
DEFAULT_ENV = Path(r"G:\Meu Drive\Opyta\Opyta_Data\.env")
DEFAULT_OUTPUT = Path(
    r"G:\Meu Drive\Opyta\Clientes\Clientes\Clientes\Itatiaia\Guanhães Energia"
    r"\Plano de resgate\PT Fauna e Ictiofauna\Dados\_secundario\_ictio"
    r"\ITAGUA001_lista_especies_ictiofauna_ultimos_dois_anos.xlsx"
)

YEAR_RE = re.compile(r"-(20\d{2})-")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
TITLE_FILL = PatternFill("solid", fgColor="0F6B5B")
WHITE_FONT = Font(color="FFFFFF", bold=True)
THIN_GRAY = Side(style="thin", color="D9E2F3")


def _year(campaign: object) -> int | None:
    match = YEAR_RE.search(str(campaign or ""))
    return int(match.group(1)) if match else None


def _clean(value: object, default: str = "Não informado") -> str:
    if value is None:
        return default
    text = str(value).strip()
    return text if text else default


def _normalize_origin(value: object) -> str:
    text = _clean(value)
    normalized = text.lower().replace("ã", "a")
    if "nao nativ" in normalized or "não nativ" in text.lower():
        return "Não nativa"
    if "nativ" in normalized:
        return "Nativa"
    return text


def load_recent_data(env_file: Path) -> tuple[list[dict], list[dict], list[int]]:
    sb = get_client(str(env_file))
    occurrences = paginate(
        sb,
        "biota_analise_consolidada",
        filters={"codigo_interno_opyta": PROJECT_CODE, "grupo_biologico": GROUP},
        select="nome_cientifico,nome_campanha,nome_empreendimento",
    )
    years = sorted({_year(row.get("nome_campanha")) for row in occurrences if _year(row.get("nome_campanha"))})
    if len(years) < 2:
        raise RuntimeError("O projeto não possui dois anos de dados de ictiofauna para o recorte solicitado.")
    selected_years = years[-2:]
    recent = [row for row in occurrences if _year(row.get("nome_campanha")) in selected_years]
    species_names = sorted({_clean(row.get("nome_cientifico"), "") for row in recent if row.get("nome_cientifico")})

    catalog = paginate(
        sb,
        "especies",
        filters={"grupo_biologico": GROUP},
        select=(
            "nome_cientifico,reino,filo,classe,ordem,familia,genero,"
            "status_ameaca_global,status_ameaca_nacional,status_copam,cites,origem"
        ),
    )
    catalog_by_name = {row.get("nome_cientifico"): row for row in catalog}
    grouped: dict[str, dict[str, set[str]]] = defaultdict(lambda: {"campanhas": set(), "empreendimentos": set()})
    for row in recent:
        name = row.get("nome_cientifico")
        if name:
            grouped[name]["campanhas"].add(_clean(row.get("nome_campanha"), ""))
            grouped[name]["empreendimentos"].add(_clean(row.get("nome_empreendimento"), ""))

    records = []
    for name in species_names:
        taxon = catalog_by_name.get(name, {})
        campaigns = sorted(grouped[name]["campanhas"])
        enterprises = sorted(grouped[name]["empreendimentos"])
        records.append(
            {
                "Nome científico": name,
                "Reino": _clean(taxon.get("reino")),
                "Filo": _clean(taxon.get("filo")),
                "Classe": _clean(taxon.get("classe")),
                "Ordem": _clean(taxon.get("ordem")),
                "Família": _clean(taxon.get("familia")),
                "Gênero": _clean(taxon.get("genero")),
                "Origem": _normalize_origin(taxon.get("origem")),
                "Status global (IUCN)": _clean(taxon.get("status_ameaca_global")),
                "Status nacional": _clean(taxon.get("status_ameaca_nacional")),
                "Status estadual (COPAM/MG)": _clean(taxon.get("status_copam")),
                "CITES": _clean(taxon.get("cites")),
                "Nº de campanhas": len(campaigns),
                "Campanhas": "; ".join(campaigns),
                "Empreendimentos": "; ".join(enterprises),
                "Fonte das ocorrências": "Supabase Opyta — biota_analise_consolidada",
                "Fonte da taxonomia/status": "Supabase Opyta — especies",
            }
        )

    campaign_rows = []
    for campaign in sorted({row.get("nome_campanha") for row in recent if row.get("nome_campanha")}):
        campaign_rows.append(
            {
                "Campanha": campaign,
                "Ano": _year(campaign),
                "Empreendimentos": "; ".join(sorted({_clean(r.get("nome_empreendimento"), "") for r in recent if r.get("nome_campanha") == campaign})),
                "Riqueza registrada": len({r.get("nome_cientifico") for r in recent if r.get("nome_campanha") == campaign and r.get("nome_cientifico")}),
            }
        )
    return records, campaign_rows, selected_years


def _write_table(ws, headers: list[str], rows: list[dict], table_name: str) -> None:
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header) for header in headers])
    header = ws[1]
    for cell in header:
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    if rows:
        table = Table(displayName=table_name, ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False)
        ws.add_table(table)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN_GRAY)
    for index, header_text in enumerate(headers, start=1):
        values = [str(ws.cell(row=row, column=index).value or "") for row in range(1, ws.max_row + 1)]
        width = min(max(len(header_text) + 2, max(map(len, values)) + 2), 48)
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.sheet_view.showGridLines = False


def build_workbook(records: list[dict], campaigns: list[dict], years: list[int], output: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Lista de espécies"
    headers = list(records[0].keys()) if records else ["Nome científico"]
    _write_table(ws, headers, records, "ListaEspeciesIctio")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["N"].width = 44
    ws.column_dimensions["O"].width = 42

    ws_campaigns = wb.create_sheet("Campanhas")
    _write_table(ws_campaigns, list(campaigns[0].keys()), campaigns, "CampanhasRecentes")

    ws_meta = wb.create_sheet("Metadados e texto")
    title = f"ITAGUA001 — Lista secundária de ictiofauna ({years[0]}–{years[1]})"
    description = (
        f"A presente base de dados secundários reúne o inventário de ictiofauna registrado no projeto "
        f"ITAGUA001, considerando exclusivamente as campanhas realizadas nos anos de {years[0]} e {years[1]}. "
        "Para cada espécie foram organizadas as classificações taxonômicas disponíveis — Reino, Filo, Classe, "
        "Ordem, Família e Gênero —, além da origem e dos status de conservação global, nacional e estadual, "
        "quando cadastrados. A listagem contempla também as campanhas e os empreendimentos em que cada táxon "
        "foi registrado, permitindo sua utilização como subsídio técnico para o Plano de Resgate de Fauna e "
        "Ictiofauna da Guanhães Energia."
    )
    metadata = [
        ("Projeto", PROJECT_CODE),
        ("ID do projeto", PROJECT_ID),
        ("Grupo biológico", GROUP),
        ("Recorte temporal", f"{years[0]}–{years[1]} (dois anos mais recentes disponíveis no banco)"),
        ("Número de campanhas", len(campaigns)),
        ("Número de espécies", len(records)),
        ("Fonte das ocorrências", "Supabase Opyta — view biota_analise_consolidada"),
        ("Fonte da taxonomia e status", "Supabase Opyta — tabela especies"),
        ("Gerado em", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Nota de uso", "Os campos 'Não informado' indicam ausência de valor cadastrado no banco; não equivalem a categoria de conservação."),
    ]
    ws_meta.merge_cells("A1:B1")
    ws_meta["A1"] = title
    ws_meta["A1"].fill = TITLE_FILL
    ws_meta["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    ws_meta["A1"].alignment = Alignment(vertical="center")
    ws_meta.row_dimensions[1].height = 28
    ws_meta["A3"] = "Texto descritivo"
    ws_meta["A3"].fill = SUBHEADER_FILL
    ws_meta["A3"].font = Font(bold=True, color="1F4E78")
    ws_meta.merge_cells("A4:B7")
    ws_meta["A4"] = description
    ws_meta["A4"].alignment = Alignment(vertical="top", wrap_text=True)
    ws_meta["A9"] = "Campo"
    ws_meta["B9"] = "Informação"
    for cell in ws_meta[9]:
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
    for key, value in metadata:
        ws_meta.append([key, value])
    for row in ws_meta.iter_rows(min_row=10, max_row=ws_meta.max_row):
        row[0].font = Font(bold=True, color="1F4E78")
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN_GRAY)
    ws_meta.column_dimensions["A"].width = 30
    ws_meta.column_dimensions["B"].width = 100
    ws_meta.sheet_view.showGridLines = False

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def validate(output: Path, expected_species: int, expected_campaigns: int) -> None:
    wb = load_workbook(output, read_only=False, data_only=False)
    required = ["Lista de espécies", "Campanhas", "Metadados e texto"]
    if wb.sheetnames != required:
        raise RuntimeError(f"Abas inesperadas: {wb.sheetnames}")
    if wb["Lista de espécies"].max_row - 1 != expected_species:
        raise RuntimeError("Quantidade de espécies divergente após a gravação.")
    if wb["Campanhas"].max_row - 1 != expected_campaigns:
        raise RuntimeError("Quantidade de campanhas divergente após a gravação.")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A")):
                    raise RuntimeError(f"Erro de fórmula em {ws.title}!{cell.coordinate}: {cell.value}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--env-file", type=Path, default=DEFAULT_ENV)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    records, campaigns, years = load_recent_data(args.env_file)
    build_workbook(records, campaigns, years, args.output)
    validate(args.output, len(records), len(campaigns))
    print(f"output={args.output}")
    print(f"years={years[0]}-{years[1]}")
    print(f"species={len(records)}")
    print(f"campaigns={len(campaigns)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
