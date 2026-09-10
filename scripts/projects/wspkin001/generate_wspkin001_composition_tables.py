"""Generate taxonomy-review composition tables for WSPKIN001.

This is a deliberately narrow Gate C pilot: one workbook per group, with no
charts or project palette. It reads the consolidated project slice and does
not write to Supabase.
"""

from __future__ import annotations

import argparse
import json
import sys
from collections.abc import Iterable
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import text


PROJECT_CODE = "WSPKIN001"
PROJECT_ID = 211
DEFAULT_DATA_ROOT = Path(r"G:\Meu Drive\Opyta\Opyta_Data")
DEFAULT_OUTPUT_ROOT = Path(r"G:\Meu Drive\Opyta\Clientes\Clientes\Clientes\WSP\BAndeirinhas_Kinross\Resultados")
GROUPS = {
    "fitoplancton": ("Fitoplâncton", "resultados_fitoplancton", "densidade"),
    "zooplancton": ("Zooplâncton", "resultados_zooplancton", "numero_de_individuos"),
    "zoobentos": ("Zoobentos", "resultados_zoobentos", "abundancia"),
}


def connect_engine(data_root: Path):
    sys.path.insert(0, str(data_root))
    from core.engine import get_engine  # noqa: PLC0415

    return get_engine()


def fetch_rows(conn, group: str) -> pd.DataFrame:
    group_name, result_table, measure_column = GROUPS[group]
    query = text(
        f"""
        SELECT c.nome_campanha, p.nome_ponto, sp.reino, sp.filo, sp.classe,
               sp.ordem, sp.familia, sp.genero, sp.nome_cientifico,
               sp.autor_e_ano, sp.bmwp_score, r.{measure_column}::numeric AS valor_registrado,
               COALESCE(r.tipo_amostragem, e.tipo_amostragem, e.tipo_de_amostragem) AS tipo_amostragem
        FROM public.{result_table} r
        JOIN public.esforcos_amostragem e ON e.id_esforco = r.id_esforco
        JOIN public.pontos_coleta p ON p.id_ponto_coleta = e.id_ponto_coleta
        JOIN public.campanhas c ON c.id_campanha = p.id_campanha
        JOIN public.especies sp ON sp.id_especie = r.id_especie
        WHERE p.id_projeto = :project_id AND e.grupo_biologico = :group_name
        ORDER BY sp.filo, sp.classe, sp.ordem, sp.familia, sp.genero, sp.nome_cientifico, c.nome_campanha, p.nome_ponto
        """
    )
    return pd.read_sql(query, conn, params={"project_id": PROJECT_ID, "group_name": group_name})


def joined(values: Iterable[object]) -> str:
    return " e ".join(sorted({str(value).strip() for value in values if pd.notna(value) and str(value).strip()}))


def first_or_na(values: Iterable[object]) -> str:
    return next((str(value).strip() for value in values if pd.notna(value) and str(value).strip()), "N.A.")


def composition_table(rows: pd.DataFrame) -> pd.DataFrame:
    taxonomy = ["reino", "filo", "classe", "ordem", "familia", "genero", "nome_cientifico", "autor_e_ano"]
    aggregations = {field: (field, first_or_na) for field in taxonomy}
    aggregations.update({
        "bmwp": ("bmwp_score", "first"),
        "ocorrencia_campanhas": ("nome_campanha", joined),
        "ocorrencia_pontos": ("nome_ponto", joined),
        "registros": ("nome_cientifico", "size"),
        "tipo_amostragem": ("tipo_amostragem", joined),
    })
    composition = rows.groupby("nome_cientifico", as_index=False).agg(**aggregations)
    composition = composition.rename(columns={
        "reino": "Reino", "filo": "Filo", "classe": "Classe", "ordem": "Ordem", "familia": "Familia",
        "genero": "Genero", "nome_cientifico": "Nome Cientifico", "autor_e_ano": "Autor e Ano",
        "ocorrencia_campanhas": "Ocorrencia (Campanhas)", "ocorrencia_pontos": "Ocorrencia (Pontos)",
        "registros": "Registros", "tipo_amostragem": "Tipo de Amostragem",
        "bmwp": "BMWP",
    })
    return composition.sort_values(["Filo", "Classe", "Ordem", "Familia", "Genero", "Nome Cientifico"], na_position="last").reset_index(drop=True)


def write_workbook(table: pd.DataFrame, rows: pd.DataFrame, group: str, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        table.to_excel(writer, sheet_name="Composicao_Taxonomica", index=False)
        pd.DataFrame([{
            "Projeto": PROJECT_CODE, "Grupo": GROUPS[group][0], "Escopo": "Piloto Gate C - conferência taxonômica",
            "Linhas de resultados": len(rows), "Taxons": len(table), "Campanhas": joined(rows["nome_campanha"]),
            "Gerado em": datetime.now().isoformat(timespec="seconds"),
        }, {
            "Projeto": PROJECT_CODE, "Grupo": GROUPS[group][0], "Escopo": "Sem paleta ou gráfico; tabela neutra para validação",
        }]).to_excel(writer, sheet_name="Metadados", index=False)
        pd.DataFrame(columns=["Nome Cientifico", "Campo", "Decisao Taxonomica", "Observacao"]).to_excel(writer, sheet_name="Conferencia", index=False)
        if group == "fitoplancton":
            audit_dir = Path("outputs/validacoes/wspkin001_fitoplancton_algaebase_r02_20260904")
            audits = sorted(audit_dir.glob("*_auditoria_fitoplancton_algaebase_r02.json"))
            if audits:
                decisions = json.loads(audits[-1].read_text(encoding="utf-8"))["decisions"]
                decision_table = pd.DataFrame(decisions).rename(columns={
                    "id_especie": "ID Especie",
                    "nome_cientifico": "Nome Cientifico",
                    "reino_anterior": "Reino Anterior",
                    "filo_anterior": "Filo Anterior",
                    "reino_algaebase": "Reino AlgaeBase",
                    "filo_algaebase": "Filo AlgaeBase",
                    "alterado": "Alterado",
                })
                decision_table["Referencia"] = "AlgaeBase"
                decision_table["Decisao"] = "Padrao definitivo WSPKIN001 aprovado em 2026-09-04"
                decision_table.to_excel(writer, sheet_name="Decisoes_Taxonomicas", index=False)
        elif group == "zooplancton":
            audit_dir = Path("outputs/validacoes/wspkin001_zooplancton_gbif_r03_20260904")
            audits = sorted(audit_dir.glob("*_aplicacao_gbif_zooplancton_r03.json"))
            if audits:
                decisions = json.loads(audits[-1].read_text(encoding="utf-8"))["decisions"]
                flat = []
                for item in decisions:
                    row = {
                        "ID Especie": item["id_especie"], "Registros WSP": item["registros_wsp"],
                        "GBIF usageKey": item["gbif_usage_key"], "GBIF Match": item["gbif_match_type"],
                        "GBIF Confianca": item["gbif_confidence"], "Acao": item["acao"],
                        "Alterado": item["alterado"],
                    }
                    for field in ("nome_cientifico", "autor_e_ano", "reino", "filo", "classe", "ordem", "familia", "genero"):
                        label = field.replace("_", " ").title()
                        row[f"{label} Anterior"] = item["antes"].get(field)
                        row[f"{label} GBIF"] = item["depois"].get(field)
                    flat.append(row)
                decision_table = pd.DataFrame(flat)
                decision_table["Referencia"] = "GBIF Backbone Taxonomy"
                decision_table["Decisao"] = "Padrao definitivo WSPKIN001 aprovado em 2026-09-04"
                decision_table.to_excel(writer, sheet_name="Decisoes_Taxonomicas", index=False)

    from openpyxl import load_workbook
    workbook = load_workbook(output_path)
    header_fill = PatternFill("solid", fgColor="595959")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 30
        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column) + 2
            sheet.column_dimensions[column[0].column_letter].width = min(max(width, 14), 42)
    workbook.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate WSPKIN001 taxonomy composition review tables.")
    parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT_ROOT)
    parser.add_argument("--opyta-data-root", type=Path, default=DEFAULT_DATA_ROOT)
    parser.add_argument("--groups", nargs="+", choices=sorted(GROUPS), default=sorted(GROUPS))
    args = parser.parse_args()

    manifest_path = args.output_root / "manifesto_gate_c_parcial_composicao_taxonomica_wspkin001.json"
    if manifest_path.exists():
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest.setdefault("groups", {})
    else:
        manifest = {"project": PROJECT_CODE, "scope": "Gate C partial - taxonomy composition review", "groups": {}}
    engine = connect_engine(args.opyta_data_root)
    try:
        with engine.connect() as conn:
            for group in args.groups:
                rows = fetch_rows(conn, group)
                if rows.empty:
                    raise RuntimeError(f"Nenhum resultado encontrado para {GROUPS[group][0]}.")
                table = composition_table(rows)
                if group == "zoobentos":
                    output_path = args.output_root / "01_tabela_composicao_taxonomica_zoobentos.xlsx"
                else:
                    output_path = args.output_root / group.capitalize() / f"01_tabela_composicao_{group}.xlsx"
                if output_path.exists():
                    raise FileExistsError(f"Saída existente preservada: {output_path}")
                write_workbook(table, rows, group, output_path)
                manifest["groups"][group] = {"output": str(output_path), "result_rows": len(rows), "taxa": len(table), "campaigns": sorted(rows["nome_campanha"].dropna().unique().tolist())}
    finally:
        engine.dispose()

    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
