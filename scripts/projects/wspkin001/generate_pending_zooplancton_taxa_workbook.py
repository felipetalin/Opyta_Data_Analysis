"""Generate the WSPKIN001 zooplankton pending-registration workbook."""

from __future__ import annotations

from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


SOURCE = Path(
    r"G:\Meu Drive\Opyta\Clientes\Clientes\Clientes\WSP\BAndeirinhas_Kinross\Migração de dados\Resultados_Migração_zooPLAN.xlsx"
)
OUTPUT = SOURCE.parent / "Cadastro_Especies_WSPKIN001_Zooplancton.xlsx"
NEW_TAXA = [
    "Arcella dentata",
    "Brachionus patulus",
    "Ciliado NI",
    "Collurella minima",
    "Difflugia kempny",
]
HEADERS = [
    "Nome_Cientifico", "Grupo_Biologico", "Situacao_Cadastro", "Reino", "Filo",
    "Classe", "Ordem", "Familia", "Genero", "Autor_e_Ano", "Fonte_Taxonomica",
    "Observacao", "Status_Preenchimento", "Registros_na_Fonte",
]


def main() -> None:
    if OUTPUT.exists():
        raise FileExistsError(f"Refusing to overwrite existing workbook: {OUTPUT}")
    source_book = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
    source_sheet = source_book["Resultados_Zooplancton"]
    counts = Counter(" ".join(str(row[7]).split()) for row in source_sheet.iter_rows(min_row=2, values_only=True) if row[7])

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Taxons_Novos"
    sheet.append(HEADERS)
    for name in NEW_TAXA:
        sheet.append([name, "Zooplâncton", "novo_no_supabase", "", "", "", "", "", "", "", "", "", "Pendente", counts[name]])

    notes = workbook.create_sheet("Orientacoes")
    notes.append(["WSPKIN001 — Zooplâncton — Pendências de Cadastro"])
    notes.append(["Preencha os campos taxonômicos e a fonte. Não altere Nome_Cientifico ou Situacao_Cadastro."])
    notes.append(["Após o preenchimento, devolver esta planilha para revalidação e aplicação no Supabase antes do Gate B."])

    header = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 34
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(max(max(len(str(c.value or "")) for c in column) + 2, 16), 34)
    notes.column_dimensions["A"].width = 120
    notes["A1"].font = Font(bold=True, color="FFFFFF")
    notes["A1"].fill = header
    for cell in notes["A"]:
        cell.alignment = Alignment(wrap_text=True)

    workbook.save(OUTPUT)
    print(f"output={OUTPUT}")
    print(f"new_taxa={len(NEW_TAXA)}")


if __name__ == "__main__":
    main()
