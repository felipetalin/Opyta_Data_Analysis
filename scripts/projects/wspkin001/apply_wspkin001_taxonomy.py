"""Apply approved WSPKIN001 phytoplankton and zooplankton taxonomy to Supabase."""

from __future__ import annotations

import os
from pathlib import Path

import openpyxl
from dotenv import load_dotenv
from sqlalchemy import bindparam, create_engine, text


ROOT = Path(r"G:\Meu Drive\Opyta\Clientes\Clientes\Clientes\WSP\BAndeirinhas_Kinross\Migração de dados")
BOOKS = (
    (ROOT / "Cadastro_Especies_WSPKIN001_Fitoplancton.xlsx", "Fitoplâncton"),
    (ROOT / "Cadastro_Especies_WSPKIN001_Zooplancton.xlsx", "Zooplâncton"),
)
FIELDS = ("reino", "filo", "classe", "ordem", "familia", "genero", "autor_e_ano")


def cell(value: object) -> str:
    if value is None or str(value).strip() in {"", "-"}:
        return "N.A."
    return str(value).strip()


def rows_from_sheet(path: Path, sheet_name: str, group: str) -> tuple[list[dict], list[dict]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    header = [str(value) if value else "" for value in next(sheet.iter_rows(max_row=1, values_only=True))]
    indexes = {name: header.index(name) for name in header}
    new_rows: list[dict] = []
    updates: list[dict] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not values[indexes["Nome_Cientifico"]]:
            continue
        payload = {
            "nome_cientifico": cell(values[indexes["Nome_Cientifico"]]),
            "grupo_biologico": group,
            **{field: cell(values[indexes[field.title().replace('_', '_')]]) for field in ()},
        }
        for field, column in {
            "reino": "Reino", "filo": "Filo", "classe": "Classe", "ordem": "Ordem",
            "familia": "Familia", "genero": "Genero", "autor_e_ano": "Autor_e_Ano",
        }.items():
            payload[field] = cell(values[indexes[column]])
        situation = cell(values[indexes["Situacao_Cadastro"]])
        if situation == "novo_no_supabase":
            new_rows.append(payload)
        elif situation == "completar_cadastro_existente":
            updates.append(payload)
    return new_rows, updates


def main() -> None:
    new_rows: list[dict] = []
    updates: list[dict] = []
    for path, group in BOOKS:
        new, update = rows_from_sheet(path, "Taxons_Novos", group)
        new_rows.extend(new)
        updates.extend(update)
        if "Complementar_Cadastro" in openpyxl.load_workbook(path, read_only=True).sheetnames:
            _, update = rows_from_sheet(path, "Complementar_Cadastro", group)
            updates.extend(update)

    load_dotenv()
    engine = create_engine(os.environ["FISICO_DB_URL"])
    names = [row["nome_cientifico"] for row in new_rows + updates]
    select_names = text("SELECT id_especie, nome_cientifico FROM public.especies WHERE nome_cientifico IN :names").bindparams(bindparam("names", expanding=True))
    insert_sql = text(
        """
        INSERT INTO public.especies
        (nome_cientifico, grupo_biologico, reino, filo, classe, ordem, familia, genero, autor_e_ano)
        VALUES
        (:nome_cientifico, :grupo_biologico, :reino, :filo, :classe, :ordem, :familia, :genero, :autor_e_ano)
        """
    )
    update_sql = text(
        """
        UPDATE public.especies SET
          reino=:reino, filo=:filo, classe=:classe, ordem=:ordem,
          familia=:familia, genero=:genero, autor_e_ano=:autor_e_ano
        WHERE nome_cientifico=:nome_cientifico
        """
    )
    with engine.begin() as conn:
        existing = {row.nome_cientifico for row in conn.execute(select_names, {"names": names})}
        duplicate_new = sorted({row["nome_cientifico"] for row in new_rows} & existing)
        missing_updates = sorted({row["nome_cientifico"] for row in updates} - existing)
        if missing_updates:
            raise RuntimeError(f"missing_updates={missing_updates}")
        insertable_rows = [row for row in new_rows if row["nome_cientifico"] not in existing]
        conn.execute(insert_sql, insertable_rows)
        conn.execute(update_sql, updates)
        final = conn.execute(select_names, {"names": names}).fetchall()
    print(f"inserted={len(insertable_rows)} reused={len(duplicate_new)} updated={len(updates)} confirmed={len(final)}")


if __name__ == "__main__":
    main()
