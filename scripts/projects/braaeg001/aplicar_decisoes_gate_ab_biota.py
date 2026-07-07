from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from sqlalchemy import text


PROJECT_CODE = "BRAAEG001"

TAXONOMY = [
    {
        "grupo_biologico": "Fitoplâncton",
        "reino": "Plantae",
        "filo": "Bacillariophyta",
        "classe": "Bacillariophyceae",
        "ordem": "Eunotiales",
        "familia": "Eunotiaceae",
        "nome_cientifico": "Eunotia pectinalis",
        "genero": "Eunotia",
        "acao": "cadastrar_se_ausente",
    },
    {
        "grupo_biologico": "Fitoplâncton",
        "reino": "Plantae",
        "filo": "Bacillariophyta",
        "classe": "Bacillariophyceae",
        "ordem": "Naviculales",
        "familia": "Pleurosigmataceae",
        "nome_cientifico": "Gyrosigma scalproides",
        "genero": "Gyrosigma",
        "acao": "cadastrar_se_ausente",
    },
    {
        "grupo_biologico": "Fitoplâncton",
        "reino": "Plantae",
        "filo": "Bacillariophyta",
        "classe": "Bacillariophyceae",
        "ordem": "Naviculales",
        "familia": "Diadesmidaceae",
        "nome_cientifico": "Humidophila contenta",
        "genero": "Humidophila",
        "acao": "cadastrar_se_ausente",
    },
    {
        "grupo_biologico": "Fitoplâncton",
        "reino": "Plantae",
        "filo": "Bacillariophyta",
        "classe": "Bacillariophyceae",
        "ordem": "Naviculales",
        "familia": "Pinnulariaceae",
        "nome_cientifico": "Pinnularia gibba",
        "genero": "Pinnularia",
        "acao": "cadastrar_se_ausente",
    },
    {
        "grupo_biologico": "Fitoplâncton",
        "reino": "Plantae",
        "filo": "Cyanobacteria",
        "classe": "Cyanophyceae",
        "ordem": "Synechococcales",
        "familia": "Pseudanabaenaceae",
        "nome_cientifico": "Anagnostidinema sp.",
        "genero": "Anagnostidinema",
        "acao": "cadastrar_se_ausente",
    },
    {
        "grupo_biologico": "Fitoplâncton",
        "reino": "Plantae",
        "filo": "Cyanobacteria",
        "classe": "Cyanophyceae",
        "ordem": "Nostocales",
        "familia": "Scytonemataceae",
        "nome_cientifico": "Scytonemataceae N.I.",
        "genero": None,
        "acao": "cadastrar_se_ausente",
    },
    {
        "grupo_biologico": "Fitoplâncton",
        "reino": "Plantae",
        "filo": "Chlorophyta",
        "classe": "Classe incerta",
        "ordem": "Ordem incerta",
        "familia": "Família incerta",
        "nome_cientifico": "Physolinum sp.",
        "genero": "Physolinum",
        "acao": "cadastrar_se_ausente",
    },
    {
        "grupo_biologico": "Zoobentos",
        "reino": "Animalia",
        "filo": "Arthropoda",
        "classe": "Insecta",
        "ordem": "Ephemeroptera",
        "familia": "Oligoneuriidae",
        "nome_cientifico": "Oligoneuriidae",
        "genero": None,
        "acao": "cadastrar_se_ausente",
    },
    {
        "grupo_biologico": "Zoobentos",
        "reino": "Animalia",
        "filo": "Mollusca",
        "classe": "Bivalvia",
        "ordem": "Veneroida",
        "familia": "Sphaeriidae",
        "nome_cientifico": "Sphaerium sp.",
        "genero": "Sphaerium",
        "acao": "usar_existente_e_renomear_resultado",
        "nome_resultado_anterior": "Sphaerium spp.",
    },
]


def norm_text(value: object) -> str:
    if value is None:
        return ""
    text_value = str(value).replace("\u00a0", " ")
    text_value = unicodedata.normalize("NFKC", text_value)
    text_value = re.sub(r"\s+", " ", text_value).strip().lower()
    return text_value


def backup_file(path: Path, stamp: str) -> Path:
    backup = path.with_name(f"{path.stem}__backup_pre_gate_ab_{stamp}{path.suffix}")
    if not backup.exists():
        shutil.copy2(path, backup)
    return backup


def find_workbook_with_sheet(input_dir: Path, sheet_name: str) -> Path:
    for path in sorted(input_dir.glob("*.xlsx")):
        if path.name.startswith("~$") or path.name.startswith("2026") or "__backup_" in path.name:
            continue
        try:
            wb = load_workbook(path, read_only=True, data_only=False)
        except Exception:
            continue
        try:
            if sheet_name in wb.sheetnames:
                return path
        finally:
            wb.close()
    raise FileNotFoundError(f"Nenhum workbook com a aba {sheet_name!r} foi encontrado em {input_dir}.")


def header_map(ws) -> dict[str, int]:
    headers: dict[str, int] = {}
    for cell in ws[1]:
        if cell.value is not None:
            headers[str(cell.value).strip()] = cell.column
    return headers


def fix_ictio_effort(input_dir: Path, stamp: str) -> dict[str, Any]:
    path = find_workbook_with_sheet(input_dir, "Resultados_Ictiofauna")
    backup = backup_file(path, stamp)
    wb = load_workbook(path)
    ws = wb["Metadados_Esforco"]
    headers = header_map(ws)
    tipo_col = headers["Tipo_de_Amostragem"]
    unidade_col = headers["Unidade_Esforco"]

    fixed_rows: list[int] = []
    for row_idx in range(2, ws.max_row + 1):
        tipo = norm_text(ws.cell(row=row_idx, column=tipo_col).value)
        unidade = norm_text(ws.cell(row=row_idx, column=unidade_col).value)
        if tipo in {"m²/100", "m2/100"} and unidade == "quantitativa":
            ws.cell(row=row_idx, column=tipo_col).value = "Quantitativa"
            ws.cell(row=row_idx, column=unidade_col).value = "m²/100"
            fixed_rows.append(row_idx)

    wb.save(path)
    wb.close()
    return {
        "arquivo": str(path),
        "backup": str(backup),
        "linhas_corrigidas": len(fixed_rows),
        "linhas": fixed_rows,
    }


def fix_sphaerium_result(input_dir: Path, stamp: str) -> dict[str, Any]:
    path = find_workbook_with_sheet(input_dir, "Resultados_Zoobentos")
    backup = backup_file(path, stamp)
    wb = load_workbook(path)
    ws = wb["Resultados_Zoobentos"]
    headers = header_map(ws)
    taxon_col = headers["Nome_Cientifico"]

    fixed_rows: list[int] = []
    for row_idx in range(2, ws.max_row + 1):
        current = str(ws.cell(row=row_idx, column=taxon_col).value or "").strip()
        if current == "Sphaerium spp.":
            ws.cell(row=row_idx, column=taxon_col).value = "Sphaerium sp."
            fixed_rows.append(row_idx)

    wb.save(path)
    wb.close()
    return {
        "arquivo": str(path),
        "backup": str(backup),
        "linhas_corrigidas": len(fixed_rows),
        "linhas": fixed_rows,
    }


def connect_engine(opyta_data_root: Path):
    sys.path.insert(0, str(opyta_data_root))
    from core.engine import get_engine  # noqa: PLC0415

    return get_engine()


def apply_taxonomy_to_db(opyta_data_root: Path) -> dict[str, Any]:
    engine = connect_engine(opyta_data_root)
    records = [
        row
        for row in TAXONOMY
        if row["acao"] == "cadastrar_se_ausente"
    ]
    inserted = []
    existing = []

    with engine.begin() as conn:
        for rec in records:
            exists = conn.execute(
                text("SELECT id_especie FROM public.especies WHERE nome_cientifico = :nome"),
                {"nome": rec["nome_cientifico"]},
            ).scalar()
            if exists:
                existing.append({"nome_cientifico": rec["nome_cientifico"], "id_especie": exists})
                continue
            row = {
                "nome_cientifico": rec["nome_cientifico"],
                "grupo_biologico": rec["grupo_biologico"],
                "reino": rec["reino"],
                "filo": rec["filo"],
                "classe": rec["classe"],
                "ordem": rec["ordem"],
                "familia": rec["familia"],
                "genero": rec["genero"],
                "observacoes": f"Cadastro Gate B {PROJECT_CODE}; taxonomia informada pelo usuario em 2026-07-03.",
            }
            new_id = conn.execute(
                text(
                    """
                    INSERT INTO public.especies (
                        nome_cientifico,
                        grupo_biologico,
                        reino,
                        filo,
                        classe,
                        ordem,
                        familia,
                        genero,
                        observacoes
                    )
                    VALUES (
                        :nome_cientifico,
                        :grupo_biologico,
                        :reino,
                        :filo,
                        :classe,
                        :ordem,
                        :familia,
                        :genero,
                        :observacoes
                    )
                    RETURNING id_especie
                    """
                ),
                row,
            ).scalar_one()
            inserted.append({"nome_cientifico": rec["nome_cientifico"], "id_especie": new_id})

    return {"inseridos": inserted, "ja_existiam": existing}


def write_taxonomy_workbook(output_dir: Path, stamp: str, db_result: dict[str, Any] | None) -> Path:
    path = output_dir / f"{stamp}_gate_b_taxonomia_biota_braaeg001.xlsx"
    records = []
    inserted_by_name = {
        row["nome_cientifico"]: row["id_especie"]
        for row in (db_result or {}).get("inseridos", [])
    }
    existing_by_name = {
        row["nome_cientifico"]: row["id_especie"]
        for row in (db_result or {}).get("ja_existiam", [])
    }
    for rec in TAXONOMY:
        row = dict(rec)
        row["id_especie_inserida"] = inserted_by_name.get(rec["nome_cientifico"])
        row["id_especie_existente"] = existing_by_name.get(rec["nome_cientifico"])
        row["origem_decisao"] = "usuario informou hierarquia taxonomica no Control Center em 2026-07-03"
        records.append(row)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame(records).to_excel(writer, sheet_name="decisao_taxonomica", index=False)
        pd.DataFrame((db_result or {}).get("inseridos", [])).to_excel(writer, sheet_name="inseridos_db", index=False)
        pd.DataFrame((db_result or {}).get("ja_existiam", [])).to_excel(writer, sheet_name="ja_existiam_db", index=False)
    return path


def main() -> int:
    parser = argparse.ArgumentParser(description="Apply approved BRAAEG001 biota Gate A/B decisions.")
    parser.add_argument("--input-dir", required=True, type=Path)
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--opyta-data-root", type=Path, default=Path.cwd().parent / "Opyta_Data")
    parser.add_argument("--apply-db", action="store_true")
    args = parser.parse_args()

    output_dir = args.output_dir or args.input_dir
    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%dT%H%M%S")

    ictio_result = fix_ictio_effort(args.input_dir, stamp)
    sphaerium_result = fix_sphaerium_result(args.input_dir, stamp)
    db_result = apply_taxonomy_to_db(args.opyta_data_root) if args.apply_db else None
    taxonomy_path = write_taxonomy_workbook(output_dir, stamp, db_result)

    result = {
        "stamp": stamp,
        "ictio": ictio_result,
        "sphaerium": sphaerium_result,
        "db": db_result,
        "taxonomy_workbook": str(taxonomy_path),
    }
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
