from __future__ import annotations

import json
import math
import re
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
import numpy as np
import pandas as pd
from matplotlib.colors import BoundaryNorm, ListedColormap
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


CLIENT_ROOT = Path(r"G:\Meu Drive\Opyta\Clientes\Clientes\Clientes\Brandt")
PROJECT_DIR = next(path for path in CLIENT_ROOT.iterdir() if path.name.startswith("A&G"))
CONSOLIDATED = (
    PROJECT_DIR
    / "resultados"
    / "Meio_fisico"
    / "migracao"
    / "consolidacao_pos_c02"
    / "20260729T183622Z_consolidado_meio_fisico_braaeg001_pos_c02.xlsx"
)
OUTPUT_DIR = PROJECT_DIR / "resultados" / "Meio_fisico" / "resultados" / "superficial"

CAMPAIGN_LABELS = {
    "C001-2026-02-CH": "Campanha-01-Chuva",
    "C002-2026-06-SC": "Campanha-02-Seca",
}
CAMPAIGN_COLORS = {
    "C001-2026-02-CH": "#1B7F22",
    "C002-2026-06-SC": "#19FF19",
}
VMP_RED = "#F04438"
VIOLATION_SHADE = "#ffcccc"

IQA_PARAM_MAP = {
    "OD": ["Oxigenio Dissolvido In Situ", "Oxigenio Dissolvido"],
    "DBO": ["Demanda Bioquimica de Oxigenio"],
    "COLI": ["Coliformes Termotolerantes", "Escherichia coli"],
    "PH": ["pH In Situ", "pH"],
    "NT": ["Nitrogenio Total"],
    "PT": ["Fosforo Total"],
    "TEMP": ["Temperatura da Amostra"],
    "TURB": ["Turbidez"],
    "ST": ["Solidos Totais"],
}
IQA_WEIGHTS = {
    "OD": 0.17,
    "COLI": 0.15,
    "PH": 0.12,
    "DBO": 0.10,
    "NT": 0.10,
    "PT": 0.10,
    "TEMP": 0.10,
    "TURB": 0.08,
    "ST": 0.08,
}
Q_CURVES = {
    "OD": [(0, 3), (10, 6), (20, 11), (30, 17), (40, 27), (50, 41), (60, 56), (70, 73), (80, 86), (90, 95), (100, 100), (110, 95), (120, 85), (130, 75), (140, 65), (150, 56)],
    "COLI_LOG": [(0, 99), (1, 90), (2, 70), (3, 45), (4, 22), (5, 7), (6, 3)],
    "PH": [(2, 2), (3, 5), (4, 11), (5, 26), (6, 60), (7, 92), (7.5, 95), (8, 88), (9, 50), (10, 22), (11, 5), (12, 2)],
    "DBO": [(0, 99), (1, 90), (2, 80), (3, 71), (4, 64), (5, 58), (8, 41), (10, 32), (15, 20), (20, 12), (30, 5)],
    "NT": [(0, 100), (1, 88), (2, 75), (3, 62), (5, 45), (10, 30), (20, 12), (50, 4), (100, 1)],
    "PT": [(0, 99), (0.1, 88), (0.2, 75), (0.5, 50), (1, 28), (2, 13), (5, 5), (10, 1)],
    "TEMP": [(-10, 50), (-5, 78), (-3, 88), (-1, 93), (0, 93), (1, 92), (3, 78), (5, 56), (8, 30), (10, 18), (15, 5)],
    "TURB": [(0, 97), (5, 84), (10, 72), (25, 47), (50, 28), (75, 15), (100, 8), (150, 3), (200, 2)],
    "ST": [(0, 80), (50, 86), (100, 90), (150, 88), (200, 82), (300, 70), (400, 56), (500, 42), (750, 22), (1000, 12), (1500, 5)],
}


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def fold(value: Any) -> str:
    text = clean_text(value).replace("ę", "e").replace("Ę", "E")
    text = unicodedata.normalize("NFKD", text)
    return "".join(ch for ch in text if not unicodedata.combining(ch))


def norm(value: Any) -> str:
    text = fold(value).lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def fix_ptbr_label(value: Any) -> str:
    out = clean_text(value)
    replacements = {
        "Arsęnio": "Arsênio",
        "Manganęs": "Manganês",
        "Nitrogęnio": "Nitrogênio",
        "Oxigęnio": "Oxigênio",
        "Bioquimica": "Bioquímica",
        "Quimica": "Química",
        "Solidos": "Sólidos",
        "Aluminio": "Alumínio",
        "Fosforo": "Fósforo",
        "Cadmio": "Cádmio",
    }
    for src, dst in replacements.items():
        out = out.replace(src, dst)
    return out


def safe_filename(value: str) -> str:
    text = fold(value)
    text = re.sub(r"[^\w\s-]", "", text)
    return re.sub(r"\s+", "_", text).strip("_")


def format_number(value: float) -> str:
    if pd.isna(value):
        return ""
    if abs(value) >= 100:
        return f"{value:.0f}"
    if abs(value) >= 10:
        return f"{value:.1f}".rstrip("0").rstrip(".")
    return f"{value:.2f}".rstrip("0").rstrip(".")


def point_order(points: list[str]) -> list[str]:
    def key(point: str) -> int:
        raw = point.split("_")[-1]
        return int(raw) if raw.isdigit() else 999

    return sorted(points, key=key)


def load_surface() -> pd.DataFrame:
    df = pd.read_excel(CONSOLIDATED, sheet_name="fisico_analise_consolidada")
    surface = df[df["matriz"].astype(str).str.contains("Superficial", na=False)].copy()
    for col in ["valor_medido", "vmp_357_cl2_min", "vmp_357_cl2_max", "vmp_amonia_dinamico"]:
        surface[col] = pd.to_numeric(surface[col], errors="coerce")
    surface["data_hora_coleta"] = pd.to_datetime(surface["data_hora_coleta"], errors="coerce")
    surface["parametro_display"] = surface["nome_parametro"].map(fix_ptbr_label)
    surface["unidade_display"] = surface["unidade_medida"].map(fix_ptbr_label)
    surface["parametro_norm"] = surface["nome_parametro"].map(norm)
    return surface.sort_values(["nome_parametro", "nome_ponto", "nome_campanha"])


def get_limits(df: pd.DataFrame) -> list[dict[str, Any]]:
    parameter_norm = norm(df["parametro_display"].iloc[0] if "parametro_display" in df.columns else df["nome_parametro"].iloc[0])
    if "manganes dissolvido" in parameter_norm or parameter_norm == "amonia":
        return []
    limits: list[dict[str, Any]] = []
    upper = df["vmp_357_cl2_max"].dropna()
    lower = df["vmp_357_cl2_min"].dropna()
    dynamic = df["vmp_amonia_dinamico"].dropna()
    if not upper.empty:
        limits.append({"kind": "max", "value": float(upper.iloc[0])})
    if not lower.empty:
        limits.append({"kind": "min", "value": float(lower.iloc[0])})
    if not dynamic.empty:
        dyn = float(dynamic.iloc[0])
        if not any(math.isclose(dyn, float(item["value"])) for item in limits):
            limits.append({"kind": "max_dynamic", "value": dyn})
    return limits


def vmp_legend_label(item: dict[str, Any]) -> str:
    kind_label = "mín" if item["kind"] == "min" else "máx"
    return f"VMP - Classe 2 ({kind_label}: {format_number(float(item['value']))})"


def classify_row(row: pd.Series, limits: list[dict[str, Any]]) -> str:
    value = row.get("valor_medido")
    sign = clean_text(row.get("sinal_limite"))
    if pd.isna(value):
        return "Sem valor numérico"
    if not limits:
        return "Sem VMP Classe 2"
    if sign in {"<", "<="}:
        return "Atende"
    for item in limits:
        limit = float(item["value"])
        if item["kind"] in {"max", "max_dynamic"} and float(value) > limit:
            return "Viola"
        if item["kind"] == "min" and float(value) < limit:
            return "Viola"
    return "Atende"


def add_violation_shade(ax, limits: list[dict[str, Any]], y_min: float, y_max: float) -> None:
    lower_limits = [float(item["value"]) for item in limits if item["kind"] == "min"]
    upper_limits = [float(item["value"]) for item in limits if item["kind"] in {"max", "max_dynamic"}]
    if lower_limits:
        ax.axhspan(y_min, max(lower_limits), color=VIOLATION_SHADE, alpha=0.28, zorder=0)
    if upper_limits:
        ax.axhspan(min(upper_limits), y_max, color=VIOLATION_SHADE, alpha=0.28, zorder=0)


def ordered_legend_items(unique_legend: dict[str, Any]) -> tuple[list[Any], list[str]]:
    ordered_labels = [label for label in CAMPAIGN_LABELS.values() if label in unique_legend]
    ordered_labels.extend(label for label in unique_legend if label.startswith("VMP - Classe 2") and label not in ordered_labels)
    ordered_labels.extend(label for label in unique_legend if label not in ordered_labels)
    return [unique_legend[label] for label in ordered_labels], ordered_labels


def y_bounds(values: np.ndarray, limits: list[dict[str, Any]], parameter: str) -> tuple[float, float]:
    clean_values = values[~np.isnan(values)]
    if len(clean_values) == 0:
        clean_values = np.array([0.0])
    y_min = float(np.nanmin(clean_values))
    y_max = float(np.nanmax(clean_values))
    for item in limits:
        y_max = max(y_max, float(item["value"]))
        y_min = min(y_min, float(item["value"]))
    if norm(parameter) == "ph in situ":
        return max(0, y_min - 0.5), y_max + 0.5
    if y_min < 0:
        pad = (y_max - y_min) * 0.15 if y_max != y_min else 1
        return y_min - pad, y_max + pad
    return -0.02 * y_max if y_max else 0, y_max * 1.22 if y_max else 1


def plot_parameter(df_param: pd.DataFrame) -> dict[str, Any]:
    source_param = clean_text(df_param["nome_parametro"].iloc[0])
    display_param = clean_text(df_param["parametro_display"].iloc[0])
    unit = clean_text(df_param["unidade_display"].dropna().iloc[0]) if not df_param["unidade_display"].dropna().empty else ""
    points = point_order(df_param["nome_ponto"].dropna().unique().tolist())
    campaigns = [campaign for campaign in CAMPAIGN_LABELS if campaign in set(df_param["nome_campanha"])]
    pivot = df_param.pivot_table(index="nome_ponto", columns="nome_campanha", values="valor_medido", aggfunc="first").reindex(points)
    limits = get_limits(df_param)
    y_min, y_max = y_bounds(pivot.to_numpy(dtype=float), limits, display_param)

    fig, ax = plt.subplots(figsize=(14, 9), dpi=600)
    x = np.arange(len(points), dtype=float)
    offsets = np.linspace(-0.14, 0.14, len(campaigns)) if len(campaigns) > 1 else [0]
    ax.set_ylim(y_min, y_max)
    add_violation_shade(ax, limits, y_min, y_max)

    for offset, campaign in zip(offsets, campaigns):
        values = pivot[campaign].to_numpy(dtype=float)
        valid = ~np.isnan(values)
        ax.scatter(
            x[valid] + offset,
            values[valid],
            s=76,
            marker="o",
            color=CAMPAIGN_COLORS[campaign],
            edgecolor="#0B4F12",
            linewidth=0.55,
            label=CAMPAIGN_LABELS[campaign],
            zorder=3,
        )
    for item in limits:
        ax.axhline(item["value"], color=VMP_RED, linewidth=2.0, label=vmp_legend_label(item), zorder=2)

    ax.set_xticks(x)
    ax.set_xticklabels(points, rotation=45, ha="right", fontsize=14)
    ax.set_ylabel(display_param if unit == "-" or not unit else f"{display_param} ({unit})", fontsize=18)
    ax.set_xlabel("Ponto", fontsize=16)
    ax.tick_params(axis="y", labelsize=12)
    ax.grid(True, alpha=0.22, linestyle="--", zorder=1)
    formatter = mticker.ScalarFormatter(useMathText=True)
    formatter.set_scientific(True)
    formatter.set_powerlimits((-3, 4))
    ax.yaxis.set_major_formatter(formatter)

    handles, labels = ax.get_legend_handles_labels()
    unique_legend = {label: handle for handle, label in zip(handles, labels)}
    ordered_handles, ordered_labels = ordered_legend_items(unique_legend)
    if ordered_labels:
        ax.legend(
            ordered_handles,
            ordered_labels,
            loc="upper center",
            bbox_to_anchor=(0.5, 1.13),
            ncol=min(4, len(ordered_labels)),
            frameon=False,
            fontsize=13,
        )
    fig.subplots_adjust(left=0.09, right=0.985, top=0.84, bottom=0.18)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out_png = OUTPUT_DIR / f"03_{safe_filename(display_param)}.png"
    out_xlsx = OUTPUT_DIR / f"03_{safe_filename(display_param)}_dados.xlsx"
    fig.savefig(out_png, facecolor="white")
    plt.close(fig)

    export_cols = [
        "nome_ponto",
        "nome_campanha",
        "data_hora_coleta",
        "nome_parametro",
        "parametro_display",
        "sinal_limite",
        "valor_medido",
        "unidade_medida",
        "vmp_357_cl2_min",
        "vmp_357_cl2_max",
        "vmp_amonia_dinamico",
        "status_conformidade",
    ]
    with pd.ExcelWriter(out_xlsx, engine="openpyxl") as writer:
        df_param[export_cols].to_excel(writer, sheet_name="dados_plotados", index=False)
    style_workbook(out_xlsx)

    return {
        "parametro": display_param,
        "parametro_fonte": source_param,
        "png": str(out_png),
        "xlsx": str(out_xlsx),
        "pontos": len(points),
        "campanhas": len(campaigns),
        "vmp": [vmp_legend_label(item) for item in limits],
        "violacoes": int((df_param["status_conformidade"] == "Viola").sum()),
    }


def plot_campaign_panel(df_param: pd.DataFrame) -> dict[str, Any]:
    display_param = clean_text(df_param["parametro_display"].iloc[0])
    unit = clean_text(df_param["unidade_display"].dropna().iloc[0]) if not df_param["unidade_display"].dropna().empty else ""
    points = point_order(df_param["nome_ponto"].dropna().unique().tolist())
    campaigns = [campaign for campaign in CAMPAIGN_LABELS if campaign in set(df_param["nome_campanha"])]
    pivot = df_param.pivot_table(index="nome_ponto", columns="nome_campanha", values="valor_medido", aggfunc="first").reindex(points)
    limits = get_limits(df_param)
    y_min, y_max = y_bounds(pivot.to_numpy(dtype=float), limits, display_param)

    fig, axes = plt.subplots(1, len(campaigns), figsize=(14, 9), dpi=600, sharey=True, squeeze=False)
    axes_flat = list(axes.ravel())
    x = np.arange(len(points), dtype=float)
    for ax, campaign in zip(axes_flat, campaigns):
        ax.set_ylim(y_min, y_max)
        add_violation_shade(ax, limits, y_min, y_max)
        values = pivot[campaign].to_numpy(dtype=float)
        valid = ~np.isnan(values)
        ax.scatter(
            x[valid],
            values[valid],
            s=88,
            marker="o",
            color=CAMPAIGN_COLORS[campaign],
            edgecolor="#0B4F12",
            linewidth=0.65,
            label=CAMPAIGN_LABELS[campaign],
            zorder=3,
        )
        for item in limits:
            ax.axhline(item["value"], color=VMP_RED, linewidth=2.0, label=vmp_legend_label(item), zorder=2)
        ax.set_title(CAMPAIGN_LABELS[campaign], fontsize=17, fontweight="bold", pad=12)
        ax.set_xticks(x)
        ax.set_xticklabels(points, rotation=45, ha="right", fontsize=12)
        ax.set_xlabel("Ponto", fontsize=14)
        ax.tick_params(axis="y", labelsize=12)
        ax.grid(True, alpha=0.22, linestyle="--", zorder=1)

    axes_flat[0].set_ylabel(display_param if unit == "-" or not unit else f"{display_param} ({unit})", fontsize=18)
    handles, labels = [], []
    for ax in axes_flat:
        h, l = ax.get_legend_handles_labels()
        handles.extend(h)
        labels.extend(l)
    unique_legend = {label: handle for handle, label in zip(handles, labels)}
    ordered_handles, ordered_labels = ordered_legend_items(unique_legend)
    fig.legend(
        ordered_handles,
        ordered_labels,
        loc="upper center",
        bbox_to_anchor=(0.5, 0.985),
        ncol=min(4, len(ordered_labels)),
        frameon=False,
        fontsize=12,
    )
    fig.subplots_adjust(left=0.08, right=0.985, top=0.82, bottom=0.18, wspace=0.08)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out_png = OUTPUT_DIR / f"03_painel_{safe_filename(display_param)}.png"
    fig.savefig(out_png, facecolor="white")
    plt.close(fig)
    return {"parametro": display_param, "png": str(out_png)}


def q_interp(key: str, x: float) -> float:
    pts = Q_CURVES[key]
    xs = [p[0] for p in pts]
    qs = [p[1] for p in pts]
    if x <= xs[0]:
        return float(qs[0])
    if x >= xs[-1]:
        return float(qs[-1])
    return float(np.interp(x, xs, qs))


def od_sat_mgl(temp_c: float) -> float:
    return 14.652 - 0.41022 * temp_c + 0.0079910 * temp_c**2 - 0.000077774 * temp_c**3


def iqa_class(value: float) -> str:
    if pd.isna(value):
        return ""
    if value > 79:
        return "Ótima"
    if value > 51:
        return "Boa"
    if value > 36:
        return "Regular"
    if value > 19:
        return "Ruim"
    return "Péssima"


def find_param_series(surface: pd.DataFrame, aliases: list[str]) -> pd.Series:
    for alias in aliases:
        alias_norm = norm(alias)
        matches = surface[surface["parametro_norm"].eq(alias_norm)]
        if not matches.empty:
            return matches.groupby(["nome_ponto", "nome_campanha"])["valor_medido"].mean()
    for alias in aliases:
        alias_norm = norm(alias)
        matches = surface[surface["parametro_norm"].str.contains(alias_norm, na=False)]
        if not matches.empty:
            return matches.groupby(["nome_ponto", "nome_campanha"])["valor_medido"].mean()
    return pd.Series(dtype=float)


def generate_iqa(surface: pd.DataFrame) -> dict[str, Any]:
    series = {key: find_param_series(surface, aliases) for key, aliases in IQA_PARAM_MAP.items()}
    keys = sorted(set().union(*(set(s.index) for s in series.values())), key=lambda item: (str(item[1]), str(item[0])))
    rows = []
    for point, campaign in keys:
        qs: dict[str, float] = {}
        od = series["OD"].get((point, campaign), np.nan)
        temp = series["TEMP"].get((point, campaign), np.nan)
        if pd.notna(od) and pd.notna(temp):
            qs["OD"] = q_interp("OD", float(od) / od_sat_mgl(float(temp)) * 100.0)
        coli = series["COLI"].get((point, campaign), np.nan)
        if pd.notna(coli) and coli > 0:
            qs["COLI"] = q_interp("COLI_LOG", float(np.log10(coli)))
        ph = series["PH"].get((point, campaign), np.nan)
        if pd.notna(ph):
            qs["PH"] = q_interp("PH", float(ph))
        for key in ["DBO", "NT", "PT", "TURB", "ST"]:
            value = series[key].get((point, campaign), np.nan)
            if pd.notna(value):
                qs[key] = q_interp(key, float(value))
        qs["TEMP"] = q_interp("TEMP", 0)
        if not qs:
            continue
        wsum = sum(IQA_WEIGHTS[key] for key in qs)
        iqa = 1.0
        for key, q_value in qs.items():
            iqa *= max(q_value, 1.0) ** (IQA_WEIGHTS[key] / wsum)
        rows.append({
            "Ponto": point,
            "Campanha": CAMPAIGN_LABELS.get(campaign, campaign),
            "IQA": float(iqa),
            "Classe": iqa_class(float(iqa)),
            "N_parametros_usados": len(qs),
            "Parametros_usados": ", ".join(sorted(qs)),
            "Observacao": "Pesos renormalizados conforme parâmetros disponíveis; temperatura avaliada como dT=0.",
        })
    df_out = pd.DataFrame(rows)
    out_xlsx = OUTPUT_DIR / "05_IQA_Tabela.xlsx"
    df_out.to_excel(out_xlsx, index=False)
    style_workbook(out_xlsx)
    out_png = OUTPUT_DIR / "05_IQA_Heatmap.png"
    if not df_out.empty:
        plot_heatmap(df_out, "Ponto", "Campanha", "IQA", out_png, "IQA", [0, 19, 36, 51, 79, 100], ["#8B0000", "#e74c3c", "#f39c12", "#3498db", "#2ecc71"])
    return {"xlsx": str(out_xlsx), "png": str(out_png) if out_png.exists() else None, "amostras": int(len(df_out))}


def to_ug_l(value: float, unit: str) -> float:
    u = norm(unit).replace(" ", "")
    if u in {"ugl", "ugpl", "microgl", "microgpl"}:
        return float(value)
    if u in {"mgl", "mgpl"}:
        return float(value) * 1000.0
    if u in {"ngl", "ngpl"}:
        return float(value) / 1000.0
    return float(value)


def iet_class(value: float) -> str:
    if pd.isna(value):
        return ""
    if value <= 47:
        return "Ultraoligotrófico"
    if value <= 52:
        return "Oligotrófico"
    if value <= 59:
        return "Mesotrófico"
    if value <= 63:
        return "Eutrófico"
    if value <= 67:
        return "Supereutrófico"
    return "Hipereutrófico"


def generate_iet(surface: pd.DataFrame) -> dict[str, Any]:
    pt = surface[surface["parametro_norm"].eq("fosforo total")].copy()
    cl = surface[surface["parametro_norm"].isin(["clorofila a", "clorofila"])]
    rows = []
    if pt.empty or cl.empty:
        out_xlsx = OUTPUT_DIR / "06_IET_Tabela.xlsx"
        out_png = OUTPUT_DIR / "06_IET_Heatmap.png"
        if out_png.exists():
            out_png.unlink()
        pd.DataFrame([{
            "Ponto": "NAO_CALCULADO",
            "Campanha": "NAO_CALCULADO",
            "IET": np.nan,
            "Classe": "NAO_CALCULADO",
            "Observacao": "IET nao calculado: o indice exige Fosforo Total e Clorofila a. Nao calcular com Fosforo Total isolado.",
        }]).to_excel(out_xlsx, index=False)
        style_workbook(out_xlsx)
        return {"xlsx": str(out_xlsx), "png": None, "amostras": 0, "observacao": "Sem Fósforo Total ou Clorofila A."}
    for df_part in [pt, cl]:
        if not df_part.empty:
            df_part["_ug_l"] = [to_ug_l(v, u) for v, u in zip(df_part["valor_medido"], df_part["unidade_medida"])]
    pt_series = pt.groupby(["nome_ponto", "nome_campanha"])["_ug_l"].mean() if not pt.empty else pd.Series(dtype=float)
    cl_series = cl.groupby(["nome_ponto", "nome_campanha"])["_ug_l"].mean() if not cl.empty else pd.Series(dtype=float)
    keys = sorted(set(pt_series.index) | set(cl_series.index), key=lambda item: (str(item[1]), str(item[0])))
    for point, campaign in keys:
        p = pt_series.get((point, campaign), np.nan)
        c = cl_series.get((point, campaign), np.nan)
        iet_p = 10 * (6 - (1.77 - 0.42 * np.log(p)) / np.log(2)) if pd.notna(p) and p > 0 else np.nan
        iet_c = 10 * (6 - (0.92 - 0.34 * np.log(c)) / np.log(2)) if pd.notna(c) and c > 0 else np.nan
        components = [value for value in [iet_p, iet_c] if pd.notna(value)]
        iet = float(np.mean(components)) if components else np.nan
        rows.append({
            "Ponto": point,
            "Campanha": CAMPAIGN_LABELS.get(campaign, campaign),
            "PT_ug_L": p,
            "Clorofila_ug_L": c,
            "IET_PT": iet_p,
            "IET_CL": iet_c,
            "IET": iet,
            "Classe": iet_class(iet),
            "Observacao": "Calculado com os componentes disponíveis; Fósforo Total convertido para µg/L quando informado em mg/L.",
        })
    df_out = pd.DataFrame(rows)
    out_xlsx = OUTPUT_DIR / "06_IET_Tabela.xlsx"
    df_out.to_excel(out_xlsx, index=False)
    style_workbook(out_xlsx)
    out_png = OUTPUT_DIR / "06_IET_Heatmap.png"
    if not df_out.empty and df_out["IET"].notna().any():
        plot_heatmap(df_out, "Ponto", "Campanha", "IET", out_png, "IET", [0, 47, 52, 59, 63, 67, 200], ["#3498db", "#5dade2", "#f1c40f", "#e67e22", "#e74c3c", "#8B0000"])
    return {"xlsx": str(out_xlsx), "png": str(out_png) if out_png.exists() else None, "amostras": int(len(df_out))}


def plot_heatmap(df: pd.DataFrame, index: str, columns: str, values: str, out_png: Path, label: str, bounds: list[float], colors: list[str]) -> None:
    pivot = df.pivot_table(index=index, columns=columns, values=values, aggfunc="mean")
    pivot = pivot.reindex(point_order(pivot.index.tolist()))
    fig, ax = plt.subplots(figsize=(12, 8), dpi=600)
    cmap = ListedColormap(colors)
    norm_obj = BoundaryNorm(bounds, cmap.N)
    im = ax.imshow(pivot.values, cmap=cmap, norm=norm_obj, aspect="auto")
    ax.set_xticks(range(len(pivot.columns)))
    ax.set_xticklabels(pivot.columns, rotation=25, ha="right", fontsize=12)
    ax.set_yticks(range(len(pivot.index)))
    ax.set_yticklabels(pivot.index, fontsize=12)
    for i in range(pivot.shape[0]):
        for j in range(pivot.shape[1]):
            value = pivot.values[i, j]
            if not np.isnan(value):
                ax.text(j, i, f"{value:.0f}", ha="center", va="center", fontsize=10, color="black")
    cbar = fig.colorbar(im, ax=ax)
    cbar.set_label(label)
    ax.set_xlabel("Campanha", fontsize=14)
    ax.set_ylabel("Ponto", fontsize=14)
    fig.tight_layout()
    fig.savefig(out_png, facecolor="white")
    plt.close(fig)


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col in ws.columns:
            values = [clean_text(cell.value) for cell in col]
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max(map(len, values), default=0) + 2, 10), 45)
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def write_tables(surface: pd.DataFrame) -> dict[str, Any]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out_xlsx = OUTPUT_DIR / "01_Conformidade_Agua_Superficial.xlsx"
    cols = [
        "nome_ponto",
        "nome_campanha",
        "data_hora_coleta",
        "parametro_display",
        "sinal_limite",
        "valor_medido",
        "unidade_display",
        "vmp_357_cl2_min",
        "vmp_357_cl2_max",
        "vmp_amonia_dinamico",
        "status_conformidade",
    ]
    summary_param = (
        surface.groupby(["parametro_display", "nome_campanha", "unidade_display"], dropna=False)
        .agg(
            registros=("valor_medido", "size"),
            valores_numericos=("valor_medido", "count"),
            violacoes=("status_conformidade", lambda s: int((s == "Viola").sum())),
            sem_vmp=("status_conformidade", lambda s: int((s == "Sem VMP Classe 2").sum())),
        )
        .reset_index()
    )
    summary_param["percentual_violacao"] = np.where(
        summary_param["valores_numericos"] > 0,
        summary_param["violacoes"] / summary_param["valores_numericos"] * 100,
        np.nan,
    )
    summary_point = (
        surface.groupby(["nome_ponto", "nome_campanha"], dropna=False)
        .agg(
            registros=("valor_medido", "size"),
            valores_numericos=("valor_medido", "count"),
            violacoes=("status_conformidade", lambda s: int((s == "Viola").sum())),
        )
        .reset_index()
    )
    summary_point["percentual_violacao"] = np.where(
        summary_point["valores_numericos"] > 0,
        summary_point["violacoes"] / summary_point["valores_numericos"] * 100,
        np.nan,
    )
    vmp_audit = (
        surface.groupby(["parametro_display", "unidade_display"], dropna=False)
        .agg(
            vmp_min=("vmp_357_cl2_min", "first"),
            vmp_max=("vmp_357_cl2_max", "first"),
            vmp_amonia_dinamico=("vmp_amonia_dinamico", "first"),
            registros=("valor_medido", "size"),
        )
        .reset_index()
    )
    with pd.ExcelWriter(out_xlsx, engine="openpyxl") as writer:
        surface[cols].to_excel(writer, sheet_name="resultados_conformidade", index=False)
        summary_param.to_excel(writer, sheet_name="sintese_parametro", index=False)
        summary_point.to_excel(writer, sheet_name="sintese_ponto", index=False)
        vmp_audit.to_excel(writer, sheet_name="auditoria_vmp", index=False)
    style_workbook(out_xlsx)

    pct_xlsx = OUTPUT_DIR / "04_Pct_Violacao.xlsx"
    summary_param.to_excel(pct_xlsx, index=False)
    style_workbook(pct_xlsx)
    pct_png = OUTPUT_DIR / "04_Pct_Violacao.png"
    plot_violation_percent(summary_param, pct_png)
    return {
        "conformidade": str(out_xlsx),
        "pct_violacao_xlsx": str(pct_xlsx),
        "pct_violacao_png": str(pct_png),
        "violacoes": int((surface["status_conformidade"] == "Viola").sum()),
    }


def plot_violation_percent(summary: pd.DataFrame, out_png: Path) -> None:
    grouped = (
        summary.groupby("parametro_display")
        .agg(violacoes=("violacoes", "sum"), valores_numericos=("valores_numericos", "sum"))
        .reset_index()
    )
    grouped["percentual_violacao"] = np.where(
        grouped["valores_numericos"] > 0,
        grouped["violacoes"] / grouped["valores_numericos"] * 100,
        0,
    )
    grouped = grouped[grouped["violacoes"] > 0].sort_values(["percentual_violacao", "violacoes"], ascending=True)
    fig, ax = plt.subplots(figsize=(14, max(7, 0.38 * max(len(grouped), 1))), dpi=600)
    if grouped.empty:
        ax.text(0.5, 0.5, "Sem violações registradas", ha="center", va="center", fontsize=16)
        ax.axis("off")
    else:
        colors = ["#C0372B" if value >= 50 else "#F04438" for value in grouped["percentual_violacao"]]
        ax.barh(grouped["parametro_display"], grouped["percentual_violacao"], color=colors)
        ax.set_xlabel("Percentual de violação (%)", fontsize=14)
        ax.tick_params(axis="y", labelsize=11)
        ax.tick_params(axis="x", labelsize=11)
        ax.grid(True, axis="x", alpha=0.22, linestyle="--")
        for i, row in enumerate(grouped.itertuples(index=False)):
            ax.text(row.percentual_violacao + 1, i, f"{row.percentual_violacao:.0f}%", va="center", fontsize=10)
        ax.set_xlim(0, max(100, grouped["percentual_violacao"].max() * 1.15))
    fig.tight_layout()
    fig.savefig(out_png, facecolor="white")
    plt.close(fig)


def write_summary(surface: pd.DataFrame, tables: dict[str, Any], figures: list[dict[str, Any]], panels: list[dict[str, Any]], iqa: dict[str, Any], iet: dict[str, Any]) -> dict[str, Any]:
    summary = {
        "timestamp": datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ"),
        "status": "OK",
        "projeto": "BRAAEG001",
        "matriz": "Água Superficial",
        "saida": str(OUTPUT_DIR),
        "fonte": str(CONSOLIDATED),
        "registros": int(len(surface)),
        "pontos": int(surface["nome_ponto"].nunique()),
        "campanhas": int(surface["nome_campanha"].nunique()),
        "parametros": int(surface["nome_parametro"].nunique()),
        "violacoes": tables["violacoes"],
        "figuras_parametros": len(figures),
        "figuras_paineis_campanha": len(panels),
        "organizacao_saida": "Todos os arquivos finais de Agua Superficial foram gravados diretamente na pasta raiz superficial, sem subpastas de teste ou divisao tecnica por tipo de grafico.",
        "tabelas": tables,
        "iqa": iqa,
        "iet": iet,
        "premissas": [
            "Somente Água Superficial.",
            "VMP operacional: Classe 2, com base em CONAMA 357 e DN COPAM-CERH/MG Nº 8/2022.",
            "Legendas das figuras apresentam apenas 'VMP - Classe 2' e o valor correspondente.",
            "Resultados com sinal < ou <= não são classificados como violação e não recebem anotação no ponto.",
            "Faixa vermelha indica região de violação para limites mínimos ou máximos.",
            "Gráficos seguem o template FERSAM aprovado pelo usuário em 2026-07-29.",
        ],
        "figuras": figures,
        "paineis": panels,
    }
    path = OUTPUT_DIR / "12_Auditoria_Execucao.json"
    path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def main() -> int:
    surface = load_surface()
    for _, group in surface.groupby("nome_parametro", sort=False):
        limits = get_limits(group)
        surface.loc[group.index, "status_conformidade"] = group.apply(lambda row: classify_row(row, limits), axis=1)

    tables = write_tables(surface)
    figures = []
    panel_params = set(surface.loc[surface["valor_medido"].notna(), "nome_parametro"].dropna().tolist())
    panels = []
    for param in sorted(panel_params, key=lambda value: norm(value)):
        group = surface[surface["nome_parametro"].eq(param)]
        if not group.empty and group["valor_medido"].notna().any():
            panels.append(plot_campaign_panel(group))

    iqa = generate_iqa(surface)
    iet = generate_iet(surface)
    summary = write_summary(surface, tables, figures, panels, iqa, iet)
    print(json.dumps({k: summary[k] for k in ["status", "registros", "pontos", "campanhas", "parametros", "violacoes", "figuras_parametros", "figuras_paineis_campanha"]}, ensure_ascii=True, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
