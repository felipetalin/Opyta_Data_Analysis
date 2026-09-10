from __future__ import annotations

import importlib.util
import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from dotenv import load_dotenv
from sqlalchemy import text
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = next(parent for parent in Path(__file__).resolve().parents if (parent / "src" / "opyta_analysis").exists())
OPYTA_DATA_ROOT = Path(r"G:\Meu Drive\Opyta\Opyta_Data")
for path in (ROOT / "src", OPYTA_DATA_ROOT):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from core.engine import get_engine  # noqa: E402


SOURCE_APPLIER = ROOT / "scripts" / "projects" / "braaeg001" / "aplicar_migracao_controlada_meio_fisico.py"
spec = importlib.util.spec_from_file_location("braaeg001_aplicar_migracao_controlada_meio_fisico", SOURCE_APPLIER)
if spec is None or spec.loader is None:
    raise RuntimeError(f"Nao foi possivel carregar {SOURCE_APPLIER}")
mf_apply = importlib.util.module_from_spec(spec)
spec.loader.exec_module(mf_apply)


PROJECT_CODE = "BRAAEG001"
CAMPAIGN = "C002-2026-06-SC"
MATRIX = "Água Superficial"
ENV_FILE = Path(r"G:\Meu Drive\Opyta\Opyta_Data\.env")
CLIENT_ROOT = Path(r"G:\Meu Drive\Opyta\Clientes\Clientes\Clientes\Brandt")
PROJECT_DIR = next(path for path in CLIENT_ROOT.iterdir() if path.name.startswith("A&G"))
OUTPUT_DIR = PROJECT_DIR / "resultados" / "Meio_fisico" / "migracao" / "campanha_02_superficial"
DRY_RUN = OUTPUT_DIR / "20260729T182456Z_dry_run_incremental_c02_agua_superficial_braaeg001.xlsx"
LOG_DIR = ROOT / "logs" / "validacao_meio_fisico"

FISICO_INSERT_COLS = mf_apply.FISICO_INSERT_COLS
cell_value = mf_apply.cell_value
insert_rows = mf_apply.insert_rows


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column_cells in ws.columns:
            values = [clean_text(cell.value) for cell in column_cells]
            width = min(max(max((len(value) for value in values), default=0) + 2, 10), 48)
            ws.column_dimensions[get_column_letter(column_cells[0].column)].width = width
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def configure_database_access() -> None:
    repo_env = ROOT / ".env"
    if repo_env.exists():
        load_dotenv(repo_env, override=False)
    if ENV_FILE.exists():
        load_dotenv(ENV_FILE, override=False)
    has_database_url = bool(os.getenv("DATABASE_URL") or os.getenv("SUPABASE_DB_URL"))
    has_db_parts = all(os.getenv(name) for name in ("DB_USER", "DB_PASSWORD", "DB_HOST", "DB_NAME"))
    has_supabase_parts = all(os.getenv(name) for name in ("SUPABASE_DB_USER", "SUPABASE_DB_PASSWORD", "SUPABASE_DB_HOST"))
    if not (has_database_url or has_db_parts or has_supabase_parts):
        fisico_url = os.getenv("FISICO_DB_URL")
        if fisico_url:
            os.environ.setdefault("DATABASE_URL", fisico_url)


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def row_payload(row: pd.Series) -> dict[str, Any]:
    return {col: cell_value(row.get(col), col) for col in FISICO_INSERT_COLS}


def main() -> int:
    ts = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    preview_df = pd.read_excel(DRY_RUN, sheet_name="preview_insert_c02", dtype=object).where(pd.notna, None)
    rows = [row_payload(row) for _, row in preview_df.iterrows()]
    if len(rows) != 588:
        raise RuntimeError(f"Total inesperado no dry-run: {len(rows)}")

    configure_database_access()
    engine = get_engine()
    backup_table = "backup_fisico_braaeg001_before_c02_" + ts.lower()
    manifest: dict[str, Any]
    try:
        with engine.begin() as conn:
            before_total = conn.execute(
                text("SELECT COUNT(*)::int FROM public.fisico_analise_consolidada WHERE codigo_interno_opyta = :code"),
                {"code": PROJECT_CODE},
            ).scalar_one()
            existing_c02 = conn.execute(
                text(
                    """
                    SELECT COUNT(*)::int
                    FROM public.fisico_analise_consolidada
                    WHERE codigo_interno_opyta = :code
                      AND nome_campanha = :campaign
                      AND matriz = :matrix
                    """
                ),
                {"code": PROJECT_CODE, "campaign": CAMPAIGN, "matrix": MATRIX},
            ).scalar_one()
            if existing_c02:
                raise RuntimeError(f"Carga bloqueada: ja existem {existing_c02} registros da C02 no banco.")

            conn.execute(
                text(
                    f"""
                    CREATE TABLE public.{backup_table} AS
                    SELECT *
                    FROM public.fisico_analise_consolidada
                    WHERE codigo_interno_opyta = :code
                    """
                ),
                {"code": PROJECT_CODE},
            )
            backup_rows = conn.execute(text(f"SELECT COUNT(*)::int FROM public.{backup_table}")).scalar_one()
            inserted = insert_rows(conn, "fisico_analise_consolidada", rows, FISICO_INSERT_COLS, batch_size=250)
            after_total = conn.execute(
                text("SELECT COUNT(*)::int FROM public.fisico_analise_consolidada WHERE codigo_interno_opyta = :code"),
                {"code": PROJECT_CODE},
            ).scalar_one()
            inserted_c02 = conn.execute(
                text(
                    """
                    SELECT COUNT(*)::int
                    FROM public.fisico_analise_consolidada
                    WHERE codigo_interno_opyta = :code
                      AND nome_campanha = :campaign
                      AND matriz = :matrix
                    """
                ),
                {"code": PROJECT_CODE, "campaign": CAMPAIGN, "matrix": MATRIX},
            ).scalar_one()
            if inserted != len(rows) or inserted_c02 != len(rows) or after_total != before_total + len(rows):
                raise RuntimeError(
                    "Conferencia pos-insert falhou: "
                    f"inserted={inserted}, inserted_c02={inserted_c02}, before={before_total}, after={after_total}"
                )

            summary = pd.read_sql(
                text(
                    """
                    SELECT matriz, nome_campanha, COUNT(*)::int AS registros,
                           COUNT(DISTINCT nome_ponto)::int AS pontos,
                           COUNT(DISTINCT nome_parametro)::int AS parametros
                    FROM public.fisico_analise_consolidada
                    WHERE codigo_interno_opyta = :code
                    GROUP BY matriz, nome_campanha
                    ORDER BY matriz, nome_campanha
                    """
                ),
                conn,
                params={"code": PROJECT_CODE},
            )

            manifest = {
                "timestamp": ts,
                "status": "APPLIED",
                "project_code": PROJECT_CODE,
                "campaign": CAMPAIGN,
                "matrix": MATRIX,
                "dry_run": str(DRY_RUN),
                "backup_table": backup_table,
                "backup_rows": int(backup_rows),
                "rows_before": int(before_total),
                "rows_inserted": int(inserted),
                "rows_c02_after": int(inserted_c02),
                "rows_after": int(after_total),
                "summary": summary.to_dict(orient="records"),
            }
    finally:
        engine.dispose()

    LOG_DIR.mkdir(parents=True, exist_ok=True)
    log_path = LOG_DIR / f"{ts}_aplicacao_incremental_c02_agua_superficial_braaeg001.json"
    log_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    audit_path = OUTPUT_DIR / f"{ts}_auditoria_pos_incremental_c02_agua_superficial_braaeg001.xlsx"
    with pd.ExcelWriter(audit_path, engine="openpyxl") as writer:
        pd.DataFrame([manifest]).drop(columns=["summary"]).to_excel(writer, sheet_name="manifesto", index=False)
        pd.DataFrame(manifest["summary"]).to_excel(writer, sheet_name="resumo_banco", index=False)
    style_workbook(audit_path)

    print(f"STATUS={manifest['status']}")
    print(f"LOG={log_path}")
    print(f"AUDITORIA={audit_path}")
    print(json.dumps({k: manifest[k] for k in ["rows_before", "rows_inserted", "rows_c02_after", "rows_after", "backup_table"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
