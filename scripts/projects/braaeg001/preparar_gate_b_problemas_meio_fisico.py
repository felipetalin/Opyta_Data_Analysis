from __future__ import annotations

import argparse
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd

from auditar_parametros_meio_fisico import fold, norm_matrix, norm_text, norm_unit, parse_number, style_workbook


SAFE_UNIT_ACTION = "aceitar_unidade_fonte"
DECISION_OPTIONS = ["APROVAR", "ALTERAR", "NAO_APROVAR", "NAO_APLICAVEL"]


def _latest_audit(output_dir: Path) -> Path:
    files = sorted(
        output_dir.glob("*_auditoria_parametros_braaeg001.xlsx"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    if not files:
        raise FileNotFoundError(f"No audit workbook found in {output_dir}")
    return files[0]


def _source_summary(source_path: Path) -> pd.DataFrame:
    source = pd.read_excel(source_path, sheet_name="Resultados_Meio_Fisico", dtype=str)
    source["_matriz_key"] = source["Matriz"].map(norm_matrix).map(norm_text)
    source["_parametro_key"] = source["Parametro"].map(norm_text)
    source["_unidade_key"] = source["Unidade_Medida"].map(norm_unit)

    rows: list[dict[str, Any]] = []
    for key, group in source.groupby(["_matriz_key", "_parametro_key", "_unidade_key"], dropna=False):
        rows.append(
            {
                "_matriz_key": key[0],
                "_parametro_key": key[1],
                "_unidade_key": key[2],
                "linhas_resultado": int(len(group)),
                "pontos_afetados": " | ".join(sorted({fold(v) for v in group["Ponto"].dropna() if fold(v)})),
                "campanhas_afetadas": " | ".join(sorted({fold(v) for v in group["Campanha"].dropna() if fold(v)})),
            }
        )
    exact = pd.DataFrame(rows)

    fallback_rows: list[dict[str, Any]] = []
    for key, group in source.groupby(["_matriz_key", "_parametro_key"], dropna=False):
        fallback_rows.append(
            {
                "_matriz_key": key[0],
                "_parametro_key": key[1],
                "linhas_resultado_parametro": int(len(group)),
                "pontos_afetados_parametro": " | ".join(
                    sorted({fold(v) for v in group["Ponto"].dropna() if fold(v)})
                ),
                "campanhas_afetadas_parametro": " | ".join(
                    sorted({fold(v) for v in group["Campanha"].dropna() if fold(v)})
                ),
            }
        )
    fallback = pd.DataFrame(fallback_rows)
    return exact.merge(fallback, on=["_matriz_key", "_parametro_key"], how="left")


def _source_vmp_text(row: pd.Series) -> str:
    values = [fold(row.get(col)) for col in ["vmp_fonte_01", "vmp_fonte_02", "vmp_fonte_03"]]
    return " | ".join(value for value in values if value)


def _has_source_vmp(row: pd.Series) -> bool:
    values = [_source_vmp_text(row)]
    return any(value and value != "-" for value in values)


def classify_problem(row: pd.Series) -> dict[str, str]:
    classification = fold(row.get("classificacao"))
    matrix = norm_text(row.get("matriz"))
    param = norm_text(row.get("parametro"))
    source_vmp = _source_vmp_text(row)
    diff_vmp = fold(row.get("diferenca_vmp"))
    master_unit = fold(row.get("unidade_mestre"))
    source_unit = fold(row.get("unidade_resultado"))

    if classification == "new_parameter":
        return {
            "prioridade_gate_b": "cadastro",
            "problema_gate_b": "parametro_matriz_ausente",
            "decisao_sugerida": "cadastrar_parametro_sem_vmp",
            "acao_no_cadastro": "Criar parametro para a matriz; manter VMP vazio porque o laudo registra '-'",
            "observacao_gate_b": "Cadastro necessario para preservar o parametro na migracao.",
        }

    if classification == "synonym_review":
        if matrix == "agua subterranea" and param == "escherichia coli":
            return {
                "prioridade_gate_b": "revisao_tecnica",
                "problema_gate_b": "alias_parametro_microbiologico",
                "decisao_sugerida": "mapear_para_parametro_mestre_existente",
                "acao_no_cadastro": "Usar o cadastro mestre de Escherichia coli por tubos multiplos para Agua Subterranea; registrar alias curto.",
                "observacao_gate_b": "O laudo usa nome curto e VMP textual 'Ausente'.",
            }
        if _has_source_vmp(row):
            return {
                "prioridade_gate_b": "cadastro",
                "problema_gate_b": "parametro_matriz_ausente_com_vmp_fonte",
                "decisao_sugerida": "cadastrar_parametro_com_vmp_fonte",
                "acao_no_cadastro": f"Criar parametro para a matriz e preencher VMP conforme laudo: {source_vmp}.",
                "observacao_gate_b": "Nao usar sugestao fuzzy como sinonimo sem decisao tecnica.",
            }
        return {
            "prioridade_gate_b": "cadastro_simples",
            "problema_gate_b": "parametro_matriz_ausente_sem_vmp",
            "decisao_sugerida": "cadastrar_parametro_sem_vmp",
            "acao_no_cadastro": "Criar parametro para a matriz e manter VMP vazio.",
            "observacao_gate_b": "Nao ha limite no laudo; entra para serie historica/apoio.",
        }

    if classification == "vmp_review":
        if matrix == "agua superficial" and param == "cobre dissolvido":
            return {
                "prioridade_gate_b": "revisao_tecnica",
                "problema_gate_b": "vmp_fonte_diverge_do_mestre",
                "decisao_sugerida": "confirmar_vmp_cobre_dissolvido",
                "acao_no_cadastro": "Conferir se deve prevalecer o VMP do laudo (0,009) ou o cadastro mestre (0,013) antes de gerar conformidade.",
                "observacao_gate_b": diff_vmp,
            }
        return {
            "prioridade_gate_b": "vmp",
            "problema_gate_b": "vmp_fonte_sem_campo_mestre",
            "decisao_sugerida": "preencher_vmp_matriz_conforme_lastro",
            "acao_no_cadastro": f"Preencher VMP da matriz conforme laudo: {source_vmp}.",
            "observacao_gate_b": diff_vmp,
        }

    if classification == "unit_review":
        if diff_vmp:
            return {
                "prioridade_gate_b": "vmp",
                "problema_gate_b": "unidade_e_vmp_divergentes",
                "decisao_sugerida": "revisar_vmp_mestre",
                "acao_no_cadastro": f"Resolver VMP antes de aceitar o mapeamento. Diferenca: {diff_vmp}.",
                "observacao_gate_b": "A classificacao veio como unidade, mas tambem ha diferenca de VMP.",
            }
        if matrix == "agua subterranea" and param == "cloreto dissolvido":
            return {
                "prioridade_gate_b": "revisao_tecnica",
                "problema_gate_b": "vmp_mestre_suspeito",
                "decisao_sugerida": "bloquear_vmp_mestre_ate_conferencia",
                "acao_no_cadastro": "Conferir VMP mestre de Cloreto Dissolvido em Agua Subterranea antes de usar em conformidade.",
                "observacao_gate_b": "Cadastro mestre traz valor ativo de 0,07, incompativel com a expectativa operacional para cloreto.",
            }
        if param in {"amonia", "nitrogenio amoniacal"}:
            return {
                "prioridade_gate_b": "revisao_tecnica",
                "problema_gate_b": "vmp_dinamico_ou_nota_legal",
                "decisao_sugerida": "tratar_vmp_amonia_com_regra_especifica",
                "acao_no_cadastro": "Manter resultado migravel, mas bloquear conformidade automatica ate definir a regra do VMP indicado como nota no laudo.",
                "observacao_gate_b": f"VMP fonte: {source_vmp}.",
            }
        if matrix == "agua subterranea" and param == "potencial redox in situ":
            return {
                "prioridade_gate_b": "cadastro",
                "problema_gate_b": "unidade_mestre_incorreta",
                "decisao_sugerida": "corrigir_unidade_mestre",
                "acao_no_cadastro": "Ajustar unidade mestre para mV ou criar cadastro especifico com mV.",
                "observacao_gate_b": f"Fonte: {source_unit}; mestre: {master_unit}.",
            }
        if matrix == "agua superficial" and param == "cobre total":
            return {
                "prioridade_gate_b": "cadastro_simples",
                "problema_gate_b": "unidade_mestre_vazia_ou_traco",
                "decisao_sugerida": "corrigir_unidade_mestre",
                "acao_no_cadastro": "Preencher unidade mestre como mg/L mantendo o parametro existente.",
                "observacao_gate_b": f"Fonte: {source_unit}; mestre: {master_unit}.",
            }
        if matrix == "agua superficial" and param == "coliformes termotolerantes":
            return {
                "prioridade_gate_b": "revisao_tecnica",
                "problema_gate_b": "unidade_microbiologica_diferente",
                "decisao_sugerida": "confirmar_equivalencia_ufc_nmp",
                "acao_no_cadastro": "Confirmar se UFC/100mL pode usar o mesmo VMP de NMP/100 mL neste produto.",
                "observacao_gate_b": f"Fonte: {source_unit}; mestre: {master_unit}.",
            }
        return {
            "prioridade_gate_b": "aceite_operacional",
            "problema_gate_b": "unidade_expressa_com_especiacao",
            "decisao_sugerida": SAFE_UNIT_ACTION,
            "acao_no_cadastro": "Aceitar o parametro mestre e preservar a unidade da fonte no resultado migrado.",
            "observacao_gate_b": f"Fonte: {source_unit}; mestre: {master_unit}.",
        }

    return {
        "prioridade_gate_b": "revisao",
        "problema_gate_b": "classificacao_nao_mapeada",
        "decisao_sugerida": "revisar_manualmente",
        "acao_no_cadastro": "Revisar manualmente antes do Gate B.",
        "observacao_gate_b": classification,
    }


def _with_user_decision_columns(df: pd.DataFrame) -> pd.DataFrame:
    output = df.copy()
    for col in reversed(
        [
            "data_decisao",
            "responsavel_decisao",
            "comentario_usuario",
            "alteracao_solicitada",
            "decisao_usuario",
        ]
    ):
        output.insert(0, col, "")
    return output


def _decision_instructions() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "campo": "aba_decisao",
                "como_preencher": "Preencher somente a aba decisao_gate_b. As demais abas sao apoio/filtro.",
            },
            {
                "campo": "decisao_usuario",
                "como_preencher": "Escolher uma opcao: APROVAR, ALTERAR, NAO_APROVAR ou NAO_APLICAVEL.",
            },
            {
                "campo": "alteracao_solicitada",
                "como_preencher": "Obrigatorio quando decisao_usuario = ALTERAR; descreva o nome, unidade, VMP ou mapeamento correto.",
            },
            {
                "campo": "comentario_usuario",
                "como_preencher": "Campo livre para justificativa tecnica, duvida ou ressalva.",
            },
            {
                "campo": "responsavel_decisao",
                "como_preencher": "Nome ou iniciais de quem revisou.",
            },
            {
                "campo": "data_decisao",
                "como_preencher": "Data da revisao no formato AAAA-MM-DD.",
            },
        ]
    )


def add_decision_validation(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = load_workbook(path)
    if "decisao_gate_b" not in wb.sheetnames:
        wb.save(path)
        return

    ws = wb["decisao_gate_b"]
    headers = {cell.value: cell.column for cell in ws[1]}
    decision_col = headers.get("decisao_usuario")
    if decision_col:
        formula = '"' + ",".join(DECISION_OPTIONS) + '"'
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        validation.error = "Escolha APROVAR, ALTERAR, NAO_APROVAR ou NAO_APLICAVEL."
        validation.errorTitle = "Decisao invalida"
        validation.prompt = "Escolha a decisao para este item do Gate B."
        validation.promptTitle = "Gate B"
        ws.add_data_validation(validation)
        validation.add(f"{ws.cell(row=2, column=decision_col).coordinate}:{ws.cell(row=max(ws.max_row, 2), column=decision_col).coordinate}")

    wb.save(path)


def build_decision_workbook(audit_path: Path, source_path: Path, output_dir: Path, log_dir: Path) -> tuple[Path, Path, dict[str, Any]]:
    output_dir.mkdir(parents=True, exist_ok=True)
    log_dir.mkdir(parents=True, exist_ok=True)

    delta = pd.read_excel(audit_path, sheet_name="delta_cadastro")
    delta["_matriz_key"] = delta["matriz"].map(norm_matrix).map(norm_text)
    delta["_parametro_key"] = delta["parametro"].map(norm_text)
    delta["_unidade_key"] = delta["unidade_resultado"].map(norm_unit)

    summary_source = _source_summary(source_path)
    decisions = delta.apply(classify_problem, axis=1, result_type="expand")
    problems = pd.concat([decisions, delta], axis=1)
    problems = problems.merge(
        summary_source,
        on=["_matriz_key", "_parametro_key", "_unidade_key"],
        how="left",
    )

    for target, fallback in [
        ("linhas_resultado", "linhas_resultado_parametro"),
        ("pontos_afetados", "pontos_afetados_parametro"),
        ("campanhas_afetadas", "campanhas_afetadas_parametro"),
    ]:
        problems[target] = problems[target].fillna(problems[fallback])

    main_cols = [
        "prioridade_gate_b",
        "problema_gate_b",
        "decisao_sugerida",
        "acao_no_cadastro",
        "observacao_gate_b",
        "classificacao",
        "matriz",
        "parametro",
        "unidade_resultado",
        "linhas_resultado",
        "pontos_afetados",
        "campanhas_afetadas",
        "id_parametro_mestre",
        "parametro_mestre",
        "matriz_mestre",
        "unidade_mestre",
        "vmp_fonte_01",
        "vmp_fonte_02",
        "vmp_fonte_03",
        "vmp_mestre",
        "diferenca_vmp",
        "sugestoes",
    ]
    problems = problems[main_cols].sort_values(
        ["prioridade_gate_b", "classificacao", "matriz", "parametro"],
        kind="stable",
    )

    cadastro = problems[
        problems["decisao_sugerida"].isin(
            [
                "cadastrar_parametro_sem_vmp",
                "cadastrar_parametro_com_vmp_fonte",
                "corrigir_unidade_mestre",
                "mapear_para_parametro_mestre_existente",
            ]
        )
    ].copy()
    vmp = problems[
        problems["prioridade_gate_b"].isin(["vmp", "revisao_tecnica"])
        | problems["problema_gate_b"].str.contains("vmp", case=False, na=False)
    ].copy()
    unidades = problems[problems["decisao_sugerida"].eq(SAFE_UNIT_ACTION)].copy()
    user_decision = _with_user_decision_columns(problems)

    count_rows = []
    for col in ["prioridade_gate_b", "classificacao", "decisao_sugerida"]:
        for key, value in sorted(Counter(problems[col]).items()):
            count_rows.append({"metrica": f"{col}_{key}", "valor": int(value)})
    count_rows.insert(0, {"metrica": "problemas_gate_b", "valor": int(len(problems))})
    count_rows.append({"metrica": "linhas_resultado_afetadas_soma", "valor": int(problems["linhas_resultado"].fillna(0).sum())})
    resumo = pd.DataFrame(count_rows)

    ts = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    output_xlsx = output_dir / f"{ts}_gate_b_problemas_parametros_braaeg001.xlsx"
    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        _decision_instructions().to_excel(writer, sheet_name="orientacao", index=False)
        resumo.to_excel(writer, sheet_name="resumo", index=False)
        user_decision.to_excel(writer, sheet_name="decisao_gate_b", index=False)
        problems.to_excel(writer, sheet_name="problemas_gate_b", index=False)
        cadastro.to_excel(writer, sheet_name="cadastro_sugerido", index=False)
        vmp.to_excel(writer, sheet_name="vmp_revisar", index=False)
        unidades.to_excel(writer, sheet_name="unidades_aceite", index=False)
    style_workbook(output_xlsx)
    add_decision_validation(output_xlsx)

    report = {
        "status": "GATE_B_PROBLEMS_PREPARED",
        "timestamp_utc": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "audit_path": str(audit_path),
        "source_path": str(source_path),
        "output_xlsx": str(output_xlsx),
        "summary": {row["metrica"]: row["valor"] for _, row in resumo.iterrows()},
    }
    output_json = log_dir / f"{ts}_gate_b_problemas_parametros_braaeg001.json"
    output_json.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return output_xlsx, output_json, report


def main() -> int:
    parser = argparse.ArgumentParser(description="Prepare the focused Gate B issue workbook for BRAAEG001.")
    parser.add_argument("--audit", type=Path, help="Audit workbook. Defaults to the newest audit in --output-dir.")
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--log-dir", default=Path("logs/validacao_meio_fisico"), type=Path)
    args = parser.parse_args()

    audit_path = args.audit or _latest_audit(args.output_dir)
    output_xlsx, output_json, report = build_decision_workbook(
        audit_path=audit_path,
        source_path=args.source,
        output_dir=args.output_dir,
        log_dir=args.log_dir,
    )

    print(f"STATUS={report['status']}")
    print(f"OUTPUT_XLSX={output_xlsx}")
    print(f"OUTPUT_JSON={output_json}")
    print("SUMMARY=" + json.dumps(report["summary"], ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
