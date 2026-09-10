from __future__ import annotations

import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
import win32com.client
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_CODE = "BRAAEG001"
PROJECT_NAME = "A&G Mineração"
CLIENT_NAME = "Brandt"
CAMPAIGN_C02 = "C002-2026-06-SC"
LABORATORY = "SGS"

BRANDT_DIR = Path(r"G:\Meu Drive\Opyta\Clientes\Clientes\Clientes\Brandt")
CLIENT_DIR = next(path for path in BRANDT_DIR.iterdir() if path.name.startswith("A&G"))
SOURCE_ROOT = CLIENT_DIR / "resultados" / "Meio_fisico" / "Campanha-02"
C01_WORKBOOK = CLIENT_DIR / "resultados" / "Meio_fisico" / "migracao" / "Resultados_Meio_Fisico.xlsx"
OUTPUT_DIR = CLIENT_DIR / "resultados" / "Meio_fisico" / "migracao" / "campanha_02_subterranea_sedimentos"

MATRIX_FOLDERS = {
    "Agua Subterranea": "Subterrânea",
    "Sedimento": "Sedimentos",
}


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).replace("\xa0", " ").strip()


def parse_date(value: Any) -> str:
    text = clean_text(value)
    if not text:
        return ""
    parsed = pd.to_datetime(text, dayfirst=True, errors="coerce")
    if pd.isna(parsed):
        return text
    return parsed.strftime("%Y-%m-%d 00:00:00")


def point_from_header(raw: str, matrix: str) -> str:
    text = clean_text(raw).upper()
    if matrix == "Sedimento":
        match = re.search(r"SED[_\s-]*(\d+)", text)
        if not match:
            return clean_text(raw)
        return f"PT_{int(match.group(1)):02d}"
    match = re.search(r"NASC[_\s-]*(\d+)", text)
    if not match:
        return clean_text(raw)
    number = int(match.group(1))
    return {
        3: "SUB_01/NASC03",
        16: "SUB_02/NASC16",
        15: "SUB_03/NASC15",
    }.get(number, f"NASC_{number:03d}")


def find_result_column(ws: Any) -> int:
    for col in range(3, 12):
        header = clean_text(ws.Cells(11, col).Text)
        if not header:
            continue
        upper = header.upper()
        if upper.startswith("VMP") or upper in {"PARAMETRO", "PARÂMETRO", "UNIDADE"}:
            continue
        return col
    raise RuntimeError("Coluna de resultado nao encontrada na linha 11.")


def read_c01_support() -> tuple[pd.DataFrame, pd.DataFrame]:
    capa = pd.read_excel(C01_WORKBOOK, sheet_name="Capa_Projeto", dtype=str).fillna("")
    points = pd.read_excel(C01_WORKBOOK, sheet_name="Pontos_e_Campanhas", dtype=str).fillna("")
    return capa, points


def read_sgs_files() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    result_rows: list[dict[str, Any]] = []
    point_rows: list[dict[str, Any]] = []
    report_rows: list[dict[str, Any]] = []

    _, c01_points = read_c01_support()
    c01_point_map = {
        clean_text(row["Ponto"]): row.to_dict()
        for _, row in c01_points.iterrows()
    }

    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        for matrix, folder_name in MATRIX_FOLDERS.items():
            source_dir = SOURCE_ROOT / folder_name
            files = sorted(source_dir.glob("*.XLS"))
            if not files:
                raise FileNotFoundError(f"Nenhum XLS encontrado em {source_dir}")
            for file in files:
                wb = excel.Workbooks.Open(str(file), 0, True)
                try:
                    ws = wb.Worksheets.Item(1)
                    report = clean_text(ws.Cells(1, 2).Text)
                    source_project = clean_text(ws.Cells(4, 2).Text)
                    result_col = find_result_column(ws)
                    raw_point = clean_text(ws.Cells(11, result_col).Text)
                    point = point_from_header(raw_point, matrix)
                    raw_date = clean_text(ws.Cells(10, result_col).Text)
                    date = parse_date(raw_date)
                    support = c01_point_map.get(point, {})

                    report_rows.append(
                        {
                            "Arquivo_XLS": file.name,
                            "Arquivo_PDF": file.with_suffix(".PDF").name if file.with_suffix(".PDF").exists() else "",
                            "Laudo": report,
                            "Projeto_Laudo": source_project,
                            "Matriz": matrix,
                            "Ponto_Laudo": raw_point,
                            "Ponto": point,
                            "Campanha": CAMPAIGN_C02,
                            "Data": date,
                            "Coluna_Resultado": result_col,
                        }
                    )
                    point_rows.append(
                        {
                            "Ponto": point,
                            "Campanha": CAMPAIGN_C02,
                            "Data": date,
                            "Latitude": support.get("Latitude", ""),
                            "Longitude": support.get("Longitude", ""),
                            "Curso_d_Agua": support.get("Curso_d_Agua", "N.A.") or "N.A.",
                            "Bacia_Hidrografica": support.get("Bacia_Hidrografica", "Bacia do rio Doce") or "Bacia do rio Doce",
                            "Municipio": support.get("Municipio", "Barão de Cocais") or "Barão de Cocais",
                            "Observacoes_Coleta": "",
                        }
                    )

                    for row_idx in range(12, 120):
                        parameter = clean_text(ws.Cells(row_idx, 1).Text)
                        if not parameter:
                            continue
                        upper = parameter.upper()
                        if upper.startswith("LEGENDA") or upper.startswith("VMP "):
                            continue
                        result = clean_text(ws.Cells(row_idx, result_col).Text)
                        unit = clean_text(ws.Cells(row_idx, 2).Text)
                        result_rows.append(
                            {
                                "Ponto": point,
                                "Campanha": CAMPAIGN_C02,
                                "Matriz": matrix,
                                "Parametro": parameter,
                                "Resultado": result,
                                "Unidade_Medida": unit,
                                "Laboratorio": LABORATORY,
                            }
                        )
                finally:
                    wb.Close(False)
    finally:
        excel.Quit()

    return (
        pd.DataFrame(result_rows),
        pd.DataFrame(point_rows).drop_duplicates(subset=["Ponto", "Campanha"]).sort_values(["Ponto", "Campanha"]),
        pd.DataFrame(report_rows),
    )


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column_cells in ws.columns:
            values = [clean_text(cell.value) for cell in column_cells]
            width = min(max(max((len(value) for value in values), default=0) + 2, 10), 55)
            ws.column_dimensions[get_column_letter(column_cells[0].column)].width = width
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main() -> int:
    ts = datetime.now().strftime("%Y%m%dT%H%M%S")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    capa, _ = read_c01_support()
    results, points, reports = read_sgs_files()

    summary_rows = [
        {"Metrica": "arquivos_xls", "Valor": len(reports)},
        {"Metrica": "registros_resultados", "Valor": len(results)},
        {"Metrica": "pontos", "Valor": points["Ponto"].nunique()},
        {"Metrica": "matrizes", "Valor": results["Matriz"].nunique()},
        {"Metrica": "parametros", "Valor": results["Parametro"].nunique()},
        {"Metrica": "resultados_vazios", "Valor": int(results["Resultado"].eq("").sum())},
    ]
    matrix_summary = (
        results.groupby("Matriz", as_index=False)
        .agg(registros=("Parametro", "size"), pontos=("Ponto", "nunique"), parametros=("Parametro", "nunique"))
        .sort_values("Matriz")
    )
    point_summary = (
        results.groupby(["Matriz", "Ponto"], as_index=False)
        .agg(registros=("Parametro", "size"), parametros=("Parametro", "nunique"))
        .sort_values(["Matriz", "Ponto"])
    )
    duplicate_counts = Counter((row.Ponto, row.Campanha, row.Matriz, row.Parametro) for row in results.itertuples(index=False))
    duplicates = pd.DataFrame(
        [
            {"Ponto": key[0], "Campanha": key[1], "Matriz": key[2], "Parametro": key[3], "ocorrencias": count}
            for key, count in duplicate_counts.items()
            if count > 1
        ]
    )

    out_xlsx = OUTPUT_DIR / f"{ts}_staging_resultados_meio_fisico_c02_subterranea_sedimentos.xlsx"
    with pd.ExcelWriter(out_xlsx, engine="openpyxl") as writer:
        capa.to_excel(writer, sheet_name="Capa_Projeto", index=False)
        points.to_excel(writer, sheet_name="Pontos_e_Campanhas", index=False)
        results.to_excel(writer, sheet_name="Resultados_Meio_Fisico", index=False)
        pd.DataFrame(summary_rows).to_excel(writer, sheet_name="Resumo", index=False)
        matrix_summary.to_excel(writer, sheet_name="Resumo_Matriz", index=False)
        point_summary.to_excel(writer, sheet_name="Resumo_Ponto", index=False)
        reports.to_excel(writer, sheet_name="Laudos_Extraidos", index=False)
        duplicates.to_excel(writer, sheet_name="Duplicidades", index=False)
    style_workbook(out_xlsx)

    print(f"STAGING={out_xlsx}")
    print(f"REGISTROS={len(results)}")
    print(f"PONTOS={points['Ponto'].nunique()}")
    print(f"MATRIZES={results['Matriz'].nunique()}")
    print(f"DUPLICIDADES={len(duplicates)}")
    print(matrix_summary.to_string(index=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
