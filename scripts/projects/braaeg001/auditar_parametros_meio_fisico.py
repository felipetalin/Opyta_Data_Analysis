from __future__ import annotations

import argparse
import json
import os
import re
import unicodedata
import urllib.parse
import urllib.request
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from dotenv import load_dotenv


VMP_COLS = [
    "vmp_357_cl1",
    "vmp_357_cl1_min",
    "vmp_357_cl1_max",
    "vmp_357_cl2",
    "vmp_357_cl2_min",
    "vmp_357_cl2_max",
    "vmp_357_cl3",
    "vmp_396_consumo_humano",
    "vmp_396_dessedentacao_animal",
    "vmp_396_irrigacao",
    "vmp_396_recreacao",
    "vmp_396_animal",
    "vmp_396_vmp",
    "vmp_396_vr",
    "vmp_430_padrao",
    "vmp_454_n1",
    "vmp_454_n2",
]

MATRIX_ALIASES = {
    "agua superficial": "Agua Superficial",
    "agua subterranea": "Agua Subterranea",
    "sedimento": "Sedimento",
    "efluente": "Efluente",
}

PARAMETER_ALIASES = {
    "dbo": "demanda bioquimica de oxigenio",
    "demanda bioquimica de oxigenio": "dbo",
    "demanda quimica de oxigenio": "dqo",
    "nitrato n": "nitrato",
    "nitrito n": "nitrito",
    "cloreto": "cloreto total",
    "cloreto dissolvido": "cloreto total",
    "ph": "ph in situ",
    "turbidez": "turbidez",
    "amonia": "nitrogenio amoniacal",
}


def fold(text: Any) -> str:
    value = "" if text is None or pd.isna(text) else str(text)
    value = value.replace("\xa0", " ").strip()
    value = unicodedata.normalize("NFKD", value)
    return "".join(ch for ch in value if not unicodedata.combining(ch))


def norm_text(text: Any) -> str:
    value = fold(text).lower()
    value = value.replace("µ", "u").replace("μ", "u")
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def norm_param(text: Any) -> str:
    value = norm_text(text)
    return PARAMETER_ALIASES.get(value, value)


def norm_matrix(text: Any) -> str:
    value = norm_text(text)
    return MATRIX_ALIASES.get(value, fold(text).strip())


def norm_unit(text: Any) -> str:
    value = fold(text).lower().replace("µ", "u").replace("μ", "u")
    value = value.replace(" ", "")
    value = value.replace("litro", "l")
    value = value.replace("kg", "kg")
    value = re.sub(r"mg[a-z]{1,3}/l$", "mg/l", value)
    value = value.replace("mg/l02", "mg/lo2")
    value = value.replace("mgo2/l", "mg/lo2")
    value = value.replace("mg/l_o2", "mg/lo2")
    value = value.replace("mg/lo2", "mg/l o2")
    value = value.replace("ufc/100ml", "ufc/100ml")
    value = value.replace("nmp/100ml", "nmp/100ml")
    return value


def parse_number(text: Any) -> float | None:
    raw = fold(text)
    if not raw:
        return None
    if re.fullmatch(r"\[[^\]]+\]", raw.strip()):
        return None
    raw = raw.replace(",", ".")
    match = re.search(r"[-+]?\d+(?:\.\d+)?", raw)
    if not match:
        return None
    try:
        return float(match.group(0))
    except ValueError:
        return None


def fetch_supabase_table(table: str, env_file: Path, page_size: int = 1000) -> list[dict[str, Any]]:
    load_dotenv(env_file, override=True)
    url = os.getenv("SUPABASE_URL")
    key = os.getenv("SUPABASE_ANON_KEY")
    if not url or not key:
        raise RuntimeError("SUPABASE_URL/SUPABASE_ANON_KEY not found")

    rows: list[dict[str, Any]] = []
    offset = 0
    while True:
        endpoint = (
            f"{url.rstrip('/')}/rest/v1/{urllib.parse.quote(table)}"
            f"?select=*&order=id_parametro.asc"
        )
        req = urllib.request.Request(
            endpoint,
            headers={
                "apikey": key,
                "Authorization": f"Bearer {key}",
                "Range": f"{offset}-{offset + page_size - 1}",
            },
        )
        with urllib.request.urlopen(req, timeout=60) as response:
            batch = json.loads(response.read().decode("utf-8"))
        rows.extend(batch)
        if len(batch) < page_size:
            return rows
        offset += page_size


def source_vmp_lookup(sgs_path: Path) -> dict[tuple[str, str], dict[str, str]]:
    if not sgs_path.exists():
        return {}
    df = pd.read_excel(sgs_path, sheet_name="Resultados_prontos", dtype=str)
    lookup: dict[tuple[str, str], dict[str, str]] = {}
    for (matrix, param), group in df.groupby(["Matriz", "Parametro"], dropna=False):
        key = (norm_matrix(matrix), norm_param(param))
        entry: dict[str, str] = {}
        for col in ["VMP_01", "VMP_02", "VMP_03", "Legenda_VMP_01", "Legenda_VMP_02", "Legenda_VMP_03"]:
            vals = sorted({fold(v) for v in group.get(col, pd.Series(dtype=str)).dropna() if fold(v)})
            entry[col] = " | ".join(vals)
        lookup[key] = entry
    return lookup


def active_master_vmp(row: dict[str, Any], matrix: str) -> dict[str, Any]:
    matrix_norm = norm_matrix(matrix)
    if matrix_norm == "Agua Superficial":
        cols = ["vmp_357_cl2_min", "vmp_357_cl2_max"]
    elif matrix_norm == "Agua Subterranea":
        cols = ["vmp_396_consumo_humano"]
    elif matrix_norm == "Sedimento":
        cols = ["vmp_454_n1", "vmp_454_n2"]
    elif matrix_norm == "Efluente":
        cols = ["vmp_430_padrao"]
    else:
        cols = VMP_COLS
    return {col: row.get(col) for col in cols if row.get(col) not in (None, "")}


def source_vmp_numeric(source: dict[str, str], matrix: str) -> dict[str, float | None]:
    matrix_norm = norm_matrix(matrix)
    if matrix_norm == "Agua Superficial":
        raw = source.get("VMP_01")
        if raw and "-" in raw and parse_number(raw) is not None:
            nums = [float(x.replace(",", ".")) for x in re.findall(r"\d+(?:[,.]\d+)?", raw)]
            if len(nums) >= 2:
                return {"vmp_357_cl2_min": nums[0], "vmp_357_cl2_max": nums[1]}
        if raw and raw.strip().startswith(">"):
            return {"vmp_357_cl2_min": parse_number(raw)}
        return {"vmp_357_cl2_max": parse_number(raw)}
    if matrix_norm == "Agua Subterranea":
        return {"vmp_396_consumo_humano": parse_number(source.get("VMP_01"))}
    if matrix_norm == "Sedimento":
        return {
            "vmp_454_n1": parse_number(source.get("VMP_02")),
            "vmp_454_n2": parse_number(source.get("VMP_03")),
        }
    return {}


def vmp_diff(source: dict[str, str], master: dict[str, Any], matrix: str) -> tuple[bool, str]:
    source_nums = source_vmp_numeric(source, matrix)
    master_active = active_master_vmp(master, matrix)
    diffs: list[str] = []
    for col, src_val in source_nums.items():
        if src_val is None:
            continue
        mst_val = parse_number(master_active.get(col))
        if mst_val is None:
            diffs.append(f"{col}: source={src_val}; master=empty")
        elif abs(src_val - mst_val) > 1e-9:
            diffs.append(f"{col}: source={src_val}; master={mst_val}")
    return bool(diffs), "; ".join(diffs)


def build_master_indexes(master_rows: list[dict[str, Any]]) -> tuple[dict[tuple[str, str], list[dict[str, Any]]], dict[str, list[dict[str, Any]]]]:
    by_key: dict[tuple[str, str], list[dict[str, Any]]] = {}
    by_param: dict[str, list[dict[str, Any]]] = {}
    for row in master_rows:
        matrix = norm_matrix(row.get("matriz")) if row.get("matriz") else ""
        param = norm_param(row.get("nome_parametro"))
        by_key.setdefault((matrix, param), []).append(row)
        by_param.setdefault(param, []).append(row)
    return by_key, by_param


def closest_suggestions(param_norm: str, candidate_params: list[str], max_items: int = 5) -> str:
    from difflib import SequenceMatcher

    scored = [
        (SequenceMatcher(None, param_norm, candidate).ratio(), candidate)
        for candidate in candidate_params
    ]
    scored.sort(reverse=True)
    return " | ".join(f"{candidate} ({score:.2f})" for score, candidate in scored[:max_items] if score >= 0.55)


def classify_row(
    row: pd.Series,
    by_key: dict[tuple[str, str], list[dict[str, Any]]],
    by_param: dict[str, list[dict[str, Any]]],
    source_vmp: dict[tuple[str, str], dict[str, str]],
    candidate_params_by_matrix: dict[str, list[str]],
) -> dict[str, Any]:
    matrix = norm_matrix(row["Matriz"])
    param = norm_param(row["Parametro"])
    unit = fold(row["Unidade_Medida"])
    unit_norm = norm_unit(unit)
    key = (matrix, param)
    source = source_vmp.get(key, {})

    matches = by_key.get(key) or by_key.get(("", param)) or []
    if matches:
        master = matches[0]
        master_unit = fold(master.get("unidade_medida"))
        master_unit_norm = norm_unit(master_unit)
        has_unit_diff = bool(master_unit_norm and unit_norm and master_unit_norm != unit_norm)
        has_vmp_diff, diff_text = vmp_diff(source, master, matrix)
        master_has_vmp = bool(active_master_vmp(master, matrix))
        source_has_vmp = any(source.get(col) for col in ["VMP_01", "VMP_02", "VMP_03"])
        if has_unit_diff:
            classification = "unit_review"
        elif has_vmp_diff:
            classification = "vmp_review"
        elif not master_has_vmp and not source_has_vmp:
            classification = "not_applicable"
        else:
            classification = "matched"
        return {
            "classificacao": classification,
            "matriz": row["Matriz"],
            "parametro": row["Parametro"],
            "unidade_resultado": unit,
            "parametro_norm": param,
            "matriz_norm": matrix,
            "id_parametro_mestre": master.get("id_parametro"),
            "parametro_mestre": master.get("nome_parametro"),
            "matriz_mestre": master.get("matriz"),
            "unidade_mestre": master_unit,
            "vmp_fonte_01": source.get("VMP_01", ""),
            "vmp_fonte_02": source.get("VMP_02", ""),
            "vmp_fonte_03": source.get("VMP_03", ""),
            "legenda_vmp_01": source.get("Legenda_VMP_01", ""),
            "legenda_vmp_02": source.get("Legenda_VMP_02", ""),
            "legenda_vmp_03": source.get("Legenda_VMP_03", ""),
            "vmp_mestre": json.dumps(active_master_vmp(master, matrix), ensure_ascii=False),
            "diferenca_vmp": diff_text,
            "sugestoes": "",
            "acao_recomendada": "reutilizar cadastro mestre" if classification in {"matched", "not_applicable"} else "revisar antes do Gate B",
        }

    same_param_other_matrix = by_param.get(param, [])
    candidates = candidate_params_by_matrix.get(matrix, []) + candidate_params_by_matrix.get("", [])
    suggestions = closest_suggestions(param, sorted(set(candidates)))
    alias_target = PARAMETER_ALIASES.get(norm_text(row["Parametro"]))
    if alias_target and (matrix, alias_target) in by_key:
        suggestions = f"{alias_target} (alias local)" + (f" | {suggestions}" if suggestions else "")
    classification = "synonym_review" if suggestions or same_param_other_matrix else "new_parameter"
    return {
        "classificacao": classification,
        "matriz": row["Matriz"],
        "parametro": row["Parametro"],
        "unidade_resultado": unit,
        "parametro_norm": param,
        "matriz_norm": matrix,
        "id_parametro_mestre": "",
        "parametro_mestre": "",
        "matriz_mestre": "",
        "unidade_mestre": "",
        "vmp_fonte_01": source.get("VMP_01", ""),
        "vmp_fonte_02": source.get("VMP_02", ""),
        "vmp_fonte_03": source.get("VMP_03", ""),
        "legenda_vmp_01": source.get("Legenda_VMP_01", ""),
        "legenda_vmp_02": source.get("Legenda_VMP_02", ""),
        "legenda_vmp_03": source.get("Legenda_VMP_03", ""),
        "vmp_mestre": "",
        "diferenca_vmp": "",
        "sugestoes": suggestions,
        "acao_recomendada": "avaliar cadastro ou sinonimo antes do Gate B",
    }


def style_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col_idx in range(1, ws.max_column + 1):
            col = get_column_letter(col_idx)
            max_len = 10
            for cell in ws[col][: min(ws.max_row, 200)]:
                max_len = max(max_len, len("" if cell.value is None else str(cell.value)))
            ws.column_dimensions[col].width = min(max_len + 2, 55)
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Audit BRAAEG001 physicochemical parameters against master registry.")
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--sgs", required=True, type=Path)
    parser.add_argument("--env-file", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--log-dir", default=Path("logs/validacao_meio_fisico"), type=Path)
    args = parser.parse_args()

    ts = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    args.log_dir.mkdir(parents=True, exist_ok=True)

    df = pd.read_excel(args.input, sheet_name="Resultados_Meio_Fisico", dtype=str)
    combos = (
        df[["Matriz", "Parametro", "Unidade_Medida"]]
        .drop_duplicates()
        .sort_values(["Matriz", "Parametro", "Unidade_Medida"])
        .reset_index(drop=True)
    )

    master_rows = fetch_supabase_table("parametros_analise", args.env_file)
    by_key, by_param = build_master_indexes(master_rows)
    candidate_params_by_matrix: dict[str, list[str]] = {}
    for matrix, param in by_key:
        candidate_params_by_matrix.setdefault(matrix, []).append(param)

    vmp_lookup = source_vmp_lookup(args.sgs)
    audit_rows = [
        classify_row(row, by_key, by_param, vmp_lookup, candidate_params_by_matrix)
        for _, row in combos.iterrows()
    ]
    audit = pd.DataFrame(audit_rows)
    delta = audit[~audit["classificacao"].isin(["matched", "not_applicable"])].copy()

    counts = Counter(audit["classificacao"])
    summary = pd.DataFrame(
        [{"metrica": "combos_auditados", "valor": int(len(audit))}]
        + [{"metrica": f"classificacao_{key}", "valor": int(value)} for key, value in sorted(counts.items())]
        + [{"metrica": "parametros_distintos_fonte", "valor": int(df["Parametro"].nunique())}]
        + [{"metrica": "linhas_resultados_fonte", "valor": int(len(df))}]
        + [{"metrica": "linhas_cadastro_mestre", "valor": int(len(master_rows))}]
    )

    master_df = pd.DataFrame(master_rows)
    output_xlsx = args.output_dir / f"{ts}_auditoria_parametros_braaeg001.xlsx"
    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="resumo", index=False)
        audit.to_excel(writer, sheet_name="auditoria_parametros", index=False)
        delta.to_excel(writer, sheet_name="delta_cadastro", index=False)
        combos.to_excel(writer, sheet_name="parametros_fonte", index=False)
        master_df.to_excel(writer, sheet_name="cadastro_mestre", index=False)
    style_workbook(output_xlsx)

    report = {
        "status": "PASS" if delta.empty else "DELTA_REVIEW_REQUIRED",
        "timestamp_utc": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "input": str(args.input),
        "sgs_support": str(args.sgs),
        "output_xlsx": str(output_xlsx),
        "master_table": "parametros_analise",
        "summary": {row["metrica"]: row["valor"] for _, row in summary.iterrows()},
        "delta_counts": dict(Counter(delta["classificacao"])) if not delta.empty else {},
    }
    output_json = args.log_dir / f"{ts}_auditoria_parametros_braaeg001.json"
    output_json.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"STATUS={report['status']}")
    print(f"OUTPUT_XLSX={output_xlsx}")
    print(f"OUTPUT_JSON={output_json}")
    print("COUNTS=" + json.dumps(report["summary"], ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
