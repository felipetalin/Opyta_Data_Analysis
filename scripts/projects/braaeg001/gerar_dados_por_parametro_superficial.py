from __future__ import annotations

import re
import unicodedata
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


CLIENT_ROOT = Path(r"G:\Meu Drive\Opyta\Clientes\Clientes\Clientes\Brandt")
PROJECT_DIR = next(path for path in CLIENT_ROOT.iterdir() if path.name.startswith("A&G"))
CONSOLIDATED = (
    PROJECT_DIR
    / "resultados"
    / "Meio_fisico"
    / "migracao"
    / "consolidacao_pos_c02"
    / "20260729T183622Z_consolidado_meio_fisico_braaeg001_pos_c02.xlsx"
)
OUTPUT = (
    PROJECT_DIR
    / "resultados"
    / "Meio_fisico"
    / "resultados"
    / "superficial"
    / "02_Dados_por_Parametro_Agua_Superficial.xlsx"
)

CAMPAIGN_LABELS = {
    "C001-2026-02-CH": "Campanha-01-Chuva",
    "C002-2026-06-SC": "Campanha-02-Seca",
}


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def fold(value: Any) -> str:
    text = clean_text(value).replace("ę", "e").replace("Ę", "E")
    text = unicodedata.normalize("NFKD", text)
    return "".join(ch for ch in text if not unicodedata.combining(ch))


def norm(value: Any) -> str:
    text = fold(value).lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def fix_ptbr_label(value: Any) -> str:
    out = clean_text(value)
    if "Ã" in out or "Â" in out:
        try:
            out = out.encode("latin1").decode("utf-8")
        except UnicodeError:
            pass
    replacements = {
        "Arsęnio": "Arsênio",
        "Manganęs": "Manganês",
        "Nitrogęnio": "Nitrogênio",
        "Oxigęnio": "Oxigênio",
        "Bioquimica": "Bioquímica",
        "Quimica": "Química",
        "Solidos": "Sólidos",
        "Aluminio": "Alumínio",
        "Fosforo": "Fósforo",
        "Cadmio": "Cádmio",
    }
    for src, dst in replacements.items():
        out = out.replace(src, dst)
    out = out.replace(chr(281), chr(234))
    return out


def safe_sheet_name(value: str, used: set[str]) -> str:
    name = re.sub(r"[:\\/?*\[\]]", "", value).strip()[:31] or "Parametro"
    base = name
    counter = 2
    while name in used:
        suffix = f"_{counter}"
        name = f"{base[:31 - len(suffix)]}{suffix}"
        counter += 1
    used.add(name)
    return name


def point_key(point: str) -> int:
    raw = clean_text(point).split("_")[-1]
    return int(raw) if raw.isdigit() else 999


def format_number(value: Any) -> str:
    if value is None or pd.isna(value):
        return "-"
    value = float(value)
    if abs(value) >= 100:
        text = f"{value:.0f}"
    elif abs(value) >= 10:
        text = f"{value:.2f}".rstrip("0").rstrip(".")
    elif abs(value) >= 1:
        text = f"{value:.3f}".rstrip("0").rstrip(".")
    else:
        text = f"{value:.5f}".rstrip("0").rstrip(".")
    return text.replace(".", ",")


def format_result(sign: Any, value: Any) -> Any:
    if value is None or pd.isna(value):
        return None
    sign_text = clean_text(sign)
    value_text = format_number(value)
    if sign_text in {"<", "<=", ">", ">="}:
        return f"{sign_text}{value_text}"
    return value_text


def effective_limits(parameter: str, legislation: str, vmin: Any, vmax: Any, vdyn: Any) -> tuple[Any, Any, Any]:
    p_norm = norm(parameter)
    if "manganes dissolvido" in p_norm:
        return None, None, None
    if p_norm == "amonia":
        return None, None, None
    if "solidos suspensos totais" in p_norm and legislation == "conama":
        return None, None, None
    if "escherichia coli" in p_norm and legislation == "conama":
        return None, None, None
    return vmin, vmax, vdyn


def format_vmp(parameter: str, legislation: str, vmin: Any, vmax: Any, vdyn: Any) -> str:
    p_norm = norm(parameter)
    vmin, vmax, vdyn = effective_limits(parameter, legislation, vmin, vmax, vdyn)
    if "oleos" in p_norm and "graxas" in p_norm:
        return "V.A."
    if p_norm == "amonia":
        return "-"
    if "nitrogenio amoniacal" in p_norm:
        return "3,7 mg/L N para pH <= 7,5; 2,0 mg/L N para 7,5 < pH <= 8,0; 1,0 mg/L N para 8,0 < pH <= 8,5; 0,5 mg/L N para pH > 8,5"
    if "escherichia coli" in p_norm and legislation == "conama":
        return "Pode substituir coliformes termotolerantes conforme critério do órgão competente"
    has_min = vmin is not None and pd.notna(vmin)
    has_max = vmax is not None and pd.notna(vmax)
    has_dyn = vdyn is not None and pd.notna(vdyn)
    if has_min and has_max:
        return f"{format_number(vmin)}-{format_number(vmax)}"
    if has_min:
        return f"≥{format_number(vmin)}"
    if has_max:
        return format_number(vmax)
    if has_dyn:
        return format_number(vdyn)
    return "-"


def status_conformidade(parameter: str, value: Any, sign: Any, vmin: Any, vmax: Any, vdyn: Any) -> str:
    if value is None or pd.isna(value):
        return "Sem valor numérico"
    vmin, vmax, vdyn = effective_limits(parameter, "copam", vmin, vmax, vdyn)
    if vmin is None and vmax is None and vdyn is None:
        return "Sem VMP Classe 2"
    if clean_text(sign) in {"<", "<="}:
        return "Atende"
    value = float(value)
    if vmin is not None and pd.notna(vmin) and value < float(vmin):
        return "Viola"
    upper_candidates = [v for v in [vmax, vdyn] if v is not None and pd.notna(v)]
    if upper_candidates and value > min(map(float, upper_candidates)):
        return "Viola"
    return "Atende"


def load_surface() -> pd.DataFrame:
    df = pd.read_excel(CONSOLIDATED, sheet_name="fisico_analise_consolidada")
    surface = df[df["matriz"].astype(str).str.contains("Superficial", na=False)].copy()
    for col in ["valor_medido", "vmp_357_cl2_min", "vmp_357_cl2_max", "vmp_amonia_dinamico"]:
        surface[col] = pd.to_numeric(surface[col], errors="coerce")
    surface["data_hora_coleta"] = pd.to_datetime(surface["data_hora_coleta"], errors="coerce")
    surface["parametro_display"] = surface["nome_parametro"].map(fix_ptbr_label)
    surface["unidade_display"] = surface["unidade_medida"].map(fix_ptbr_label)
    surface.loc[surface["parametro_display"].map(norm).eq("nitrogenio amoniacal"), "unidade_display"] = "mg N/L"
    surface["campanha_display"] = surface["nome_campanha"].map(lambda x: CAMPAIGN_LABELS.get(clean_text(x), clean_text(x)))
    return surface


def build_parameter_table(group: pd.DataFrame) -> pd.DataFrame:
    parameter = clean_text(group["parametro_display"].iloc[0])
    unit = clean_text(group["unidade_display"].dropna().iloc[0]) if not group["unidade_display"].dropna().empty else ""
    vmin = group["vmp_357_cl2_min"].dropna().iloc[0] if not group["vmp_357_cl2_min"].dropna().empty else None
    vmax = group["vmp_357_cl2_max"].dropna().iloc[0] if not group["vmp_357_cl2_max"].dropna().empty else None
    vdyn = group["vmp_amonia_dinamico"].dropna().iloc[0] if not group["vmp_amonia_dinamico"].dropna().empty else None

    rows = []
    ordered = group.sort_values(
        by=["nome_ponto", "nome_campanha"],
        key=lambda s: s.map(point_key) if s.name == "nome_ponto" else s,
    )
    for rec in ordered.itertuples(index=False):
        rows.append({
            "Parâmetro": parameter,
            "Unidade": unit,
            "VMP - CONAMA n° 357 (2005) - Classe 2": format_vmp(parameter, "conama", vmin, vmax, vdyn),
            "VMP - COPAM n° 8 (2022) - Classe 2": format_vmp(parameter, "copam", vmin, vmax, vdyn),
            "Campanha": rec.campanha_display,
            "Ponto amostral": rec.nome_ponto,
            "Data da coleta": rec.data_hora_coleta,
            "Sinal": clean_text(rec.sinal_limite) or None,
            "Valor numérico": rec.valor_medido,
            "Resultado formatado": format_result(rec.sinal_limite, rec.valor_medido),
            "Status COPAM Classe 2": status_conformidade(parameter, rec.valor_medido, rec.sinal_limite, vmin, vmax, vdyn),
            "Observações": clean_text(rec.observacoes_resultado) or None,
        })
    return pd.DataFrame(rows)


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    white_font = Font(color="FFFFFF", bold=True)
    violation_fill = PatternFill("solid", fgColor="F4CCCC")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = white_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            status_cell = row[10] if len(row) > 10 else None
            if status_cell and status_cell.value == "Viola":
                for cell in row:
                    cell.fill = violation_fill
        for col in ws.columns:
            values = [clean_text(cell.value) for cell in col]
            width = min(max(max(map(len, values), default=0) + 2, 12), 55)
            ws.column_dimensions[get_column_letter(col[0].column)].width = width
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main() -> int:
    surface = load_surface()
    used: set[str] = set()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        summary = []
        for _, group in surface.groupby("parametro_display", sort=True):
            table = build_parameter_table(group)
            sheet = safe_sheet_name(clean_text(group["parametro_display"].iloc[0]), used)
            table.to_excel(writer, sheet_name=sheet, index=False)
            summary.append({
                "Parâmetro": clean_text(group["parametro_display"].iloc[0]),
                "Aba": sheet,
                "Registros": len(table),
                "Campanhas": table["Campanha"].nunique(),
                "Pontos": table["Ponto amostral"].nunique(),
                "Violações": int((table["Status COPAM Classe 2"] == "Viola").sum()),
            })
        pd.DataFrame(summary).to_excel(writer, sheet_name=safe_sheet_name("Indice", used), index=False)
    style_workbook(OUTPUT)
    print(f"OK: {OUTPUT}")
    print(f"abas_parametros={surface['parametro_display'].nunique()} registros={len(surface)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
