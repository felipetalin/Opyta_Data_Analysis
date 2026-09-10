from __future__ import annotations

import json
import os
import sys
import importlib.util
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import text


ROOT = next(parent for parent in Path(__file__).resolve().parents if (parent / "src" / "opyta_analysis").exists())
OPYTA_DATA_ROOT = Path(r"G:\Meu Drive\Opyta\Opyta_Data")
for path in (ROOT / "src", OPYTA_DATA_ROOT):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from core.engine import get_engine  # noqa: E402
SOURCE_MIGRATOR = ROOT / "scripts" / "projects" / "braaeg001" / "preparar_migracao_controlada_meio_fisico.py"
spec = importlib.util.spec_from_file_location("braaeg001_preparar_migracao_controlada_meio_fisico", SOURCE_MIGRATOR)
if spec is None or spec.loader is None:
    raise RuntimeError(f"Nao foi possivel carregar {SOURCE_MIGRATOR}")
mf = importlib.util.module_from_spec(spec)
spec.loader.exec_module(mf)

PROJECT_CODE = mf.PROJECT_CODE
VMP_COLS = mf.VMP_COLS
RestClient = mf.RestClient
build_lookup = mf.build_lookup
build_master_indexes = mf.build_master_indexes
build_migration_preview = mf.build_migration_preview
canonical_matrix = mf.canonical_matrix
clean_text = mf.clean_text
fetch_master = mf.fetch_master
norm_param = mf.norm_param
style_workbook = mf.style_workbook


ENV_FILE = Path(r"G:\Meu Drive\Opyta\Opyta_Data\.env")
CLIENT_ROOT = Path(r"G:\Meu Drive\Opyta\Clientes\Clientes\Clientes\Brandt")
PROJECT_DIR = next(path for path in CLIENT_ROOT.iterdir() if path.name.startswith("A&G"))
OUTPUT_DIR = PROJECT_DIR / "resultados" / "Meio_fisico" / "migracao" / "campanha_02_superficial"
STAGING_C02 = OUTPUT_DIR / "20260729T151930_staging_resultados_meio_fisico_c02_agua_superficial.xlsx"
C01_DRY_RUN = PROJECT_DIR / "resultados" / "Meio_fisico" / "migracao" / "20260702T193527Z_dry_run_migracao_controlada_braaeg001.xlsx"
LOG_DIR = ROOT / "logs" / "validacao_meio_fisico"

FISICO_INSERT_COLS = [
    "codigo_interno_opyta",
    "nome_projeto",
    "nome_empresa",
    "nome_campanha",
    "nome_ponto",
    "data_hora_coleta",
    "latitude",
    "longitude",
    "bacia_hidrografica",
    "matriz",
    "nome_parametro",
    "sinal_limite",
    "valor_medido",
    "unidade_medida",
    "unidade_original_laudo",
    "laboratorio_responsavel",
    "observacoes_resultado",
    "vmp_357_cl1_min",
    "vmp_357_cl1_max",
    "vmp_357_cl2_min",
    "vmp_357_cl2_max",
    "vmp_amonia_dinamico",
    "vmp_454_n1",
    "vmp_454_n2",
    "vmp_396_consumo_humano",
    "vmp_396_dessedentacao_animal",
    "vmp_396_irrigacao",
    "vmp_396_recreacao",
    "vmp_430_padrao",
]


def configure_database_access() -> None:
    repo_env = ROOT / ".env"
    if repo_env.exists():
        load_dotenv(repo_env, override=False)
    if ENV_FILE.exists():
        load_dotenv(ENV_FILE, override=False)
    has_database_url = bool(os.getenv("DATABASE_URL") or os.getenv("SUPABASE_DB_URL"))
    has_db_parts = all(os.getenv(name) for name in ("DB_USER", "DB_PASSWORD", "DB_HOST", "DB_NAME"))
    has_supabase_parts = all(os.getenv(name) for name in ("SUPABASE_DB_USER", "SUPABASE_DB_PASSWORD", "SUPABASE_DB_HOST"))
    if not (has_database_url or has_db_parts or has_supabase_parts):
        fisico_url = os.getenv("FISICO_DB_URL")
        if fisico_url:
            os.environ.setdefault("DATABASE_URL", fisico_url)


def infer_aliases_from_c01(master_by_id: dict[int, dict[str, Any]]) -> pd.DataFrame:
    preview = pd.read_excel(C01_DRY_RUN, sheet_name="preview_migracao", dtype=str).fillna("")
    rows: list[dict[str, Any]] = []
    for _, row in preview.iterrows():
        source_param = clean_text(row.get("nome_parametro"))
        source_matrix = canonical_matrix(row.get("matriz"))
        mapped_id = clean_text(row.get("id_parametro_mapeado"))
        if not source_param or not mapped_id:
            continue
        id_param = int(float(mapped_id))
        master = master_by_id.get(id_param, {})
        if not master:
            continue
        if norm_param(source_param) == norm_param(master.get("nome_parametro")) and source_matrix == canonical_matrix(master.get("matriz")):
            continue
        rows.append(
            {
                "matriz_fonte": source_matrix,
                "parametro_fonte": source_param,
                "parametro_mestre": master.get("nome_parametro"),
                "id_parametro_mestre": id_param,
                "regra": "alias inferido do dry-run aprovado da C01",
            }
        )
    if not rows:
        return pd.DataFrame(columns=["matriz_fonte", "parametro_fonte", "parametro_mestre", "id_parametro_mestre", "regra"])
    return pd.DataFrame(rows).drop_duplicates(subset=["matriz_fonte", "parametro_fonte", "id_parametro_mestre"])


def fetch_existing_summary() -> tuple[pd.DataFrame, pd.DataFrame]:
    configure_database_access()
    engine = get_engine()
    try:
        with engine.begin() as conn:
            existing_summary = pd.read_sql(
                text(
                    """
                    SELECT matriz, nome_campanha, COUNT(*)::int AS registros,
                           COUNT(DISTINCT nome_ponto)::int AS pontos,
                           COUNT(DISTINCT nome_parametro)::int AS parametros
                    FROM public.fisico_analise_consolidada
                    WHERE codigo_interno_opyta = :code
                    GROUP BY matriz, nome_campanha
                    ORDER BY matriz, nome_campanha
                    """
                ),
                conn,
                params={"code": PROJECT_CODE},
            )
            existing_keys = pd.read_sql(
                text(
                    """
                    SELECT nome_campanha, nome_ponto, matriz, nome_parametro, laboratorio_responsavel
                    FROM public.fisico_analise_consolidada
                    WHERE codigo_interno_opyta = :code
                    """
                ),
                conn,
                params={"code": PROJECT_CODE},
            )
    finally:
        engine.dispose()
    return existing_summary, existing_keys


def normalize_key_df(df: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    out = df.copy()
    for col in columns:
        out[col] = out[col].map(clean_text)
    return out


def sql_literal(value: Any) -> str:
    if value is None:
        return "NULL"
    try:
        if pd.isna(value):
            return "NULL"
    except (TypeError, ValueError):
        pass
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return str(float(value)).rstrip("0").rstrip(".")
    text_value = str(value)
    if text_value == "":
        return "NULL"
    return "'" + text_value.replace("'", "''") + "'"


def build_insert_sql(preview: pd.DataFrame) -> list[str]:
    lines = [
        "-- BRAAEG001 C02 Agua Superficial dry-run incremental.",
        "-- Revisar e aprovar antes de executar; gerado sem escrita no banco.",
    ]
    for _, row in preview.iterrows():
        values = ", ".join(sql_literal(row.get(col)) for col in FISICO_INSERT_COLS)
        lines.append(f"INSERT INTO public.fisico_analise_consolidada ({', '.join(FISICO_INSERT_COLS)}) VALUES ({values});")
    return lines


def build_rollback_sql() -> list[str]:
    return [
        "-- Rollback planejado para remover somente a C02 Agua Superficial do BRAAEG001.",
        "-- Executar apenas se a carga incremental C02 tiver sido aplicada.",
        "DELETE FROM public.fisico_analise_consolidada",
        "WHERE codigo_interno_opyta = 'BRAAEG001'",
        "  AND nome_campanha = 'C002-2026-06-SC'",
        "  AND matriz = 'Agua Superficial';",
    ]


def write_excel(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name[:31], index=False)
    style_workbook(path)


def main() -> int:
    ts = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    client = RestClient(ENV_FILE)
    master = fetch_master(client)
    _, _, master_by_id = build_master_indexes(master)
    aliases = infer_aliases_from_c01(master_by_id)
    lookup = build_lookup(master, aliases)
    preview, mapping_issues = build_migration_preview(STAGING_C02, lookup)

    existing_summary, existing_keys = fetch_existing_summary()
    key_cols = ["nome_campanha", "nome_ponto", "matriz", "nome_parametro", "laboratorio_responsavel"]
    planned_keys = normalize_key_df(preview[key_cols], key_cols)
    existing_keys_norm = normalize_key_df(existing_keys, key_cols)
    duplicate_planned = planned_keys.merge(existing_keys_norm.drop_duplicates(), on=key_cols, how="inner")

    summary = pd.DataFrame(
        [
            {"Metrica": "registros_planejados_c02", "Valor": len(preview)},
            {"Metrica": "pontos_planejados_c02", "Valor": preview["nome_ponto"].nunique()},
            {"Metrica": "parametros_planejados_c02", "Valor": preview["nome_parametro"].nunique()},
            {"Metrica": "problemas_mapeamento", "Valor": len(mapping_issues)},
            {"Metrica": "duplicidades_com_banco", "Valor": len(duplicate_planned)},
            {"Metrica": "registros_banco_antes_braaeg001", "Valor": int(existing_summary["registros"].sum()) if not existing_summary.empty else 0},
            {"Metrica": "registros_banco_apos_planejado", "Valor": int(existing_summary["registros"].sum()) + len(preview) if not existing_summary.empty else len(preview)},
            {"Metrica": "parametros_com_vmp_agua_superficial", "Valor": int(preview[[col for col in VMP_COLS if col in preview.columns]].notna().any(axis=1).sum())},
            {"Metrica": "status", "Valor": "READY_FOR_APPLY_APPROVAL" if mapping_issues.empty and duplicate_planned.empty else "REVIEW_REQUIRED"},
        ]
    )

    dry_run_xlsx = OUTPUT_DIR / f"{ts}_dry_run_incremental_c02_agua_superficial_braaeg001.xlsx"
    apply_sql = OUTPUT_DIR / f"{ts}_apply_incremental_c02_agua_superficial_braaeg001.sql"
    rollback_sql = OUTPUT_DIR / f"{ts}_rollback_incremental_c02_agua_superficial_braaeg001.sql"
    log_path = LOG_DIR / f"{ts}_dry_run_incremental_c02_agua_superficial_braaeg001.json"

    write_excel(
        dry_run_xlsx,
        {
            "resumo": summary,
            "existente_banco": existing_summary,
            "preview_insert_c02": preview,
            "problemas_mapeamento": mapping_issues,
            "duplicidades_com_banco": duplicate_planned,
            "aliases_c01_reutilizados": aliases,
        },
    )
    apply_sql.write_text("\n".join(build_insert_sql(preview)) + "\n", encoding="utf-8")
    rollback_sql.write_text("\n".join(build_rollback_sql()) + "\n", encoding="utf-8")

    manifest = {
        "timestamp": ts,
        "status": clean_text(summary.loc[summary["Metrica"] == "status", "Valor"].iloc[0]),
        "staging": str(STAGING_C02),
        "dry_run_xlsx": str(dry_run_xlsx),
        "apply_sql": str(apply_sql),
        "rollback_sql": str(rollback_sql),
        "records_planned": int(len(preview)),
        "mapping_issues": int(len(mapping_issues)),
        "duplicate_existing_keys": int(len(duplicate_planned)),
        "existing_rows_before": int(existing_summary["registros"].sum()) if not existing_summary.empty else 0,
        "expected_rows_after": int(existing_summary["registros"].sum()) + len(preview) if not existing_summary.empty else len(preview),
    }
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    log_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"STATUS={manifest['status']}")
    print(f"DRY_RUN={dry_run_xlsx}")
    print(f"APPLY_SQL={apply_sql}")
    print(f"ROLLBACK_SQL={rollback_sql}")
    print(f"LOG={log_path}")
    print(json.dumps({k: manifest[k] for k in ['records_planned', 'mapping_issues', 'duplicate_existing_keys', 'existing_rows_before', 'expected_rows_after']}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
