from __future__ import annotations

import re
import sys
import unicodedata
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
import win32com.client
import fitz
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_CODE = "BRAAEG001"
PROJECT_NAME = "A&G Mineração"
CLIENT_NAME = "Brandt Meio Ambiente Ltda."
CAMPAIGN_C02 = "C002-2026-06-SC"
MATRIX = "Agua Superficial"
LABORATORY = "SGS"


def project_root() -> Path:
    return next(parent for parent in Path(__file__).resolve().parents if (parent / "src" / "opyta_analysis").exists())


ROOT = project_root()
BRANDT_DIR = Path(r"G:\Meu Drive\Opyta\Clientes\Clientes\Clientes\Brandt")
CLIENT_DIR = next(path for path in BRANDT_DIR.iterdir() if path.name.startswith("A&G"))
SOURCE_DIR = CLIENT_DIR / "resultados" / "Meio_fisico" / "Campanha-02" / "Resultados-Agua_superficial"
C01_WORKBOOK = CLIENT_DIR / "resultados" / "Meio_fisico" / "migracao" / "Resultados_Meio_Fisico.xlsx"
OUTPUT_DIR = CLIENT_DIR / "resultados" / "Meio_fisico" / "migracao" / "campanha_02_superficial"


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).replace("\xa0", " ").strip()


def fold(value: Any) -> str:
    text = clean_text(value)
    text = unicodedata.normalize("NFKD", text)
    return "".join(ch for ch in text if not unicodedata.combining(ch))


def norm_text(value: Any) -> str:
    text = fold(value).lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def normalize_point(raw_point: str) -> str:
    match = re.search(r"(\d+)$", clean_text(raw_point))
    if not match:
        return clean_text(raw_point)
    return f"PT_{int(match.group(1)):02d}"


def extract_pdf_sampling_date(pdf_path: Path) -> str:
    if not pdf_path.exists():
        return ""
    text = "\n".join(page.get_text() for page in fitz.open(pdf_path))
    dates_section = text.split("DATAS", 1)[1] if "DATAS" in text else text
    match = re.search(r"\d{2}/\d{2}/\d{4}\s+\d{2}:\d{2}", dates_section)
    if match:
        return re.sub(r"\s+", " ", match.group(0)).strip()
    match = re.search(r"\d{2}/\d{2}/\d{4}", dates_section)
    return match.group(0) if match else ""


def read_sgs_xls_files() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    files = sorted(SOURCE_DIR.glob("*.XLS"))
    if not files:
        raise FileNotFoundError(f"Nenhum XLS encontrado em {SOURCE_DIR}")

    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False

    result_rows: list[dict[str, Any]] = []
    vmp_rows: list[dict[str, Any]] = []
    report_rows: list[dict[str, Any]] = []

    try:
        for file in files:
            wb = excel.Workbooks.Open(str(file), 0, True)
            try:
                ws = wb.Worksheets.Item(1)
                report = clean_text(ws.Cells(1, 2).Text)
                source_project = clean_text(ws.Cells(4, 2).Text)
                xls_date = clean_text(ws.Cells(10, 5).Text)
                pdf_date = extract_pdf_sampling_date(file.with_suffix(".PDF"))
                raw_date = xls_date or pdf_date
                raw_point = clean_text(ws.Cells(11, 5).Text)
                point = normalize_point(raw_point)

                report_rows.append(
                    {
                        "Arquivo_XLS": file.name,
                        "Arquivo_PDF": file.with_suffix(".PDF").name if file.with_suffix(".PDF").exists() else "",
                        "Laudo": report,
                        "Projeto_Laudo": source_project,
                        "Ponto_Laudo": raw_point,
                        "Ponto": point,
                        "Campanha": CAMPAIGN_C02,
                        "Data": raw_date,
                        "Data_XLS": xls_date,
                        "Data_PDF": pdf_date,
                    }
                )

                for row in range(12, 61):
                    parameter = clean_text(ws.Cells(row, 1).Text)
                    if not parameter:
                        continue
                    unit = clean_text(ws.Cells(row, 2).Text)
                    vmp_01 = clean_text(ws.Cells(row, 3).Text)
                    vmp_02 = clean_text(ws.Cells(row, 4).Text)
                    result = clean_text(ws.Cells(row, 5).Text)

                    result_rows.append(
                        {
                            "Ponto": point,
                            "Campanha": CAMPAIGN_C02,
                            "Matriz": MATRIX,
                            "Parametro": parameter,
                            "Resultado": result,
                            "Unidade_Medida": unit,
                            "Laboratorio": LABORATORY,
                        }
                    )
                    vmp_rows.append(
                        {
                            "Ponto": point,
                            "Campanha": CAMPAIGN_C02,
                            "Laudo": report,
                            "Parametro": parameter,
                            "Unidade_Medida": unit,
                            "VMP_01_CONAMA_357_Classe_2": vmp_01,
                            "VMP_02_COPAM_08_Classe_2": vmp_02,
                            "Resultado": result,
                        }
                    )
            finally:
                wb.Close(False)
    finally:
        excel.Quit()

    return pd.DataFrame(result_rows), pd.DataFrame(vmp_rows), pd.DataFrame(report_rows)


def load_c01() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    capa = pd.read_excel(C01_WORKBOOK, sheet_name="Capa_Projeto", dtype=str).fillna("")
    points = pd.read_excel(C01_WORKBOOK, sheet_name="Pontos_e_Campanhas", dtype=str).fillna("")
    results = pd.read_excel(C01_WORKBOOK, sheet_name="Resultados_Meio_Fisico", dtype=str).fillna("")
    return capa, points, results


def build_points_c02(points_c01: pd.DataFrame, reports: pd.DataFrame) -> pd.DataFrame:
    c01_surface = points_c01[points_c01["Ponto"].isin(sorted(reports["Ponto"].unique()))].copy()
    c01_surface = c01_surface.drop_duplicates(subset=["Ponto"], keep="first")
    dates = reports.set_index("Ponto")["Data"].to_dict()
    c02 = c01_surface.copy()
    c02["Campanha"] = CAMPAIGN_C02
    c02["Data"] = c02["Ponto"].map(dates).fillna("")
    return c02[points_c01.columns]


def make_audit(results_c02: pd.DataFrame, reports: pd.DataFrame, points_c02: pd.DataFrame, results_c01: pd.DataFrame) -> dict[str, pd.DataFrame]:
    c01_surface = results_c01[norm_series(results_c01["Matriz"]) == norm_text(MATRIX)].copy()

    c01_params = pd.DataFrame(
        sorted({(norm_text(p), clean_text(p)) for p in c01_surface["Parametro"]}),
        columns=["Parametro_norm", "Parametro_C01"],
    )
    c02_params = pd.DataFrame(
        sorted({(norm_text(p), clean_text(p)) for p in results_c02["Parametro"]}),
        columns=["Parametro_norm", "Parametro_C02"],
    )
    param_delta = c01_params.merge(c02_params, on="Parametro_norm", how="outer")
    param_delta["Situacao"] = param_delta.apply(
        lambda row: "somente_C01"
        if clean_text(row.get("Parametro_C02")) == ""
        else ("somente_C02" if clean_text(row.get("Parametro_C01")) == "" else "presente_C01_C02"),
        axis=1,
    )

    units_c01 = (
        c01_surface.assign(Parametro_norm=norm_series(c01_surface["Parametro"]))
        .groupby("Parametro_norm")["Unidade_Medida"]
        .apply(lambda values: " | ".join(sorted({clean_text(v) for v in values if clean_text(v)})))
        .reset_index(name="Unidades_C01")
    )
    units_c02 = (
        results_c02.assign(Parametro_norm=norm_series(results_c02["Parametro"]))
        .groupby("Parametro_norm")["Unidade_Medida"]
        .apply(lambda values: " | ".join(sorted({clean_text(v) for v in values if clean_text(v)})))
        .reset_index(name="Unidades_C02")
    )
    units = units_c01.merge(units_c02, on="Parametro_norm", how="outer").fillna("")
    units["Situacao"] = units.apply(
        lambda row: "unidade_divergente" if row["Unidades_C01"] and row["Unidades_C02"] and row["Unidades_C01"] != row["Unidades_C02"] else "ok",
        axis=1,
    )

    grain = ["Ponto", "Campanha", "Matriz", "Parametro", "Laboratorio"]
    duplicates = results_c02[results_c02.duplicated(subset=grain, keep=False)].sort_values(grain)

    expected_points = {f"PT_{i:02d}" for i in range(1, 13)}
    actual_points = set(results_c02["Ponto"].unique())
    point_rows = []
    for point in sorted(expected_points | actual_points):
        point_results = results_c02[results_c02["Ponto"] == point]
        point_rows.append(
            {
                "Ponto": point,
                "Esperado": "sim" if point in expected_points else "nao",
                "Presente_C02": "sim" if point in actual_points else "nao",
                "Registros_Resultados": len(point_results),
                "Parametros_Unicos": point_results["Parametro"].nunique(),
                "Data": clean_text(points_c02.loc[points_c02["Ponto"] == point, "Data"].iloc[0]) if point in set(points_c02["Ponto"]) else "",
            }
        )

    report_project_counts = Counter(reports["Projeto_Laudo"])
    date_mismatch = reports[
        (reports["Data_XLS"].map(clean_text) != "")
        & (reports["Data_PDF"].map(clean_text) != "")
        & (reports["Data_XLS"].map(lambda value: clean_text(value)[:10]) != reports["Data_PDF"].map(lambda value: clean_text(value)[:10]))
    ]
    summary = pd.DataFrame(
        [
            {"Item": "Arquivos XLS lidos", "Valor": len(reports)},
            {"Item": "Arquivos PDF pareados", "Valor": int((reports["Arquivo_PDF"] != "").sum())},
            {"Item": "Registros C02 extraidos", "Valor": len(results_c02)},
            {"Item": "Pontos C02", "Valor": results_c02["Ponto"].nunique()},
            {"Item": "Parametros C02", "Valor": results_c02["Parametro"].nunique()},
            {"Item": "Duplicidades no grao tecnico", "Valor": len(duplicates)},
            {"Item": "Parametros somente C02", "Valor": int((param_delta["Situacao"] == "somente_C02").sum())},
            {"Item": "Parametros somente C01", "Valor": int((param_delta["Situacao"] == "somente_C01").sum())},
            {"Item": "Unidades divergentes C01 x C02", "Valor": int((units["Situacao"] == "unidade_divergente").sum())},
            {"Item": "Datas XLS x PDF divergentes", "Valor": len(date_mismatch)},
            {"Item": "Projeto informado nos laudos", "Valor": " | ".join(f"{k}: {v}" for k, v in sorted(report_project_counts.items()))},
            {"Item": "Criterio de data", "Valor": "Usar Data_XLS quando preenchida; usar Data_PDF apenas como complemento quando Data_XLS estiver vazia."},
            {"Item": "Observacao operacional", "Valor": "Staging sem carga no Supabase; requer aprovacao do Gate A/C02 antes da migracao incremental."},
        ]
    )

    return {
        "Resumo": summary,
        "Laudos_C02": reports,
        "Pontos_C02": pd.DataFrame(point_rows),
        "Parametros_Delta": param_delta,
        "Unidades_C01_C02": units,
        "Duplicidades": duplicates,
    }


def norm_series(series: pd.Series) -> pd.Series:
    return series.map(norm_text)


def write_workbook(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet, df in sheets.items():
            df.to_excel(writer, sheet_name=sheet[:31], index=False)
    style_workbook(path)


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column_cells in ws.columns:
            values = [clean_text(cell.value) for cell in column_cells]
            width = min(max(max((len(v) for v in values), default=0) + 2, 10), 48)
            ws.column_dimensions[get_column_letter(column_cells[0].column)].width = width
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main() -> int:
    timestamp = datetime.now().strftime("%Y%m%dT%H%M%S")
    results_c02, vmps_c02, reports = read_sgs_xls_files()
    capa_c01, points_c01, results_c01 = load_c01()
    points_c02 = build_points_c02(points_c01, reports)

    staging_path = OUTPUT_DIR / f"{timestamp}_staging_resultados_meio_fisico_c02_agua_superficial.xlsx"
    audit_path = OUTPUT_DIR / f"{timestamp}_auditoria_c02_vs_c01_agua_superficial.xlsx"

    write_workbook(
        staging_path,
        {
            "Capa_Projeto": capa_c01,
            "Pontos_e_Campanhas": points_c02,
            "Resultados_Meio_Fisico": results_c02,
            "Laudos_C02": reports,
            "VMP_Laudos_C02": vmps_c02,
        },
    )
    audit = make_audit(results_c02, reports, points_c02, results_c01)
    write_workbook(audit_path, audit)

    print(f"STAGING={staging_path}")
    print(f"AUDITORIA={audit_path}")
    print(f"REGISTROS_C02={len(results_c02)}")
    print(f"PONTOS_C02={results_c02['Ponto'].nunique()}")
    print(f"PARAMETROS_C02={results_c02['Parametro'].nunique()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
