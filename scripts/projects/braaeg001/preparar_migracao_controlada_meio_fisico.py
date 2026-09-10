from __future__ import annotations

import argparse
import json
import os
import re
import sys
import unicodedata
import urllib.parse
import urllib.request
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from dotenv import load_dotenv

ROOT = next(parent for parent in Path(__file__).resolve().parents if (parent / "src" / "opyta_analysis").exists())
SRC_PATH = ROOT / "src"
if str(SRC_PATH) not in sys.path:
    sys.path.insert(0, str(SRC_PATH))

from opyta_analysis.meio_fisico.rules import parse_resultado


PROJECT_CODE = "BRAAEG001"
PROJECT_NAME = "A&G Mineração"
CLIENT_NAME = "Brandt Meio Ambiente Ltda."
CLIENT_ID = 1
ENV_FILE = Path(r"G:\Meu Drive\Opyta\Opyta_Data\.env")

VMP_COLS = [
    "vmp_357_cl1_min",
    "vmp_357_cl1_max",
    "vmp_357_cl2_min",
    "vmp_357_cl2_max",
    "vmp_amonia_dinamico",
    "vmp_454_n1",
    "vmp_454_n2",
    "vmp_396_consumo_humano",
    "vmp_396_dessedentacao_animal",
    "vmp_396_irrigacao",
    "vmp_396_recreacao",
    "vmp_430_padrao",
]

PARAM_TABLE_VMP_COLS = [col for col in VMP_COLS if col != "vmp_amonia_dinamico"]

PARAM_TABLE_COLS = [
    "nome_parametro",
    "unidade_medida",
    "matriz",
    "vmp_357_cl1",
    "vmp_357_cl2",
    "vmp_357_cl3",
    "vmp_396_vmp",
    "vmp_396_vr",
    "vmp_454_n1",
    "vmp_454_n2",
    "vmp_430_padrao",
    "vmp_396_consumo_humano",
    "vmp_396_animal",
    "vmp_396_irrigacao",
    "vmp_396_recreacao",
    "vmp_357_cl2_min",
    "vmp_357_cl2_max",
    "vmp_357_cl1_min",
    "vmp_357_cl1_max",
    "vmp_396_dessedentacao_animal",
]

MATRIX_CANONICAL = {
    "agua superficial": "Água Superficial",
    "agua subterranea": "Água Subterrânea",
    "sedimento": "Sedimento",
    "efluente": "Efluente",
}

UNIT_CANONICAL = {
    "mg/l": "mg/L",
    "mgb/l": "mg/L",
    "mgf/l": "mg/L",
    "mgf-/l": "mg/L",
    "mgmo/l": "mg/L",
    "mgv/l": "mg/L",
    "mgco/l": "mg/L",
    "mgca/l": "mg/L",
    "mgk/l": "mg/L",
    "mgmg/l": "mg/L",
    "mgna/l": "mg/L",
    "mgal/l": "mg/L",
    "mgas/l": "mg/L",
    "mgba/l": "mg/L",
    "mgcd/l": "mg/L",
    "mgcr/l": "mg/L",
    "mgfe/l": "mg/L",
    "mghg/l": "mg/L",
    "mgmn/l": "mg/L",
    "mgpb/l": "mg/L",
    "mgse/l": "mg/L",
    "mgcaco3/l": "mg/L",
    "mgcl2/l": "mg/L",
    "mgcu/l": "mg/L",
    "mgzn/l": "mg/L",
    "mgni/l": "mg/L",
    "mgcl/l": "mg/L",
    "mgcl-/l": "mg/L",
    "mgcn-/l": "mg/L",
    "mgo2/l": "mg/L",
    "mgn_no3/l": "mg/L",
    "mgn_no2/l": "mg/L",
    "mgn_nh3/l": "mg/L",
    "mgso4/l": "mg/L",
    "mg/kg": "mg/kg",
    "mg hg/kg": "mg/kg",
    "nmp/100ml": "NMP/100mL",
    "ufc/100ml": "NMP/100mL",
    "mv": "mV",
    "oc": "°C",
    "şc": "°C",
    "ºc": "°C",
    "°c": "°C",
}

PARAMETER_ALIASES = {
    "nitrato n": "nitrato",
    "nitrito n": "nitrito",
    "amonia": "nitrogenio amoniacal",
    "cloreto": "cloretos",
}

GROUNDWATER_USER_LIMITS_MG_L = {
    "boro total": {
        "vmp_396_consumo_humano": 0.5,
        "vmp_396_dessedentacao_animal": 5.0,
        "vmp_396_irrigacao": 1.0,
    },
    "fluoreto": {
        "vmp_396_consumo_humano": 1.5,
        "vmp_396_dessedentacao_animal": 2.0,
        "vmp_396_irrigacao": 1.0,
    },
    "molibdenio total": {
        "vmp_396_consumo_humano": 0.07,
        "vmp_396_dessedentacao_animal": 0.15,
        "vmp_396_irrigacao": 0.01,
    },
    "vanadio total": {
        "vmp_396_consumo_humano": 0.1,
        "vmp_396_dessedentacao_animal": 0.1,
        "vmp_396_irrigacao": 0.1,
    },
}

EXISTING_UPDATE_RULES = {
    72: {"vmp_396_consumo_humano": 250.0},
    40: {"vmp_357_cl2_max": 0.009},
    3: {"vmp_396_consumo_humano": 0.7, "vmp_357_cl2_max": 0.7},
    92: {"unidade_medida": "mV"},
    155: {"unidade_medida": "mg/L"},
}


def fold(text: Any) -> str:
    value = "" if text is None or pd.isna(text) else str(text)
    value = value.replace("\xa0", " ").strip()
    value = unicodedata.normalize("NFKD", value)
    return "".join(ch for ch in value if not unicodedata.combining(ch))


def clean_text(text: Any) -> str:
    if text is None:
        return ""
    try:
        if pd.isna(text):
            return ""
    except (TypeError, ValueError):
        pass
    return str(text).replace("\xa0", " ").strip()


def norm_text(text: Any) -> str:
    value = fold(text).lower()
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def norm_param(text: Any) -> str:
    value = norm_text(text)
    return PARAMETER_ALIASES.get(value, value)


def norm_unit(text: Any) -> str:
    value = fold(text).lower().replace(" ", "")
    return UNIT_CANONICAL.get(value, clean_text(text))


def canonical_matrix(value: Any) -> str:
    return MATRIX_CANONICAL.get(norm_text(value), fold(value).strip())


def sql_literal(value: Any) -> str:
    if value is None or value == "":
        return "NULL"
    try:
        if pd.isna(value):
            return "NULL"
    except (TypeError, ValueError):
        pass
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return str(float(value)).rstrip("0").rstrip(".")
    text = str(value).replace("'", "''")
    return f"'{text}'"


def style_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col_idx in range(1, ws.max_column + 1):
            col = get_column_letter(col_idx)
            max_len = 10
            for cell in ws[col][: min(ws.max_row, 200)]:
                max_len = max(max_len, len("" if cell.value is None else str(cell.value)))
            ws.column_dimensions[col].width = min(max_len + 2, 65)
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


class RestClient:
    def __init__(self, env_file: Path) -> None:
        load_dotenv(env_file, override=True)
        self.url = os.getenv("SUPABASE_URL", "").rstrip("/")
        self.key = os.getenv("SUPABASE_ANON_KEY", "")
        if not self.url or not self.key:
            raise RuntimeError("SUPABASE_URL/SUPABASE_ANON_KEY ausentes.")

    def get(self, table: str, query: str, page_size: int = 1000) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        offset = 0
        while True:
            endpoint = f"{self.url}/rest/v1/{urllib.parse.quote(table)}?{query}"
            req = urllib.request.Request(
                endpoint,
                headers={
                    "apikey": self.key,
                    "Authorization": f"Bearer {self.key}",
                    "Accept": "application/json",
                    "Range": f"{offset}-{offset + page_size - 1}",
                },
            )
            with urllib.request.urlopen(req, timeout=60) as response:
                batch = json.loads(response.read().decode("utf-8"))
            rows.extend(batch)
            if len(batch) < page_size:
                return rows
            offset += page_size


def read_capa(path: Path) -> dict[str, Any]:
    raw = pd.read_excel(path, sheet_name="Capa_Projeto", dtype=str)
    if raw.empty:
        return {}
    row = raw.iloc[0].to_dict()
    return {str(k): v for k, v in row.items()}


def project_payload(capa: dict[str, Any]) -> dict[str, Any]:
    return {
        "id_cliente": CLIENT_ID,
        "nome_projeto": clean_text(capa.get("Nome_do_Projeto")) or PROJECT_NAME,
        "codigo_interno_opyta": PROJECT_CODE,
        "data_inicio": str(capa.get("Data_Inicio")).split(" ")[0] if fold(capa.get("Data_Inicio")) else None,
        "data_fim_prevista": str(capa.get("Data_Fim_Prevista")).split(" ")[0] if fold(capa.get("Data_Fim_Prevista")) else None,
        "descricao_projeto": "Meio físico - migração inicial A&G Mineração",
        "status_projeto": "ativo",
    }


def fetch_master(client: RestClient) -> list[dict[str, Any]]:
    return client.get("parametros_analise", "select=*&order=id_parametro.asc")


def build_master_indexes(master: list[dict[str, Any]]) -> tuple[dict[tuple[str, str], dict[str, Any]], dict[str, list[dict[str, Any]]], dict[int, dict[str, Any]]]:
    by_matrix_param: dict[tuple[str, str], dict[str, Any]] = {}
    by_param: dict[str, list[dict[str, Any]]] = {}
    by_id: dict[int, dict[str, Any]] = {}
    for row in master:
        matrix = canonical_matrix(row.get("matriz")) if row.get("matriz") else ""
        param = norm_param(row.get("nome_parametro"))
        by_matrix_param.setdefault((matrix, param), row)
        by_param.setdefault(param, []).append(row)
        by_id[int(row["id_parametro"])] = row
    return by_matrix_param, by_param, by_id


def find_row(by_matrix_param: dict[tuple[str, str], dict[str, Any]], by_param: dict[str, list[dict[str, Any]]], matrix: str, param: str) -> dict[str, Any] | None:
    key = (canonical_matrix(matrix), norm_param(param))
    if key in by_matrix_param:
        return by_matrix_param[key]
    global_rows = by_param.get(norm_param(param), [])
    return global_rows[0] if global_rows else None


def parameter_insert_payload(row: pd.Series) -> dict[str, Any]:
    matrix = canonical_matrix(row["matriz"])
    param_norm = norm_param(row["parametro"])
    payload = {col: None for col in PARAM_TABLE_COLS}
    payload.update(
        {
            "nome_parametro": clean_text(row["parametro"]),
            "unidade_medida": norm_unit(row["unidade_resultado"]),
            "matriz": matrix,
        }
    )
    if matrix == "Água Subterrânea":
        payload.update(GROUNDWATER_USER_LIMITS_MG_L.get(param_norm, {}))
    return payload


def build_actions(decisions: pd.DataFrame, master_by_id: dict[int, dict[str, Any]]) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    inserts: list[dict[str, Any]] = []
    updates: list[dict[str, Any]] = []
    backup: list[dict[str, Any]] = []
    aliases: list[dict[str, Any]] = []
    seen_update_ids: set[int] = set()

    for _, row in decisions.iterrows():
        id_raw = fold(row.get("id_parametro_mestre"))
        id_param = int(float(id_raw)) if id_raw else None
        matrix = canonical_matrix(row.get("matriz"))
        param_norm = norm_param(row.get("parametro"))

        if not id_param:
            if param_norm == "escherichia coli" and matrix == "Água Subterrânea":
                aliases.append(
                    {
                        "matriz_fonte": row["matriz"],
                        "parametro_fonte": row["parametro"],
                        "parametro_mestre": "Escherichia Coli por tubos múltiplos (substrato enzimático) - NMP",
                        "id_parametro_mestre": 237,
                        "regra": "alias local de migracao; nao existe tabela de alias no banco",
                    }
                )
                continue
            payload = parameter_insert_payload(row)
            inserts.append(
                {
                    **payload,
                    "matriz_fonte": row["matriz"],
                    "parametro_fonte": row["parametro"],
                    "instrucao_usuario": row.get("instrucao_para_execucao", ""),
                    "acao": "insert_parametros_analise",
                }
            )
            continue

        if id_param in EXISTING_UPDATE_RULES and id_param not in seen_update_ids:
            seen_update_ids.add(id_param)
            original = master_by_id.get(id_param, {})
            changes = EXISTING_UPDATE_RULES[id_param]
            backup.append({**original, "acao": "backup_update_parametros_analise"})
            updates.append(
                {
                    "id_parametro": id_param,
                    "nome_parametro": original.get("nome_parametro"),
                    "matriz": original.get("matriz"),
                    "instrucao_usuario": row.get("instrucao_para_execucao", ""),
                    "acao": "update_parametros_analise",
                    **changes,
                }
            )
        elif id_param == 54:
            aliases.append(
                {
                    "matriz_fonte": row["matriz"],
                    "parametro_fonte": row["parametro"],
                    "parametro_mestre": "Nitrogênio Amoniacal",
                    "id_parametro_mestre": 54,
                    "regra": "VMP dinamico por pH na carga consolidada; parametros_analise nao possui coluna vmp_amonia_dinamico",
                }
            )
        elif id_param == 41:
            aliases.append(
                {
                    "matriz_fonte": row["matriz"],
                    "parametro_fonte": row["parametro"],
                    "parametro_mestre": row.get("parametro_mestre", ""),
                    "id_parametro_mestre": id_param,
                    "regra": "unidade microbiologica ajustada para NMP/100mL na carga consolidada",
                }
            )

    aliases.extend(
        [
            {
                "matriz_fonte": "Agua Superficial",
                "parametro_fonte": "Cloreto",
                "parametro_mestre": "Cloretos",
                "id_parametro_mestre": 190,
                "regra": "usar cadastro especifico de Agua Superficial com VMP 357=250",
            }
        ]
    )
    return pd.DataFrame(inserts), pd.DataFrame(updates), pd.DataFrame(backup), pd.DataFrame(aliases)


def apply_actions_to_master(master: list[dict[str, Any]], inserts: pd.DataFrame, updates: pd.DataFrame) -> list[dict[str, Any]]:
    proposed = [dict(row) for row in master]
    by_id = {int(row["id_parametro"]): row for row in proposed}
    next_id = max(by_id) + 1 if by_id else 1
    for _, upd in updates.iterrows():
        row = by_id.get(int(upd["id_parametro"]))
        if not row:
            continue
        for col in PARAM_TABLE_COLS:
            if col in upd and fold(upd[col]):
                row[col] = upd[col]
    for _, ins in inserts.iterrows():
        row = {col: None for col in ["id_parametro"] + PARAM_TABLE_COLS}
        row["id_parametro"] = next_id
        next_id += 1
        for col in PARAM_TABLE_COLS:
            if col in ins:
                row[col] = None if not fold(ins[col]) else ins[col]
        proposed.append(row)
    return proposed


def build_lookup(proposed_master: list[dict[str, Any]], aliases: pd.DataFrame) -> dict[tuple[str, str], dict[str, Any]]:
    by_matrix_param, by_param, by_id = build_master_indexes(proposed_master)
    lookup: dict[tuple[str, str], dict[str, Any]] = {}
    for row in proposed_master:
        matrix = canonical_matrix(row.get("matriz")) if row.get("matriz") else ""
        param = norm_param(row.get("nome_parametro"))
        lookup[(matrix, param)] = row
    for _, alias in aliases.iterrows():
        target = by_id.get(int(alias["id_parametro_mestre"]))
        if target:
            lookup[(canonical_matrix(alias["matriz_fonte"]), norm_param(alias["parametro_fonte"]))] = target
    return lookup


def active_vmp(row: dict[str, Any], matrix: str) -> dict[str, Any]:
    matrix = canonical_matrix(matrix)
    if matrix == "Água Superficial":
        cols = ["vmp_357_cl1_min", "vmp_357_cl1_max", "vmp_357_cl2_min", "vmp_357_cl2_max"]
    elif matrix == "Água Subterrânea":
        cols = ["vmp_396_consumo_humano", "vmp_396_dessedentacao_animal", "vmp_396_irrigacao", "vmp_396_recreacao"]
    elif matrix == "Sedimento":
        cols = ["vmp_454_n1", "vmp_454_n2"]
    else:
        cols = PARAM_TABLE_VMP_COLS
    return {col: row.get(col) for col in cols if row.get(col) not in (None, "")}


def ammonia_limit(ph: float | None) -> float | None:
    if ph is None:
        return None
    if ph <= 7.5:
        return 3.7
    if ph <= 8.0:
        return 2.0
    if ph <= 8.5:
        return 1.0
    return 0.5


def build_migration_preview(source_path: Path, proposed_lookup: dict[tuple[str, str], dict[str, Any]]) -> tuple[pd.DataFrame, pd.DataFrame]:
    results = pd.read_excel(source_path, sheet_name="Resultados_Meio_Fisico", dtype=str).fillna("")
    points = pd.read_excel(source_path, sheet_name="Pontos_e_Campanhas", dtype=str).fillna("")
    point_map = {
        (clean_text(row["Ponto"]), clean_text(row["Campanha"])): row.to_dict()
        for _, row in points.iterrows()
    }

    ph_by_point: dict[tuple[str, str], float | None] = {}
    for _, row in results.iterrows():
        if norm_text(row["Parametro"]) == "ph in situ":
            value, sign = parse_resultado(row["Resultado"])
            ph_by_point[(clean_text(row["Ponto"]), clean_text(row["Campanha"]))] = None if sign in {"<", "<="} else value

    rows: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    for idx, row in results.iterrows():
        matrix = canonical_matrix(row["Matriz"])
        param_key = norm_param(row["Parametro"])
        lookup_key = (matrix, param_key)
        master = proposed_lookup.get(lookup_key) or proposed_lookup.get(("", param_key))
        if not master:
            issues.append(
                {
                    "linha_fonte": int(idx) + 2,
                    "matriz": row["Matriz"],
                    "parametro": row["Parametro"],
                    "problema": "parametro_sem_mapeamento_pos_gate_b",
                }
            )
            master = {}
        point = point_map.get((clean_text(row["Ponto"]), clean_text(row["Campanha"])), {})
        value, sign = parse_resultado(row["Resultado"])
        out = {
            "codigo_interno_opyta": PROJECT_CODE,
            "nome_projeto": PROJECT_NAME,
            "nome_empresa": CLIENT_NAME,
            "nome_campanha": clean_text(row["Campanha"]),
            "nome_ponto": clean_text(row["Ponto"]),
            "data_hora_coleta": clean_text(point.get("Data")) or None,
            "latitude": pd.to_numeric(point.get("Latitude"), errors="coerce"),
            "longitude": pd.to_numeric(point.get("Longitude"), errors="coerce"),
            "bacia_hidrografica": clean_text(point.get("Bacia_Hidrografica")) or None,
            "matriz": matrix,
            "nome_parametro": clean_text(row["Parametro"]),
            "sinal_limite": sign or None,
            "valor_medido": value,
            "unidade_medida": norm_unit(row["Unidade_Medida"]),
            "unidade_original_laudo": clean_text(row["Unidade_Medida"]) or None,
            "laboratorio_responsavel": clean_text(row["Laboratorio"]) or None,
            "observacoes_resultado": None,
            "id_parametro_mapeado": master.get("id_parametro"),
            "nome_parametro_mapeado": master.get("nome_parametro"),
        }
        for col in VMP_COLS:
            out[col] = None
        for col, vmp in active_vmp(master, matrix).items():
            out[col] = vmp
        if matrix == "Água Superficial" and param_key == "nitrogenio amoniacal":
            out["vmp_amonia_dinamico"] = ammonia_limit(ph_by_point.get((clean_text(row["Ponto"]), clean_text(row["Campanha"]))))
        rows.append(out)
    return pd.DataFrame(rows), pd.DataFrame(issues)


def insert_sql(table: str, row: dict[str, Any], cols: list[str]) -> str:
    values = ", ".join(sql_literal(row.get(col)) for col in cols)
    return f"INSERT INTO public.{table} ({', '.join(cols)}) VALUES ({values});"


def update_sql(table: str, key_col: str, key_value: Any, changes: dict[str, Any]) -> str:
    set_clause = ", ".join(f"{col} = {sql_literal(value)}" for col, value in changes.items())
    return f"UPDATE public.{table} SET {set_clause} WHERE {key_col} = {sql_literal(key_value)};"


def build_sql(inserts: pd.DataFrame, updates: pd.DataFrame, backup: pd.DataFrame, project: dict[str, Any]) -> tuple[list[str], list[str]]:
    apply_lines = ["-- BRAAEG001 Gate B dry-run SQL; revisar antes de executar."]
    rollback_lines = ["-- BRAAEG001 Gate B rollback SQL; executar apenas se o apply tiver sido aplicado."]

    project_cols = ["id_cliente", "nome_projeto", "codigo_interno_opyta", "data_inicio", "data_fim_prevista", "descricao_projeto", "status_projeto"]
    apply_lines.append(insert_sql("projetos", project, project_cols))
    rollback_lines.insert(1, "DELETE FROM public.projetos WHERE codigo_interno_opyta = 'BRAAEG001';")

    for _, row in inserts.iterrows():
        payload = {col: row.get(col) for col in PARAM_TABLE_COLS}
        apply_lines.append(insert_sql("parametros_analise", payload, PARAM_TABLE_COLS))
        rollback_lines.append(
            "DELETE FROM public.parametros_analise "
            f"WHERE nome_parametro = {sql_literal(payload.get('nome_parametro'))} "
            f"AND matriz = {sql_literal(payload.get('matriz'))};"
        )

    for _, row in updates.iterrows():
        changes = {col: row[col] for col in PARAM_TABLE_COLS if col in row and fold(row[col])}
        changes.pop("nome_parametro", None)
        changes.pop("matriz", None)
        apply_lines.append(update_sql("parametros_analise", "id_parametro", row["id_parametro"], changes))

    if not backup.empty:
        for _, row in backup.iterrows():
            restored = {col: row.get(col) for col in PARAM_TABLE_COLS}
            rollback_lines.append(update_sql("parametros_analise", "id_parametro", row["id_parametro"], restored))

    return apply_lines, rollback_lines


def main() -> int:
    parser = argparse.ArgumentParser(description="Prepare BRAAEG001 controlled migration dry-run artifacts.")
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--decisions", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--log-dir", default=Path("logs/validacao_meio_fisico"), type=Path)
    parser.add_argument("--env-file", default=ENV_FILE, type=Path)
    args = parser.parse_args()

    ts = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    args.log_dir.mkdir(parents=True, exist_ok=True)

    client = RestClient(args.env_file)
    master = fetch_master(client)
    _, _, master_by_id = build_master_indexes(master)
    projects = client.get("projetos", "select=*&codigo_interno_opyta=eq.BRAAEG001")
    existing_rows = client.get("fisico_analise_consolidada", "select=id_consolidated&codigo_interno_opyta=eq.BRAAEG001&limit=1")

    decisions = pd.read_excel(args.decisions, sheet_name="decisoes_normalizadas", dtype=str).fillna("")
    inserts, updates, backup, aliases = build_actions(decisions, master_by_id)
    proposed_master = apply_actions_to_master(master, inserts, updates)
    lookup = build_lookup(proposed_master, aliases)
    migration_preview, mapping_issues = build_migration_preview(args.source, lookup)

    capa = read_capa(args.source)
    project = project_payload(capa)
    apply_lines, rollback_lines = build_sql(inserts, updates, backup, project)

    apply_sql = args.output_dir / f"{ts}_braaeg001_gate_b_apply_dry_run.sql"
    rollback_sql = args.output_dir / f"{ts}_braaeg001_gate_b_rollback.sql"
    apply_sql.write_text("\n".join(apply_lines) + "\n", encoding="utf-8")
    rollback_sql.write_text("\n".join(rollback_lines) + "\n", encoding="utf-8")

    migration_summary = pd.DataFrame(
        [{"metrica": "linhas_preview_migracao", "valor": int(len(migration_preview))}]
        + [{"metrica": f"matriz_{key}", "valor": int(value)} for key, value in sorted(Counter(migration_preview["matriz"]).items())]
        + [{"metrica": "parametros_sem_mapeamento", "valor": int(len(mapping_issues))}]
        + [{"metrica": "valores_numericos_nulos", "valor": int(migration_preview["valor_medido"].isna().sum())}]
        + [{"metrica": "projeto_existente_supabase", "valor": int(bool(projects))}]
        + [{"metrica": "registros_existentes_braaeg001", "valor": int(bool(existing_rows))}]
        + [{"metrica": "parametros_insert_planejados", "valor": int(len(inserts))}]
        + [{"metrica": "parametros_update_planejados", "valor": int(len(updates))}]
        + [{"metrica": "aliases_migracao", "valor": int(len(aliases))}]
    )

    output_xlsx = args.output_dir / f"{ts}_dry_run_migracao_controlada_braaeg001.xlsx"
    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        migration_summary.to_excel(writer, sheet_name="resumo", index=False)
        pd.DataFrame([project]).to_excel(writer, sheet_name="projeto_payload", index=False)
        pd.DataFrame(projects).to_excel(writer, sheet_name="projeto_existente", index=False)
        inserts.to_excel(writer, sheet_name="parametros_insert", index=False)
        updates.to_excel(writer, sheet_name="parametros_update", index=False)
        backup.to_excel(writer, sheet_name="parametros_backup", index=False)
        aliases.to_excel(writer, sheet_name="alias_migracao", index=False)
        mapping_issues.to_excel(writer, sheet_name="problemas_mapeamento", index=False)
        migration_preview.to_excel(writer, sheet_name="preview_migracao", index=False)
        pd.DataFrame({"sql_apply": apply_lines}).to_excel(writer, sheet_name="sql_apply", index=False)
        pd.DataFrame({"sql_rollback": rollback_lines}).to_excel(writer, sheet_name="sql_rollback", index=False)
    style_workbook(output_xlsx)

    report = {
        "status": "DRY_RUN_READY" if mapping_issues.empty and not projects and not existing_rows else "DRY_RUN_REVIEW_REQUIRED",
        "timestamp_utc": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "source": str(args.source),
        "decisions": str(args.decisions),
        "output_xlsx": str(output_xlsx),
        "apply_sql": str(apply_sql),
        "rollback_sql": str(rollback_sql),
        "summary": {row["metrica"]: row["valor"] for _, row in migration_summary.iterrows()},
    }
    output_json = args.log_dir / f"{ts}_dry_run_migracao_controlada_braaeg001.json"
    output_json.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"STATUS={report['status']}")
    print(f"OUTPUT_XLSX={output_xlsx}")
    print(f"OUTPUT_JSON={output_json}")
    print(f"APPLY_SQL={apply_sql}")
    print(f"ROLLBACK_SQL={rollback_sql}")
    print("SUMMARY=" + json.dumps(report["summary"], ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
