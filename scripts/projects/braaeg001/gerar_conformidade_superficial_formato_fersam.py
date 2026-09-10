from __future__ import annotations

import re
import shutil
import unicodedata
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


CLIENT_ROOT = Path(r"G:\Meu Drive\Opyta\Clientes\Clientes\Clientes\Brandt")
PROJECT_DIR = next(path for path in CLIENT_ROOT.iterdir() if path.name.startswith("A&G"))
CONSOLIDATED = (
    PROJECT_DIR
    / "resultados"
    / "Meio_fisico"
    / "migracao"
    / "consolidacao_pos_c02"
    / "20260729T183622Z_consolidado_meio_fisico_braaeg001_pos_c02.xlsx"
)
OUTPUT = (
    PROJECT_DIR
    / "resultados"
    / "Meio_fisico"
    / "resultados"
    / "superficial"
    / "01_Conformidade_Agua_Superficial.xlsx"
)
BACKUP = OUTPUT.with_name("01_Conformidade_Agua_Superficial_detalhada.xlsx")

CAMPAIGN_LABELS = {
    "C001-2026-02-CH": "Campanha-01-Chuva",
    "C002-2026-06-SC": "Campanha-02-Seca",
}


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def fold(value: Any) -> str:
    text = clean_text(value).replace("Ä™", "e").replace("Ä˜", "E")
    text = unicodedata.normalize("NFKD", text)
    return "".join(ch for ch in text if not unicodedata.combining(ch))


def norm(value: Any) -> str:
    text = fold(value).lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def fix_ptbr_label(value: Any) -> str:
    out = clean_text(value)
    if "Ã" in out or "Â" in out:
        try:
            out = out.encode("latin1").decode("utf-8")
        except UnicodeError:
            pass
    replacements = {
        "ArsÄ™nio": "ArsÃªnio",
        "ManganÄ™s": "ManganÃªs",
        "NitrogÄ™nio": "NitrogÃªnio",
        "OxigÄ™nio": "OxigÃªnio",
        "Bioquimica": "BioquÃ­mica",
        "Quimica": "QuÃ­mica",
        "Solidos": "SÃ³lidos",
        "Aluminio": "AlumÃ­nio",
        "Fosforo": "FÃ³sforo",
        "Cadmio": "CÃ¡dmio",
    }
    for src, dst in replacements.items():
        out = out.replace(src, dst)
    out = out.replace("S" + chr(195) + chr(179) + "lidos", "S" + chr(243) + "lidos")
    out = out.replace("Bioqu" + chr(195) + chr(173) + "mica", "Bioqu" + chr(237) + "mica")
    out = out.replace("Qu" + chr(195) + chr(173) + "mica", "Qu" + chr(237) + "mica")
    out = out.replace(chr(281), chr(234))
    return out


def point_order(points: list[str]) -> list[str]:
    def key(point: str) -> int:
        raw = point.split("_")[-1]
        return int(raw) if raw.isdigit() else 999

    return sorted(points, key=key)


def format_number(value: Any) -> str:
    if value is None or pd.isna(value):
        return "-"
    value = float(value)
    if abs(value) >= 100:
        text = f"{value:.0f}"
    elif abs(value) >= 10:
        text = f"{value:.2f}".rstrip("0").rstrip(".")
    elif abs(value) >= 1:
        text = f"{value:.3f}".rstrip("0").rstrip(".")
    else:
        text = f"{value:.5f}".rstrip("0").rstrip(".")
    return text.replace(".", ",")


def format_result(sign: Any, value: Any) -> Any:
    if value is None or pd.isna(value):
        return None
    formatted = format_number(value)
    sign_text = clean_text(sign)
    if sign_text in {"<", "<=", ">", ">="}:
        return f"{sign_text}{formatted}"
    return formatted


def effective_limits(parameter: str, legislation: str, vmin: Any, vmax: Any, vdyn: Any) -> tuple[Any, Any, Any]:
    p_norm = norm(parameter)
    if "manganes dissolvido" in p_norm:
        return None, None, None
    if p_norm == "amonia":
        return None, None, None
    if "solidos suspensos totais" in p_norm and legislation == "conama":
        return None, None, None
    if "escherichia coli" in p_norm and legislation == "conama":
        return None, None, None
    return vmin, vmax, vdyn


def format_vmp(parameter: str, legislation: str, vmin: Any, vmax: Any, vdyn: Any) -> str:
    p_norm = norm(parameter)
    vmin, vmax, vdyn = effective_limits(parameter, legislation, vmin, vmax, vdyn)
    if "oleos" in p_norm and "graxas" in p_norm:
        return "V.A."
    if p_norm == "amonia":
        return "-"
    if "nitrogenio amoniacal" in p_norm:
        return "3,7 mg/L N para pH <= 7,5; 2,0 mg/L N para 7,5 < pH <= 8,0; 1,0 mg/L N para 8,0 < pH <= 8,5; 0,5 mg/L N para pH > 8,5"
    if "escherichia coli" in p_norm and legislation == "conama":
        return "Pode substituir coliformes termotolerantes conforme critério do órgão competente"
    has_min = vmin is not None and pd.notna(vmin)
    has_max = vmax is not None and pd.notna(vmax)
    has_dyn = vdyn is not None and pd.notna(vdyn)
    if has_min and has_max:
        return f"{format_number(vmin)}-{format_number(vmax)}"
    if has_min:
        return f"≥{format_number(vmin)}"
    if has_max:
        return format_number(vmax)
    if has_dyn:
        return format_number(vdyn)
    return "-"

def format_vmp_numeric(value: Any) -> Any:
    if value is None or pd.isna(value):
        return None
    value = float(value)
    if abs(value) >= 100:
        return round(value, 0)
    if abs(value) >= 10:
        return round(value, 2)
    if abs(value) >= 1:
        return round(value, 3)
    return round(value, 5)


def violates(parameter: str, value: Any, sign: Any, vmin: Any, vmax: Any, vdyn: Any) -> bool:
    vmin, vmax, vdyn = effective_limits(parameter, "copam", vmin, vmax, vdyn)
    if value is None or pd.isna(value) or clean_text(sign) in {"<", "<="}:
        return False
    value = float(value)
    if vmin is not None and pd.notna(vmin) and value < float(vmin):
        return True
    upper_candidates = [v for v in [vmax, vdyn] if v is not None and pd.notna(v)]
    if upper_candidates and value > min(map(float, upper_candidates)):
        return True
    return False


def load_surface() -> pd.DataFrame:
    df = pd.read_excel(CONSOLIDATED, sheet_name="fisico_analise_consolidada")
    surface = df[df["matriz"].astype(str).str.contains("Superficial", na=False)].copy()
    numeric_cols = [
        "valor_medido",
        "vmp_357_cl1_min",
        "vmp_357_cl1_max",
        "vmp_357_cl2_min",
        "vmp_357_cl2_max",
        "vmp_amonia_dinamico",
    ]
    for col in numeric_cols:
        surface[col] = pd.to_numeric(surface[col], errors="coerce")
    surface["parametro_display"] = surface["nome_parametro"].map(fix_ptbr_label)
    surface["unidade_display"] = surface["unidade_medida"].map(fix_ptbr_label)
    surface.loc[surface["parametro_display"].map(norm).eq("nitrogenio amoniacal"), "unidade_display"] = "mg N/L"
    surface["parametro_norm"] = surface["parametro_display"].map(norm)
    return surface


def build_rows(surface: pd.DataFrame, campaign: str) -> tuple[list[str], list[list[Any]], set[tuple[int, int]]]:
    points = point_order(surface["nome_ponto"].dropna().unique().tolist())
    headers = [
        "ParÃ¢metros analisados",
        "Unidade",
        "VMP - CONAMA nÂ° 357 (2005) - Classe 2",
        "VMP - COPAM nÂ° 8 (2022) - Classe 2",
    ]
    headers.extend(points)

    rows: list[list[Any]] = []
    violation_cells: set[tuple[int, int]] = set()
    grouped = surface.sort_values(["parametro_norm", "nome_ponto", "nome_campanha"]).groupby("parametro_display", sort=False)
    for row_idx, (parameter, group) in enumerate(grouped, start=3):
        unit = group["unidade_display"].dropna().iloc[0] if not group["unidade_display"].dropna().empty else None
        v1_min = group["vmp_357_cl1_min"].dropna().iloc[0] if not group["vmp_357_cl1_min"].dropna().empty else None
        v1_max = group["vmp_357_cl1_max"].dropna().iloc[0] if not group["vmp_357_cl1_max"].dropna().empty else None
        v2_min = group["vmp_357_cl2_min"].dropna().iloc[0] if not group["vmp_357_cl2_min"].dropna().empty else None
        v2_max = group["vmp_357_cl2_max"].dropna().iloc[0] if not group["vmp_357_cl2_max"].dropna().empty else None
        vdyn = group["vmp_amonia_dinamico"].dropna().iloc[0] if not group["vmp_amonia_dinamico"].dropna().empty else None
        vmp_conama = format_vmp(parameter, "conama", v2_min, v2_max, vdyn)
        vmp_copam = format_vmp(parameter, "copam", v2_min, v2_max, vdyn)
        row = [
            parameter,
            unit,
            vmp_conama,
            vmp_copam,
        ]
        col_idx = 5
        for point in points:
            match = group[(group["nome_campanha"].eq(campaign)) & (group["nome_ponto"].eq(point))]
            if match.empty:
                row.append(None)
            else:
                rec = match.iloc[0]
                row.append(format_result(rec["sinal_limite"], rec["valor_medido"]))
                if violates(parameter, rec["valor_medido"], rec["sinal_limite"], v2_min, v2_max, vdyn):
                    violation_cells.add((row_idx, col_idx))
            col_idx += 1
        rows.append(row)
    return headers, rows, violation_cells


def write_workbook(surface: pd.DataFrame) -> None:
    if OUTPUT.exists() and not BACKUP.exists():
        shutil.copy2(OUTPUT, BACKUP)

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E78")
    campaign_fill = PatternFill("solid", fgColor="D9EAD3")
    violation_fill = PatternFill("solid", fgColor="F4CCCC")
    white_font = Font(color="FFFFFF", bold=True)
    bold_font = Font(bold=True)

    campaigns = [campaign for campaign in CAMPAIGN_LABELS if campaign in set(surface["nome_campanha"])]
    for idx, campaign in enumerate(campaigns):
        ws = wb.active if idx == 0 else wb.create_sheet()
        ws.title = CAMPAIGN_LABELS[campaign]
        headers, rows, violation_cells = build_rows(surface, campaign)

        for col, value in enumerate(headers[:4], start=1):
            ws.cell(1, col, value)
            ws.cell(1, col).fill = header_fill
            ws.cell(1, col).font = white_font
            ws.cell(1, col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.cell(2, col, None)
            ws.cell(2, col).fill = header_fill

        ws.merge_cells(start_row=1, start_column=5, end_row=1, end_column=len(headers))
        cell = ws.cell(1, 5, "Pontos amostrais")
        cell.fill = campaign_fill
        cell.font = bold_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        for col, value in enumerate(headers[4:], start=5):
            ws.cell(2, col, value)
            ws.cell(2, col).fill = header_fill
            ws.cell(2, col).font = white_font
            ws.cell(2, col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row in rows:
            ws.append(row)

        for row_idx, col_idx in violation_cells:
            ws.cell(row_idx, col_idx).fill = violation_fill

        ws.freeze_panes = "E3"
        ws.auto_filter.ref = ws.dimensions
        for row in ws.iter_rows(min_row=3, min_col=1, max_col=4):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=3, min_col=5):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions["A"].width = 36
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["C"].width = 36
        ws.column_dimensions["D"].width = 36
        for col in range(5, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 14
        ws.row_dimensions[1].height = 42
        ws.row_dimensions[2].height = 30

    conv = wb.create_sheet("Conversoes_Unidade")
    conv.append(["Parametro", "Unidade_Cadastro", "Unidade_Dados", "Fator_Conversao_VMP"])
    for cell in conv[1]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    conv.append(["FÃ³sforo Total", "mg/L", "Âµg/L", 1000])
    for col in range(1, conv.max_column + 1):
        conv.column_dimensions[get_column_letter(col)].width = 24
    conv.auto_filter.ref = conv.dimensions

    obs = wb.create_sheet("Observacoes_Normativas")
    obs.append(["Parâmetro", "Observação"])
    for cell in obs[1]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    obs_rows = [
        ["Fósforo Total", "Valor de 0,1 mg P/L mantido para ambiente lótico ou tributário de ambiente intermediário; revisar caso algum ponto represente ambiente lêntico ou tributário direto de ambiente lêntico."],
        ["Cloro Residual Livre", "Não foi aplicada comparação direta, pois as normas tratam cloro residual total, não especificamente cloro residual livre."],
        ["Coliformes Termotolerantes", "O valor de referência é apresentado, mas a conclusão legal de conformidade anual depende de atendimento em 80% ou mais de pelo menos seis amostras ao longo de um ano, com frequência bimestral."],
        ["Escherichia coli", "Na CONAMA, o parâmetro pode substituir coliformes termotolerantes conforme critério do órgão competente; na norma mineira foi mantido o valor de 1.000 NMP/100 mL."],
        ["Amônia", "A linha em mg NH3/L não foi comparada diretamente ao VMP de nitrogênio amoniacal total sem conversão adequada."],
        ["Manganês Dissolvido", "Sem VMP aplicado; o padrão de 0,1 mg/L refere-se a manganês total."],
    ]
    for row in obs_rows:
        obs.append(row)
    obs.column_dimensions["A"].width = 28
    obs.column_dimensions["B"].width = 120
    for row in obs.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    obs.auto_filter.ref = obs.dimensions

    wb.save(OUTPUT)


def main() -> int:
    surface = load_surface()
    write_workbook(surface)
    print(f"OK: {OUTPUT}")
    print(f"registros={len(surface)} parametros={surface['parametro_display'].nunique()} pontos={surface['nome_ponto'].nunique()} campanhas={surface['nome_campanha'].nunique()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())


