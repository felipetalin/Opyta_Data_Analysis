from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from opyta_analysis.pipelines.diagnostico.zoobentos import _load_zoobentos_df
from opyta_analysis.supabase_client import get_client, paginate


PROJECT_ID = 30
GROUP = "Zoobentos"
END_YEAR = 2021
OUTPUT_DIR = Path("outputs/consultas/GEOHER001_bentos_ate_2021")


LABELS = {
    "family": "Fam\u00edlia",
    "taxon": "T\u00e1xon",
    "abundance": "Abund\u00e2ncia",
    "last_campaign": "\u00daltima campanha",
    "efforts": "Esfor\u00e7os",
    "efforts_with_result": "Esfor\u00e7os com resultado",
    "biological_records": "Registros biol\u00f3gicos",
    "taxa": "T\u00e1xons",
    "status": "Situa\u00e7\u00e3o",
    "pending": "Pend\u00eancia",
}


def clean(value):
    if pd.isna(value) or not str(value).strip():
        return None
    return re.sub(r"\s+", " ", str(value).strip())


def campaign_number(value) -> int:
    match = re.match(r"^C0*(\d+)", str(value or "").strip(), flags=re.I)
    return int(match.group(1)) if match else 9999


def load_taxa() -> tuple[pd.DataFrame, pd.DataFrame]:
    raw = _load_zoobentos_df(PROJECT_ID, GROUP, None)
    raw["data_dt"] = pd.to_datetime(raw["data_hora_coleta"], errors="coerce")
    raw["year"] = raw["data_dt"].dt.year
    raw["abundance_num"] = pd.to_numeric(raw["contagem"], errors="coerce").fillna(0)
    data = raw[raw["year"] <= END_YEAR].copy()

    for column in [
        "filo",
        "classe",
        "ordem",
        "familia",
        "nome_cientifico",
        "taxon_final",
        "nome_campanha",
    ]:
        data[column] = data[column].map(clean)

    records = []
    for taxon, group in data.groupby("taxon_final", dropna=False):
        representative = group.iloc[0]
        campaigns = sorted(
            group["nome_campanha"].dropna().unique().tolist(),
            key=lambda value: (
                group.loc[group["nome_campanha"].eq(value), "year"].min(),
                campaign_number(value),
                value,
            ),
        )
        records.append(
            {
                "Filo": clean(representative["filo"]) or "N\u00e3o informado",
                "Classe": clean(representative["classe"]) or "N\u00e3o informada",
                "Ordem": clean(representative["ordem"]) or "N\u00e3o informada",
                LABELS["family"]: (
                    clean(representative["familia"]) or "Indeterminada (NA)"
                ),
                LABELS["taxon"]: taxon or "T\u00e1xon n\u00e3o identificado",
                "Campanhas com registro": len(campaigns),
                "Abund\u00e2ncia total": int(group["abundance_num"].sum()),
                "Primeira campanha": campaigns[0] if campaigns else "",
                LABELS["last_campaign"]: campaigns[-1] if campaigns else "",
                "Lista de campanhas": ", ".join(campaigns),
            }
        )

    taxa = pd.DataFrame(records).sort_values(
        ["Filo", "Classe", "Ordem", LABELS["family"], LABELS["taxon"]]
    )
    return data, taxa.reset_index(drop=True)


def load_campaigns(data: pd.DataFrame) -> pd.DataFrame:
    sb = get_client(None)
    points = paginate(
        sb,
        "pontos_coleta",
        filters={"id_projeto": PROJECT_ID},
        select="id_ponto_coleta,id_campanha,nome_ponto,data_hora_coleta",
    )
    point_ids = {point["id_ponto_coleta"] for point in points}
    point_map = {point["id_ponto_coleta"]: point for point in points}

    efforts = paginate(
        sb,
        "esforcos_amostragem",
        filters={"grupo_biologico": GROUP},
        select="id_esforco,id_ponto_coleta,grupo_biologico",
    )
    efforts = [
        effort for effort in efforts if effort.get("id_ponto_coleta") in point_ids
    ]

    results = paginate(
        sb,
        "resultados_zoobentos",
        select="id_resultado_bento,id_esforco,id_especie,abundancia",
    )
    result_effort_ids = {
        result.get("id_esforco")
        for result in results
        if result.get("id_esforco") is not None
    }

    campaigns_db = paginate(
        sb,
        "campanhas",
        select="id_campanha,nome_campanha",
    )
    campaign_map = {
        campaign["id_campanha"]: campaign["nome_campanha"]
        for campaign in campaigns_db
    }

    effort_rows = []
    for effort in efforts:
        point = point_map[effort["id_ponto_coleta"]]
        date = pd.to_datetime(point.get("data_hora_coleta"), errors="coerce")
        effort_rows.append(
            {
                "id_campaign": point.get("id_campanha"),
                "campaign": campaign_map.get(
                    point.get("id_campanha"), "Campanha desconhecida"
                ),
                "year": int(date.year) if not pd.isna(date) else None,
                "id_point": point.get("id_ponto_coleta"),
                "id_effort": effort.get("id_esforco"),
                "has_result": effort.get("id_esforco") in result_effort_ids,
            }
        )

    effort_df = pd.DataFrame(effort_rows)
    effort_df = effort_df[effort_df["year"] <= END_YEAR].copy()

    result_summary = (
        data.groupby("nome_campanha", dropna=False)
        .agg(
            result_records=("id_resultado_pk", "count"),
            taxa=("taxon_final", "nunique"),
            abundance=("abundance_num", "sum"),
        )
        .reset_index()
        .rename(columns={"nome_campanha": "campaign"})
    )

    campaigns = (
        effort_df.groupby(["id_campaign", "campaign", "year"], dropna=False)
        .agg(
            points=("id_point", "nunique"),
            efforts=("id_effort", "nunique"),
            efforts_with_result=("has_result", "sum"),
        )
        .reset_index()
        .merge(result_summary, on="campaign", how="left")
    )
    for column in ["result_records", "taxa", "abundance"]:
        campaigns[column] = campaigns[column].fillna(0).astype(int)

    campaigns["status"] = campaigns["result_records"].gt(0).map(
        {
            True: "Com resultados biol\u00f3gicos",
            False: "Executada sem resultados cadastrados",
        }
    )
    campaigns = campaigns.sort_values(["year", "id_campaign"]).reset_index(drop=True)
    return campaigns.rename(
        columns={
            "id_campaign": "ID campanha",
            "campaign": "Campanha",
            "year": "Ano",
            "points": "Pontos",
            "efforts": LABELS["efforts"],
            "efforts_with_result": LABELS["efforts_with_result"],
            "result_records": LABELS["biological_records"],
            "taxa": LABELS["taxa"],
            "abundance": LABELS["abundance"],
            "status": LABELS["status"],
        }
    )


def build_summary(data: pd.DataFrame, taxa: pd.DataFrame, campaigns: pd.DataFrame) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"Indicador": "Per\u00edodo avaliado", "Valor": "At\u00e9 31/12/2021"},
            {
                "Indicador": "Campanhas executadas com esfor\u00e7o de Zoobentos",
                "Valor": int(campaigns["ID campanha"].nunique()),
            },
            {
                "Indicador": "Campanhas com resultados biol\u00f3gicos",
                "Valor": int((campaigns[LABELS["biological_records"]] > 0).sum()),
            },
            {
                "Indicador": "Campanhas executadas sem resultados cadastrados",
                "Valor": int((campaigns[LABELS["biological_records"]] == 0).sum()),
            },
            {
                "Indicador": "T\u00e1xons distintos registrados",
                "Valor": int(taxa[LABELS["taxon"]].nunique()),
            },
            {"Indicador": "Filos", "Valor": int(taxa["Filo"].nunique())},
            {"Indicador": "Classes", "Valor": int(taxa["Classe"].nunique())},
            {"Indicador": "Ordens", "Valor": int(taxa["Ordem"].nunique())},
            {
                "Indicador": "Fam\u00edlias informadas",
                "Valor": int(
                    taxa.loc[
                        taxa[LABELS["family"]].ne("Indeterminada (NA)"),
                        LABELS["family"],
                    ].nunique()
                ),
            },
            {
                "Indicador": "Registros biol\u00f3gicos",
                "Valor": int(len(data)),
            },
            {
                "Indicador": "Abund\u00e2ncia acumulada",
                "Valor": int(data["abundance_num"].sum()),
            },
        ]
    )


def format_workbook(writer) -> None:
    for worksheet in writer.book.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
        for column_cells in worksheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            width = min(max(max_length + 2, 12), 55)
            worksheet.column_dimensions[
                get_column_letter(column_cells[0].column)
            ].width = width
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_outputs(
    data: pd.DataFrame,
    taxa: pd.DataFrame,
    campaigns: pd.DataFrame,
    summary: pd.DataFrame,
) -> tuple[Path, Path]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    family = LABELS["family"]
    taxon = LABELS["taxon"]
    biological_records = LABELS["biological_records"]

    pending = taxa[taxa[family].eq("Indeterminada (NA)")][
        ["Filo", "Classe", "Ordem", family, taxon]
    ].copy()
    pending[LABELS["pending"]] = (
        "Fam\u00edlia indeterminada conforme decis\u00e3o taxon\u00f4mica aprovada."
    )
    applied_corrections = pd.DataFrame(
        [
            {
                "Item": "Rhynchobdellida",
                "Decis\u00e3o": "Ordem; fam\u00edlia Indeterminada (NA)",
            },
            {
                "Item": "Tubificida",
                "Decis\u00e3o": "Ordem; fam\u00edlia Naididae",
            },
            {
                "Item": "Ostracoda",
                "Decis\u00e3o": "Classe; fam\u00edlia Indeterminada (NA)",
            },
            {
                "Item": "Atopsyche",
                "Decis\u00e3o": "Consolidado no nome can\u00f4nico Atopsyche sp.",
            },
            {
                "Item": "Atopsyche sp.",
                "Decis\u00e3o": "G\u00eanero; fam\u00edlia Hydrobiosidae",
            }
        ]
    )

    xlsx = OUTPUT_DIR / "lista_taxonomica_bentos_GEOHER001_ate_2021.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as writer:
        summary.to_excel(writer, index=False, sheet_name="Resumo")
        taxa.to_excel(writer, index=False, sheet_name="Lista_taxonomica")
        campaigns.to_excel(writer, index=False, sheet_name="Campanhas_ate_2021")
        pending.to_excel(writer, index=False, sheet_name="Taxonomia_pendente")
        applied_corrections.to_excel(
            writer,
            index=False,
            sheet_name="Correcoes_aplicadas",
        )
        format_workbook(writer)

    md = OUTPUT_DIR / "lista_taxonomica_bentos_GEOHER001_ate_2021.md"
    md_lines = [
        "# GEOHER001 - lista taxon\u00f4mica de bentos at\u00e9 2021",
        "",
        "## Resumo",
        "",
        summary.to_markdown(index=False),
        "",
        "## Campanhas",
        "",
        campaigns[
            [
                "Campanha",
                "Ano",
                "Pontos",
                LABELS["efforts"],
                biological_records,
                LABELS["taxa"],
                LABELS["abundance"],
                LABELS["status"],
            ]
        ].to_markdown(index=False),
        "",
        "## Lista taxon\u00f4mica",
        "",
        taxa[
            [
                "Filo",
                "Classe",
                "Ordem",
                family,
                taxon,
                "Campanhas com registro",
                "Abund\u00e2ncia total",
            ]
        ].to_markdown(index=False),
        "",
        "## Pend\u00eancias",
        "",
        (
            "- Rhynchobdellida e Ostracoda foram mantidos com fam\u00edlia "
            "Indeterminada (NA)."
        ),
        (
            "- Tubificida foi registrado na fam\u00edlia Naididae."
        ),
        (
            "- Atopsyche foi consolidado no nome can\u00f4nico Atopsyche sp., "
            "fam\u00edlia Hydrobiosidae."
        ),
    ]
    md.write_text("\n".join(md_lines) + "\n", encoding="utf-8")
    return xlsx, md


def main() -> None:
    data, taxa = load_taxa()
    campaigns = load_campaigns(data)
    summary = build_summary(data, taxa, campaigns)
    xlsx, md = write_outputs(data, taxa, campaigns, summary)

    print(summary.to_string(index=False))
    print(f"taxa={len(taxa)}")
    print(f"xlsx={xlsx.resolve()}")
    print(f"markdown={md.resolve()}")


if __name__ == "__main__":
    main()
