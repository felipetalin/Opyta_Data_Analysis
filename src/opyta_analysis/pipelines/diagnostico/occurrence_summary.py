from __future__ import annotations

import re
from pathlib import Path

import matplotlib.colors as mcolors
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _clean_text(value) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except TypeError:
        pass
    text = str(value).strip()
    return "" if text.lower() in {"", "nan", "none"} else text


def _mode_or_first(series: pd.Series) -> str:
    values = series.dropna().astype(str).str.strip()
    values = values[(values != "") & (~values.str.lower().isin(["nan", "none"]))]
    if values.empty:
        return ""
    modes = values.mode()
    return str(modes.iloc[0] if not modes.empty else values.iloc[0])


def _campaign_short(campaign: str) -> str:
    match = re.match(r"^C0*(\d+)", str(campaign).strip(), flags=re.IGNORECASE)
    if match:
        return f"C{int(match.group(1)):02d}"
    return str(campaign).strip()


def _campaign_year(campaign: str) -> int | None:
    match = re.search(r"(19\d{2}|20\d{2})", str(campaign))
    return int(match.group(1)) if match else None


def _campaign_season(campaign: str) -> str:
    text = str(campaign).upper()
    if re.search(r"(^|[-_\s])CH($|[-_\s])", text):
        return "CH"
    if re.search(r"(^|[-_\s])SC($|[-_\s])", text):
        return "SC"
    return ""


def _campaign_sort_key(campaign: str) -> tuple[int, int, str]:
    text = str(campaign)
    number = re.match(r"^C0*(\d+)", text, flags=re.IGNORECASE)
    year = _campaign_year(text)
    return (year or 9999, int(number.group(1)) if number else 9999, text)


def _point_sort_key(point: str) -> tuple[str, int, str]:
    text = str(point).strip()
    match = re.search(r"^(.*?)(\d+)$", text)
    if not match:
        return (text.lower(), 9999, text)
    return (match.group(1).lower(), int(match.group(2)), text)


def _hex_without_hash(value: str) -> str:
    text = str(value).strip().lstrip("#")
    return text if re.fullmatch(r"[0-9A-Fa-f]{6}", text) else "002060"


def _chunk_taxa(
    taxa: list[str],
    metadata: pd.DataFrame,
    *,
    group_column: str | None,
    max_rows: int,
) -> list[list[str]]:
    if len(taxa) <= max_rows:
        return [taxa]

    if not group_column or group_column not in metadata.columns:
        return [taxa[i : i + max_rows] for i in range(0, len(taxa), max_rows)]

    meta = metadata.set_index("Táxon").reindex(taxa)
    grouped: list[list[str]] = []
    for _group, frame in meta.groupby(group_column, sort=False, dropna=False):
        group_taxa = frame.index.astype(str).tolist()
        if len(group_taxa) <= max_rows:
            grouped.append(group_taxa)
        else:
            grouped.extend(group_taxa[i : i + max_rows] for i in range(0, len(group_taxa), max_rows))

    pages: list[list[str]] = []
    current: list[str] = []
    for group_taxa in grouped:
        if current and len(current) + len(group_taxa) > max_rows:
            pages.append(current)
            current = []
        current.extend(group_taxa)
    if current:
        pages.append(current)
    return pages


def _render_frequency_heatmaps(
    matrix: pd.DataFrame,
    *,
    metadata: pd.DataFrame,
    group_column: str | None,
    output_dir: Path,
    filename_prefix: str,
    xlabel: str,
    theme: dict,
    generated_files: list[str],
    campaign_columns: bool,
) -> list[str]:
    if matrix.empty:
        return []

    max_rows = max(8, int(theme.get("occurrence_heatmap_rows_per_page", 28)))
    pages = _chunk_taxa(
        matrix.index.astype(str).tolist(),
        metadata,
        group_column=group_column,
        max_rows=max_rows,
    )
    primary = str(theme.get("primary_hex", "#002060"))
    cmap = mcolors.LinearSegmentedColormap.from_list(
        "occurrence_frequency",
        ["#FFFFFF", primary],
    )
    size = theme.get("figsize_standard", [11.69, 8.27])
    base_width = float(size[0])
    tick_size = int(theme.get("heatmap_tick_size", theme.get("font_size_base", 11)))
    annotation_size = int(theme.get("heatmap_annotation_size", theme.get("annotation_size", 10)))
    label_size = int(theme.get("label_size", theme.get("font_size_base", 15)))
    outputs: list[str] = []

    for page_number, taxa_page in enumerate(pages, start=1):
        page = matrix.reindex(taxa_page).fillna(0) * 100
        fig_height = max(5.5, min(float(size[1]), 2.0 + 0.23 * len(taxa_page)))
        fig, ax = plt.subplots(
            figsize=(base_width, fig_height),
            dpi=int(theme.get("dpi", 600)),
        )
        values = page.to_numpy(dtype=float)
        image = ax.imshow(values, cmap=cmap, vmin=0, vmax=100, aspect="auto")

        x_labels = [
            _campaign_short(column) if campaign_columns else str(column)
            for column in page.columns
        ]
        x_rotation = 0 if campaign_columns else 90
        ax.set_xticks(np.arange(len(page.columns)))
        ax.set_xticklabels(x_labels, rotation=x_rotation, ha="center", fontsize=tick_size)
        ax.set_yticks(np.arange(len(taxa_page)))
        ax.set_yticklabels(taxa_page, fontsize=tick_size, fontstyle="italic")
        ax.set_xlabel(xlabel, fontsize=label_size)
        ax.set_ylabel("Táxon", fontsize=label_size)

        ax.set_xticks(np.arange(-0.5, len(page.columns), 1), minor=True)
        ax.set_yticks(np.arange(-0.5, len(taxa_page), 1), minor=True)
        ax.grid(which="minor", color="#D9D9D9", linestyle="-", linewidth=0.7)
        ax.tick_params(which="minor", bottom=False, left=False)

        for row in range(values.shape[0]):
            for col in range(values.shape[1]):
                value = float(values[row, col])
                if value <= 0:
                    continue
                ax.text(
                    col,
                    row,
                    f"{value:.0f}",
                    ha="center",
                    va="center",
                    fontsize=annotation_size,
                    color="white" if value >= 65 else "black",
                )

        colorbar = fig.colorbar(image, ax=ax, fraction=0.025, pad=0.015)
        colorbar.set_label("Frequência de ocorrência (%)", fontsize=label_size)
        colorbar.ax.tick_params(labelsize=tick_size)
        fig.tight_layout()

        suffix = f"_parte_{page_number:02d}" if len(pages) > 1 else ""
        output = output_dir / f"{filename_prefix}{suffix}.png"
        fig.savefig(output, dpi=int(theme.get("dpi", 600)), bbox_inches="tight")
        plt.close(fig)
        generated_files.append(str(output))
        outputs.append(str(output))

    return outputs


def _style_workbook(
    workbook_path: Path,
    *,
    metadata_columns: list[str],
    theme: dict,
) -> None:
    workbook = load_workbook(workbook_path)
    primary = _hex_without_hash(theme.get("primary_hex", "#002060"))
    secondary = _hex_without_hash(theme.get("secondary_hex", "#5B9BD5"))
    header_fill = PatternFill("solid", fgColor=primary)
    header_font = Font(color="FFFFFF", bold=True)
    presence_fill = PatternFill("solid", fgColor=secondary)
    not_sampled_fill = PatternFill("solid", fgColor="D9D9D9")

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False

        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for column_cells in sheet.columns:
            max_length = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in list(column_cells)[:250]
            )
            letter = get_column_letter(column_cells[0].column)
            header = str(column_cells[0].value or "")
            if header in {"Táxon", "Nome popular"}:
                width = min(max(max_length + 2, 22), 42)
            else:
                width = min(max(max_length + 2, 10), 22)
            sheet.column_dimensions[letter].width = width

        if sheet.title in {"Freq_Campanha", "Freq_Ponto"} and sheet.max_row >= 2:
            data_start = len(metadata_columns) + 2
            if sheet.max_column >= data_start:
                start_letter = get_column_letter(data_start)
                end_letter = get_column_letter(sheet.max_column)
                data_range = f"{start_letter}2:{end_letter}{sheet.max_row}"
                sheet.conditional_formatting.add(
                    data_range,
                    ColorScaleRule(
                        start_type="num",
                        start_value=0,
                        start_color="FFFFFF",
                        end_type="num",
                        end_value=1,
                        end_color=primary,
                    ),
                )
                for row in sheet.iter_rows(
                    min_row=2,
                    max_row=sheet.max_row,
                    min_col=data_start,
                    max_col=sheet.max_column,
                ):
                    for cell in row:
                        cell.number_format = "0%"
                        cell.alignment = Alignment(horizontal="center")

        if sheet.title == "Resumo_Geral":
            percent_headers = {
                "% campanhas",
                "% pontos",
                "% ocorrência global",
            }
            for cell in sheet[1]:
                if str(cell.value) in percent_headers:
                    for row in range(2, sheet.max_row + 1):
                        sheet.cell(row=row, column=cell.column).number_format = "0.0%"

        if sheet.title.startswith("Ocorrencia_"):
            for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
                for cell in row:
                    if cell.value == "X":
                        cell.fill = presence_fill
                        cell.font = Font(color="FFFFFF", bold=True)
                        cell.alignment = Alignment(horizontal="center")
                    elif cell.value == "NA":
                        cell.fill = not_sampled_fill
                        cell.font = Font(color="666666")
                        cell.alignment = Alignment(horizontal="center")
            for cell in sheet[1]:
                if str(cell.value).endswith(" - %OC"):
                    for row in range(2, sheet.max_row + 1):
                        sheet.cell(row=row, column=cell.column).number_format = "0%"

    workbook.save(workbook_path)


def export_occurrence_summary(
    *,
    records: pd.DataFrame,
    sampling_units: pd.DataFrame,
    taxon_col: str,
    campaign_col: str,
    point_col: str,
    abundance_col: str,
    metadata_map: dict[str, str],
    group_slug: str,
    output_dir: Path,
    theme: dict,
    generated_files: list[str],
) -> dict:
    required_records = [taxon_col, campaign_col, point_col, abundance_col]
    missing_records = [column for column in required_records if column not in records.columns]
    if missing_records:
        raise RuntimeError(
            "Colunas ausentes para síntese de ocorrência: "
            + ", ".join(missing_records)
        )
    required_units = [campaign_col, point_col]
    missing_units = [column for column in required_units if column not in sampling_units.columns]
    if missing_units:
        raise RuntimeError(
            "Colunas ausentes nas unidades amostrais: "
            + ", ".join(missing_units)
        )

    work = records.copy()
    units = sampling_units.copy()
    for column in [taxon_col, campaign_col, point_col]:
        work[column] = work[column].map(_clean_text)
    for column in [campaign_col, point_col]:
        units[column] = units[column].map(_clean_text)
    work[abundance_col] = pd.to_numeric(work[abundance_col], errors="coerce").fillna(0)
    work = work[
        (work[taxon_col] != "")
        & (work[campaign_col] != "")
        & (work[point_col] != "")
    ].copy()
    units = units[(units[campaign_col] != "") & (units[point_col] != "")].copy()

    observed_units = work[[campaign_col, point_col]].drop_duplicates()
    units = (
        pd.concat([units[[campaign_col, point_col]], observed_units], ignore_index=True)
        .drop_duplicates()
        .reset_index(drop=True)
    )
    campaigns = sorted(units[campaign_col].unique().tolist(), key=_campaign_sort_key)
    points = sorted(units[point_col].unique().tolist(), key=_point_sort_key)

    abundance = (
        work.groupby([taxon_col, campaign_col, point_col], as_index=False)[abundance_col]
        .sum()
    )
    abundance["Presença"] = (abundance[abundance_col] > 0).astype(int)
    taxa = sorted(
        abundance.loc[abundance["Presença"] > 0, taxon_col].unique().tolist()
    )
    if not taxa:
        return {"taxa": 0, "warning": "sem ocorrências positivas"}

    metadata_internal = [column for column in metadata_map if column in work.columns]
    if metadata_internal:
        metadata = (
            work.groupby(taxon_col, as_index=False)
            .agg(**{column: (column, _mode_or_first) for column in metadata_internal})
            .rename(columns={taxon_col: "Táxon", **metadata_map})
        )
    else:
        metadata = pd.DataFrame({"Táxon": taxa})
    metadata_columns = [
        metadata_map[column]
        for column in metadata_internal
        if metadata_map[column] in metadata.columns
    ]
    metadata = metadata[metadata_columns + ["Táxon"]]
    metadata = metadata[metadata["Táxon"].isin(taxa)].copy()
    sort_columns = metadata_columns + ["Táxon"]
    metadata = metadata.sort_values(sort_columns, na_position="last").reset_index(drop=True)
    taxa_order = metadata["Táxon"].astype(str).tolist()

    taxa_frame = pd.DataFrame({taxon_col: taxa_order})
    complete = taxa_frame.merge(units, how="cross")
    complete = complete.merge(
        abundance[[taxon_col, campaign_col, point_col, abundance_col, "Presença"]],
        on=[taxon_col, campaign_col, point_col],
        how="left",
    )
    complete[abundance_col] = complete[abundance_col].fillna(0)
    complete["Presença"] = complete["Presença"].fillna(0).astype(int)

    campaign_denominator = units.groupby(campaign_col)[point_col].nunique()
    campaign_presence = (
        complete.groupby([taxon_col, campaign_col])["Presença"]
        .sum()
        .unstack(fill_value=0)
        .reindex(index=taxa_order, columns=campaigns, fill_value=0)
    )
    frequency_campaign = campaign_presence.div(
        campaign_denominator.reindex(campaigns),
        axis=1,
    ).fillna(0)

    point_denominator = units.groupby(point_col)[campaign_col].nunique()
    point_presence = (
        complete.groupby([taxon_col, point_col])["Presença"]
        .sum()
        .unstack(fill_value=0)
        .reindex(index=taxa_order, columns=points, fill_value=0)
    )
    frequency_point = point_presence.div(
        point_denominator.reindex(points),
        axis=1,
    ).fillna(0)

    global_presence = complete.groupby(taxon_col)["Presença"].sum().reindex(taxa_order)
    campaigns_present = (campaign_presence > 0).sum(axis=1)
    points_present = (point_presence > 0).sum(axis=1)
    campaign_names = (
        complete[complete["Presença"] > 0]
        .groupby(taxon_col)[campaign_col]
        .apply(lambda values: ", ".join(sorted(set(values), key=_campaign_sort_key)))
        .reindex(taxa_order)
        .fillna("")
    )
    point_names = (
        complete[complete["Presença"] > 0]
        .groupby(taxon_col)[point_col]
        .apply(lambda values: ", ".join(sorted(set(values), key=_point_sort_key)))
        .reindex(taxa_order)
        .fillna("")
    )

    summary = metadata.set_index("Táxon").reindex(taxa_order).reset_index()
    summary["Campanhas com ocorrência"] = campaigns_present.reindex(taxa_order).to_numpy()
    summary["% campanhas"] = summary["Campanhas com ocorrência"] / max(len(campaigns), 1)
    summary["Pontos com ocorrência"] = points_present.reindex(taxa_order).to_numpy()
    summary["% pontos"] = summary["Pontos com ocorrência"] / max(len(points), 1)
    summary["Ocorrências campanha-ponto"] = global_presence.to_numpy()
    summary["% ocorrência global"] = summary["Ocorrências campanha-ponto"] / max(len(units), 1)
    summary["Campanhas observadas"] = campaign_names.to_numpy()
    summary["Pontos observados"] = point_names.to_numpy()
    summary = summary[metadata_columns + ["Táxon"] + [
        "Campanhas com ocorrência",
        "% campanhas",
        "Pontos com ocorrência",
        "% pontos",
        "Ocorrências campanha-ponto",
        "% ocorrência global",
        "Campanhas observadas",
        "Pontos observados",
    ]]

    frequency_campaign_export = (
        metadata.set_index("Táxon")
        .reindex(taxa_order)
        .join(frequency_campaign)
        .reset_index()
    )
    frequency_campaign_export = frequency_campaign_export[
        metadata_columns + ["Táxon"] + campaigns
    ]
    frequency_point_export = (
        metadata.set_index("Táxon")
        .reindex(taxa_order)
        .join(frequency_point)
        .reset_index()
    )
    frequency_point_export = frequency_point_export[
        metadata_columns + ["Táxon"] + points
    ]

    long_base = complete.rename(
        columns={
            taxon_col: "Táxon",
            campaign_col: "Campanha",
            point_col: "Ponto",
            abundance_col: "Abundância",
        }
    )
    long_base["Campanha curta"] = long_base["Campanha"].map(_campaign_short)
    long_base["Ano"] = long_base["Campanha"].map(_campaign_year)
    long_base["Estação"] = long_base["Campanha"].map(_campaign_season)
    long_base = long_base.merge(metadata, on="Táxon", how="left")
    long_base = long_base[
        metadata_columns
        + ["Táxon", "Campanha", "Campanha curta", "Ano", "Estação", "Ponto", "Abundância", "Presença"]
    ]
    long_base["_campaign_order"] = long_base["Campanha"].map(
        {campaign: index for index, campaign in enumerate(campaigns)}
    )
    long_base["_point_order"] = long_base["Ponto"].map(
        {point: index for index, point in enumerate(points)}
    )
    long_base["_taxon_order"] = long_base["Táxon"].map(
        {taxon: index for index, taxon in enumerate(taxa_order)}
    )
    long_base = (
        long_base.sort_values(["_taxon_order", "_campaign_order", "_point_order"])
        .drop(columns=["_taxon_order", "_campaign_order", "_point_order"])
        .reset_index(drop=True)
    )

    annual_tables: dict[int, pd.DataFrame] = {}
    for year in sorted({year for year in map(_campaign_year, campaigns) if year is not None}):
        year_campaigns = [campaign for campaign in campaigns if _campaign_year(campaign) == year]
        table = metadata.set_index("Táxon").reindex(taxa_order).reset_index()
        for campaign in year_campaigns:
            sampled_points = set(
                units.loc[units[campaign_col] == campaign, point_col].astype(str).tolist()
            )
            presence_campaign = (
                complete[complete[campaign_col] == campaign]
                .pivot_table(
                    index=taxon_col,
                    columns=point_col,
                    values="Presença",
                    aggfunc="max",
                    fill_value=0,
                )
                .reindex(index=taxa_order, columns=points, fill_value=0)
            )
            for point in points:
                column = f"{_campaign_short(campaign)} - {point}"
                if point not in sampled_points:
                    table[column] = "NA"
                else:
                    table[column] = presence_campaign[point].map({1: "X", 0: ""}).to_numpy()
            oc = presence_campaign[list(sampled_points)].sum(axis=1) if sampled_points else 0
            table[f"{_campaign_short(campaign)} - OC"] = (
                oc.to_numpy() if hasattr(oc, "to_numpy") else 0
            )
            table[f"{_campaign_short(campaign)} - %OC"] = (
                (oc / len(sampled_points)).to_numpy()
                if sampled_points and hasattr(oc, "to_numpy")
                else 0
            )
        annual_tables[year] = table

    workbook_path = output_dir / f"04A_tabela_sintese_ocorrencia_{group_slug}.xlsx"
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Resumo_Geral", index=False)
        frequency_campaign_export.to_excel(writer, sheet_name="Freq_Campanha", index=False)
        frequency_point_export.to_excel(writer, sheet_name="Freq_Ponto", index=False)
        for year, table in annual_tables.items():
            table.to_excel(writer, sheet_name=f"Ocorrencia_{year}", index=False)
        long_base.to_excel(writer, sheet_name="Base_Longa", index=False)
    _style_workbook(
        workbook_path,
        metadata_columns=metadata_columns,
        theme=theme,
    )
    generated_files.append(str(workbook_path))

    group_column = "Ordem" if "Ordem" in metadata.columns else None
    campaign_outputs = _render_frequency_heatmaps(
        frequency_campaign,
        metadata=metadata,
        group_column=group_column,
        output_dir=output_dir,
        filename_prefix=f"04B_grafico_frequencia_ocorrencia_por_campanha_{group_slug}",
        xlabel="Campanha",
        theme=theme,
        generated_files=generated_files,
        campaign_columns=True,
    )
    point_outputs = _render_frequency_heatmaps(
        frequency_point,
        metadata=metadata,
        group_column=group_column,
        output_dir=output_dir,
        filename_prefix=f"04C_grafico_frequencia_ocorrencia_por_ponto_{group_slug}",
        xlabel="Ponto amostral",
        theme=theme,
        generated_files=generated_files,
        campaign_columns=False,
    )

    return {
        "taxa": int(len(taxa_order)),
        "campaigns": campaigns,
        "points": points,
        "sampling_units": int(len(units)),
        "annual_sheets": sorted(annual_tables),
        "campaign_heatmaps": campaign_outputs,
        "point_heatmaps": point_outputs,
        "workbook": str(workbook_path),
    }
