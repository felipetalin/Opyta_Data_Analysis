from __future__ import annotations

import re
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl import Workbook


SAMPLING_EVENT_COLUMNS = [
    "eventID",
    "samplingProtocol",
    "samplingEffort",
    "sampleSizeValue",
    "sampleSizeUnit",
    "eventDate",
    "eventRemarks",
    "county",
    "municipality",
    "waterBody",
    "locality",
    "decimalLatitude",
    "decimalLongitude",
    "geodeticDatum",
]

ASSOCIATED_OCCURRENCE_COLUMNS = [
    "eventID",
    "occurrenceID",
    "basisOfRecord",
    "scientificName",
    "kingdom",
    "phylum",
    "class",
    "order",
    "family",
    "taxonRank",
    "identificationQualifier",
    "recordedBy",
    "individualCount",
    "sex",
    "lifeStage",
    "reproductiveCondition",
    "preparations",
    "occurrenceRemarks",
]

FISH_BIOMETRIC_COLUMNS = [
    "eventID",
    "occurrenceID",
    "scientificName",
    "individualCount",
    "Weight",
    "StandardLength",
    "TotalLength",
    "Sex",
    "GonadalStage",
    "GonadWeight",
]

ORIENTATION_SHEET = "Orienta\u00e7\u00f5es"
COUNTY = "Gr\u00e3o Mogol"
MUNICIPALITY = "Gr\u00e3o Mogol"
GEODETIC_DATUM = "WGS84"
BASIS_OF_RECORD = "Esp\u00e9cime vivo"
RECORDED_BY = "Opyta"
TAXON_RANK = "Esp\u00e9cie"
DEGREE = "\u00b0"


def _normalize_text(value: str) -> str:
    txt = str(value).strip().lower()
    txt = unicodedata.normalize("NFKD", txt).encode("ascii", "ignore").decode("ascii")
    txt = re.sub(r"[^a-z0-9]+", " ", txt)
    return re.sub(r"\s+", " ", txt).strip()


def _slug_title(value: str) -> str:
    normalized = _normalize_text(value).title()
    return re.sub(r"[^A-Za-z0-9]+", "_", normalized).strip("_") or "Projeto"


def _get_col(df: pd.DataFrame, *cands: str) -> str | None:
    for col in cands:
        if col in df.columns:
            return col
    return None


def _clean(value) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except TypeError:
        pass
    text = str(value).strip()
    return "" if text.lower() == "nan" else text


def _numeric_or_blank(value):
    num = pd.to_numeric(value, errors="coerce")
    if pd.isna(num):
        return ""
    num = float(num)
    return int(num) if num.is_integer() else round(num, 6)


def _date_or_blank(value):
    dt = pd.to_datetime(value, errors="coerce")
    if pd.isna(dt):
        return _clean(value)
    return dt.date()


def _coord_or_blank(value) -> str:
    num = pd.to_numeric(value, errors="coerce")
    if pd.isna(num):
        return ""
    return f"{float(num):.6f}{DEGREE}"


def _effort_text(effort, unit) -> str:
    effort_value = _numeric_or_blank(effort)
    unit_text = _clean(unit)
    if effort_value == "" and unit_text == "":
        return ""
    return f"{effort_value} {unit_text}".strip()


def _mode_or_blank(series: pd.Series) -> str:
    values = series.dropna().astype(str).str.strip()
    values = values[(values != "") & (values.str.lower() != "nan")]
    if values.empty:
        return ""
    modes = values.mode()
    return str(modes.iloc[0] if not modes.empty else values.iloc[0])


def _project_slug(df: pd.DataFrame) -> str:
    if "nome_projeto" not in df.columns or df["nome_projeto"].dropna().empty:
        return "Projeto"
    return _slug_title(_mode_or_blank(df["nome_projeto"]))


def _write_dataframe(ws, df: pd.DataFrame) -> None:
    for col_idx, col_name in enumerate(df.columns, start=1):
        ws.cell(row=1, column=col_idx, value=col_name)
    for row_idx, row in enumerate(df.itertuples(index=False, name=None), start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=None if value == "" else value)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        header = str(column_cells[0].value or "")
        max_len = max([len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells[:200]] + [len(header)])
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 12), 42)


def _orientation_sheet(wb: Workbook) -> None:
    ws = wb.active
    ws.title = ORIENTATION_SHEET
    ws["E4"] = "GOVERNO DO ESTADO DE MINAS GERAIS"
    ws["C6"] = "Modelo DarwinCore IEF"


def export_darwincore_ief(
    *,
    df: pd.DataFrame,
    group: str,
    output_dir: Path,
    generated_files: list[str],
    include_fish_biometrics: bool = False,
) -> dict:
    if df.empty:
        return {"rows": 0, "warning": "dataset vazio"}

    c_id = _get_col(df, "id_resultado_pk", "id_resultado_bento")
    c_ponto = _get_col(df, "nome_ponto", "codigo_ponto", "ponto", "locality")
    c_date = _get_col(df, "data_hora_coleta", "data_campanha", "data_coleta", "eventDate")
    c_lat = _get_col(df, "latitude", "lat", "decimalLatitude")
    c_lon = _get_col(df, "longitude", "lon", "decimalLongitude")
    c_method = _get_col(df, "metodo_de_captura", "metodo", "metodo_amostragem", "samplingProtocol")
    c_effort = _get_col(df, "esforco", "samplingEffort")
    c_unit = _get_col(df, "unidade_esforco", "sampleSizeUnit")
    c_sci = _get_col(df, "nome_cientifico", "scientificName", "taxon_final")
    c_reino = _get_col(df, "reino", "kingdom")
    c_filo = _get_col(df, "filo", "phylum")
    c_classe = _get_col(df, "classe", "class")
    c_ordem = _get_col(df, "ordem", "order")
    c_familia = _get_col(df, "familia", "family")
    c_count = _get_col(df, "contagem", "numero_de_individuos", "abundancia", "individualCount")
    c_water = _get_col(df, "bacia_hidrografica", "curso_d_agua", "waterBody")
    c_weight = _get_col(df, "biomassa", "peso", "weight")
    c_standard_length = _get_col(df, "medida_1", "comprimento_padrao", "standard_length", "StandardLength")
    c_total_length = _get_col(df, "medida_2", "comprimento_total", "total_length", "TotalLength")
    c_sex = _get_col(df, "sexo", "sex", "Sex")
    c_gonadal_stage = _get_col(df, "estagio_gonadal", "estadio_gonadal", "gonadal_stage", "GonadalStage")
    c_gonad_weight = _get_col(df, "peso_gonada", "gonad_weight", "GonadWeight")

    sort_cols = [c for c in ["nome_campanha", "nome_ponto", "nome_cientifico", "id_resultado_pk"] if c in df.columns]
    df_work = df.sort_values(sort_cols, na_position="last").reset_index(drop=True) if sort_cols else df.reset_index(drop=True)

    sampling_rows: list[dict] = []
    occurrence_rows: list[dict] = []
    fish_rows: list[dict] = []

    for idx, row in enumerate(df_work.to_dict(orient="records"), start=1):
        locality = _clean(row.get(c_ponto)) if c_ponto else ""
        event_id = f"{idx}-{locality}" if locality else str(idx)
        occurrence_id = event_id
        individual_count = _numeric_or_blank(row.get(c_count)) if c_count else ""

        sampling_rows.append(
            {
                "eventID": event_id,
                "samplingProtocol": _clean(row.get(c_method)) if c_method else "",
                "samplingEffort": _effort_text(row.get(c_effort), row.get(c_unit)) if c_effort or c_unit else "",
                "sampleSizeValue": _numeric_or_blank(row.get(c_effort)) if c_effort else "",
                "sampleSizeUnit": _clean(row.get(c_unit)) if c_unit else "",
                "eventDate": _date_or_blank(row.get(c_date)) if c_date else "",
                "eventRemarks": "",
                "county": COUNTY,
                "municipality": MUNICIPALITY,
                "waterBody": _clean(row.get(c_water)) if c_water else "",
                "locality": locality,
                "decimalLatitude": _coord_or_blank(row.get(c_lat)) if c_lat else "",
                "decimalLongitude": _coord_or_blank(row.get(c_lon)) if c_lon else "",
                "geodeticDatum": GEODETIC_DATUM,
            }
        )

        occurrence_rows.append(
            {
                "eventID": event_id,
                "occurrenceID": occurrence_id,
                "basisOfRecord": BASIS_OF_RECORD,
                "scientificName": _clean(row.get(c_sci)) if c_sci else "",
                "kingdom": _clean(row.get(c_reino)) if c_reino else "",
                "phylum": _clean(row.get(c_filo)) if c_filo else "",
                "class": _clean(row.get(c_classe)) if c_classe else "",
                "order": _clean(row.get(c_ordem)) if c_ordem else "",
                "family": _clean(row.get(c_familia)) if c_familia else "",
                "taxonRank": TAXON_RANK,
                "identificationQualifier": "",
                "recordedBy": RECORDED_BY,
                "individualCount": individual_count,
                "sex": "",
                "lifeStage": "",
                "reproductiveCondition": "",
                "preparations": "",
                "occurrenceRemarks": "",
            }
        )

        if include_fish_biometrics:
            fish_rows.append(
                {
                    "eventID": event_id,
                    "occurrenceID": occurrence_id,
                    "scientificName": _clean(row.get(c_sci)) if c_sci else "",
                    "individualCount": individual_count,
                    "Weight": _numeric_or_blank(row.get(c_weight)) if c_weight else "",
                    "StandardLength": _numeric_or_blank(row.get(c_standard_length)) if c_standard_length else "",
                    "TotalLength": _numeric_or_blank(row.get(c_total_length)) if c_total_length else "",
                    "Sex": _clean(row.get(c_sex)) if c_sex else "",
                    "GonadalStage": _clean(row.get(c_gonadal_stage)) if c_gonadal_stage else "",
                    "GonadWeight": _numeric_or_blank(row.get(c_gonad_weight)) if c_gonad_weight else "",
                }
            )

    wb = Workbook()
    _orientation_sheet(wb)

    ws_sampling = wb.create_sheet("Sampling Events")
    sampling_df = pd.DataFrame(sampling_rows, columns=SAMPLING_EVENT_COLUMNS)
    _write_dataframe(ws_sampling, sampling_df)

    ws_occurrences = wb.create_sheet("Associated Occurrences")
    occurrence_df = pd.DataFrame(occurrence_rows, columns=ASSOCIATED_OCCURRENCE_COLUMNS)
    _write_dataframe(ws_occurrences, occurrence_df)

    if include_fish_biometrics:
        ws_fish = wb.create_sheet("Fish Biometric data")
        fish_df = pd.DataFrame(fish_rows, columns=FISH_BIOMETRIC_COLUMNS)
        _write_dataframe(ws_fish, fish_df)

    output_dir.mkdir(parents=True, exist_ok=True)
    out_file = output_dir / f"DarwinCore_IEF_{_slug_title(group)}_{_project_slug(df_work)}.xlsx"
    wb.save(out_file)
    generated_files.append(str(out_file))

    result = {
        "rows": int(len(occurrence_df)),
        "sampling_events": int(len(sampling_df)),
        "associated_occurrences": int(len(occurrence_df)),
        "file": str(out_file),
    }
    if include_fish_biometrics:
        result["fish_biometric_rows"] = int(len(fish_rows))
    return result
