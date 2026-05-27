from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from PIL import Image


def _xlsx_shape(path: Path) -> list[dict[str, Any]]:
    shapes: list[dict[str, Any]] = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            headers = []
            if sheet.max_row >= 1:
                headers = [str(cell.value or "")[:80] for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            shapes.append(
                {
                    "arquivo": path.name,
                    "sheet": sheet.title,
                    "rows": sheet.max_row,
                    "columns": sheet.max_column,
                    "headers": " | ".join(headers),
                }
            )
    finally:
        workbook.close()
    return shapes


def _png_check(path: Path) -> dict[str, Any]:
    with Image.open(path) as image:
        gray = image.convert("L")
        extrema = gray.getextrema()
        is_blank = extrema[0] == extrema[1]
        return {
            "arquivo": path.name,
            "width": image.size[0],
            "height": image.size[1],
            "mode": image.mode,
            "blank": is_blank,
            "small": image.size[0] < 300 or image.size[1] < 250,
        }


def inventory_results_dirs(result_dirs: dict[str, str | Path]) -> dict[str, Any]:
    files: list[dict[str, Any]] = []
    xlsx_sheets: list[dict[str, Any]] = []
    png_checks: list[dict[str, Any]] = []
    summary: list[dict[str, Any]] = []

    for group, raw_dir in result_dirs.items():
        base = Path(raw_dir)
        group_files = [
            path
            for path in base.rglob("*")
            if path.is_file() and path.name.lower() != "desktop.ini"
        ] if base.exists() else []

        ext_counter = Counter(path.suffix.lower() or "<sem_ext>" for path in group_files)
        duplicate_names = [
            name for name, count in Counter(path.name for path in group_files).items() if count > 1
        ]
        summary.append(
            {
                "grupo": group,
                "pasta": str(base),
                "exists": base.exists(),
                "arquivos": len(group_files),
                "png": ext_counter.get(".png", 0),
                "xlsx": ext_counter.get(".xlsx", 0),
                "duplicated_file_names": len(duplicate_names),
                "duplicated_names_sample": ", ".join(sorted(duplicate_names)[:10]),
            }
        )

        for path in sorted(group_files, key=lambda p: str(p).lower()):
            rel = path.relative_to(base)
            record = {
                "grupo": group,
                "arquivo": path.name,
                "relative_path": str(rel),
                "extension": path.suffix.lower() or "<sem_ext>",
                "size_bytes": path.stat().st_size,
            }
            files.append(record)

            if path.suffix.lower() == ".xlsx":
                try:
                    for shape in _xlsx_shape(path):
                        xlsx_sheets.append({"grupo": group, "relative_path": str(rel), **shape})
                except Exception as exc:  # noqa: BLE001
                    xlsx_sheets.append(
                        {
                            "grupo": group,
                            "relative_path": str(rel),
                            "arquivo": path.name,
                            "sheet": "<erro>",
                            "rows": None,
                            "columns": None,
                            "headers": repr(exc),
                        }
                    )
            elif path.suffix.lower() == ".png":
                try:
                    png_checks.append({"grupo": group, "relative_path": str(rel), **_png_check(path)})
                except Exception as exc:  # noqa: BLE001
                    png_checks.append(
                        {
                            "grupo": group,
                            "relative_path": str(rel),
                            "arquivo": path.name,
                            "width": None,
                            "height": None,
                            "mode": "<erro>",
                            "blank": True,
                            "small": True,
                            "erro": repr(exc),
                        }
                    )

    return {
        "summary": summary,
        "files": files,
        "xlsx_sheets": xlsx_sheets,
        "png_checks": png_checks,
    }

