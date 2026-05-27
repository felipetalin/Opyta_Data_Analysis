from __future__ import annotations

import re
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter
from pathlib import Path
from typing import Any

from docx import Document


WORD_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def _norm_style(value: object) -> str:
    return str(value or "").strip().lower()


def _is_heading(style_name: str) -> bool:
    style = _norm_style(style_name)
    return style.startswith("heading") or "titulo" in style or "título" in style


def _is_toc_style(style_name: str) -> bool:
    style = _norm_style(style_name)
    return style.startswith("toc") or "table of figures" in style or "sumario" in style


def _is_caption_style(style_name: str) -> bool:
    style = _norm_style(style_name)
    return "caption" in style or "legenda" in style or "table of figures" in style


def _caption_match(text: str, style_name: str = "") -> re.Match[str] | None:
    match = re.match(
        r"^\s*(Figura|Tabela|Quadro|Gr[aá]fico)\s+([0-9]+(?:[.\-][0-9A-Za-z]+)*)\b",
        text,
        flags=re.IGNORECASE,
    )
    if not match:
        return None
    if _is_caption_style(style_name) or _is_toc_style(style_name):
        return match

    suffix = text[match.end(): match.end() + 4]
    if re.match(r"\s*[-–—:]", suffix):
        return match
    return None


def _extract_comments(docx_path: Path) -> list[dict[str, Any]]:
    comments: list[dict[str, Any]] = []
    with zipfile.ZipFile(docx_path) as zf:
        if "word/comments.xml" not in zf.namelist():
            return comments
        root = ET.fromstring(zf.read("word/comments.xml"))
        for comment in root.findall(".//w:comment", WORD_NS):
            text = "".join(t.text or "" for t in comment.findall(".//w:t", WORD_NS)).strip()
            comments.append(
                {
                    "id": comment.get(f"{{{WORD_NS['w']}}}id"),
                    "author": comment.get(f"{{{WORD_NS['w']}}}author"),
                    "date": comment.get(f"{{{WORD_NS['w']}}}date"),
                    "text": text,
                }
            )
    return comments


def _extract_xml_paragraphs(docx_path: Path) -> list[dict[str, Any]]:
    with zipfile.ZipFile(docx_path) as zf:
        names = set(zf.namelist())
        style_map: dict[str, str] = {}
        if "word/styles.xml" in names:
            styles_root = ET.fromstring(zf.read("word/styles.xml"))
            for style in styles_root.findall(".//w:style", WORD_NS):
                style_id = style.get(f"{{{WORD_NS['w']}}}styleId")
                name = style.find("w:name", WORD_NS)
                if style_id and name is not None:
                    style_map[style_id] = name.get(f"{{{WORD_NS['w']}}}val") or style_id

        root = ET.fromstring(zf.read("word/document.xml"))

    records: list[dict[str, Any]] = []
    for idx, paragraph in enumerate(root.findall(".//w:p", WORD_NS), start=1):
        texts = [text.text for text in paragraph.findall(".//w:t", WORD_NS) if text.text]
        value = "".join(texts).strip()
        if not value:
            continue

        style_id = ""
        ppr = paragraph.find("w:pPr", WORD_NS)
        if ppr is not None:
            pstyle = ppr.find("w:pStyle", WORD_NS)
            if pstyle is not None:
                style_id = pstyle.get(f"{{{WORD_NS['w']}}}val") or ""
        style_name = style_map.get(style_id, style_id or "<sem estilo>")
        records.append(
            {
                "paragraph_index": idx,
                "style": style_name,
                "text": value,
                "text_len": len(value),
                "is_toc": _is_toc_style(style_name),
            }
        )
    return records


def _extract_zip_metadata(docx_path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(docx_path) as zf:
        names = set(zf.namelist())
        doc_xml = ET.fromstring(zf.read("word/document.xml"))
        media = [name for name in names if name.startswith("word/media/")]
        return {
            "media_files_count": len(media),
            "comments_xml": "word/comments.xml" in names,
            "footnotes_xml": "word/footnotes.xml" in names,
            "endnotes_xml": "word/endnotes.xml" in names,
            "track_changes_insertions": len(doc_xml.findall(".//w:ins", WORD_NS)),
            "track_changes_deletions": len(doc_xml.findall(".//w:del", WORD_NS)),
        }


def extract_docx_audit(docx_path: str | Path) -> dict[str, Any]:
    path = Path(docx_path)
    document = Document(str(path))

    paragraphs = _extract_xml_paragraphs(path)
    headings: list[dict[str, Any]] = []
    captions: list[dict[str, Any]] = []
    references: list[dict[str, Any]] = []
    style_counter: Counter[str] = Counter()

    ref_pattern = re.compile(
        r"\b(Figura|Tabela|Quadro|Gr[aá]fico)\s+([0-9]+(?:[.\-][0-9A-Za-z]+)*)",
        flags=re.IGNORECASE,
    )

    for record in paragraphs:
        text = str(record.get("text", ""))
        style_name = str(record.get("style", ""))
        style_counter[style_name or "<sem estilo>"] += 1

        if _is_heading(style_name):
            headings.append(record.copy())

        cap = _caption_match(text, style_name)
        if cap:
            captions.append(
                {
                    "paragraph_index": record.get("paragraph_index"),
                    "style": style_name,
                    "tipo": cap.group(1).capitalize(),
                    "numero": cap.group(2),
                    "text": text,
                    "is_toc": _is_toc_style(style_name),
                }
            )

        if not _is_toc_style(style_name) and cap is None:
            for match in ref_pattern.finditer(text):
                start = max(0, match.start() - 120)
                end = min(len(text), match.end() + 160)
                references.append(
                    {
                        "paragraph_index": record.get("paragraph_index"),
                        "tipo": match.group(1).capitalize(),
                        "numero": match.group(2),
                        "contexto": text[start:end],
                    }
                )

    tables: list[dict[str, Any]] = []
    for idx, table in enumerate(document.tables, start=1):
        first_row: list[str] = []
        if table.rows:
            first_row = [" ".join(cell.text.split())[:120] for cell in table.rows[0].cells]
        tables.append(
            {
                "table_index": idx,
                "rows": len(table.rows),
                "columns": len(table.columns),
                "first_row": " | ".join(first_row),
            }
        )

    text_all = "\n".join(p["text"] for p in paragraphs)
    words_count = len(re.findall(r"\w+", text_all, flags=re.UNICODE))

    zip_meta = _extract_zip_metadata(path)
    comments = _extract_comments(path)

    return {
        "docx_path": str(path),
        "file_name": path.name,
        "size_bytes": path.stat().st_size,
        "paragraphs_count": len(paragraphs),
        "words_count_estimate": words_count,
        "tables_count": len(document.tables),
        "inline_shapes_count": len(document.inline_shapes),
        "sections_count": len(document.sections),
        "styles_count": dict(style_counter.most_common()),
        "paragraphs": paragraphs,
        "headings": headings,
        "captions": captions,
        "references": references,
        "tables": tables,
        "comments": comments,
        "zip_metadata": zip_meta,
    }


def render_extracted_markdown(doc: dict[str, Any]) -> str:
    lines = [
        f"# Texto extraido - {doc.get('file_name', '')}",
        "",
        f"- Paragrafos: {doc.get('paragraphs_count', 0)}",
        f"- Palavras estimadas: {doc.get('words_count_estimate', 0)}",
        f"- Tabelas: {doc.get('tables_count', 0)}",
        f"- Figuras/objetos inline: {doc.get('inline_shapes_count', 0)}",
        "",
    ]
    for para in doc.get("paragraphs", []):
        style = para.get("style") or "<sem estilo>"
        text = para.get("text") or ""
        if _is_heading(style):
            level_match = re.search(r"(\d+)", style)
            level = int(level_match.group(1)) if level_match else 2
            level = max(1, min(level, 6))
            lines.append(f"{'#' * level} {text}")
        else:
            lines.append(text)
        lines.append("")
    return "\n".join(lines)
