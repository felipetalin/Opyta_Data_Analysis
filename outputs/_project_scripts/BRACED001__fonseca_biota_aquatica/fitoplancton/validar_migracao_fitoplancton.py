from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import text


TAXON_MAP = {"Scytonemataceae n.i.": "Scytonemataceae"}


def get_engine(opyta_data_root: Path):
    sys.path.insert(0, str(opyta_data_root))
    from core.engine import get_engine as build_engine

    return build_engine()


def clean_series(series: pd.Series) -> pd.Series:
    return series.fillna("").astype(str).str.strip()


def compare_values(source: pd.DataFrame, database: pd.DataFrame, keys: list[str], source_value: str, database_value: str, label: str) -> pd.DataFrame:
    left = source.groupby(keys, dropna=False)[source_value].sum().reset_index(name="valor_fonte")
    right = database.groupby(keys, dropna=False)[database_value].sum().reset_index(name="valor_banco")
    merged = left.merge(right, on=keys, how="outer", indicator=True)
    merged["valor_fonte"] = pd.to_numeric(merged["valor_fonte"], errors="coerce")
    merged["valor_banco"] = pd.to_numeric(merged["valor_banco"], errors="coerce")
    merged["diferenca"] = merged["valor_banco"] - merged["valor_fonte"]
    merged["tipo_comparacao"] = label
    return merged[(merged["_merge"] != "both") | merged["diferenca"].abs().gt(1e-9)].copy()


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        for column in worksheet.columns:
            width = min(70, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            worksheet.column_dimensions[column[0].column_letter].width = width
    workbook.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate BRACED001 phytoplankton migration and consolidation.")
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--opyta-data-root", required=True, type=Path)
    args = parser.parse_args()

    source_results = pd.read_excel(args.source, sheet_name="Resultados_Fitoplancton")
    source_points = pd.read_excel(args.source, sheet_name="Pontos_e_Campanhas")
    source_efforts = pd.read_excel(args.source, sheet_name="Metadados_Esforco")
    source_results["Nome_Cientifico"] = clean_series(source_results["Nome_Cientifico"]).replace(TAXON_MAP)
    for frame, columns in (
        (source_results, ["Campanha", "Ponto", "Metodo_de_Captura", "Tipo_de_Amostragem"]),
        (source_efforts, ["Campanha", "Ponto", "Metodo_de_Captura", "Tipo_de_Amostragem"]),
        (source_points, ["Campanha", "Ponto"]),
    ):
        for column in columns:
            frame[column] = clean_series(frame[column])
    source_results["Numero_de_Individuos"] = pd.to_numeric(source_results["Numero_de_Individuos"], errors="coerce").fillna(0)
    source_efforts["Esforco"] = pd.to_numeric(source_efforts["Esforco"], errors="coerce")

    engine = get_engine(args.opyta_data_root)
    with engine.connect() as connection:
        db_results = pd.read_sql(
            text(
                """
                SELECT c.nome_campanha AS "Campanha", p.nome_ponto AS "Ponto",
                       e.metodo_de_captura AS "Metodo_de_Captura",
                       COALESCE(rf.tipo_amostragem, e.tipo_amostragem, e.tipo_de_amostragem) AS "Tipo_de_Amostragem",
                       sp.nome_cientifico AS "Nome_Cientifico", rf.densidade AS densidade
                FROM public.resultados_fitoplancton rf
                JOIN public.esforcos_amostragem e ON e.id_esforco = rf.id_esforco
                JOIN public.pontos_coleta p ON p.id_ponto_coleta = e.id_ponto_coleta
                JOIN public.campanhas c ON c.id_campanha = p.id_campanha
                JOIN public.especies sp ON sp.id_especie = rf.id_especie
                WHERE p.id_projeto = 133 AND e.grupo_biologico = 'Fitoplâncton'
                """
            ),
            connection,
        )
        db_efforts = pd.read_sql(
            text(
                """
                SELECT c.nome_campanha AS "Campanha", p.nome_ponto AS "Ponto",
                       e.metodo_de_captura AS "Metodo_de_Captura",
                       COALESCE(e.tipo_amostragem, e.tipo_de_amostragem) AS "Tipo_de_Amostragem",
                       e.esforco, e.unidade_esforco AS "Unidade_Esforco"
                FROM public.esforcos_amostragem e
                JOIN public.pontos_coleta p ON p.id_ponto_coleta = e.id_ponto_coleta
                JOIN public.campanhas c ON c.id_campanha = p.id_campanha
                WHERE p.id_projeto = 133 AND e.grupo_biologico = 'Fitoplâncton'
                """
            ),
            connection,
        )
        db_points = pd.read_sql(
            text(
                """
                SELECT c.nome_campanha AS "Campanha", p.nome_ponto AS "Ponto", p.latitude, p.longitude, p.data_hora_coleta
                FROM public.pontos_coleta p
                JOIN public.campanhas c ON c.id_campanha = p.id_campanha
                WHERE p.id_projeto = 133 AND c.nome_campanha IN ('C001-2026-03-CH', 'C002-2026-06-SC')
                """
            ),
            connection,
        )
        consolidated = pd.read_sql(
            text(
                """
                SELECT nome_campanha AS "Campanha", nome_ponto AS "Ponto", metodo_de_captura AS "Metodo_de_Captura",
                       tipo_amostragem AS "Tipo_de_Amostragem", nome_cientifico AS "Nome_Cientifico", contagem
                FROM public.biota_analise_consolidada
                WHERE id_projeto = 133 AND grupo_biologico = 'Fitoplâncton'
                """
            ),
            connection,
        )

    keys = ["Campanha", "Ponto", "Metodo_de_Captura", "Tipo_de_Amostragem", "Nome_Cientifico"]
    result_diff = compare_values(source_results, db_results, keys, "Numero_de_Individuos", "densidade", "fonte_vs_base")
    consolidated_diff = compare_values(source_results, consolidated, keys, "Numero_de_Individuos", "contagem", "fonte_vs_consolidado")

    effort_keys = ["Campanha", "Ponto", "Metodo_de_Captura", "Tipo_de_Amostragem"]
    effort_diff = source_efforts[effort_keys + ["Esforco", "Unidade_Esforco"]].merge(
        db_efforts[effort_keys + ["esforco", "Unidade_Esforco"]], on=effort_keys, how="outer", indicator=True
    )
    effort_diff["diferenca"] = pd.to_numeric(effort_diff["esforco"], errors="coerce") - pd.to_numeric(effort_diff["Esforco"], errors="coerce")
    effort_diff = effort_diff[
        (effort_diff["_merge"] != "both")
        | effort_diff["diferenca"].abs().gt(1e-9)
        | (clean_series(effort_diff["Unidade_Esforco_x"]) != clean_series(effort_diff["Unidade_Esforco_y"]))
    ].copy()

    point_compare = source_points[["Campanha", "Ponto", "Latitude", "Longitude", "Data"]].merge(
        db_points, on=["Campanha", "Ponto"], how="outer", indicator=True
    )
    point_compare["lat_diff"] = pd.to_numeric(point_compare["latitude"], errors="coerce") - pd.to_numeric(point_compare["Latitude"], errors="coerce")
    point_compare["lon_diff"] = pd.to_numeric(point_compare["longitude"], errors="coerce") - pd.to_numeric(point_compare["Longitude"], errors="coerce")
    point_compare["data_fonte"] = pd.to_datetime(point_compare["Data"], errors="coerce", dayfirst=True).dt.date
    point_compare["data_banco"] = pd.to_datetime(point_compare["data_hora_coleta"], errors="coerce").dt.date
    point_diff = point_compare[
        (point_compare["_merge"] != "both")
        | point_compare["lat_diff"].abs().gt(1e-9)
        | point_compare["lon_diff"].abs().gt(1e-9)
        | (point_compare["data_fonte"] != point_compare["data_banco"])
    ].copy()

    summary = {
        "projeto": "BRACED001",
        "grupo": "Fitoplancton",
        "fonte_resultados": len(source_results),
        "banco_resultados": len(db_results),
        "consolidado_resultados": len(consolidated),
        "fonte_esforcos": len(source_efforts),
        "banco_esforcos": len(db_efforts),
        "fonte_taxons_apos_alias": source_results["Nome_Cientifico"].nunique(),
        "banco_taxons": db_results["Nome_Cientifico"].nunique(),
        "divergencias_resultados_base": len(result_diff),
        "divergencias_resultados_consolidado": len(consolidated_diff),
        "divergencias_esforcos": len(effort_diff),
        "divergencias_pontos_datas_coordenadas": len(point_diff),
        "status": "PASS" if all(frame.empty for frame in (result_diff, consolidated_diff, effort_diff, point_diff)) else "BLOCKED",
    }

    args.output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%dT%H%M%S")
    base = args.output_dir / f"{stamp}_auditoria_pos_migracao_fitoplancton_braced001"
    with pd.ExcelWriter(base.with_suffix(".xlsx"), engine="openpyxl") as writer:
        pd.DataFrame([{"indicador": key, "valor": value} for key, value in summary.items()]).to_excel(writer, sheet_name="00_resumo", index=False)
        result_diff.to_excel(writer, sheet_name="01_diff_base", index=False)
        consolidated_diff.to_excel(writer, sheet_name="02_diff_consolidado", index=False)
        effort_diff.to_excel(writer, sheet_name="03_diff_esforcos", index=False)
        point_diff.to_excel(writer, sheet_name="04_diff_pontos", index=False)
        db_results.to_excel(writer, sheet_name="05_resultados_banco", index=False)
    style_workbook(base.with_suffix(".xlsx"))
    base.with_suffix(".json").write_text(json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(json.dumps({"workbook": str(base.with_suffix('.xlsx')), "json": str(base.with_suffix('.json')), **summary}, ensure_ascii=False, indent=2))
    return 0 if summary["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
