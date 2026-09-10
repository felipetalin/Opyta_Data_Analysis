from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import bindparam, text


FIELD_MAP = {
    "Nome_Cientifico": "nome_cientifico",
    "Nome_Popular": "nome_popular",
    "Grupo_Biologico": "grupo_biologico",
    "Reino": "reino",
    "Filo": "filo",
    "Classe": "classe",
    "Ordem": "ordem",
    "Familia": "familia",
    "Genero": "genero",
    "Autor_e_Ano": "autor_e_ano",
    "Status_Ameaca_Estadual": "status_estadual",
    "Status_Ameaca_Nacional": "status_ameaca_nacional",
    "Status_Ameaca_Global": "status_ameaca_global",
    "Origem": "origem",
    "Habito_Alimentar": "habito_alimentar",
    "Estrategia_Reprodutiva": "estrategia_reprodutiva",
    "Valor_Economico": "valor_economico",
    "Observacoes": "observacoes",
    "Cinegetica": "cinegetica",
    "Xerimbabo": "xerimbabo",
}
BOOL_FIELDS = {"cinegetica", "xerimbabo"}


def clean(value):
    if value is None or pd.isna(value):
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def bool_or_none(value):
    value = clean(value)
    if value is None or isinstance(value, bool):
        return value
    normalized = str(value).strip().casefold()
    if normalized in {"sim", "s", "true", "1", "x"}:
        return True
    if normalized in {"nao", "não", "n", "false", "0"}:
        return False
    return None


def read_source(path: Path) -> pd.DataFrame:
    frame = pd.read_excel(path, sheet_name="Cadastro_Especies")
    rename = {source: target for source, target in FIELD_MAP.items() if source in frame.columns}
    frame = frame.rename(columns=rename)
    frame = frame.loc[frame["nome_cientifico"].notna(), list(dict.fromkeys(rename.values()))].copy()
    for column in frame.columns:
        frame[column] = frame[column].map(clean)
    for column in BOOL_FIELDS & set(frame.columns):
        frame[column] = frame[column].map(bool_or_none)
    return frame.drop_duplicates("nome_cientifico", keep="last").set_index("nome_cientifico")


def get_engine(opyta_data_root: Path):
    sys.path.insert(0, str(opyta_data_root))
    from core.engine import get_engine as build_engine

    return build_engine()


def database_columns(connection) -> set[str]:
    rows = connection.execute(
        text(
            """
            SELECT column_name
            FROM information_schema.columns
            WHERE table_schema = 'public' AND table_name = 'especies'
            """
        )
    )
    return {row[0] for row in rows}


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        for column in worksheet.columns:
            width = min(70, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            worksheet.column_dimensions[column[0].column_letter].width = width
    workbook.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Apply the approved BRACED001 phytoplankton Gate B package.")
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--decision-json", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--opyta-data-root", required=True, type=Path)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--approved-gate-b", action="store_true")
    args = parser.parse_args()
    if not args.apply or not args.approved_gate_b:
        raise SystemExit("Use --apply --approved-gate-b somente apos aprovacao explicita do Gate B.")

    source = read_source(args.source)
    package = json.loads(args.decision_json.read_text(encoding="utf-8"))
    actions = pd.DataFrame(package["actions"])
    expected = package["summary"]
    if expected["taxons_novos"] != 92 or expected["taxons_existentes_reutilizados"] != 106:
        raise RuntimeError("Pacote Gate B nao corresponde ao escopo aprovado de 92 novos e 106 existentes.")

    stamp = datetime.now().strftime("%Y%m%dT%H%M%S")
    backup_table = f"backup_especies_braced001_fitoplancton_gate_b_{stamp.lower()}"
    if not re.fullmatch(r"[a-z0-9_]+", backup_table):
        raise RuntimeError("Nome de backup invalido.")

    applied: list[dict[str, object]] = []
    engine = get_engine(args.opyta_data_root)
    with engine.begin() as connection:
        columns = database_columns(connection)
        connection.execute(text(f"CREATE TABLE public.{backup_table} AS TABLE public.especies"))

        for action in actions.to_dict("records"):
            name = action["nome_cientifico"]
            operation = action["acao"]
            if operation == "REUSE_ALIAS":
                applied.append(
                    {
                        "nome_cientifico": name,
                        "acao": operation,
                        "id_especie": int(action["id_especie"]),
                        "campos_alterados": 0,
                    }
                )
                continue

            row = source.loc[name].to_dict()
            values = {
                field: clean(value)
                for field, value in row.items()
                if field in columns and field != "nome_cientifico" and clean(value) is not None
            }
            values["nome_cientifico"] = name

            if operation == "INSERT_NEW":
                insert_columns = list(values)
                new_id = connection.execute(
                    text(
                        f"INSERT INTO public.especies ({', '.join(insert_columns)}) "
                        f"VALUES ({', '.join(':' + column for column in insert_columns)}) "
                        "ON CONFLICT (nome_cientifico) DO NOTHING RETURNING id_especie"
                    ),
                    values,
                ).scalar()
                if new_id is None:
                    raise RuntimeError(f"Taxon novo ja existe ou nao foi inserido: {name}")
                applied.append({"nome_cientifico": name, "acao": operation, "id_especie": int(new_id), "campos_alterados": len(values) - 1})
                continue

            id_especie = int(action["id_especie"])
            assignments = [f"{field} = COALESCE({field}, :{field})" for field in values if field != "nome_cientifico"]
            parameters = {**values, "id_especie": id_especie}
            if operation == "FILL_NULLS_AND_FIX_GROUP_ENCODING":
                assignments = [item for item in assignments if not item.startswith("grupo_biologico =")]
                assignments.append("grupo_biologico = :grupo_biologico")
            result = connection.execute(
                text(f"UPDATE public.especies SET {', '.join(assignments)} WHERE id_especie = :id_especie"),
                parameters,
            )
            if result.rowcount != 1:
                raise RuntimeError(f"Registro existente nao atualizado de forma unica: {name} / {id_especie}")
            applied.append(
                {
                    "nome_cientifico": name,
                    "acao": operation,
                    "id_especie": id_especie,
                    "campos_alterados": int(action.get("campos_a_preencher") or 0) + (1 if operation == "FILL_NULLS_AND_FIX_GROUP_ENCODING" else 0),
                }
            )

        target_names = [row["nome_banco"] for row in actions.to_dict("records")]
        verify = connection.execute(
            text(
                "SELECT id_especie, nome_cientifico, grupo_biologico FROM public.especies "
                "WHERE nome_cientifico IN :names"
            ).bindparams(bindparam("names", expanding=True)),
            {"names": sorted(set(target_names))},
        ).mappings().all()

    verify_df = pd.DataFrame(verify)
    missing = sorted(set(target_names) - set(verify_df["nome_cientifico"]))
    invalid_groups = verify_df[verify_df["grupo_biologico"].astype(str).str.casefold() != "fitoplâncton".casefold()]
    summary = {
        "projeto": "BRACED001",
        "grupo": "Fitoplancton",
        "modo": "apply_approved",
        "backup_table": f"public.{backup_table}",
        "taxons_novos_inseridos": sum(row["acao"] == "INSERT_NEW" for row in applied),
        "taxons_existentes_processados": sum(row["acao"] != "INSERT_NEW" for row in applied),
        "registros_aplicados": len(applied),
        "taxons_verificados_pos_acao": len(verify_df),
        "taxons_ausentes_pos_acao": len(missing),
        "grupos_invalidos_pos_acao": len(invalid_groups),
        "status": "PASS" if not missing and invalid_groups.empty and len(applied) == 198 else "BLOCKED",
    }

    args.output_dir.mkdir(parents=True, exist_ok=True)
    base = args.output_dir / f"{stamp}_apply_gate_b_fitoplancton_braced001"
    with pd.ExcelWriter(base.with_suffix(".xlsx"), engine="openpyxl") as writer:
        pd.DataFrame([{"indicador": key, "valor": value} for key, value in summary.items()]).to_excel(
            writer, sheet_name="00_resumo", index=False
        )
        pd.DataFrame(applied).to_excel(writer, sheet_name="01_aplicados", index=False)
        verify_df.to_excel(writer, sheet_name="02_validacao_pos", index=False)
        pd.DataFrame({"nome_cientifico": missing}).to_excel(writer, sheet_name="03_ausentes", index=False)
        invalid_groups.to_excel(writer, sheet_name="04_grupos_invalidos", index=False)
    style_workbook(base.with_suffix(".xlsx"))
    base.with_suffix(".json").write_text(
        json.dumps({"summary": summary, "applied": applied}, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )
    print(json.dumps({"workbook": str(base.with_suffix('.xlsx')), "json": str(base.with_suffix('.json')), **summary}, ensure_ascii=False, indent=2))
    return 0 if summary["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
