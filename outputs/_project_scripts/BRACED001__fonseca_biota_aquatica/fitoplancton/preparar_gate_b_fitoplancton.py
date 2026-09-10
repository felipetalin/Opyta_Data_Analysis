from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import bindparam, text


GROUP_FIXES = {"Frustulia sp.", "Iconella delicatissima"}
ALIAS_MAP = {"Scytonemataceae n.i.": "Scytonemataceae"}


def get_engine(opyta_data_root: Path):
    sys.path.insert(0, str(opyta_data_root))
    from core.engine import get_engine as build_engine

    return build_engine()


def read_database(engine, names: list[str]) -> pd.DataFrame:
    query = text(
        """
        SELECT id_especie, nome_cientifico, grupo_biologico
        FROM public.especies
        WHERE nome_cientifico IN :names
        ORDER BY nome_cientifico
        """
    ).bindparams(bindparam("names", expanding=True))
    with engine.connect() as connection:
        return pd.read_sql(query, connection, params={"names": names})


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        for column in worksheet.columns:
            width = min(70, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            worksheet.column_dimensions[column[0].column_letter].width = width
    workbook.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Prepare the BRACED001 phytoplankton Gate B decision package without database writes.")
    parser.add_argument("--audit-workbook", required=True, type=Path)
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--opyta-data-root", required=True, type=Path)
    args = parser.parse_args()

    actions = pd.read_excel(args.audit_workbook, sheet_name="01_acoes")
    comparisons = pd.read_excel(args.audit_workbook, sheet_name="02_comparacao_campos")
    remaining = pd.read_excel(args.audit_workbook, sheet_name="03_pendencias_pos_acao")
    conciliation = pd.read_excel(args.source, sheet_name="CONCILIACAO_BANCO")

    target_names = sorted(set(actions["nome_cientifico"]) | set(ALIAS_MAP.values()))
    database = read_database(get_engine(args.opyta_data_root), target_names)
    db_by_name = database.set_index("nome_cientifico").to_dict("index")

    actions = actions.copy()
    actions["nome_banco"] = actions["nome_cientifico"]
    actions["ajuste_especial"] = None
    for source_name, database_name in ALIAS_MAP.items():
        mask = actions["nome_cientifico"].eq(source_name)
        target = db_by_name.get(database_name)
        if not mask.any() or target is None:
            raise RuntimeError(f"Alias sem correspondencia validada: {source_name} -> {database_name}")
        actions.loc[mask, "acao"] = "REUSE_ALIAS"
        actions.loc[mask, "id_especie"] = target["id_especie"]
        actions.loc[mask, "nome_banco"] = database_name
        actions.loc[mask, "campos_a_preencher"] = 0
        actions.loc[mask, "ajuste_especial"] = "De/para aprovado no cadastro final e usado em BRAAEG001."

    for name in GROUP_FIXES:
        mask = actions["nome_cientifico"].eq(name)
        if not mask.any() or pd.isna(actions.loc[mask, "id_especie"]).all():
            raise RuntimeError(f"Registro existente nao localizado para correcao de grupo: {name}")
        actions.loc[mask, "acao"] = "FILL_NULLS_AND_FIX_GROUP_ENCODING"
        actions.loc[mask, "ajuste_especial"] = "Corrigir Fitopl?ncton para Fitoplâncton; preservar demais valores nao nulos."

    gomphonema = actions["nome_cientifico"].eq("Gomphonema lagenula")
    gomphonema_db = db_by_name.get("Gomphonema lagenula")
    if not gomphonema.any() or gomphonema_db is None:
        raise RuntimeError("Gomphonema lagenula nao localizado por correspondencia exata no banco.")
    actions.loc[gomphonema, "id_especie"] = gomphonema_db["id_especie"]
    actions.loc[gomphonema, "nome_banco"] = "Gomphonema lagenula"
    actions.loc[gomphonema, "ajuste_especial"] = "Reutilizar correspondencia exata; preservar taxonomia nao nula do banco."

    new_taxa = actions[actions["acao"].eq("INSERT_NEW")].copy()
    existing_taxa = actions[~actions["acao"].eq("INSERT_NEW")].copy()
    group_differences = comparisons[
        comparisons["campo"].eq("grupo_biologico")
        & comparisons["decisao"].eq("PRESERVE_DB_DIFFERENCE")
    ]
    expected_group_fixes = set(group_differences["nome_cientifico"])
    if expected_group_fixes != GROUP_FIXES:
        raise RuntimeError(f"Conflitos de grupo inesperados: {sorted(expected_group_fixes)}")

    fills = int((comparisons["decisao"] == "FILL_NULL").sum())
    preserved = int((comparisons["decisao"] == "PRESERVE_DB_DIFFERENCE").sum()) - len(GROUP_FIXES)
    special = pd.DataFrame(
        [
            {
                "taxon_fonte": "Scytonemataceae n.i.",
                "decisao": "REUSE_ALIAS",
                "taxon_banco": "Scytonemataceae",
                "id_especie": db_by_name["Scytonemataceae"]["id_especie"],
                "motivo": "Evitar duplicata e manter o padrao usado em BRAAEG001.",
            },
            {
                "taxon_fonte": "Frustulia sp.",
                "decisao": "FIX_GROUP_ENCODING",
                "taxon_banco": "Frustulia sp.",
                "id_especie": db_by_name["Frustulia sp."]["id_especie"],
                "motivo": "Corrigir Fitopl?ncton para Fitoplâncton; preservar demais divergencias.",
            },
            {
                "taxon_fonte": "Iconella delicatissima",
                "decisao": "FIX_GROUP_ENCODING",
                "taxon_banco": "Iconella delicatissima",
                "id_especie": db_by_name["Iconella delicatissima"]["id_especie"],
                "motivo": "Corrigir Fitopl?ncton para Fitoplâncton; preservar demais divergencias.",
            },
            {
                "taxon_fonte": "Gomphonema lagenula",
                "decisao": "REUSE_EXACT",
                "taxon_banco": "Gomphonema lagenula",
                "id_especie": gomphonema_db["id_especie"],
                "motivo": "Correspondencia exata encontrada; nao criar duplicata nem sobrescrever taxonomia nao nula.",
            },
        ]
    )

    summary = {
        "projeto": "BRACED001",
        "grupo": "Fitoplancton",
        "modo": "audit_only",
        "taxons_resultados": 198,
        "taxons_cadastro_canonico": len(actions),
        "linhas_duplicadas_descartadas": 87,
        "estrategia_duplicados": "keep_last",
        "taxons_novos": len(new_taxa),
        "taxons_existentes_reutilizados": len(existing_taxa),
        "campos_nulos_a_preencher": fills,
        "correcoes_grupo_codificacao": len(GROUP_FIXES),
        "diferencas_nao_nulas_preservadas": preserved,
        "aliases": len(ALIAS_MAP),
        "pendencias_obrigatorias_pos_acao": len(remaining),
        "registros_aplicados": 0,
        "status_gate_b": "AWAITING_APPROVAL",
    }

    args.output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%dT%H%M%S")
    base = args.output_dir / f"{stamp}_gate_b_decisao_fitoplancton_braced001"
    with pd.ExcelWriter(base.with_suffix(".xlsx"), engine="openpyxl") as writer:
        pd.DataFrame([{"indicador": key, "valor": value} for key, value in summary.items()]).to_excel(
            writer, sheet_name="00_resumo", index=False
        )
        actions.to_excel(writer, sheet_name="01_acoes_consolidadas", index=False)
        special.to_excel(writer, sheet_name="02_ajustes_especiais", index=False)
        new_taxa.to_excel(writer, sheet_name="03_taxons_novos", index=False)
        existing_taxa.to_excel(writer, sheet_name="04_taxons_existentes", index=False)
        comparisons.to_excel(writer, sheet_name="05_comparacao_campos", index=False)
        conciliation.to_excel(writer, sheet_name="06_conciliacao_fonte", index=False)
        remaining.to_excel(writer, sheet_name="07_pendencias", index=False)
    style_workbook(base.with_suffix(".xlsx"))
    base.with_suffix(".json").write_text(
        json.dumps(
            {
                "summary": summary,
                "special_decisions": special.to_dict("records"),
                "actions": actions.to_dict("records"),
            },
            ensure_ascii=False,
            indent=2,
            default=str,
        ),
        encoding="utf-8",
    )
    print(json.dumps({"workbook": str(base.with_suffix('.xlsx')), "json": str(base.with_suffix('.json')), **summary}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
