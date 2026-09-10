from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import bindparam, text


FIELD_MAP = {
    "Nome_Cientifico": "nome_cientifico",
    "Nome_Popular": "nome_popular",
    "Grupo_Biologico": "grupo_biologico",
    "Reino": "reino",
    "Filo": "filo",
    "Classe": "classe",
    "Ordem": "ordem",
    "Familia": "familia",
    "Genero": "genero",
    "Autor_e_Ano": "autor_e_ano",
    "Status_Ameaca_Nacional": "status_ameaca_nacional",
    "Status_Ameaca_Global": "status_ameaca_global",
    "Origem": "origem",
    "Habito_Alimentar": "habito_alimentar",
    "Estrategia_Reprodutiva": "estrategia_reprodutiva",
    "Valor_Economico": "valor_economico",
    "Cinegetica": "cinegetica",
    "Xerimbabo": "xerimbabo",
    "Observacoes": "observacoes",
    "Status_Estadual": "status_estadual",
    "Status_Ameaca_Estadual": "status_estadual",
    "bmwp_score": "bmwp_score",
}

REQUIRED_ANALYTICAL_FIELDS = (
    "status_ameaca_nacional",
    "status_ameaca_global",
    "habito_alimentar",
    "estrategia_reprodutiva",
    "valor_economico",
)

PRESERVE_NA_TEXT = False


def clean(value):
    if value is None or pd.isna(value):
        return None
    if isinstance(value, str):
        value = value.strip()
        null_markers = {"(vazio)"} if PRESERVE_NA_TEXT else {"(vazio)", "n.a."}
        if value.casefold() in null_markers:
            return None
        return value or None
    return value


def bool_or_none(value):
    value = clean(value)
    if value is None or isinstance(value, bool):
        return value
    normalized = str(value).strip().lower()
    if normalized in {"sim", "s", "true", "1", "x"}:
        return True
    if normalized in {"nao", "não", "n", "false", "0"}:
        return False
    return None


def read_source(path: Path) -> pd.DataFrame:
    sheet = "Cadastro_Especies"
    xls = pd.ExcelFile(path)
    if sheet not in xls.sheet_names:
        sheet = "Especies"
    frame = pd.read_excel(path, sheet_name=sheet)
    rename = {source: target for source, target in FIELD_MAP.items() if source in frame.columns}
    frame = frame.rename(columns=rename)
    if "nome_cientifico" not in frame.columns:
        raise ValueError("Coluna Nome_Cientifico nao encontrada no cadastro.")
    frame = frame.loc[frame["nome_cientifico"].notna(), list(dict.fromkeys(rename.values()))].copy()
    for column in frame.columns:
        frame[column] = frame[column].map(clean)
    for column in ("cinegetica", "xerimbabo"):
        if column in frame.columns:
            frame[column] = frame[column].map(bool_or_none)
    if "bmwp_score" in frame.columns:
        frame["bmwp_score"] = pd.to_numeric(frame["bmwp_score"], errors="coerce")
        frame["bmwp_score"] = frame["bmwp_score"].where(frame["bmwp_score"].notna(), None)
    source_rows = len(frame)
    duplicated = frame.duplicated("nome_cientifico", keep=False)
    duplicate_taxa = int(frame.loc[duplicated, "nome_cientifico"].nunique())
    frame = frame.drop_duplicates("nome_cientifico", keep="last").copy()
    frame.attrs.update(
        {
            "source_rows": source_rows,
            "duplicate_rows_removed": source_rows - len(frame),
            "duplicate_taxa": duplicate_taxa,
            "duplicate_strategy": "keep_last",
        }
    )
    return frame


def apply_taxon_overrides(frame: pd.DataFrame, path: Path | None) -> tuple[pd.DataFrame, list[dict]]:
    if path is None:
        return frame, []
    payload = json.loads(path.read_text(encoding="utf-8"))
    applied: list[dict] = []
    for item in payload.get("taxon_overrides", []):
        match = str(item["match"]).strip()
        values = item.get("values", {})
        mask = frame["nome_cientifico"].eq(match)
        if int(mask.sum()) != 1:
            raise ValueError(f"Override taxonomico exige uma linha para {match!r}; encontradas {int(mask.sum())}.")
        for field, value in values.items():
            if field not in frame.columns:
                raise ValueError(f"Campo de override ausente no cadastro: {field}")
            frame.loc[mask, field] = clean(value)
        applied.append({"match": match, "values": values, "reason": item.get("reason", "")})
    if frame["nome_cientifico"].duplicated().any():
        duplicates = sorted(frame.loc[frame["nome_cientifico"].duplicated(False), "nome_cientifico"].unique())
        raise ValueError(f"Overrides geraram nomes cientificos duplicados: {duplicates}")
    return frame, applied


def get_engine(opyta_data_root: Path):
    sys.path.insert(0, str(opyta_data_root))
    from core.engine import get_engine as build_engine

    return build_engine()


def database_columns(connection) -> list[str]:
    rows = connection.execute(
        text(
            """
            SELECT column_name
            FROM information_schema.columns
            WHERE table_schema = 'public' AND table_name = 'especies'
            ORDER BY ordinal_position
            """
        )
    )
    return [row[0] for row in rows]


def read_database(connection, names: list[str], columns: list[str]) -> pd.DataFrame:
    selected = [column for column in columns if column in {"id_especie", *FIELD_MAP.values()}]
    query = text(
        f"SELECT {', '.join(selected)} FROM public.especies "
        "WHERE nome_cientifico IN :names ORDER BY nome_cientifico"
    ).bindparams(bindparam("names", expanding=True))
    return pd.read_sql(query, connection, params={"names": names})


def build_audit(
    source: pd.DataFrame,
    database: pd.DataFrame,
    available: set[str],
    required_fields: tuple[str, ...] = REQUIRED_ANALYTICAL_FIELDS,
):
    db_by_name = {
        row["nome_cientifico"]: row
        for row in database.to_dict("records")
    }
    actions = []
    comparisons = []
    proposed_rows = []

    for source_row in source.to_dict("records"):
        name = source_row["nome_cientifico"]
        db_row = db_by_name.get(name)
        if db_row is None:
            actions.append(
                {
                    "nome_cientifico": name,
                    "acao": "INSERT_NEW",
                    "id_especie": None,
                    "campos_a_preencher": sum(
                        clean(value) is not None
                        for field, value in source_row.items()
                        if field in available and field != "nome_cientifico"
                    ),
                }
            )
            proposed = dict(source_row)
        else:
            proposed = dict(db_row)
            fill_count = 0
            difference_count = 0
            for field, source_value in source_row.items():
                if field == "nome_cientifico" or field not in available:
                    continue
                source_value = clean(source_value)
                db_value = clean(db_row.get(field))
                if db_value is None and source_value is not None:
                    decision = "FILL_NULL"
                    proposed[field] = source_value
                    fill_count += 1
                elif source_value is not None and str(source_value) != str(db_value):
                    decision = "PRESERVE_DB_DIFFERENCE"
                    difference_count += 1
                else:
                    decision = "UNCHANGED"
                comparisons.append(
                    {
                        "nome_cientifico": name,
                        "id_especie": db_row.get("id_especie"),
                        "campo": field,
                        "valor_fonte": source_value,
                        "valor_banco": db_value,
                        "decisao": decision,
                    }
                )
            actions.append(
                {
                    "nome_cientifico": name,
                    "acao": "FILL_NULLS_ONLY" if fill_count else "PRESERVE_EXISTING",
                    "id_especie": db_row.get("id_especie"),
                    "campos_a_preencher": fill_count,
                    "diferencas_preservadas": difference_count,
                }
            )
        proposed_rows.append(proposed)

    remaining = []
    for row in proposed_rows:
        missing = [field for field in required_fields if clean(row.get(field)) is None]
        if missing:
            remaining.append({"nome_cientifico": row["nome_cientifico"], "campos_ausentes": ", ".join(missing)})
    return pd.DataFrame(actions), pd.DataFrame(comparisons), pd.DataFrame(remaining)


def apply_approved(connection, source: pd.DataFrame, actions: pd.DataFrame, available: set[str]):
    applied = []
    source_by_name = {row["nome_cientifico"]: row for row in source.to_dict("records")}
    for action in actions.to_dict("records"):
        name = action["nome_cientifico"]
        row = source_by_name[name]
        if action["acao"] == "INSERT_NEW":
            values = {field: clean(value) for field, value in row.items() if field in available}
            columns = list(values)
            query = text(
                f"INSERT INTO public.especies ({', '.join(columns)}) "
                f"VALUES ({', '.join(':' + column for column in columns)}) "
                "ON CONFLICT (nome_cientifico) DO NOTHING RETURNING id_especie"
            )
            new_id = connection.execute(query, values).scalar()
            applied.append({"nome_cientifico": name, "acao": "INSERT_NEW", "id_especie": new_id})
        elif action["acao"] == "FILL_NULLS_ONLY":
            values = {
                field: clean(value)
                for field, value in row.items()
                if field in available and field != "nome_cientifico" and clean(value) is not None
            }
            assignments = ", ".join(f"{field} = COALESCE({field}, :{field})" for field in values)
            connection.execute(
                text(f"UPDATE public.especies SET {assignments} WHERE nome_cientifico = :nome_cientifico"),
                {**values, "nome_cientifico": name},
            )
            applied.append({"nome_cientifico": name, "acao": "FILL_NULLS_ONLY", "id_especie": action["id_especie"]})
    return applied


def write_workbook(
    path: Path,
    summary: dict,
    actions: pd.DataFrame,
    comparisons: pd.DataFrame,
    remaining: pd.DataFrame,
    missing_in_source: pd.DataFrame | None = None,
    extra_in_source: pd.DataFrame | None = None,
):
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame([{"indicador": key, "valor": value} for key, value in summary.items()]).to_excel(
            writer, sheet_name="00_resumo", index=False
        )
        actions.to_excel(writer, sheet_name="01_acoes", index=False)
        comparisons.to_excel(writer, sheet_name="02_comparacao_campos", index=False)
        remaining.to_excel(writer, sheet_name="03_pendencias_pos_acao", index=False)
        if missing_in_source is not None:
            missing_in_source.to_excel(writer, sheet_name="04_resultados_sem_cadastro", index=False)
        if extra_in_source is not None:
            extra_in_source.to_excel(writer, sheet_name="05_cadastro_sem_resultados", index=False)

    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        for column in worksheet.columns:
            width = min(60, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            worksheet.column_dimensions[column[0].column_letter].width = width
    workbook.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Audita e aplica, apos aprovacao, o Gate B de especies do BRACED001.")
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--results-file", type=Path)
    parser.add_argument("--results-sheet", default="Resultados_Ictiofauna")
    parser.add_argument("--group", default="Ictiofauna")
    parser.add_argument("--output-prefix", default="gate_b_especies_braced001")
    parser.add_argument("--required-fields", nargs="*", default=list(REQUIRED_ANALYTICAL_FIELDS))
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--overrides-json", type=Path)
    parser.add_argument("--preserve-na-text", action="store_true")
    parser.add_argument("--opyta-data-root", type=Path, default=Path.cwd().parent / "Opyta_Data")
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--approved-gate-b", action="store_true")
    args = parser.parse_args()
    global PRESERVE_NA_TEXT
    PRESERVE_NA_TEXT = args.preserve_na_text
    if args.apply and not args.approved_gate_b:
        raise SystemExit("Use --approved-gate-b somente apos aprovacao explicita do Gate B.")

    source = read_source(args.source)
    source, taxon_overrides = apply_taxon_overrides(source, args.overrides_json)
    result_taxa: set[str] = set()
    result_aliases: dict[str, str] = {}
    if args.overrides_json:
        override_payload = json.loads(args.overrides_json.read_text(encoding="utf-8"))
        result_aliases = {str(key).strip(): str(value).strip() for key, value in override_payload.get("result_aliases", {}).items()}
    if args.results_file:
        results = pd.read_excel(args.results_file, sheet_name=args.results_sheet)
        if "Nome_Cientifico" not in results.columns:
            raise ValueError(f"Coluna Nome_Cientifico nao encontrada em {args.results_sheet}.")
        result_taxa = {
            result_aliases.get(str(value).strip(), str(value).strip())
            for value in results["Nome_Cientifico"].dropna()
            if str(value).strip()
        }
    source_taxa = set(source["nome_cientifico"])
    missing_taxa = sorted(result_taxa - source_taxa)
    extra_taxa = sorted(source_taxa - result_taxa) if result_taxa else []
    missing_in_source = pd.DataFrame({"nome_cientifico": missing_taxa})
    extra_in_source = pd.DataFrame({"nome_cientifico": extra_taxa})
    engine = get_engine(args.opyta_data_root)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%dT%H%M%S")

    with engine.begin() as connection:
        columns = database_columns(connection)
        database = read_database(connection, source["nome_cientifico"].tolist(), columns)
        actions, comparisons, remaining = build_audit(source, database, set(columns), tuple(args.required_fields))
        applied = apply_approved(connection, source, actions, set(columns)) if args.apply else []

    summary = {
        "projeto": "BRACED001",
        "grupo": args.group,
        "modo": "apply" if args.apply else "audit_only",
        "taxons_resultados": len(result_taxa),
        "taxons_resultados_sem_cadastro": len(missing_taxa),
        "taxons_cadastro_sem_resultados": len(extra_taxa),
        "linhas_cadastro_origem": source.attrs.get("source_rows", len(source)),
        "linhas_duplicadas_removidas": source.attrs.get("duplicate_rows_removed", 0),
        "taxons_duplicados": source.attrs.get("duplicate_taxa", 0),
        "estrategia_duplicados": source.attrs.get("duplicate_strategy", "none"),
        "especies_fonte": len(source),
        "especies_novas": int((actions["acao"] == "INSERT_NEW").sum()),
        "especies_com_nulos_a_preencher": int((actions["acao"] == "FILL_NULLS_ONLY").sum()),
        "campos_nulos_a_preencher": int(actions["campos_a_preencher"].fillna(0).sum()),
        "diferencas_preservadas": int((comparisons["decisao"] == "PRESERVE_DB_DIFFERENCE").sum()),
        "pendencias_obrigatorias_pos_acao": len(remaining),
        "registros_aplicados": len(applied),
        "normalizacoes_taxonomicas": len(taxon_overrides),
        "aliases_resultados": len(result_aliases),
    }
    base = args.output_dir / f"{args.output_prefix}_{stamp}"
    write_workbook(base.with_suffix(".xlsx"), summary, actions, comparisons, remaining, missing_in_source, extra_in_source)
    base.with_suffix(".json").write_text(
        json.dumps(
            {
                "summary": summary,
                "actions": actions.to_dict("records"),
                "resultados_sem_cadastro": missing_taxa,
                "cadastro_sem_resultados": extra_taxa,
                "applied": applied,
                "taxon_overrides": taxon_overrides,
                "result_aliases": result_aliases,
            },
            ensure_ascii=False,
            indent=2,
            default=str,
        ),
        encoding="utf-8",
    )
    print(json.dumps({"workbook": str(base.with_suffix('.xlsx')), "json": str(base.with_suffix('.json')), **summary}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
