"""Generate the WSPKIN001 phytoplankton taxon-registration workbook."""

from __future__ import annotations

import os
from collections import Counter
from pathlib import Path

import openpyxl
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from supabase import create_client


SOURCE = Path(
    r"G:\Meu Drive\Opyta\Clientes\Clientes\Clientes\WSP\BAndeirinhas_Kinross\Migração de dados\Resultados_Migração _Fito_wsp.xlsx"
)
OUTPUT = SOURCE.parent / "Cadastro_Especies_WSPKIN001_Fitoplancton.xlsx"
GROUP = "Fitoplâncton"
FIELDS = [
    "Nome_Cientifico",
    "Grupo_Biologico",
    "Situacao_Cadastro",
    "Reino",
    "Filo",
    "Classe",
    "Ordem",
    "Familia",
    "Genero",
    "Autor_e_Ano",
    "Fonte_Taxonomica",
    "Observacao",
    "Status_Preenchimento",
    "Registros_na_Fonte",
]


def style_sheet(sheet) -> None:
    header = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 34
    for column in sheet.columns:
        letter = column[0].column_letter
        width = max(len(str(cell.value or "")) for cell in column) + 2
        sheet.column_dimensions[letter].width = min(max(width, 16), 34)


def source_taxa() -> Counter[str]:
    workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
    sheet = workbook["Resultados_Fitoplancton"]
    return Counter(
        row[7]
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[7]
    )


def registered_taxa(names: list[str]) -> dict[str, dict]:
    load_dotenv()
    client = create_client(os.environ["SUPABASE_URL"], os.environ["SUPABASE_SERVICE_ROLE_KEY"])
    rows: list[dict] = []
    for start in range(0, len(names), 50):
        rows.extend(
            client.table("especies")
            .select("nome_cientifico,reino,filo,classe,ordem,familia,genero,autor_e_ano")
            .in_("nome_cientifico", names[start : start + 50])
            .execute()
            .data
        )
    return {row["nome_cientifico"]: row for row in rows}


def add_rows(sheet, rows: list[list[object]]) -> None:
    sheet.append(FIELDS)
    for row in rows:
        sheet.append(row)
    style_sheet(sheet)


def main() -> None:
    if OUTPUT.exists():
        raise FileExistsError(f"Refusing to overwrite existing workbook: {OUTPUT}")

    counts = source_taxa()
    names = sorted(counts)
    registered = registered_taxa(names)
    new_taxa = sorted(set(names) - set(registered))
    incomplete = sorted(
        name
        for name, row in registered.items()
        if any(not row.get(field) for field in ("reino", "filo", "classe", "ordem", "familia", "genero"))
    )

    workbook = Workbook()
    new_sheet = workbook.active
    new_sheet.title = "Taxons_Novos"
    add_rows(
        new_sheet,
        [
            [name, GROUP, "novo_no_supabase", "", "", "", "", "", "", "", "", "", "Pendente", counts[name]]
            for name in new_taxa
        ],
    )

    incomplete_sheet = workbook.create_sheet("Complementar_Cadastro")
    add_rows(
        incomplete_sheet,
        [
            [
                name,
                GROUP,
                "completar_cadastro_existente",
                row.get("reino") or "",
                row.get("filo") or "",
                row.get("classe") or "",
                row.get("ordem") or "",
                row.get("familia") or "",
                row.get("genero") or "",
                row.get("autor_e_ano") or "",
                "",
                "Completar os campos taxonômicos vazios.",
                "Pendente",
                counts[name],
            ]
            for name in incomplete
            for row in [registered[name]]
        ],
    )

    notes = workbook.create_sheet("Orientacoes")
    notes.append(["WSPKIN001 — Fitoplâncton — Pendências de Cadastro"])
    notes.append(["Preencha ou corrija somente os campos taxonômicos e a fonte. Não altere Nome_Cientifico ou Situacao_Cadastro."])
    notes.append(["Após o preenchimento, devolver esta planilha para revalidação e aplicação no Supabase antes do Gate B."])
    notes.column_dimensions["A"].width = 120
    notes["A1"].font = Font(bold=True, color="FFFFFF")
    notes["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    notes["A2"].alignment = Alignment(wrap_text=True)
    notes["A3"].alignment = Alignment(wrap_text=True)

    workbook.save(OUTPUT)
    print(f"output={OUTPUT}")
    print(f"new_taxa={len(new_taxa)} incomplete_taxa={len(incomplete)}")


if __name__ == "__main__":
    main()
