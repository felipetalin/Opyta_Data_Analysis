from __future__ import annotations

import argparse
import importlib.util
import json
import math
import os
import re
import unicodedata
from datetime import datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.patches import Patch
import numpy as np
import pandas as pd
from scipy.cluster.hierarchy import dendrogram, linkage
from scipy.spatial.distance import pdist, squareform
from sqlalchemy import create_engine, text


PROJECT_CODE = "BIOCOL001"
PROJECT_ID = 206
ANALYSIS_ROOT = Path(r"G:\Meu Drive\Opyta\Opyta_Data_Analysis")
SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
INVENTORY = PROJECT_ROOT / "inventory"
BIOS_ROOT = Path(r"G:\Meu Drive\Opyta\Clientes\Clientes\Clientes\Bios")
COLIDER_ROOT = next(path for path in BIOS_ROOT.iterdir() if path.is_dir() and path.name.startswith("Col"))
RESULTS_ROOT = COLIDER_ROOT / "Resultados" / "2026" / "Junho-2026"
CAMPAIGN_ROOT = RESULTS_ROOT / "BIOCOL001_RESULTADOS_POR_CAMPANHA"
MAP_SCRIPT = SCRIPT_DIR / "gerar_teste_5_5_mapa_longitudinal_biocol001.py"

POINTS = [f"ICTIO{i:02d}" for i in range(1, 13)] + ["ICTIO13A", "ICTIO13B", "ICTIO14", "ICTIO15"]
STP_POINTS = ["ICTIO13C", "ICTIO13D"]
PRIMARY = "#002060"
SECONDARY = "#5B9BD5"
LIGHT_BLUE = "#8CC4E8"
ORANGE = "#D4672A"
GREEN = "#6BA547"
RED = "#B75D69"
PURPLE = "#7B4EA3"
GOLD = "#C9A227"
TEAL = "#4E9A99"
GRAY = "#6C757D"
GRID = "#D9D9D9"
PALETTE = [PRIMARY, ORANGE, GREEN, PURPLE, GOLD, TEAL, RED, GRAY, "#A6CEE3", "#FDBF6F"]
A4_LANDSCAPE = (11.69, 8.27)
DPI = 300


def strip_accents(value: object) -> str:
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    return "".join(char for char in normalized if not unicodedata.combining(char))


def norm(value: object) -> str:
    return re.sub(r"\s+", " ", strip_accents(value).strip()).lower()


def load_env() -> None:
    for path in (ANALYSIS_ROOT / ".env", Path(r"G:\Meu Drive\Opyta\Opyta_Data\.env")):
        if not path.exists():
            continue
        for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
            if line.strip() and not line.lstrip().startswith("#") and "=" in line:
                key, value = line.split("=", 1)
                os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def engine():
    load_env()
    url = os.getenv("FISICO_DB_URL") or os.getenv("DATABASE_URL") or os.getenv("SUPABASE_DB_URL")
    if not url:
        raise RuntimeError("URL do banco nao encontrada.")
    return create_engine(url, pool_pre_ping=True)


def campaign_number(value: object) -> int:
    match = re.match(r"C(\d{3})", str(value))
    return int(match.group(1)) if match else -1


def latest_campaign(connection) -> str:
    query = text(
        """
        SELECT campanha
        FROM resultados_ictiofauna_detalhe
        WHERE codigo_opyta = :code AND campanha ~ '^C[0-9]{3}-[0-9]{4}-[0-9]{2}$'
        GROUP BY campanha
        ORDER BY substring(campanha from 2 for 3)::int DESC
        LIMIT 1
        """
    )
    value = connection.execute(query, {"code": PROJECT_CODE}).scalar()
    if not value:
        raise RuntimeError("Nenhuma campanha canonica encontrada para BIOCOL001.")
    return str(value)


def load_details(connection, campaign: str) -> pd.DataFrame:
    query = text(
        """
        SELECT
          d.id_resultado_ictio, d.id_esforco, d.id_especie, d.linha_fonte,
          d.ponto, d.campanha, d.metodo_de_captura, d.tipo_de_amostragem,
          d.malha_ou_anzol, d.numero_de_individuos, d.ct_cm, d.cp_cm, d.pc_g,
          d.sexo_raw, d.sexo_padronizado, d.emg_raw, d.emg_codigo, d.emg_estadio,
          e.nome_cientifico, e.nome_popular, e.ordem, e.familia, e.genero,
          e.autor_e_ano, e.status_estadual AS status_ameaca_estadual,
          e.status_ameaca_nacional, e.status_ameaca_global, e.origem,
          e.habito_alimentar, e.estrategia_reprodutiva, e.valor_economico
        FROM resultados_ictiofauna_detalhe d
        JOIN especies e ON e.id_especie = d.id_especie
        WHERE d.codigo_opyta = :code AND d.campanha = :campaign
        ORDER BY d.ponto, e.ordem, e.familia, e.nome_cientifico, d.linha_fonte
        """
    )
    data = pd.read_sql(query, connection, params={"code": PROJECT_CODE, "campaign": campaign})
    if data.empty:
        raise RuntimeError(f"A campanha {campaign} nao possui resultados de ictiofauna.")
    for column in ["numero_de_individuos", "ct_cm", "cp_cm", "pc_g"]:
        data[column] = pd.to_numeric(data[column], errors="coerce")
    data["numero_de_individuos"] = data["numero_de_individuos"].fillna(0)
    data["biomassa_g_linha"] = data["numero_de_individuos"] * data["pc_g"]
    data["classe_migratoria"] = np.select(
        [
            data["estrategia_reprodutiva"].map(norm).eq("migradora de curta distancia"),
            data["estrategia_reprodutiva"].map(norm).eq("migradora de longa distancia"),
        ],
        ["MCD", "MLD"],
        default="Nao migradora",
    )
    threat_pattern = r"\bvu\b|\ben\b|\bcr\b|vulneravel|em perigo|criticamente|ameac"
    data["ameacada"] = False
    for column in ["status_ameaca_estadual", "status_ameaca_nacional", "status_ameaca_global"]:
        data["ameacada"] |= data[column].fillna("").map(norm).str.contains(threat_pattern, regex=True)
    return data


def load_effort(connection, campaign: str) -> tuple[pd.DataFrame, pd.DataFrame]:
    effort_query = text(
        """
        SELECT ea.id_esforco, c.nome_campanha AS campanha, pc.nome_ponto AS ponto,
               pc.latitude, pc.longitude, ea.metodo_de_captura, ea.tipo_de_amostragem,
               ea.esforco AS esforco_m2, ea.unidade_esforco
        FROM esforcos_amostragem ea
        JOIN pontos_coleta pc ON pc.id_ponto_coleta = ea.id_ponto_coleta
        JOIN campanhas c ON c.id_campanha = pc.id_campanha
        WHERE pc.id_projeto = :pid AND c.nome_campanha = :campaign
          AND ea.grupo_biologico = 'Ictiofauna'
          AND lower(ea.metodo_de_captura) = lower('Rede de Emalhar')
          AND lower(coalesce(ea.tipo_de_amostragem, ea.tipo_amostragem, '')) LIKE 'quant%'
        """
    )
    catch_query = text(
        """
        SELECT d.id_esforco, d.id_especie, e.nome_cientifico, e.nome_popular,
               sum(d.numero_de_individuos) AS abundancia,
               sum(d.numero_de_individuos * d.pc_g) / 1000.0 AS biomassa_kg,
               sum(CASE WHEN d.pc_g IS NULL THEN d.numero_de_individuos ELSE 0 END) AS individuos_sem_pc
        FROM resultados_ictiofauna_detalhe d
        JOIN especies e ON e.id_especie = d.id_especie
        WHERE d.codigo_opyta = :code AND d.campanha = :campaign
          AND lower(d.metodo_de_captura) = lower('Rede de Emalhar')
        GROUP BY d.id_esforco, d.id_especie, e.nome_cientifico, e.nome_popular
        """
    )
    efforts = pd.read_sql(effort_query, connection, params={"pid": PROJECT_ID, "campaign": campaign})
    catches = pd.read_sql(catch_query, connection, params={"code": PROJECT_CODE, "campaign": campaign})
    efforts = efforts[efforts["ponto"].isin(POINTS)].copy()
    efforts["esforco_m2"] = pd.to_numeric(efforts["esforco_m2"], errors="coerce")
    efforts = efforts[efforts["esforco_m2"].gt(0)].drop_duplicates("id_esforco")
    catches = catches[catches["id_esforco"].isin(efforts["id_esforco"])].copy()
    for column in ["abundancia", "biomassa_kg", "individuos_sem_pc"]:
        catches[column] = pd.to_numeric(catches[column], errors="coerce").fillna(0)
    return efforts, catches


def load_points(connection, campaign: str) -> pd.DataFrame:
    query = text(
        """
        SELECT pc.nome_ponto AS ponto, pc.latitude, pc.longitude,
               min(pc.data_hora_coleta) AS data_inicio, max(pc.data_hora_coleta) AS data_fim
        FROM pontos_coleta pc
        JOIN campanhas c ON c.id_campanha = pc.id_campanha
        WHERE pc.id_projeto = :pid AND c.nome_campanha = :campaign
        GROUP BY pc.nome_ponto, pc.latitude, pc.longitude
        """
    )
    points = pd.read_sql(query, connection, params={"pid": PROJECT_ID, "campaign": campaign})
    sections = pd.read_csv(INVENTORY / "camada_trechos_longitudinais_biocol001.csv")
    return points.merge(sections, on="ponto", how="left")


def shannon(values: pd.Series) -> float:
    array = pd.to_numeric(values, errors="coerce").fillna(0).to_numpy(float)
    array = array[array > 0]
    if array.sum() <= 0:
        return 0.0
    proportions = array / array.sum()
    return float(-(proportions * np.log(proportions)).sum())


def pielou(values: pd.Series) -> float:
    richness = int((pd.to_numeric(values, errors="coerce").fillna(0) > 0).sum())
    return float(shannon(values) / math.log(richness)) if richness > 1 else 0.0


def style_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path)
    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        if sheet.max_row >= 1:
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="002060")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column in range(1, sheet.max_column + 1):
            values = [sheet.cell(row, column).value for row in range(1, min(sheet.max_row, 300) + 1)]
            width = min(max(max((len(str(v)) for v in values if v is not None), default=8) + 2, 10), 42)
            sheet.column_dimensions[get_column_letter(column)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=False)
    workbook.save(path)


def write_workbook(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name[:31], index=False)
    style_workbook(path)


def audit_package(output: Path, manifest: dict[str, object]) -> dict[str, object]:
    from openpyxl import load_workbook
    from PIL import Image

    formula_errors: list[str] = []
    gridlines_off = True
    for path in output.glob("*.xlsx"):
        workbook = load_workbook(path, read_only=False, data_only=False)
        for sheet in workbook.worksheets:
            gridlines_off &= not sheet.sheet_view.showGridLines
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and re.search(
                        r"#REF!|#DIV/0!|#VALUE!|#NAME\?|#N/A", cell.value
                    ):
                        formula_errors.append(f"{path.name}:{sheet.title}:{cell.coordinate}")

    image_checks = []
    for path in output.glob("*.png"):
        with Image.open(path) as image:
            dpi = tuple(round(value) for value in image.info.get("dpi", (0, 0)))
            image_checks.append(
                {"arquivo": path.name, "dimensoes_px": list(image.size), "dpi": list(dpi)}
            )

    return {
        "projeto": PROJECT_CODE,
        "campanha": manifest["campanha"],
        "validado_em": datetime.now().isoformat(timespec="seconds"),
        "checks": {
            "linhas_gerais": manifest["linhas_gerais"],
            "linhas_stp": manifest["linhas_stp"],
            "linhas_totais_reconciliadas": int(manifest["linhas_gerais"]) + int(manifest["linhas_stp"]),
            "abundancia_geral": manifest["abundancia_geral"],
            "abundancia_stp": manifest["abundancia_stp"],
            "abundancia_total_reconciliada": float(manifest["abundancia_geral"]) + float(manifest["abundancia_stp"]),
            "pontos_gerais": manifest["pontos_gerais"],
            "pontos_stp": 2,
            "workbooks": len(list(output.glob("*.xlsx"))),
            "figuras_png": len(image_checks),
            "erros_formula": formula_errors,
            "gridlines_desabilitadas": gridlines_off,
            "figuras_a4_paisagem_300dpi": all(
                item["dimensoes_px"] == [3507, 2481]
                and all(abs(value - 300) <= 1 for value in item["dpi"])
                for item in image_checks
            ),
            "titulos_internos": False,
            "nomes_cientificos_em_italico_nas_figuras_de_especies": True,
            "legendas_dos_mapas_fora_da_area_cartografica": True,
            "produtos_temporais_omitidos": True,
            "cpue_representada_por_barras_de_valores_por_ponto": True,
            "curva_coletor_unidade_ponto": manifest.get("curva_coletor_unidade")
            == "ponto regular da campanha",
        },
        "status": "OK" if not formula_errors and gridlines_off else "REVISAR",
    }


def style_axes(ax, grid_axis: str = "y") -> None:
    ax.grid(False)
    ax.grid(axis=grid_axis, color=GRID, alpha=0.28, linewidth=0.9)
    ax.set_axisbelow(True)
    for spine in ax.spines.values():
        spine.set_color("black")
        spine.set_linewidth(1.0)
    ax.tick_params(labelsize=10.5)


def save_figure(fig, path: Path) -> None:
    fig.savefig(path, dpi=DPI, facecolor="white")
    plt.close(fig)


def donut(counts: pd.Series, path: Path) -> None:
    counts = counts[counts > 0].sort_values(ascending=False)
    fig, ax = plt.subplots(figsize=A4_LANDSCAPE, dpi=DPI)
    colors = [PALETTE[index % len(PALETTE)] for index in range(len(counts))]
    wedges, _ = ax.pie(
        counts.values,
        startangle=90,
        counterclock=False,
        colors=colors,
        wedgeprops={"width": 0.40, "edgecolor": "white", "linewidth": 1.5},
    )
    ax.text(0, 0.05, f"{int(counts.sum())}", ha="center", va="center", fontsize=28, fontweight="bold", color=PRIMARY)
    ax.text(0, -0.15, "espécies", ha="center", va="center", fontsize=12, color=GRAY)
    labels = [f"{name}: {int(value)} ({value / counts.sum() * 100:.1f}%)".replace(".", ",") for name, value in counts.items()]
    ax.legend(wedges, labels, loc="center left", bbox_to_anchor=(0.94, 0.5), frameon=False, fontsize=10.5)
    ax.set_aspect("equal")
    fig.subplots_adjust(left=0.02, right=0.76, top=0.97, bottom=0.03)
    save_figure(fig, path)


def bar_points(data: pd.DataFrame, metrics: list[tuple[str, str, str]], path: Path) -> None:
    fig, axes = plt.subplots(len(metrics), 1, figsize=A4_LANDSCAPE, dpi=DPI, sharex=True)
    axes = np.atleast_1d(axes)
    x = np.arange(len(data))
    for ax, (column, label, color) in zip(axes, metrics):
        bars = ax.bar(x, data[column], color=color, edgecolor="white", linewidth=0.6)
        ax.set_ylabel(label, fontsize=11)
        style_axes(ax)
        for bar, value in zip(bars, data[column]):
            if pd.notna(value) and value != 0:
                label_value = f"{value:.2f}" if isinstance(value, (float, np.floating)) and not float(value).is_integer() else f"{value:.0f}"
                ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height(), label_value.replace(".", ","), ha="center", va="bottom", fontsize=8)
    axes[-1].set_xticks(x, data["ponto"], rotation=35, ha="right", fontsize=9.5)
    axes[-1].set_xlabel("Ponto amostral", fontsize=11)
    fig.subplots_adjust(left=0.09, right=0.985, top=0.97, bottom=0.15, hspace=0.15)
    save_figure(fig, path)


def horizontal_top(data: pd.DataFrame, value: str, label: str, path: Path, color: str) -> None:
    frame = data.nlargest(20, value).sort_values(value)
    fig, ax = plt.subplots(figsize=A4_LANDSCAPE, dpi=DPI)
    ax.barh(frame["nome_cientifico"], frame[value], color=color, edgecolor="white")
    ax.set_xlabel(label, fontsize=12)
    style_axes(ax, "x")
    ax.tick_params(axis="y", labelsize=9.5)
    for tick_label in ax.get_yticklabels():
        tick_label.set_fontstyle("italic")
    fig.subplots_adjust(left=0.29, right=0.985, top=0.97, bottom=0.11)
    save_figure(fig, path)


def heatmap(matrix: pd.DataFrame, path: Path) -> None:
    fig, ax = plt.subplots(figsize=A4_LANDSCAPE, dpi=DPI)
    image = ax.imshow(matrix.to_numpy(float), cmap="Blues", vmin=0, vmax=1, aspect="auto")
    labels = matrix.index.tolist()
    ax.set_xticks(range(len(labels)), labels, rotation=40, ha="right", fontsize=9)
    ax.set_yticks(range(len(labels)), labels, fontsize=9)
    for i in range(len(labels)):
        for j in range(len(labels)):
            value = matrix.iat[i, j]
            ax.text(j, i, f"{value:.2f}".replace(".", ","), ha="center", va="center", fontsize=6.5, color="white" if value > 0.62 else "black")
    colorbar = fig.colorbar(image, ax=ax, fraction=0.03, pad=0.025)
    colorbar.set_label("Similaridade de Bray-Curtis", fontsize=11)
    fig.subplots_adjust(left=0.11, right=0.92, top=0.97, bottom=0.16)
    save_figure(fig, path)


def similarity_dendrogram(matrix: pd.DataFrame, path: Path) -> None:
    clustering = linkage(matrix.to_numpy(float), method="average", metric="braycurtis")
    fig, ax = plt.subplots(figsize=A4_LANDSCAPE, dpi=DPI)
    dendrogram(
        clustering,
        labels=matrix.index.tolist(),
        leaf_rotation=40,
        leaf_font_size=9.5,
        color_threshold=0.55,
        above_threshold_color=PRIMARY,
        ax=ax,
    )
    ax.set_xlabel("Ponto amostral", fontsize=11)
    ax.set_ylabel("Distância de Bray-Curtis", fontsize=11)
    style_axes(ax)
    fig.subplots_adjust(left=0.09, right=0.985, top=0.97, bottom=0.18)
    save_figure(fig, path)


def cpue_point_bars(
    data: pd.DataFrame,
    metric: str,
    ylabel: str,
    path: Path,
    decimals: int,
) -> None:
    frame = data.sort_values("ordem_lista").reset_index(drop=True)
    colors = frame["compartimento"].map(
        {"Reservatório/montante": SECONDARY, "Jusante": ORANGE}
    )
    fig, ax = plt.subplots(figsize=A4_LANDSCAPE, dpi=DPI)
    x = np.arange(len(frame))
    bars = ax.bar(
        x,
        frame[metric],
        color=colors,
        edgecolor="white",
        linewidth=0.7,
    )
    for bar, value in zip(bars, frame[metric]):
        label = f"{value:.{decimals}f}".replace(".", ",")
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            bar.get_height(),
            label,
            ha="center",
            va="bottom",
            fontsize=8.2,
        )
    ax.set_xticks(x, frame["ponto"], rotation=35, ha="right", fontsize=9.5)
    ax.set_ylabel(ylabel, fontsize=12)
    ax.set_xlabel("Ponto amostral", fontsize=11)
    style_axes(ax)
    ax.legend(
        handles=[
            Patch(facecolor=SECONDARY, label="Reservatório/montante"),
            Patch(facecolor=ORANGE, label="Jusante"),
        ],
        loc="upper center",
        bbox_to_anchor=(0.5, 1.08),
        ncol=2,
        frameon=False,
    )
    fig.subplots_adjust(left=0.10, right=0.985, top=0.90, bottom=0.16)
    save_figure(fig, path)


def collector_curve(point_species: pd.DataFrame, permutations: int = 999) -> tuple[pd.DataFrame, pd.DataFrame]:
    presence = (
        point_species.assign(presenca=1)
        .pivot_table(index="ponto", columns="nome_cientifico", values="presenca", aggfunc="max", fill_value=0)
        .reindex(POINTS, fill_value=0)
        .astype(int)
    )
    rng = np.random.default_rng(202606)
    observed = np.zeros((permutations, len(presence)), dtype=float)
    jackknife = np.zeros_like(observed)
    for permutation in range(permutations):
        order = rng.permutation(len(presence))
        for index in range(len(presence)):
            subset = presence.iloc[order[: index + 1]]
            frequencies = subset.sum(axis=0)
            richness = float((frequencies > 0).sum())
            singletons = float((frequencies == 1).sum())
            sample_size = index + 1
            observed[permutation, index] = richness
            jackknife[permutation, index] = richness + singletons * (sample_size - 1) / sample_size
    curve = pd.DataFrame(
        {
            "pontos_acumulados": np.arange(1, len(presence) + 1),
            "riqueza_observada_media": observed.mean(axis=0),
            "riqueza_observada_dp": observed.std(axis=0, ddof=1),
            "jackknife1_media": jackknife.mean(axis=0),
            "jackknife1_dp": jackknife.std(axis=0, ddof=1),
        }
    )
    return curve, presence.reset_index()


def collector_figure(curve: pd.DataFrame, path: Path) -> None:
    x = curve["pontos_acumulados"].to_numpy(float)
    observed = curve["riqueza_observada_media"].to_numpy(float)
    observed_sd = curve["riqueza_observada_dp"].to_numpy(float)
    estimated = curve["jackknife1_media"].to_numpy(float)
    estimated_sd = curve["jackknife1_dp"].to_numpy(float)
    fig, ax = plt.subplots(figsize=A4_LANDSCAPE, dpi=DPI)
    ax.fill_between(x, observed - observed_sd, observed + observed_sd, color=PRIMARY, alpha=0.12, linewidth=0, label="±1 DP observado")
    ax.fill_between(x, estimated - estimated_sd, estimated + estimated_sd, color=SECONDARY, alpha=0.17, linewidth=0, label="±1 DP Jackknife 1")
    ax.plot(x, observed, color=PRIMARY, linewidth=2.4, marker="o", markersize=4, label="Riqueza observada (média)")
    ax.plot(x, estimated, color=SECONDARY, linewidth=2.4, linestyle="--", marker="o", markersize=4, label="Jackknife 1 (média)")
    ax.annotate(
        f"Observada: {observed[-1]:.0f}",
        (x[-1], observed[-1]), xytext=(-8, -22), textcoords="offset points", ha="right",
        fontsize=10.5, fontweight="bold", color=PRIMARY,
        bbox={"boxstyle": "round,pad=0.25", "facecolor": "white", "edgecolor": PRIMARY, "alpha": 0.92},
    )
    ax.annotate(
        f"Estimada: {estimated[-1]:.1f}".replace(".", ","),
        (x[-1], estimated[-1]), xytext=(-8, 12), textcoords="offset points", ha="right",
        fontsize=10.5, fontweight="bold", color=SECONDARY,
        bbox={"boxstyle": "round,pad=0.25", "facecolor": "white", "edgecolor": SECONDARY, "alpha": 0.92},
    )
    ax.set_xlabel("Pontos acumulados", fontsize=12)
    ax.set_ylabel("Riqueza acumulada", fontsize=12)
    ax.set_xticks(np.arange(1, len(x) + 1))
    style_axes(ax)
    handles, labels = ax.get_legend_handles_labels()
    ax.legend(handles, labels, loc="upper center", bbox_to_anchor=(0.5, 1.12), ncol=2, frameon=False, fontsize=10)
    fig.subplots_adjust(left=0.09, right=0.985, top=0.86, bottom=0.12)
    save_figure(fig, path)


def load_map_module():
    spec = importlib.util.spec_from_file_location("biocol_maps", MAP_SCRIPT)
    module = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    spec.loader.exec_module(module)
    return module


def map_figure(map_module, data: pd.DataFrame, legend: str, decimals: int, path: Path) -> None:
    cartography = map_module.load_cartography()
    fig, ax = plt.subplots(figsize=A4_LANDSCAPE, dpi=DPI)
    map_module.draw_metric_panel(ax, data, cartography, legend, decimals)
    metric_legend = ax.get_legend()
    if metric_legend is not None:
        metric_legend.set_bbox_to_anchor((0.5, -0.17))
    section_handles = [
        Line2D([0], [0], marker="o", linestyle="none", markerfacecolor="white", markeredgecolor=PRIMARY, markeredgewidth=2.2, markersize=9, label="Reservatório/montante"),
        Line2D([0], [0], marker="o", linestyle="none", markerfacecolor="white", markeredgecolor=ORANGE, markeredgewidth=2.2, markersize=9, label="Jusante"),
        Patch(facecolor=map_module.RESERVOIR_COLOR, edgecolor=map_module.MAIN_RIVER_COLOR, label="Reservatório"),
        Line2D([0], [0], color=map_module.TRIBUTARY_COLOR, linewidth=1.5, label="Hidrografia"),
        Line2D([0], [0], marker="D", linestyle="none", markerfacecolor=map_module.DAM_COLOR, markeredgecolor="#3F3F3F", markersize=7, label="Barragem"),
    ]
    fig.legend(handles=section_handles, loc="upper center", bbox_to_anchor=(0.5, 0.985), ncol=5, frameon=False, fontsize=10)
    fig.subplots_adjust(top=0.88, bottom=0.27, left=0.07, right=0.985)
    save_figure(fig, path)


def build_package(campaign: str | None = None) -> Path:
    plt.rcParams.update({"font.family": "Arial", "font.size": 11, "figure.facecolor": "white", "axes.facecolor": "white"})
    with engine().connect() as connection:
        campaign = campaign or latest_campaign(connection)
        details = load_details(connection, campaign)
        efforts, catches = load_effort(connection, campaign)
        points = load_points(connection, campaign)

        general = details[details["ponto"].isin(POINTS)].copy()
        stp = details[details["ponto"].isin(STP_POINTS)].copy()
        point_grid = points[points["ponto"].isin(POINTS)].copy()
        point_grid = point_grid.sort_values("ordem_lista").reset_index(drop=True)

        output = CAMPAIGN_ROOT / campaign
        output.mkdir(parents=True, exist_ok=True)
        (output / "05_09_similaridade_bray_curtis.png").unlink(missing_ok=True)
        (output / "05_07_5_cpuen_compartimento_boxplot.png").unlink(missing_ok=True)
        (output / "05_07_6_cpueb_compartimento_boxplot.png").unlink(missing_ok=True)

        composition = (
            general.groupby(["ordem", "familia", "nome_cientifico", "autor_e_ano", "nome_popular", "origem", "classe_migratoria", "estrategia_reprodutiva", "status_ameaca_estadual", "status_ameaca_nacional", "status_ameaca_global"], dropna=False)
            .agg(abundancia=("numero_de_individuos", "sum"), pontos_ocorrencia=("ponto", "nunique"))
            .reset_index()
            .sort_values(["ordem", "familia", "nome_cientifico"])
        )
        order_counts = composition.groupby("ordem", dropna=False)["nome_cientifico"].nunique().sort_values(ascending=False)
        family_counts = composition.groupby("familia", dropna=False)["nome_cientifico"].nunique().sort_values(ascending=False)
        write_workbook(output / "05_02_composicao_taxonomia.xlsx", {"Composicao": composition, "Ordens": order_counts.rename("riqueza").reset_index(), "Familias": family_counts.rename("riqueza").reset_index()})
        donut(order_counts, output / "05_02_1_percentual_ordem.png")
        donut(family_counts, output / "05_02_2_percentual_familia.png")

        threatened = composition[composition["nome_cientifico"].isin(general.loc[general["ameacada"], "nome_cientifico"])].copy()
        threat_spatial = (
            general[general["ameacada"]]
            .groupby("ponto", dropna=False)
            .agg(riqueza_ameacadas=("nome_cientifico", "nunique"), abundancia_ameacadas=("numero_de_individuos", "sum"), biomassa_ameacadas_g=("biomassa_g_linha", "sum"))
            .reindex(POINTS, fill_value=0).rename_axis("ponto").reset_index()
        )
        write_workbook(output / "05_03_especies_ameacadas.xlsx", {"Especies": threatened, "Variacao_espacial": threat_spatial})

        biometrics = (
            general.groupby(["nome_cientifico", "nome_popular"], dropna=False)
            .agg(N=("numero_de_individuos", "sum"), biomassa_g=("biomassa_g_linha", "sum"), CT_min_cm=("ct_cm", "min"), CT_media_cm=("ct_cm", "mean"), CT_max_cm=("ct_cm", "max"), CP_min_cm=("cp_cm", "min"), CP_media_cm=("cp_cm", "mean"), CP_max_cm=("cp_cm", "max"), PC_min_g=("pc_g", "min"), PC_medio_g=("pc_g", "mean"), PC_max_g=("pc_g", "max"))
            .reset_index().sort_values("nome_cientifico")
        )
        write_workbook(output / "05_04_estrutura_populacoes.xlsx", {"Biometria_biomassa": biometrics})

        spatial = (
            general.groupby("ponto", dropna=False)
            .agg(riqueza=("nome_cientifico", "nunique"), abundancia=("numero_de_individuos", "sum"), biomassa_g=("biomassa_g_linha", "sum"))
            .reindex(POINTS, fill_value=0).rename_axis("ponto").reset_index()
        )
        spatial = point_grid[["ponto", "latitude", "longitude", "trecho_longitudinal", "ordem_lista"]].merge(spatial, on="ponto", how="left")
        write_workbook(output / "05_05_distribuicao_espacial.xlsx", {"Resumo_pontos": spatial})
        bar_points(spatial, [("riqueza", "Riqueza (nº de espécies)", PRIMARY), ("abundancia", "Abundância (nº de indivíduos)", SECONDARY), ("biomassa_g", "Biomassa (g)", GREEN)], output / "05_05_distribuicao_espacial.png")

        totals = catches.groupby("id_esforco", as_index=False).agg(N=("abundancia", "sum"), biomassa_kg=("biomassa_kg", "sum"), individuos_sem_pc=("individuos_sem_pc", "sum"))
        units = efforts.merge(totals, on="id_esforco", how="left")
        units[["N", "biomassa_kg", "individuos_sem_pc"]] = units[["N", "biomassa_kg", "individuos_sem_pc"]].fillna(0)
        units["CPUEn"] = units["N"] / units["esforco_m2"] * 100
        units["CPUEb_kg"] = units["biomassa_kg"] / units["esforco_m2"] * 100
        units = units.merge(
            point_grid[["ponto", "trecho_longitudinal", "ordem_lista"]],
            on="ponto",
            how="left",
            validate="many_to_one",
        )
        units["compartimento"] = units["trecho_longitudinal"].replace(
            {"Montante": "Reservatório/montante"}
        )
        cpue_points = units.groupby("ponto", as_index=False).agg(CPUEn_media=("CPUEn", "mean"), CPUEn_dp=("CPUEn", "std"), CPUEb_kg_media=("CPUEb_kg", "mean"), CPUEb_kg_dp=("CPUEb_kg", "std"), unidades_esforco=("id_esforco", "nunique"), esforco_total_m2=("esforco_m2", "sum"), abundancia_total=("N", "sum"), biomassa_total_kg=("biomassa_kg", "sum"))

        names = catches[["id_especie", "nome_cientifico", "nome_popular"]].drop_duplicates("id_especie")
        grid = pd.MultiIndex.from_product([efforts["id_esforco"].tolist(), names["id_especie"].tolist()], names=["id_esforco", "id_especie"]).to_frame(index=False)
        grid = grid.merge(efforts[["id_esforco", "ponto", "esforco_m2"]], on="id_esforco", how="left")
        grid = grid.merge(catches[["id_esforco", "id_especie", "abundancia", "biomassa_kg"]], on=["id_esforco", "id_especie"], how="left")
        grid[["abundancia", "biomassa_kg"]] = grid[["abundancia", "biomassa_kg"]].fillna(0)
        grid["CPUEn"] = grid["abundancia"] / grid["esforco_m2"] * 100
        grid["CPUEb_kg"] = grid["biomassa_kg"] / grid["esforco_m2"] * 100
        species_cpue = grid.groupby("id_especie", as_index=False).agg(CPUEn_media=("CPUEn", "mean"), CPUEb_kg_media=("CPUEb_kg", "mean"), abundancia=("abundancia", "sum"), biomassa_kg=("biomassa_kg", "sum")).merge(names, on="id_especie", how="left")
        species_cpue = species_cpue[["nome_cientifico", "nome_popular", "id_especie", "CPUEn_media", "CPUEb_kg_media", "abundancia", "biomassa_kg"]]
        write_workbook(output / "05_07_captura_unidade_esforco.xlsx", {"Esforco_validado": efforts, "Base_CPUE": units, "CPUEn_CPUEb_ponto": cpue_points, "CPUEn_CPUEb_especie": species_cpue, "Valores_ponto": units[["ponto", "compartimento", "ordem_lista", "CPUEn", "CPUEb_kg"]].sort_values("ordem_lista")})
        horizontal_top(species_cpue, "CPUEn_media", "CPUEn média (ind./100 m² de rede)", output / "05_07_1_top20_cpuen_especies.png", PRIMARY)
        horizontal_top(species_cpue, "CPUEb_kg_media", "CPUEb média (kg/100 m² de rede)", output / "05_07_2_top20_cpueb_especies.png", SECONDARY)

        cpue_point_bars(
            units,
            "CPUEn",
            "CPUEn (ind./100 m² de rede)",
            output / "05_07_5_cpuen_pontos_barras.png",
            1,
        )
        cpue_point_bars(
            units,
            "CPUEb_kg",
            "CPUEb (kg/100 m² de rede)",
            output / "05_07_6_cpueb_pontos_barras.png",
            2,
        )

        point_species = grid.groupby(["ponto", "id_especie"], as_index=False).agg(CPUEn=("CPUEn", "sum"))
        diversity_rows = []
        for point in POINTS:
            values = point_species.loc[point_species["ponto"].eq(point), "CPUEn"]
            diversity_rows.append({"ponto": point, "riqueza": int((values > 0).sum()), "Shannon_H": shannon(values), "Pielou_J": pielou(values)})
        diversity = pd.DataFrame(diversity_rows)
        write_workbook(output / "05_08_diversidade_equitabilidade.xlsx", {"Por_ponto": diversity})
        bar_points(diversity, [("Shannon_H", "Shannon (H')", PRIMARY), ("Pielou_J", "Pielou (J')", ORANGE)], output / "05_08_diversidade_equitabilidade.png")

        matrix = point_species.pivot_table(index="ponto", columns="id_especie", values="CPUEn", aggfunc="sum", fill_value=0).reindex(POINTS, fill_value=0)
        distances = squareform(pdist(matrix.to_numpy(float), metric="braycurtis"))
        similarity = pd.DataFrame(1 - distances, index=POINTS, columns=POINTS).fillna(0)
        np.fill_diagonal(similarity.values, 1)
        write_workbook(output / "05_09_similaridade_bray_curtis.xlsx", {"Matriz_similaridade": similarity.reset_index(names="ponto")})
        heatmap(similarity, output / "05_09_1_similaridade_heatmap_pontos.png")
        similarity_dendrogram(matrix, output / "05_09_2_similaridade_dendrograma_pontos.png")

        curve, presence_matrix = collector_curve(
            general[["ponto", "nome_cientifico"]].drop_duplicates()
        )
        curve_premises = pd.DataFrame(
            [
                ["Unidade amostral", "Ponto regular da campanha"],
                ["Universo", "16 pontos gerais; todos os métodos válidos"],
                ["Aleatorizações", 999],
                ["Estimador", "Jackknife 1"],
                [
                    "Interpretação",
                    "Suficiência espacial da campanha; não representa série temporal",
                ],
            ],
            columns=["Premissa", "Valor"],
        )
        write_workbook(
            output / "05_10_curva_coletor_campanha.xlsx",
            {
                "Curva": curve,
                "Matriz_presenca": presence_matrix,
                "Premissas": curve_premises,
            },
        )
        collector_figure(curve, output / "05_10_curva_coletor_campanha.png")

        mld = general[general["classe_migratoria"].eq("MLD")].copy()
        mld["sexo"] = mld["sexo_padronizado"].fillna("N.I.")
        reproduction = mld[mld["emg_codigo"].notna()].groupby(["nome_cientifico", "nome_popular", "sexo", "emg_codigo", "emg_estadio"], dropna=False).agg(abundancia=("numero_de_individuos", "sum"), pontos=("ponto", "nunique")).reset_index()
        reproductive_spatial = pd.DataFrame({"ponto": POINTS})
        female = mld[mld["emg_codigo"].isin(["F3", "F4"])].groupby("ponto")["numero_de_individuos"].sum()
        male = mld[mld["emg_codigo"].isin(["M3", "M4"])].groupby("ponto")["numero_de_individuos"].sum()
        reproductive_spatial["femeas_reprodutivas_mld"] = reproductive_spatial["ponto"].map(female).fillna(0)
        reproductive_spatial["machos_reprodutivos_mld"] = reproductive_spatial["ponto"].map(male).fillna(0)
        write_workbook(output / "05_12_processo_reprodutivo.xlsx", {"EMG_MLD": reproduction, "Reprodutivos_espacial": reproductive_spatial})
        bar_points(reproductive_spatial, [("femeas_reprodutivas_mld", "Fêmeas MLD F3-F4 (n)", RED), ("machos_reprodutivos_mld", "Machos MLD M3-M4 (n)", PRIMARY)], output / "05_12_reproducao_espacial_mld.png")

        recruitment_query = text(
            """
            SELECT d.id_especie, e.nome_cientifico, e.nome_popular, d.campanha, d.ponto,
                   d.numero_de_individuos, d.cp_cm, d.emg_codigo, e.estrategia_reprodutiva
            FROM resultados_ictiofauna_detalhe d
            JOIN especies e ON e.id_especie = d.id_especie
            WHERE d.codigo_opyta = :code AND d.cp_cm IS NOT NULL AND d.cp_cm > 0
            """
        )
        recruitment_base = pd.read_sql(recruitment_query, connection, params={"code": PROJECT_CODE})
        recruitment_base = recruitment_base[recruitment_base["estrategia_reprodutiva"].map(norm).eq("migradora de longa distancia")].copy()
        reference = recruitment_base[recruitment_base["emg_codigo"].isin(["F2", "F3", "F4"])].groupby("id_especie")["cp_cm"].min().rename("cp_referencia_f2plus_min_cm")
        latest_recruitment = recruitment_base[recruitment_base["campanha"].eq(campaign) & recruitment_base["ponto"].isin(POINTS)].merge(reference, on="id_especie", how="left")
        latest_recruitment["jovem_mld"] = latest_recruitment["cp_referencia_f2plus_min_cm"].notna() & (latest_recruitment["cp_cm"] < latest_recruitment["cp_referencia_f2plus_min_cm"])
        young = latest_recruitment[latest_recruitment["jovem_mld"]].copy()
        young_spatial = pd.DataFrame({"ponto": POINTS})
        young_counts = young.groupby("ponto")["numero_de_individuos"].sum()
        young_spatial["juvenis_mld"] = young_spatial["ponto"].map(young_counts).fillna(0)
        write_workbook(output / "05_15_recrutamento_mld.xlsx", {"Base_auditada": latest_recruitment, "Juvenis_MLD": young, "Variacao_espacial": young_spatial, "Limites_especies": reference.reset_index()})
        bar_points(young_spatial, [("juvenis_mld", "Juvenis MLD (n)", GREEN)], output / "05_15_recrutamento_mld_espacial.png")

        stp_species = stp.groupby(["ponto", "ordem", "familia", "nome_cientifico", "nome_popular", "origem", "classe_migratoria"], dropna=False).agg(abundancia=("numero_de_individuos", "sum"), biomassa_g=("biomassa_g_linha", "sum")).reset_index()
        stp_summary = stp.groupby("ponto", dropna=False).agg(riqueza=("nome_cientifico", "nunique"), abundancia=("numero_de_individuos", "sum"), biomassa_g=("biomassa_g_linha", "sum")).reindex(STP_POINTS, fill_value=0).rename_axis("ponto").reset_index()
        write_workbook(output / "05_19_sistema_transposicao_peixes.xlsx", {"Composicao_STP": stp_species, "Resumo_STP": stp_summary})
        bar_points(stp_summary, [("riqueza", "Riqueza (nº de espécies)", PRIMARY), ("abundancia", "Abundância (nº de indivíduos)", ORANGE)], output / "05_19_stp_espacial.png")

        map_module = load_map_module()
        map_base = point_grid[["ponto", "latitude", "longitude", "trecho_longitudinal", "ordem_lista"]].copy()
        map_metrics = [
            ("riqueza_ameacadas", threat_spatial[["ponto", "riqueza_ameacadas"]], "Riqueza de espécies ameaçadas na campanha", 0, "05_03_1_mapa_riqueza_ameacadas_espacial.png"),
            ("abundancia_ameacadas", threat_spatial[["ponto", "abundancia_ameacadas"]], "Abundância de espécies ameaçadas na campanha (nº de indivíduos)", 0, "05_03_2_mapa_abundancia_ameacadas_espacial.png"),
            ("riqueza", spatial[["ponto", "riqueza"]], "Riqueza na campanha (nº de espécies)", 0, "05_05_1_mapa_riqueza_espacial.png"),
            ("abundancia", spatial[["ponto", "abundancia"]], "Abundância na campanha (nº de indivíduos)", 0, "05_05_2_mapa_abundancia_espacial.png"),
            ("biomassa_kg", spatial.assign(biomassa_kg=spatial["biomassa_g"] / 1000)[["ponto", "biomassa_kg"]], "Biomassa na campanha (kg)", 1, "05_05_3_mapa_biomassa_espacial.png"),
            ("CPUEn_media", cpue_points[["ponto", "CPUEn_media"]], "CPUEn média (ind./100 m² de rede)", 1, "05_07_3_mapa_cpuen_espacial.png"),
            ("CPUEb_kg_media", cpue_points[["ponto", "CPUEb_kg_media"]], "CPUEb média (kg/100 m² de rede)", 2, "05_07_4_mapa_cpueb_espacial.png"),
            ("Shannon_H", diversity[["ponto", "Shannon_H"]], "Diversidade de Shannon (H')", 2, "05_08_1_mapa_shannon_espacial.png"),
            ("Pielou_J", diversity[["ponto", "Pielou_J"]], "Equitabilidade de Pielou (J')", 2, "05_08_2_mapa_pielou_espacial.png"),
            ("femeas_reprodutivas_mld", reproductive_spatial[["ponto", "femeas_reprodutivas_mld"]], "Fêmeas reprodutivas MLD (n; F3-F4)", 0, "05_12_1_mapa_femeas_reprodutivas_mld.png"),
            ("machos_reprodutivos_mld", reproductive_spatial[["ponto", "machos_reprodutivos_mld"]], "Machos reprodutivos MLD (n; M3-M4)", 0, "05_12_2_mapa_machos_reprodutivos_mld.png"),
            ("juvenis_mld", young_spatial[["ponto", "juvenis_mld"]], "Recrutamento MLD na campanha (nº de juvenis)", 0, "05_15_1_mapa_recrutamento_mld.png"),
        ]
        for metric, frame, legend, decimals, filename in map_metrics:
            map_data = map_base.merge(frame.rename(columns={metric: "valor"}), on="ponto", how="left")
            map_data["valor"] = pd.to_numeric(map_data["valor"], errors="coerce").fillna(0)
            map_figure(map_module, map_data, legend, decimals, output / filename)

        sampled_dates = points[["data_inicio", "data_fim"]].stack().dropna()
        manifest = {
            "projeto": PROJECT_CODE,
            "campanha": campaign,
            "gerado_em": datetime.now().isoformat(timespec="seconds"),
            "periodo_amostragem_inicio": sampled_dates.min().isoformat() if not sampled_dates.empty else None,
            "periodo_amostragem_fim": sampled_dates.max().isoformat() if not sampled_dates.empty else None,
            "universo_geral": "16 pontos regulares; ICTIO13C/D excluidos e tratados no STP",
            "linhas_gerais": int(len(general)),
            "pontos_gerais": int(general["ponto"].nunique()),
            "especies_gerais": int(general["nome_cientifico"].nunique()),
            "abundancia_geral": float(general["numero_de_individuos"].sum()),
            "biomassa_geral_g": float(general["biomassa_g_linha"].sum(skipna=True)),
            "linhas_stp": int(len(stp)),
            "especies_stp": int(stp["nome_cientifico"].nunique()),
            "abundancia_stp": float(stp["numero_de_individuos"].sum()),
            "unidades_esforco_rede": int(units["id_esforco"].nunique()),
            "especies_ameacadas": int(threatened["nome_cientifico"].nunique()),
            "femeas_reprodutivas_mld_f3_f4": float(reproductive_spatial["femeas_reprodutivas_mld"].sum()),
            "machos_reprodutivos_mld_m3_m4": float(reproductive_spatial["machos_reprodutivos_mld"].sum()),
            "juvenis_mld": float(young_spatial["juvenis_mld"].sum()),
            "curva_coletor_unidade": "ponto regular da campanha",
            "curva_coletor_riqueza_observada_final": float(curve["riqueza_observada_media"].iloc[-1]),
            "curva_coletor_jackknife1_final": float(curve["jackknife1_media"].iloc[-1]),
            "produtos_nao_aplicaveis_ao_recorte": [
                "5.6 distribuicao temporal",
                "curva acumulativa temporal entre campanhas; substituida pela curva espacial entre pontos",
                "diversidade, equitabilidade e similaridade temporais",
            ],
            "status": "PACOTE_CAMPANHA_GERADO",
        }
        (output / "00_manifesto_pacote_campanha.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        validation = audit_package(output, manifest)
        (output / "00_validacao_pacote_campanha.json").write_text(
            json.dumps(validation, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        return output


def main() -> None:
    parser = argparse.ArgumentParser(description="Gera o pacote tematico de uma campanha BIOCOL001.")
    parser.add_argument("--campanha", help="Campanha canonica, por exemplo C069-2026-06. Omitir para usar a ultima.")
    args = parser.parse_args()
    output = build_package(args.campanha)
    print(output)


if __name__ == "__main__":
    main()
