from __future__ import annotations

import os
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from matplotlib.patches import Patch
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(
    "G:/Meu Drive/Opyta/Clientes/Clientes/Clientes/Bios/Colíder/Resultados/2026/Junho-2026/"
    "BIOCOL001_RESULTADOS_ICTIOFAUNA_FINAL_R02"
)
BASE_PATH = ROOT / "base_analitica_geral_biocol001.xlsx"
SPATIAL_PATH = ROOT / "05_05_distribuicao_espacial_riqueza_abundancia.xlsx"
TEMPORAL_PATH = ROOT / "05_06_riqueza_temporal_campanha.xlsx"
SPATIAL_FIGURE_PATH = ROOT / "05_05_distribuicao_espacial_riqueza_abundancia.png"

PRIMARY = "#002060"
SECONDARY = "#5B9BD5"
GRID = "#D9D9D9"


def matrix(data: pd.DataFrame, species_names: pd.DataFrame, dimension: str, value: str) -> pd.DataFrame:
    result = data.pivot_table(
        index="nome_cientifico",
        columns=dimension,
        values=value,
        aggfunc="sum",
        fill_value=0,
    )
    result["Total"] = result.sum(axis=1)
    result = result.reset_index().merge(species_names, on="nome_cientifico", how="left")
    columns = ["nome_cientifico", "nome_popular"] + [column for column in result.columns if column not in {"nome_cientifico", "nome_popular"}]
    return result[columns].sort_values("nome_cientifico")


def occurrence_matrix(data: pd.DataFrame, species_names: pd.DataFrame, dimension: str) -> pd.DataFrame:
    occurrence = data.assign(ocorrencia=(data["numero_de_individuos"] > 0).astype(int))
    result = occurrence.pivot_table(
        index="nome_cientifico",
        columns=dimension,
        values="ocorrencia",
        aggfunc="max",
        fill_value=0,
    )
    result["Total_unidades_ocorrencia"] = result.sum(axis=1)
    result = result.reset_index().merge(species_names, on="nome_cientifico", how="left")
    columns = ["nome_cientifico", "nome_popular"] + [column for column in result.columns if column not in {"nome_cientifico", "nome_popular"}]
    return result[columns].sort_values("nome_cientifico")


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "B2" if sheet.max_column > 5 else "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 30
        for column in range(1, sheet.max_column + 1):
            values = [sheet.cell(row, column).value for row in range(1, min(sheet.max_row, 150) + 1)]
            width = max(len(str(value or "")) for value in values) + 2
            sheet.column_dimensions[get_column_letter(column)].width = min(max(width, 11), 42)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
    temporary = path.with_name(f"{path.stem}.tmp.xlsx")
    workbook.save(temporary)
    os.replace(temporary, path)


def write_workbook(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, data in sheets.items():
            data.to_excel(writer, sheet_name=name, index=False)
    style_workbook(path)


def plot_spatial_summary(spatial_summary: pd.DataFrame) -> None:
    plt.rcParams.update(
        {
            "font.family": "Arial",
            "font.size": 17,
            "axes.labelsize": 17,
            "xtick.labelsize": 13,
            "ytick.labelsize": 13,
            "legend.fontsize": 15,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "savefig.facecolor": "white",
        }
    )
    spatial = spatial_summary.sort_values("ponto").reset_index(drop=True)
    x = list(range(len(spatial)))
    fig, ax = plt.subplots(figsize=(18, 10.2), dpi=300)
    ax2 = ax.twinx()

    ax.bar([value - 0.2 for value in x], spatial["riqueza"], width=0.4, color=PRIMARY)
    ax2.bar(
        [value + 0.2 for value in x],
        spatial["abundancia"],
        width=0.4,
        color=SECONDARY,
        alpha=0.75,
    )
    ax.set_xticks(x)
    ax.set_xticklabels(spatial["ponto"], rotation=45, ha="right")
    ax.set_ylabel("Riqueza acumulada (nº de espécies)")
    ax2.set_ylabel("Abundância acumulada (nº de indivíduos)")
    ax.grid(axis="y", color=GRID, alpha=0.25, linewidth=1)
    ax2.grid(False)
    for axis in (ax, ax2):
        for spine in axis.spines.values():
            spine.set_visible(True)
            spine.set_color("black")

    fig.legend(
        handles=[
            Patch(facecolor=PRIMARY, edgecolor="none", label="Riqueza acumulada (nº de espécies)"),
            Patch(facecolor=SECONDARY, edgecolor="none", alpha=0.75, label="Abundância acumulada (nº de indivíduos)"),
        ],
        loc="upper center",
        bbox_to_anchor=(0.5, 0.975),
        ncol=2,
        frameon=False,
    )
    fig.subplots_adjust(left=0.08, right=0.92, bottom=0.22, top=0.88)
    fig.savefig(SPATIAL_FIGURE_PATH, dpi=300, facecolor="white")
    plt.close(fig)


def main() -> None:
    base = pd.read_excel(BASE_PATH, sheet_name="base")
    numeric = ["numero_de_individuos", "biomassa_g_linha"]
    for column in numeric:
        base[column] = pd.to_numeric(base[column], errors="coerce").fillna(0)
    base = base[base["nome_cientifico"].notna() & base["ponto"].notna() & base["campanha"].notna()].copy()
    species_names = (
        base.assign(nome_popular=base["nome_popular"].fillna("").astype(str).str.strip())
        .sort_values(["nome_cientifico", "nome_popular"], key=lambda column: column.eq("") if column.name == "nome_popular" else column)
        .drop_duplicates("nome_cientifico")[["nome_cientifico", "nome_popular"]]
    )

    spatial_summary = (
        base.groupby("ponto", dropna=False)
        .agg(
            riqueza=("nome_cientifico", "nunique"),
            abundancia=("numero_de_individuos", "sum"),
            biomassa_g=("biomassa_g_linha", "sum"),
            campanhas_amostradas=("campanha", "nunique"),
        )
        .reset_index()
        .sort_values("ponto")
    )
    spatial_sheets = {
        "Resumo_pontos": spatial_summary,
        "Abundancia_especie": matrix(base, species_names, "ponto", "numero_de_individuos"),
        "Biomassa_especie": matrix(base, species_names, "ponto", "biomassa_g_linha"),
        "Ocorrencia_especie": occurrence_matrix(base, species_names, "ponto"),
    }

    temporal_summary = (
        base.groupby("campanha", dropna=False)
        .agg(
            riqueza=("nome_cientifico", "nunique"),
            abundancia=("numero_de_individuos", "sum"),
            biomassa_g=("biomassa_g_linha", "sum"),
            pontos_amostrados=("ponto", "nunique"),
            fase_reservatorio=("fase_reservatorio", "first"),
            evento_reservatorio=("evento_reservatorio", "first"),
            periodo_rebaixamento_parcial=("periodo_rebaixamento_parcial", "first"),
            marco_reenchimento=("marco_reenchimento", "first"),
        )
        .reset_index()
        .sort_values("campanha")
    )
    temporal_sheets = {
        "Resumo_campanhas": temporal_summary,
        "Abundancia_especie": matrix(base, species_names, "campanha", "numero_de_individuos"),
        "Biomassa_especie": matrix(base, species_names, "campanha", "biomassa_g_linha"),
        "Ocorrencia_especie": occurrence_matrix(base, species_names, "campanha"),
    }

    write_workbook(SPATIAL_PATH, spatial_sheets)
    write_workbook(TEMPORAL_PATH, temporal_sheets)
    plot_spatial_summary(spatial_summary)

    print(f"pontos={base['ponto'].nunique()} campanhas={base['campanha'].nunique()} especies={base['nome_cientifico'].nunique()}")
    print(f"abundancia={base['numero_de_individuos'].sum():.0f} biomassa_g={base['biomassa_g_linha'].sum():.4f}")
    print(SPATIAL_PATH)
    print(TEMPORAL_PATH)
    print(SPATIAL_FIGURE_PATH)


if __name__ == "__main__":
    main()
