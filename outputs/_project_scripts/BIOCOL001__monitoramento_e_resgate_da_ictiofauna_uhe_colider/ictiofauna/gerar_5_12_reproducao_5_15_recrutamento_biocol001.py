from __future__ import annotations

import os
import re
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.lines import Line2D
from matplotlib.patches import Patch
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows


ROOT = Path(
    "G:/Meu Drive/Opyta/Clientes/Clientes/Clientes/Bios/Col\u00edder/Resultados/2026/Junho-2026/"
    "BIOCOL001_RESULTADOS_ICTIOFAUNA_FINAL_R02"
)
REPRO_BASE = ROOT / "base_reprodutiva_biometria_biocol001.xlsx"
RECRUIT_BASE = ROOT / "base_recrutamento_mld_biocol001.xlsx"
GENERAL_BASE = ROOT / "base_analitica_sem_marcacao_biocol001.xlsx"
REPRO_BOOK = ROOT / "05_12_processo_reprodutivo.xlsx"
RECRUIT_BOOK = ROOT / "05_15_recrutamento_mld_auditoria.xlsx"
FIG_REPRO_SPATIAL = ROOT / "05_12_1_reproducao_espacial.png"
FIG_REPRO_TEMPORAL = ROOT / "05_12_2_reproducao_temporal.png"
FIG_RECRUIT_SPATIAL = ROOT / "05_15_1_recrutamento_mld_espacial.png"
FIG_RECRUIT_TEMPORAL = ROOT / "05_15_2_recrutamento_mld_temporal.png"

PRIMARY = "#002060"
SECONDARY = "#5B9BD5"
LIGHT = "#DBE5F1"
ORANGE = "#D4672A"
GRID = "#D9D9D9"
PRE = "#F0F0F0"
STAGE_COLORS = {1: "#DBE5F1", 2: "#9DC3E6", 3: "#5B9BD5", 4: "#002060"}
STAGE_LABELS = {
    1: "1 - Repouso",
    2: "2 - Matura\u00e7\u00e3o inicial",
    3: "3 - Matura\u00e7\u00e3o avan\u00e7ada/maduro",
    4: "4 - Desovado/esgotado",
}
A4_LANDSCAPE = (11.69, 8.27)
POINT_ORDER = [f"ICTIO{i:02d}" for i in range(1, 13)] + ["ICTIO13A", "ICTIO13B", "ICTIO14", "ICTIO15"]
EMG_ORDER = ["F1", "F2", "F3", "F4", "M1", "M2", "M3", "M4"]


def campaign_order(value: str) -> int:
    match = re.match(r"C(\d{3})", str(value))
    return int(match.group(1)) if match else 999


def prepare_reproduction():
    raw = pd.read_excel(REPRO_BASE, sheet_name="base")
    data = raw[
        raw["camada_operacional"].eq("malha_regular")
        & raw["sexo_padronizado"].isin(["F", "M"])
        & raw["emg_codigo"].isin(EMG_ORDER)
    ].copy()
    data["numero_de_individuos"] = pd.to_numeric(data["numero_de_individuos"], errors="coerce").fillna(0)
    data["ordem_campanha"] = data["campanha"].map(campaign_order)
    data["estagio_ordem"] = data["emg_codigo"].str.extract(r"(\d)", expand=False).astype(int)

    species = (
        data.pivot_table(
            index=["nome_cientifico", "nome_popular"],
            columns="emg_codigo",
            values="numero_de_individuos",
            aggfunc="sum",
            fill_value=0,
        )
        .reindex(columns=EMG_ORDER, fill_value=0)
        .reset_index()
    )
    species["Total_femeas"] = species[["F1", "F2", "F3", "F4"]].sum(axis=1)
    species["Total_machos"] = species[["M1", "M2", "M3", "M4"]].sum(axis=1)
    species["Total_EMG"] = species["Total_femeas"] + species["Total_machos"]
    species = species.sort_values(["Total_EMG", "nome_cientifico"], ascending=[False, True]).reset_index(drop=True)

    spatial_abs = (
        data.pivot_table(index="ponto", columns="emg_codigo", values="numero_de_individuos", aggfunc="sum", fill_value=0)
        .reindex(index=POINT_ORDER, columns=EMG_ORDER, fill_value=0)
        .reset_index()
    )
    campaigns = (
        pd.read_excel(GENERAL_BASE, sheet_name="base")
        [["campanha", "fase_reservatorio", "evento_reservatorio"]]
        .drop_duplicates("campanha")
    )
    campaigns["ordem_campanha"] = campaigns["campanha"].map(campaign_order)
    campaigns = campaigns.sort_values("ordem_campanha").query("ordem_campanha <= 69")
    temporal_abs = data.pivot_table(
        index="campanha", columns="emg_codigo", values="numero_de_individuos", aggfunc="sum", fill_value=0
    ).reindex(columns=EMG_ORDER, fill_value=0).reset_index()
    temporal_abs = campaigns.merge(temporal_abs, on="campanha", how="left")
    temporal_abs[EMG_ORDER] = temporal_abs[EMG_ORDER].fillna(0)

    spatial_pct = relative_by_sex(spatial_abs, "ponto")
    temporal_pct = relative_by_sex(temporal_abs, "campanha")

    reproductive_migrants = data[
        data["classe_migratoria"].eq("MLD")
        & data["emg_codigo"].isin(["F3", "F4", "M3", "M4"])
    ].copy()
    migrant_spatial = (
        reproductive_migrants.pivot_table(
            index="ponto",
            columns="emg_codigo",
            values="numero_de_individuos",
            aggfunc="sum",
            fill_value=0,
        )
        .reindex(index=POINT_ORDER, columns=["F3", "F4", "M3", "M4"], fill_value=0)
        .reset_index()
    )
    migrant_spatial["Femeas_reprodutivas"] = migrant_spatial[["F3", "F4"]].sum(axis=1)
    migrant_spatial["Machos_reprodutivos"] = migrant_spatial[["M3", "M4"]].sum(axis=1)
    migrant_species = (
        reproductive_migrants.groupby(
            ["classe_migratoria", "nome_cientifico", "nome_popular", "emg_codigo"],
            dropna=False,
        )["numero_de_individuos"]
        .sum()
        .unstack("emg_codigo", fill_value=0)
        .reindex(columns=["F3", "F4", "M3", "M4"], fill_value=0)
        .reset_index()
    )
    migrant_species["Femeas_reprodutivas"] = migrant_species[["F3", "F4"]].sum(axis=1)
    migrant_species["Machos_reprodutivos"] = migrant_species[["M3", "M4"]].sum(axis=1)
    return (
        data,
        species,
        spatial_abs,
        spatial_pct,
        temporal_abs,
        temporal_pct,
        migrant_spatial,
        migrant_species,
    )


def relative_by_sex(data: pd.DataFrame, id_col: str):
    output = data.copy()
    for prefix in ["F", "M"]:
        columns = [f"{prefix}{stage}" for stage in range(1, 5)]
        total = output[columns].sum(axis=1)
        for column in columns:
            output[f"{column}_pct"] = np.where(total.gt(0), output[column] / total * 100, 0)
    leading = [id_col]
    metadata = [c for c in ["ordem_campanha", "fase_reservatorio", "evento_reservatorio"] if c in output.columns]
    values = [f"{code}_pct" for code in EMG_ORDER]
    return output[leading + metadata + values]


def prepare_recruitment():
    source = pd.read_excel(RECRUIT_BASE, sheet_name="base")
    source = source[
        source["universo_geral"].eq(True)
        & source["classe_migratoria"].eq("MLD")
    ].copy()
    source["numero_de_individuos"] = pd.to_numeric(
        source["numero_de_individuos"], errors="coerce"
    ).fillna(0)

    audit_rows = []
    for (species_id, species), group in source.groupby(
        ["id_especie", "nome_cientifico"], dropna=False
    ):
        reference = group[group["emg_codigo"].isin(["F2", "F3", "F4"])]
        young = group[group["jovem_preliminar"].eq(True)]
        audit_rows.append(
            {
                "id_especie": species_id,
                "nome_cientifico": species,
                "classe_recrutamento": "MLD",
                "estrategia_reprodutiva": group["estrategia_reprodutiva"].iloc[0],
                "cp_referencia_f2plus_min_cm": group["cp_referencia_f2plus_min_cm"].iloc[0],
                "linhas_referencia_f2plus": len(reference),
                "individuos_referencia_f2plus": reference["numero_de_individuos"].sum(),
                "linhas_com_cp_mld": len(group),
                "individuos_com_cp_mld": group["numero_de_individuos"].sum(),
                "linhas_jovens_preliminar": len(young),
                "individuos_jovens_preliminar": young["numero_de_individuos"].sum(),
                "status_criterio_recrutamento": group["status_criterio_recrutamento"].iloc[0],
            }
        )
    audit = pd.DataFrame(audit_rows).sort_values("nome_cientifico").reset_index(drop=True)

    data = source[source["jovem_preliminar"].eq(True)].copy()
    data["numero_de_individuos"] = pd.to_numeric(data["numero_de_individuos"], errors="coerce").fillna(0)
    data["ordem_campanha"] = data["campanha"].map(campaign_order)

    general = (
        data.groupby(["nome_cientifico", "nome_popular", "cp_referencia_f2plus_min_cm", "status_criterio_recrutamento"], dropna=False)
        .agg(abundancia_juvenis=("numero_de_individuos", "sum"), campanhas_ocorrencia=("campanha", "nunique"), pontos_ocorrencia=("ponto", "nunique"))
        .reset_index()
        .sort_values(["abundancia_juvenis", "nome_cientifico"], ascending=[False, True])
    )
    spatial = (
        data.pivot_table(index="ponto", columns="nome_cientifico", values="numero_de_individuos", aggfunc="sum", fill_value=0)
        .reindex(POINT_ORDER, fill_value=0)
    )
    spatial["Total_juvenis_MLD"] = spatial.sum(axis=1)
    spatial = spatial.reset_index()

    campaigns = (
        pd.read_excel(GENERAL_BASE, sheet_name="base")
        [["campanha", "fase_reservatorio", "evento_reservatorio"]]
        .drop_duplicates("campanha")
    )
    campaigns["ordem_campanha"] = campaigns["campanha"].map(campaign_order)
    campaigns = campaigns.sort_values("ordem_campanha").query("ordem_campanha <= 69")
    temporal = data.pivot_table(
        index="campanha", columns="nome_cientifico", values="numero_de_individuos", aggfunc="sum", fill_value=0
    ).reset_index()
    temporal = campaigns.merge(temporal, on="campanha", how="left")
    species_cols = [c for c in temporal.columns if c not in ["campanha", "fase_reservatorio", "evento_reservatorio", "ordem_campanha"]]
    temporal[species_cols] = temporal[species_cols].fillna(0)
    temporal["Total_juvenis_MLD"] = temporal[species_cols].sum(axis=1)

    long = data.sort_values(["ordem_campanha", "ponto", "nome_cientifico"])[[
        "campanha", "ponto", "nome_cientifico", "nome_popular", "numero_de_individuos", "cp_cm",
        "cp_referencia_f2plus_min_cm", "status_criterio_recrutamento", "fase_reservatorio", "evento_reservatorio"
    ]].reset_index(drop=True)
    return data, audit, general, spatial, temporal, long


def write_new_workbook(path: Path, sheets: dict[str, pd.DataFrame]):
    workbook = Workbook()
    workbook.remove(workbook.active)
    append_sheets(workbook, sheets, replace=False)
    workbook.save(path)


def append_sheets(workbook, sheets: dict[str, pd.DataFrame], replace=True):
    for name, data in sheets.items():
        if replace and name in workbook.sheetnames:
            del workbook[name]
        sheet = workbook.create_sheet(name)
        export = data.reset_index() if data.index.name else data
        for row in dataframe_to_rows(export, index=False, header=True):
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column in range(1, sheet.max_column + 1):
            values = [sheet.cell(row, column).value for row in range(1, min(sheet.max_row, 300) + 1)]
            width = max(len(str(value or "")) for value in values) + 2
            sheet.column_dimensions[get_column_letter(column)].width = min(max(width, 12), 40)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"


def update_workbook(path: Path, sheets: dict[str, pd.DataFrame]):
    workbook = load_workbook(path)
    append_sheets(workbook, sheets)
    temporary = path.with_name(f"{path.stem}.tmp.xlsx")
    workbook.save(temporary)
    os.replace(temporary, path)


def style_axes(ax, grid_axis="y"):
    ax.grid(False)
    ax.grid(axis=grid_axis, color=GRID, alpha=0.25, linewidth=1)
    for spine in ax.spines.values():
        spine.set_visible(True)
        spine.set_color("black")
        spine.set_linewidth(1.2)
    ax.tick_params(colors="black", labelsize=13)


def shade_events(ax, temporal: pd.DataFrame):
    pre = temporal["fase_reservatorio"].eq("pre_enchimento").to_numpy()
    lowering = temporal["evento_reservatorio"].eq("rebaixamento_parcial_reservatorio").to_numpy()
    refill = temporal["evento_reservatorio"].eq("reenchimento_reservatorio").to_numpy()
    if pre.any():
        idx = np.where(pre)[0]
        ax.axvspan(idx.min() - 0.5, idx.max() + 0.5, color=PRE, alpha=0.75, linewidth=0, zorder=0)
    if lowering.any():
        idx = np.where(lowering)[0]
        ax.axvspan(idx.min() - 0.5, idx.max() + 0.5, color=LIGHT, alpha=0.72, linewidth=0, zorder=0)
    if refill.any():
        ax.axvline(np.where(refill)[0][0], color=ORANGE, linewidth=1.6, linestyle="--", zorder=3)


def stacked_panel(ax, data: pd.DataFrame, prefix: str, labels: list[str], temporal=False):
    x = np.arange(len(data))
    bottom = np.zeros(len(data))
    if temporal:
        shade_events(ax, data)
    for stage in range(1, 5):
        values = data[f"{prefix}{stage}_pct"].to_numpy(float)
        ax.bar(x, values, bottom=bottom, width=0.82, color=STAGE_COLORS[stage], edgecolor="white", linewidth=0.35)
        bottom += values
    ax.set_ylabel("Abund\u00e2ncia relativa (%)")
    ax.set_ylim(0, 100)
    style_axes(ax)
    ax.tick_params(axis="y", labelsize=10.5)
    ax.yaxis.label.set_size(12)
    if temporal:
        tick_indices = np.unique(np.linspace(0, len(data) - 1, 12, dtype=int))
        tick_labels = []
        for idx in tick_indices:
            parts = str(labels[idx]).split("-")
            tick_labels.append(f"{parts[0]}\n{'-'.join(parts[1:])}" if len(parts) > 1 else parts[0])
        ax.set_xticks(tick_indices, tick_labels, rotation=0, ha="center", fontsize=9.5)
    else:
        ax.set_xticks(x, labels, rotation=35, ha="right", fontsize=10)


def plot_reproduction(spatial_pct: pd.DataFrame, temporal_pct: pd.DataFrame):
    legend = [Patch(facecolor=STAGE_COLORS[i], edgecolor="white", label=STAGE_LABELS[i]) for i in range(1, 5)]
    event_legend = [
        Patch(facecolor=PRE, edgecolor="none", alpha=0.75, label="Pr\u00e9-enchimento"),
        Patch(facecolor=LIGHT, edgecolor="none", alpha=0.72, label="Rebaixamento parcial"),
        Line2D([0], [0], color=ORANGE, linewidth=1.6, linestyle="--", label="Reenchimento"),
    ]

    fig, axes = plt.subplots(2, 1, figsize=A4_LANDSCAPE, dpi=300, sharex=True)
    stacked_panel(axes[0], spatial_pct, "F", spatial_pct["ponto"].tolist())
    stacked_panel(axes[1], spatial_pct, "M", spatial_pct["ponto"].tolist())
    axes[0].text(0.01, 0.94, "F\u00eameas", transform=axes[0].transAxes, fontsize=14, fontweight="bold", va="top")
    axes[1].text(0.01, 0.94, "Machos", transform=axes[1].transAxes, fontsize=14, fontweight="bold", va="top")
    axes[1].set_xlabel("Trecho amostral")
    axes[1].xaxis.label.set_size(12)
    fig.legend(
        handles=[legend[0], legend[2], legend[1], legend[3]],
        loc="upper center",
        bbox_to_anchor=(0.5, 0.985),
        ncol=2,
        frameon=False,
        fontsize=10.5,
        columnspacing=1.8,
        handlelength=1.8,
    )
    fig.subplots_adjust(top=0.84, bottom=0.17, left=0.085, right=0.985, hspace=0.16)
    fig.savefig(FIG_REPRO_SPATIAL, dpi=300, facecolor="white")
    plt.close(fig)

    fig, axes = plt.subplots(2, 1, figsize=A4_LANDSCAPE, dpi=300, sharex=True)
    labels = temporal_pct["campanha"].tolist()
    stacked_panel(axes[0], temporal_pct, "F", labels, temporal=True)
    stacked_panel(axes[1], temporal_pct, "M", labels, temporal=True)
    axes[0].text(0.01, 0.94, "F\u00eameas", transform=axes[0].transAxes, fontsize=14, fontweight="bold", va="top")
    axes[1].text(0.01, 0.94, "Machos", transform=axes[1].transAxes, fontsize=14, fontweight="bold", va="top")
    axes[1].set_xlabel("Campanha")
    axes[1].xaxis.label.set_size(12)
    fig.legend(
        handles=legend,
        loc="upper center",
        bbox_to_anchor=(0.5, 0.985),
        ncol=4,
        frameon=False,
        fontsize=9.5,
        columnspacing=1.25,
        handlelength=1.7,
    )
    fig.legend(
        handles=event_legend,
        loc="upper center",
        bbox_to_anchor=(0.5, 0.94),
        ncol=3,
        frameon=False,
        fontsize=9.5,
        columnspacing=1.5,
        handlelength=1.7,
    )
    fig.subplots_adjust(top=0.84, bottom=0.14, left=0.085, right=0.985, hspace=0.16)
    fig.savefig(FIG_REPRO_TEMPORAL, dpi=300, facecolor="white")
    plt.close(fig)


def plot_recruitment(spatial: pd.DataFrame, temporal: pd.DataFrame):
    x = np.arange(len(spatial))
    fig, ax = plt.subplots(figsize=(18, 10.2), dpi=300)
    ax.bar(x, spatial["Total_juvenis_MLD"], width=0.72, color=PRIMARY)
    ax.set_xticks(x, spatial["ponto"], rotation=45, ha="right")
    ax.set_xlabel("Trecho amostral")
    ax.set_ylabel("Abund\u00e2ncia total de juvenis MLD")
    ax.set_ylim(bottom=0)
    style_axes(ax)
    fig.tight_layout()
    fig.savefig(FIG_RECRUIT_SPATIAL, dpi=300, facecolor="white", bbox_inches="tight")
    plt.close(fig)

    x = np.arange(len(temporal))
    values = temporal["Total_juvenis_MLD"].to_numpy(float)
    fig, ax = plt.subplots(figsize=(18, 10.2), dpi=300)
    shade_events(ax, temporal)
    line = ax.plot(x, values, color=PRIMARY, linewidth=2.8, label="Total de juvenis MLD")[0]
    ax.fill_between(x, values, color=PRIMARY, alpha=0.12)
    step = max(1, len(temporal) // 12)
    ax.set_xticks(x[::step], temporal["campanha"].iloc[::step], rotation=45, ha="right")
    ax.set_xlabel("Campanha")
    ax.set_ylabel("Abund\u00e2ncia total de juvenis MLD")
    ax.set_ylim(bottom=0)
    style_axes(ax)
    handles = [
        line,
        Patch(facecolor=PRE, edgecolor="none", alpha=0.75, label="Pr\u00e9-enchimento"),
        Patch(facecolor=LIGHT, edgecolor="none", alpha=0.72, label="Rebaixamento parcial"),
        Line2D([0], [0], color=ORANGE, linewidth=1.6, linestyle="--", label="Reenchimento"),
    ]
    ax.legend(handles=handles, loc="upper center", bbox_to_anchor=(0.5, 1.08), ncol=4, frameon=False)
    fig.tight_layout()
    fig.savefig(FIG_RECRUIT_TEMPORAL, dpi=300, facecolor="white", bbox_inches="tight")
    plt.close(fig)

def main():
    plt.rcParams.update(
        {
            "font.family": "Arial",
            "font.size": 17,
            "axes.labelsize": 17,
            "xtick.labelsize": 13,
            "ytick.labelsize": 13,
            "legend.fontsize": 13,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "savefig.facecolor": "white",
        }
    )
    (
        repro,
        species,
        spatial_abs,
        spatial_pct,
        temporal_abs,
        temporal_pct,
        migrant_spatial,
        migrant_species,
    ) = prepare_reproduction()
    repro_premises = pd.DataFrame(
        {
            "premissa": ["Universo", "Per\u00edodo", "Categorias", "Abund\u00e2ncia relativa", "Aus\u00eancia de EMG"],
            "regra": [
                "16 pontos gerais; marca\u00e7\u00e3o e STP exclu\u00eddos",
                "C001-C069; 49 campanhas possuem EMG de machos ou f\u00eameas",
                "F1-F4 e M1-M4; est\u00e1gios 1-4 usam a paleta EMG de BIOPOR001",
                "Percentual calculado separadamente dentro de cada sexo por ponto/campanha",
                "Registros F/M sem EMG n\u00e3o entram na distribui\u00e7\u00e3o por categoria",
            ],
        }
    )
    emg_dictionary = pd.DataFrame(
        [
            {
                "EMG_codigo": f"{sex}{stage}",
                "sexo": "F\u00eamea" if sex == "F" else "Macho",
                "estagio": STAGE_LABELS[stage],
            }
            for sex in ["F", "M"]
            for stage in range(1, 5)
        ]
    )
    write_new_workbook(
        REPRO_BOOK,
        {
            "Geral_especie_EMG": species,
            "Dicionario_EMG": emg_dictionary,
            "Espacial_absoluta": spatial_abs,
            "Espacial_percentual": spatial_pct,
            "Temporal_absoluta": temporal_abs,
            "Temporal_percentual": temporal_pct,
            "Migradoras_reprod_espacial": migrant_spatial,
            "Migradoras_reprod_especies": migrant_species,
            "Premissas": repro_premises,
            "Premissas_mapas_reprod": pd.DataFrame(
                {
                    "premissa": ["Universo", "Grupo", "Femeas", "Machos", "Estadio 2"],
                    "regra": [
                        "16 pontos gerais; marcacao e STP excluidos",
                        "Somente migradoras MLD",
                        "F3 + F4",
                        "M3 + M4",
                        "Exclusivo da regra de recrutamento; nao entra nos mapas reprodutivos",
                    ],
                }
            ),
        },
    )
    plot_reproduction(spatial_pct, temporal_pct)

    recruits, audit, general, spatial, temporal, long = prepare_recruitment()
    recruit_premises = pd.DataFrame(
        {
            "premissa": ["Universo", "Grupo", "Crit\u00e9rio", "Espacial", "Temporal", "Limita\u00e7\u00e3o"],
            "regra": [
                "16 pontos gerais; marca\u00e7\u00e3o e STP exclu\u00eddos",
                "Somente migradores de longa dist\u00e2ncia (MLD)",
                "Jovem preliminar = CP_cm inferior ao menor CP_cm F2/F3/F4 da esp\u00e9cie",
                "Abund\u00e2ncia de juvenis por ponto e esp\u00e9cie, integrada entre campanhas",
                "Abund\u00e2ncia de juvenis por campanha e esp\u00e9cie, integrada entre pontos",
                "Esp\u00e9cies sem refer\u00eancia F2+ n\u00e3o s\u00e3o classificadas como juvenis",
            ],
        }
    )
    write_new_workbook(
        RECRUIT_BOOK,
        {
            "Recrutamento": audit,
            "Juvenis_MLD_geral": general,
            "Juvenis_MLD_espacial": spatial,
            "Juvenis_MLD_temporal": temporal,
            "Juvenis_MLD_registros": long,
            "Premissas_MLD": recruit_premises,
        },
    )
    plot_recruitment(spatial, temporal)
    print(
        {
            "repro_especies": int(species.shape[0]),
            "repro_individuos": float(repro["numero_de_individuos"].sum()),
            "repro_campanhas": int(repro["campanha"].nunique()),
            "repro_pontos": int(repro["ponto"].nunique()),
            "juvenis_mld": float(recruits["numero_de_individuos"].sum()),
            "especies_mld": int(recruits["nome_cientifico"].nunique()),
            "campanhas_mld": int(recruits["campanha"].nunique()),
            "pontos_mld": int(recruits["ponto"].nunique()),
        }
    )


if __name__ == "__main__":
    main()
