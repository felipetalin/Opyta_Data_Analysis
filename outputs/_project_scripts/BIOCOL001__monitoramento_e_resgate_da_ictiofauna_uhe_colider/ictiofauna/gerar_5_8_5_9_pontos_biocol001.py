from __future__ import annotations

import importlib.util
import os
import sys
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.patches import Patch
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from scipy.cluster.hierarchy import dendrogram, linkage
from scipy.spatial.distance import pdist, squareform


HERE = Path(__file__).resolve().parent
ROOT = Path(
    "G:/Meu Drive/Opyta/Clientes/Clientes/Clientes/Bios/Col\u00edder/Resultados/2026/Junho-2026/"
    "BIOCOL001_RESULTADOS_ICTIOFAUNA_FINAL_R02"
)
BOOK_DIVERSITY = ROOT / "05_08_dados_diversidade_equitabilidade.xlsx"
BOOK_SIMILARITY = ROOT / "05_09_dados_similaridade_bray_curtis.xlsx"
FIG_DIVERSITY = ROOT / "05_08_diversidade_equitabilidade_pontos.png"
FIG_SIMILARITY = ROOT / "05_09_similaridade_bray_curtis_pontos.png"
FIG_DIVERSITY_TEMPORAL = ROOT / "05_08_diversidade_equitabilidade.png"
FIG_SIMILARITY_TEMPORAL = ROOT / "05_09_similaridade_bray_curtis_campanhas.png"

PRIMARY = "#002060"
SECONDARY = "#5B9BD5"
GRID = "#D9D9D9"
PRE_FILL = "#F0F0F0"
DRAWDOWN = "#DBE5F1"
REFILL = "#D4672A"
POINT_ORDER = [f"ICTIO{i:02d}" for i in range(1, 13)] + ["ICTIO13A", "ICTIO13B", "ICTIO14", "ICTIO15"]


def load_cpue_module():
    path = HERE / "gerar_5_7_cpue_biocol001.py"
    spec = importlib.util.spec_from_file_location("biocol001_cpue", path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def shannon(values: np.ndarray) -> float:
    values = values[values > 0]
    proportions = values / values.sum()
    return float(-(proportions * np.log(proportions)).sum())


def build_point_matrix(efforts: pd.DataFrame, catches: pd.DataFrame):
    species = catches[["id_especie", "nome_cientifico", "nome_popular"]].drop_duplicates("id_especie")
    grid = pd.MultiIndex.from_product(
        [efforts["id_esforco"].tolist(), species["id_especie"].tolist()],
        names=["id_esforco", "id_especie"],
    ).to_frame(index=False)
    grid = grid.merge(
        efforts[["id_esforco", "campanha", "ponto", "esforco_m2"]],
        on="id_esforco",
        how="left",
    ).merge(
        catches[["id_esforco", "id_especie", "abundancia"]],
        on=["id_esforco", "id_especie"],
        how="left",
    )
    grid["abundancia"] = pd.to_numeric(grid["abundancia"], errors="coerce").fillna(0)
    grid["CPUEn"] = grid["abundancia"] / grid["esforco_m2"] * 100

    long = (
        grid.groupby(["ponto", "id_especie"], as_index=False)
        .agg(CPUEn_media=("CPUEn", "mean"), CPUEn_soma=("CPUEn", "sum"), ocorrencias=("abundancia", lambda x: int((x > 0).sum())))
        .merge(species, on="id_especie", how="left")
    )
    long["nome_popular"] = long["nome_popular"].fillna("N.I.")
    long["ordem_ponto"] = long["ponto"].map({p: i for i, p in enumerate(POINT_ORDER)})
    long = long.sort_values(["ordem_ponto", "nome_cientifico"]).drop(columns="ordem_ponto").reset_index(drop=True)

    matrix = long.pivot(index="ponto", columns="nome_cientifico", values="CPUEn_media").fillna(0)
    matrix = matrix.reindex(POINT_ORDER).fillna(0)
    summary_rows = []
    effort_count = efforts.groupby("ponto")["id_esforco"].nunique()
    for point, row in matrix.iterrows():
        values = row.to_numpy(float)
        richness = int((values > 0).sum())
        h = shannon(values)
        summary_rows.append(
            {
                "ponto": point,
                "n_unidades_esforco": int(effort_count.get(point, 0)),
                "riqueza": richness,
                "CPUEn_total_media": float(values.sum()),
                "Shannon_H": h,
                "Pielou_J": float(h / np.log(richness)) if richness > 1 else 0.0,
            }
        )
    return matrix, long, pd.DataFrame(summary_rows)


def similarity_outputs(matrix: pd.DataFrame):
    distance = squareform(pdist(matrix.to_numpy(float), metric="braycurtis"))
    similarity = 1 - distance
    sim = pd.DataFrame(similarity, index=matrix.index, columns=matrix.index)
    dist = pd.DataFrame(distance, index=matrix.index, columns=matrix.index)
    pairs = []
    for i, point_a in enumerate(matrix.index):
        for j in range(i + 1, len(matrix.index)):
            point_b = matrix.index[j]
            pairs.append(
                {
                    "ponto_A": point_a,
                    "ponto_B": point_b,
                    "similaridade_Bray_Curtis": float(sim.iloc[i, j]),
                    "distancia_Bray_Curtis": float(dist.iloc[i, j]),
                }
            )
    pairs = pd.DataFrame(pairs).sort_values("similaridade_Bray_Curtis", ascending=False).reset_index(drop=True)
    return sim, dist, pairs


def replace_sheets(path: Path, sheets: dict[str, pd.DataFrame]):
    workbook = load_workbook(path)
    for name, data in sheets.items():
        if name in workbook.sheetnames:
            del workbook[name]
        sheet = workbook.create_sheet(name)
        export = data.reset_index() if data.index.name or not isinstance(data.index, pd.RangeIndex) else data
        for row in dataframe_to_rows(export, index=False, header=True):
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column in range(1, sheet.max_column + 1):
            values = [sheet.cell(row, column).value for row in range(1, min(sheet.max_row, 250) + 1)]
            width = max(len(str(value or "")) for value in values) + 2
            sheet.column_dimensions[get_column_letter(column)].width = min(max(width, 12), 38)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
    temporary = path.with_name(f"{path.stem}.tmp.xlsx")
    workbook.save(temporary)
    os.replace(temporary, path)


def style_axes(ax, grid_axis="y"):
    ax.grid(False)
    ax.grid(axis=grid_axis, color=GRID, alpha=0.25, linewidth=1)
    for spine in ax.spines.values():
        spine.set_visible(True)
        spine.set_color("black")
        spine.set_linewidth(1.2)
    ax.tick_params(colors="black", labelsize=13)


def plot_diversity(summary: pd.DataFrame):
    x = np.arange(len(summary))
    fig, axes = plt.subplots(2, 1, figsize=(18, 10.2), dpi=300, sharex=True)
    axes[0].bar(x, summary["Shannon_H"], color=PRIMARY, width=0.72)
    axes[0].set_ylabel("Shannon H'")
    style_axes(axes[0])
    axes[1].bar(x, summary["Pielou_J"], color=SECONDARY, width=0.72)
    axes[1].set_ylabel("Pielou J'")
    axes[1].set_xlabel("Trecho amostral")
    axes[1].set_xticks(x, summary["ponto"], rotation=45, ha="right")
    axes[1].set_ylim(0, 1)
    style_axes(axes[1])
    fig.tight_layout()
    fig.savefig(FIG_DIVERSITY, dpi=300, facecolor="white", bbox_inches="tight")
    plt.close(fig)


def plot_similarity(matrix: pd.DataFrame):
    distances = pdist(matrix.to_numpy(float), metric="braycurtis")
    clusters = linkage(distances, method="average")
    fig, ax = plt.subplots(figsize=(18, 10.2), dpi=300)
    dendrogram(clusters, labels=matrix.index.tolist(), ax=ax, leaf_rotation=45, leaf_font_size=13)
    ax.set_ylabel("Distância Bray-Curtis")
    ax.set_xlabel("Trecho amostral")
    style_axes(ax)
    fig.tight_layout()
    fig.savefig(FIG_SIMILARITY, dpi=300, facecolor="white", bbox_inches="tight")
    plt.close(fig)


def shade_reservoir_events(ax, campaigns: pd.Series):
    numbers = campaigns.str.extract(r"C(\d{3})", expand=False).astype(int)
    pre_positions = np.flatnonzero(numbers.le(20).to_numpy())
    drawdown_positions = np.flatnonzero(numbers.between(62, 67).to_numpy())
    refill_positions = np.flatnonzero(numbers.eq(68).to_numpy())
    if len(pre_positions):
        ax.axvspan(pre_positions.min() - 0.5, pre_positions.max() + 0.5, color=PRE_FILL, zorder=0)
    if len(drawdown_positions):
        ax.axvspan(drawdown_positions.min() - 0.5, drawdown_positions.max() + 0.5, color=DRAWDOWN, zorder=0)
    if len(refill_positions):
        ax.axvline(refill_positions[0], color=REFILL, linestyle="--", linewidth=1.8, zorder=3)


def plot_diversity_temporal(diversity: pd.DataFrame):
    campaign = (
        diversity.groupby("campanha", as_index=False)
        .agg(Shannon_H=("Shannon_H", "mean"), Pielou_J=("Pielou_J", "mean"))
    )
    campaign["ordem"] = campaign["campanha"].str.extract(r"C(\d{3})", expand=False).astype(int)
    campaign = campaign.sort_values("ordem").reset_index(drop=True)
    x = np.arange(len(campaign))
    fig, ax = plt.subplots(figsize=(18, 10.2), dpi=300)
    shade_reservoir_events(ax, campaign["campanha"])
    line_h = ax.plot(x, campaign["Shannon_H"], color=PRIMARY, linewidth=2.8, label="Shannon H'")[0]
    ax2 = ax.twinx()
    line_j = ax2.plot(x, campaign["Pielou_J"], color=SECONDARY, linewidth=2.8, label="Pielou J'")[0]
    step = max(1, len(campaign) // 12)
    ticks = np.arange(0, len(campaign), step)
    ax.set_xticks(ticks, campaign.loc[ticks, "campanha"], rotation=45, ha="right")
    ax.set_ylabel("Shannon H'")
    ax2.set_ylabel("Pielou J'")
    style_axes(ax)
    style_axes(ax2)
    handles = [
        line_h,
        line_j,
        Patch(facecolor=PRE_FILL, edgecolor="none", label="Pre-enchimento"),
        Patch(facecolor=DRAWDOWN, edgecolor="none", label="Rebaixamento parcial"),
        Line2D([0], [0], color=REFILL, linestyle="--", linewidth=1.8, label="Reenchimento"),
    ]
    ax.legend(handles=handles, loc="upper center", bbox_to_anchor=(0.5, 1.08), ncol=5, frameon=False)
    fig.tight_layout()
    fig.savefig(FIG_DIVERSITY_TEMPORAL, dpi=300, facecolor="white", bbox_inches="tight")
    plt.close(fig)


def plot_similarity_temporal(similarity: pd.DataFrame):
    similarity = similarity.apply(pd.to_numeric, errors="coerce")
    distance = 1 - similarity.to_numpy(float)
    np.fill_diagonal(distance, 0)
    clusters = linkage(squareform(distance, checks=False), method="average")
    fig, ax = plt.subplots(figsize=(18, 10.2), dpi=300)
    dendrogram(
        clusters,
        labels=similarity.index.astype(str).tolist(),
        ax=ax,
        leaf_rotation=90,
        leaf_font_size=8,
    )
    ax.set_ylabel("Distancia Bray-Curtis")
    style_axes(ax)
    fig.tight_layout()
    fig.savefig(FIG_SIMILARITY_TEMPORAL, dpi=300, facecolor="white", bbox_inches="tight")
    plt.close(fig)


def restore_temporal_figures():
    diversity = pd.read_excel(BOOK_DIVERSITY, sheet_name="Diversidade")
    similarity = pd.read_excel(BOOK_SIMILARITY, sheet_name="Similaridade", index_col=0)
    plot_diversity_temporal(diversity)
    plot_similarity_temporal(similarity)


def main():
    plt.rcParams.update(
        {
            "font.family": "Arial",
            "font.size": 17,
            "axes.labelsize": 17,
            "xtick.labelsize": 13,
            "ytick.labelsize": 13,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "savefig.facecolor": "white",
        }
    )
    cpue = load_cpue_module()
    efforts, catches, _, _ = cpue.load_data()
    matrix, long, summary = build_point_matrix(efforts, catches)
    sim, dist, pairs = similarity_outputs(matrix)

    premises = pd.DataFrame(
        {
            "premissa": [
                "Universo",
                "Período",
                "Matriz comunitária",
                "Diversidade",
                "Equitabilidade",
                "Similaridade",
            ],
            "regra": [
                "16 pontos gerais; marcação e STP excluídos",
                "C001-C069, sem diferenciação entre campanhas",
                "CPUEn média por espécie e ponto, com ausências iguais a zero",
                "Shannon calculado sobre a matriz CPUEn integrada",
                "Pielou = Shannon / ln(riqueza)",
                "Bray-Curtis sobre a mesma matriz CPUEn; agrupamento UPGMA",
            ],
        }
    )
    replace_sheets(
        BOOK_DIVERSITY,
        {
            "Geral_por_ponto": summary,
            "Matriz_CPUEn_longa": long[["ponto", "nome_cientifico", "nome_popular", "CPUEn_media", "CPUEn_soma", "ocorrencias"]],
            "Premissas_pontos": premises,
        },
    )
    replace_sheets(
        BOOK_SIMILARITY,
        {
            "Geral_por_ponto": sim,
            "Distancia_por_ponto": dist,
            "Pares_pontos": pairs,
            "Premissas_pontos": premises,
        },
    )
    plot_diversity(summary)
    plot_similarity(matrix)
    restore_temporal_figures()
    print(
        {
            "pontos": len(summary),
            "especies_rede": matrix.shape[1],
            "unidades_esforco": len(efforts),
            "pares": len(pairs),
            "shannon_min": summary["Shannon_H"].min(),
            "shannon_max": summary["Shannon_H"].max(),
            "pielou_min": summary["Pielou_J"].min(),
            "pielou_max": summary["Pielou_J"].max(),
        }
    )


if __name__ == "__main__":
    if "--temporal-only" in sys.argv:
        plt.rcParams.update(
            {
                "font.family": "Arial",
                "font.size": 17,
                "axes.labelsize": 17,
                "xtick.labelsize": 13,
                "ytick.labelsize": 13,
                "legend.fontsize": 15,
                "figure.facecolor": "white",
                "axes.facecolor": "white",
                "savefig.facecolor": "white",
            }
        )
        restore_temporal_figures()
        print({"temporal_figures": [str(FIG_DIVERSITY_TEMPORAL), str(FIG_SIMILARITY_TEMPORAL)]})
    else:
        main()
