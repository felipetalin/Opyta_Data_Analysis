from __future__ import annotations

import re
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.lines import Line2D
from matplotlib.patches import Patch
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy.cluster.hierarchy import fcluster, linkage
from scipy.spatial import ConvexHull, QhullError
from scipy.spatial.distance import squareform
from matplotlib.transforms import Bbox

from gerar_5_7_cpue_biocol001 import POINTS, load_data


ANALYSIS_ROOT = Path("G:/Meu Drive/Opyta/Opyta_Data_Analysis")
OPYTA_ROOT = ANALYSIS_ROOT.parent
BIOS_ROOT = OPYTA_ROOT / "Clientes/Clientes/Clientes/Bios"
COLIDER_ROOT = next(path for path in BIOS_ROOT.glob("Col*der") if path.is_dir())
RESULTS_ROOT = COLIDER_ROOT / "Resultados/2026/Junho-2026"
FINAL_ROOT = RESULTS_ROOT / "BIOCOL001_RESULTADOS_ICTIOFAUNA_FINAL_R02"
OUTPUT_ROOT = RESULTS_ROOT / "BIOCOL001_ANALISES_EXPLORATORIAS_R01"
BASE_ANALYTICAL = FINAL_ROOT / "base_analitica_sem_marcacao_biocol001.xlsx"

BETA_FIGURE = OUTPUT_ROOT / "EXP_01_diversidade_beta_temporal_componentes_geral.png"
BETA_WORKBOOK = OUTPUT_ROOT / "EXP_01_diversidade_beta_temporal_componentes_geral.xlsx"
PCOA_FIGURE = OUTPUT_ROOT / "EXP_02_pcoa_biplot_bray_curtis_vetores_especies.png"
PCOA_WORKBOOK = OUTPUT_ROOT / "EXP_02_pcoa_biplot_bray_curtis_vetores_especies.xlsx"
TEMPORAL_PCOA_FIGURE = OUTPUT_ROOT / "EXP_03_pcoa_temporal_conjunta_fases_vetores_especies.png"
TEMPORAL_PCOA_WORKBOOK = OUTPUT_ROOT / "EXP_03_pcoa_temporal_conjunta_fases_vetores_especies.xlsx"

FIGSIZE = (18, 10.2)
DPI = 300
CUT_HEIGHT = 0.55
VECTOR_MIN_STRENGTH = 0.55
VECTOR_MAX_SPECIES = 4
VECTOR_MIN_POINTS = 3

COLORS = {
    "primary": "#002060",
    "secondary": "#5B9BD5",
    "nestedness": "#9DC3E6",
    "pre": "#F0F0F0",
    "lowering": "#DBE5F1",
    "refill": "#D4672A",
    "grid": "#D9D9D9",
}
GROUP_COLORS = ["#0B7A3B", "#6A8F2F", "#2A78B8", "#B1782C", "#7A4FA3", "#C43B3B", "#4C4C4C"]
PHASE_STYLES = {
    "pre": ("Pré-enchimento", "#7F7F7F", "o"),
    "pos": ("Pós-enchimento", "#002060", "o"),
    "rebaixamento": ("Rebaixamento parcial", "#5B9BD5", "s"),
    "reenchimento": ("Reenchimento", "#D4672A", "D"),
}

PERIODS = [
    ("pre", "Pré-enchimento", 1, 20),
    ("pos", "Pós-enchimento", 21, 61),
    ("rebaixamento", "Rebaixamento parcial", 62, 67),
    ("reenchimento", "Reenchimento", 68, 69),
]


def campaign_number(value: object) -> int:
    match = re.match(r"C(\d{3})", str(value))
    return int(match.group(1)) if match else 999


def period_key(sequence: int) -> str | None:
    for key, _label, start, end in PERIODS:
        if start <= sequence <= end:
            return key
    return None


def period_label(key: str) -> str:
    return next(label for item_key, label, _start, _end in PERIODS if item_key == key)


def style_axes(ax: plt.Axes, grid_axis: str = "y") -> None:
    ax.grid(False)
    ax.grid(axis=grid_axis, color=COLORS["grid"], alpha=0.35, linewidth=0.8, zorder=0)
    for spine in ax.spines.values():
        spine.set_visible(True)
        spine.set_color("black")
        spine.set_linewidth(1.0)
    ax.tick_params(colors="black", labelsize=12)


def beta_components(previous: np.ndarray, current: np.ndarray) -> tuple[float, float, float, int, int, int]:
    x = np.asarray(previous, dtype=bool)
    y = np.asarray(current, dtype=bool)
    shared = int(np.sum(x & y))
    lost = int(np.sum(x & ~y))
    gained = int(np.sum(~x & y))
    denominator_sorensen = 2 * shared + lost + gained
    beta_sorensen = (lost + gained) / denominator_sorensen if denominator_sorensen else 0.0
    denominator_turnover = shared + min(lost, gained)
    beta_turnover = min(lost, gained) / denominator_turnover if denominator_turnover else 0.0
    beta_nestedness = max(beta_sorensen - beta_turnover, 0.0)
    return beta_sorensen, beta_turnover, beta_nestedness, shared, lost, gained


def build_beta_temporal() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    base = pd.read_excel(BASE_ANALYTICAL, sheet_name="base")
    base["ordem_campanha"] = base["campanha"].map(campaign_number)
    base["numero_de_individuos"] = pd.to_numeric(base["numero_de_individuos"], errors="coerce").fillna(0)
    observed = base[
        base["ponto"].isin(POINTS)
        & base["ordem_campanha"].between(1, 69)
        & base["nome_cientifico"].notna()
        & base["numero_de_individuos"].gt(0)
    ].copy()
    observed["nome_cientifico"] = observed["nome_cientifico"].astype(str).str.strip()
    observed = observed[observed["nome_cientifico"].ne("")]

    campaigns = (
        observed[["campanha", "ordem_campanha"]]
        .drop_duplicates()
        .sort_values("ordem_campanha")
        .reset_index(drop=True)
    )
    species = sorted(observed["nome_cientifico"].unique())
    matrix = (
        observed.assign(Presenca=1)
        .pivot_table(index="campanha", columns="nome_cientifico", values="Presenca", aggfunc="max", fill_value=0)
        .reindex(index=campaigns["campanha"], columns=species, fill_value=0)
        .astype(int)
    )

    rows: list[dict[str, object]] = []
    for idx in range(1, len(campaigns)):
        previous = campaigns.iloc[idx - 1]
        current = campaigns.iloc[idx]
        beta_sor, beta_sim, beta_nes, shared, lost, gained = beta_components(
            matrix.iloc[idx - 1].to_numpy(), matrix.iloc[idx].to_numpy()
        )
        rows.append(
            {
                "campanha_anterior": previous["campanha"],
                "ordem_anterior": int(previous["ordem_campanha"]),
                "campanha": current["campanha"],
                "ordem_campanha": int(current["ordem_campanha"]),
                "periodo_destino": period_label(period_key(int(current["ordem_campanha"])) or "pos"),
                "riqueza_anterior": int(matrix.iloc[idx - 1].sum()),
                "riqueza_atual": int(matrix.iloc[idx].sum()),
                "especies_compartilhadas": shared,
                "especies_perdidas": lost,
                "especies_ganhas": gained,
                "Beta_Sorensen": beta_sor,
                "Turnover_BetaSim": beta_sim,
                "Nestedness_BetaNes": beta_nes,
            }
        )
    beta = pd.DataFrame(rows)
    matrix_output = matrix.reset_index().merge(campaigns, on="campanha", how="left")
    leading = ["campanha", "ordem_campanha"]
    matrix_output = matrix_output[leading + [column for column in matrix_output.columns if column not in leading]]
    audit = pd.DataFrame(
        [
            {
                "campanhas": campaigns["campanha"].nunique(),
                "comparacoes_consecutivas": len(beta),
                "pontos_regulares": observed["ponto"].nunique(),
                "especies_no_universo": len(species),
                "metodos_incidencia": observed["metodo_de_captura"].nunique(),
                "identidade_beta_max_erro": float(
                    np.max(np.abs(beta["Beta_Sorensen"] - beta["Turnover_BetaSim"] - beta["Nestedness_BetaNes"]))
                ),
            }
        ]
    )
    return beta, matrix_output, audit


def plot_beta_temporal(beta: pd.DataFrame) -> None:
    fig, ax = plt.subplots(figsize=FIGSIZE, dpi=DPI)
    ax.axvspan(1.5, 20.5, color=COLORS["pre"], alpha=0.85, linewidth=0, zorder=0)
    ax.axvspan(61.5, 67.5, color=COLORS["lowering"], alpha=0.85, linewidth=0, zorder=0)
    ax.axvline(67.5, color=COLORS["refill"], linewidth=1.7, linestyle="--", zorder=1)

    lines = [
        ("Beta_Sorensen", "β-Sørensen (β-sor)", COLORS["primary"]),
        ("Turnover_BetaSim", "Turnover (β-sim)", COLORS["secondary"]),
        ("Nestedness_BetaNes", "Nestedness (β-nes)", COLORS["nestedness"]),
    ]
    for column, label, color in lines:
        ax.plot(
            beta["ordem_campanha"],
            beta[column],
            color=color,
            marker="o",
            markersize=4.8,
            linewidth=2.4,
            alpha=0.92,
            label=label,
            zorder=3,
        )

    ticks = sorted(set(range(2, 70, 4)) | {20, 21, 62, 68, 69})
    labels = beta.set_index("ordem_campanha")["campanha"].to_dict()
    ax.set_xticks(ticks)
    ax.set_xticklabels([str(labels.get(tick, f"C{tick:03d}" )).split("-")[0] for tick in ticks], rotation=45, ha="right")
    ax.set_xlim(1.5, 69.5)
    ax.set_ylim(0, 1.02)
    ax.set_xlabel("Campanha")
    ax.set_ylabel("Dissimilaridade")
    style_axes(ax)

    handles = [
        Line2D([0], [0], color=color, marker="o", linewidth=2.4, label=label)
        for _column, label, color in lines
    ]
    handles.extend(
        [
            Patch(facecolor=COLORS["pre"], edgecolor="none", label="Pré-enchimento"),
            Patch(facecolor=COLORS["lowering"], edgecolor="none", label="Rebaixamento parcial"),
            Line2D([0], [0], color=COLORS["refill"], linestyle="--", linewidth=1.7, label="Reenchimento"),
        ]
    )
    ax.legend(
        handles=handles,
        loc="upper center",
        bbox_to_anchor=(0.5, 1.13),
        ncol=6,
        frameon=False,
        fontsize=13,
        handlelength=2.4,
        columnspacing=1.3,
    )
    fig.subplots_adjust(left=0.075, right=0.985, bottom=0.14, top=0.86)
    fig.savefig(BETA_FIGURE, dpi=DPI, facecolor="white")
    plt.close(fig)


def bray_curtis_matrix(values: np.ndarray) -> np.ndarray:
    n_samples = values.shape[0]
    distance = np.zeros((n_samples, n_samples), dtype=float)
    for i in range(n_samples):
        for j in range(i + 1, n_samples):
            denominator = float(np.sum(values[i] + values[j]))
            value = 0.0 if denominator == 0 else float(np.sum(np.abs(values[i] - values[j])) / denominator)
            distance[i, j] = value
            distance[j, i] = value
    return distance


def pcoa(distance: np.ndarray) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    n_samples = distance.shape[0]
    squared = distance**2
    centering = np.eye(n_samples) - np.ones((n_samples, n_samples)) / n_samples
    matrix_b = -0.5 * centering @ squared @ centering
    eigenvalues, eigenvectors = np.linalg.eigh(matrix_b)
    order = np.argsort(eigenvalues)[::-1]
    eigenvalues = eigenvalues[order]
    eigenvectors = eigenvectors[:, order]
    positive = np.maximum(eigenvalues[:2], 0)
    coordinates = eigenvectors[:, :2] * np.sqrt(positive)
    positive_total = float(eigenvalues[eigenvalues > 0].sum())
    explained = positive / positive_total if positive_total > 0 else np.array([0.0, 0.0])
    return coordinates, explained, eigenvalues


def build_period_tables() -> tuple[dict[str, pd.DataFrame], pd.DataFrame, pd.DataFrame]:
    efforts, catches, _units, _species = load_data()
    efforts = efforts.copy()
    efforts["periodo"] = efforts["ordem_campanha"].map(period_key)
    efforts = efforts[efforts["periodo"].notna()].copy()
    catches = catches[catches["id_esforco"].isin(efforts["id_esforco"])].copy()

    names = (
        catches[["id_especie", "nome_cientifico", "nome_popular"]]
        .drop_duplicates("id_especie")
        .sort_values("nome_cientifico")
    )
    species_ids = names["id_especie"].tolist()
    grid = pd.MultiIndex.from_product(
        [efforts["id_esforco"].tolist(), species_ids], names=["id_esforco", "id_especie"]
    ).to_frame(index=False)
    grid = grid.merge(
        efforts[["id_esforco", "campanha", "ordem_campanha", "ponto", "esforco_m2", "periodo"]],
        on="id_esforco",
        how="left",
    )
    grid = grid.merge(catches[["id_esforco", "id_especie", "abundancia"]], on=["id_esforco", "id_especie"], how="left")
    grid["abundancia"] = pd.to_numeric(grid["abundancia"], errors="coerce").fillna(0)
    grid["CPUEn"] = grid["abundancia"] / grid["esforco_m2"] * 100
    grid = grid.merge(names, on="id_especie", how="left")

    period_species = (
        grid.groupby(["periodo", "ponto", "nome_cientifico"], as_index=False)["CPUEn"]
        .mean()
        .rename(columns={"CPUEn": "CPUEn_media"})
    )
    species_names = names["nome_cientifico"].tolist()
    tables: dict[str, pd.DataFrame] = {}
    for key, _label, _start, _end in PERIODS:
        table = (
            period_species[period_species["periodo"].eq(key)]
            .pivot_table(index="ponto", columns="nome_cientifico", values="CPUEn_media", aggfunc="mean", fill_value=0)
            .reindex(index=POINTS, columns=species_names, fill_value=0)
        )
        tables[key] = table

    audit = (
        efforts.groupby("periodo", as_index=False)
        .agg(campanhas=("campanha", "nunique"), pontos=("ponto", "nunique"), unidades_esforco=("id_esforco", "size"))
    )
    audit["periodo_label"] = audit["periodo"].map(period_label)
    audit["ordem_periodo"] = audit["periodo"].map({key: idx for idx, (key, *_rest) in enumerate(PERIODS, start=1)})
    audit = audit.sort_values("ordem_periodo")
    return tables, audit, names


def species_vectors(
    transformed: pd.DataFrame,
    raw: pd.DataFrame,
    coordinates: np.ndarray,
    min_occurrence: int = VECTOR_MIN_POINTS,
) -> pd.DataFrame:
    records: list[dict[str, object]] = []
    for species in transformed.columns:
        values = transformed[species].to_numpy(dtype=float)
        occurrence = int((raw[species] > 0).sum())
        if occurrence < min_occurrence or np.nanstd(values) == 0:
            continue
        corr1 = float(np.corrcoef(values, coordinates[:, 0])[0, 1]) if np.nanstd(coordinates[:, 0]) > 0 else 0.0
        corr2 = float(np.corrcoef(values, coordinates[:, 1])[0, 1]) if np.nanstd(coordinates[:, 1]) > 0 else 0.0
        corr1 = 0.0 if np.isnan(corr1) else corr1
        corr2 = 0.0 if np.isnan(corr2) else corr2
        strength = float(np.hypot(corr1, corr2))
        records.append(
            {
                "nome_cientifico": species,
                "n_pontos_ocorrencia": occurrence,
                "corr_pcoa1": corr1,
                "corr_pcoa2": corr2,
                "forca_associacao": strength,
            }
        )
    vectors = pd.DataFrame(records).sort_values("forca_associacao", ascending=False)
    selected = vectors[vectors["forca_associacao"].ge(VECTOR_MIN_STRENGTH)].head(VECTOR_MAX_SPECIES)
    if selected.empty:
        selected = vectors.head(min(3, len(vectors)))
    return selected.reset_index(drop=True)


def build_pcoa_products(
    tables: dict[str, pd.DataFrame]
) -> tuple[dict[str, dict[str, object]], pd.DataFrame]:
    products: dict[str, dict[str, object]] = {}
    audit_rows: list[dict[str, object]] = []
    for key, label, start, end in PERIODS:
        raw = tables[key]
        transformed = np.sqrt(raw)
        distance = bray_curtis_matrix(transformed.to_numpy(dtype=float))
        coordinates, explained, eigenvalues = pcoa(distance)
        hierarchy = linkage(squareform(distance, checks=False), method="average")
        groups = fcluster(hierarchy, t=CUT_HEIGHT, criterion="distance")
        vectors = species_vectors(transformed, raw, coordinates)
        scores = pd.DataFrame(
            {
                "ponto": raw.index,
                "grupo_bray": groups.astype(int),
                "PCoA1": coordinates[:, 0],
                "PCoA2": coordinates[:, 1],
                "PCoA1_variancia_pct": explained[0] * 100,
                "PCoA2_variancia_pct": explained[1] * 100,
            }
        )
        products[key] = {
            "label": label,
            "start": start,
            "end": end,
            "raw": raw,
            "transformed": transformed,
            "distance": distance,
            "scores": scores,
            "vectors": vectors,
        }
        negative_sum = float(np.abs(eigenvalues[eigenvalues < 0]).sum())
        positive_sum = float(eigenvalues[eigenvalues > 0].sum())
        audit_rows.append(
            {
                "periodo": label,
                "campanhas": f"C{start:03d}-C{end:03d}",
                "pontos": len(raw),
                "especies_com_cpuen": int((raw.sum(axis=0) > 0).sum()),
                "grupos_bray_corte_0_55": int(len(np.unique(groups))),
                "PCoA1_variancia_pct": explained[0] * 100,
                "PCoA2_variancia_pct": explained[1] * 100,
                "razao_modulo_autovalores_negativos": negative_sum / positive_sum if positive_sum else np.nan,
                "vetores_exibidos": len(vectors),
            }
        )
    return products, pd.DataFrame(audit_rows)


def overlap_area(first, second) -> float:
    width = max(0.0, min(first.x1, second.x1) - max(first.x0, second.x0))
    height = max(0.0, min(first.y1, second.y1) - max(first.y0, second.y0))
    return width * height


def place_annotation(
    ax: plt.Axes,
    xy: tuple[float, float],
    label: str,
    occupied: list,
    *,
    fontsize: float,
    fontstyle: str = "normal",
    color: str = "#222222",
    leader: bool = False,
) -> None:
    preferred = [
        (8, 8),
        (-8, 8),
        (8, -12),
        (-8, -12),
        (16, 1),
        (-16, 1),
        (1, 16),
        (1, -18),
        (22, 12),
        (-22, 12),
        (22, -16),
        (-22, -16),
        (32, 2),
        (-32, 2),
    ]
    extended = []
    for radius in (42, 56, 72, 90):
        extended.extend(
            [
                (radius, 0),
                (-radius, 0),
                (0, radius),
                (0, -radius),
                (radius, radius // 2),
                (-radius, radius // 2),
                (radius, -radius // 2),
                (-radius, -radius // 2),
            ]
        )
    candidates = preferred + extended
    figure = ax.figure
    renderer = figure.canvas.get_renderer()
    axes_box = ax.get_window_extent(renderer=renderer)
    points_to_pixels = figure.dpi / 72.0
    anchor_x, anchor_y = ax.transData.transform(xy)
    estimated_width = max(fontsize * 0.60 * len(label) + 6, fontsize * 2.5) * points_to_pixels
    estimated_height = fontsize * 1.35 * points_to_pixels

    best_candidate = candidates[0]
    best_score = float("inf")
    for candidate in candidates:
        center_x = anchor_x + candidate[0] * points_to_pixels
        center_y = anchor_y + candidate[1] * points_to_pixels
        candidate_box = Bbox.from_extents(
            center_x - estimated_width / 2,
            center_y - estimated_height / 2,
            center_x + estimated_width / 2,
            center_y + estimated_height / 2,
        )
        score = sum(overlap_area(candidate_box, box) for box in occupied)
        if (
            candidate_box.x0 < axes_box.x0 + 2
            or candidate_box.x1 > axes_box.x1 - 2
            or candidate_box.y0 < axes_box.y0 + 2
            or candidate_box.y1 > axes_box.y1 - 2
        ):
            score += 1_000_000_000
        if score < best_score:
            best_score = score
            best_candidate = candidate
            if score == 0:
                break

    center_x = anchor_x + best_candidate[0] * points_to_pixels
    center_y = anchor_y + best_candidate[1] * points_to_pixels
    occupied.append(
        Bbox.from_extents(
            center_x - estimated_width / 2,
            center_y - estimated_height / 2,
            center_x + estimated_width / 2,
            center_y + estimated_height / 2,
        )
    )
    ax.annotate(
        label,
        xy=xy,
        xytext=best_candidate,
        textcoords="offset points",
        fontsize=fontsize,
        fontstyle=fontstyle,
        color=color,
        ha="center",
        va="center",
        arrowprops={"arrowstyle": "-", "color": "#777777", "lw": 0.55, "alpha": 0.75} if leader else None,
        zorder=5,
    )


def plot_pcoa(products: dict[str, dict[str, object]]) -> None:
    fig, axes = plt.subplots(2, 2, figsize=FIGSIZE, dpi=100)
    axes = axes.ravel()
    fig.subplots_adjust(left=0.065, right=0.985, bottom=0.075, top=0.94, wspace=0.20, hspace=0.30)

    for ax, (key, _label, start, end) in zip(axes, PERIODS, strict=True):
        product = products[key]
        scores = product["scores"]
        assert isinstance(scores, pd.DataFrame)
        vectors = product["vectors"]
        assert isinstance(vectors, pd.DataFrame)
        point_labels: list[tuple[float, float, str]] = []
        vector_labels: list[tuple[float, float, str]] = []

        for group_id in sorted(scores["grupo_bray"].unique()):
            subset = scores[scores["grupo_bray"].eq(group_id)]
            color = GROUP_COLORS[(int(group_id) - 1) % len(GROUP_COLORS)]
            ax.scatter(
                subset["PCoA1"],
                subset["PCoA2"],
                s=95,
                color=color,
                edgecolor="#1F1F1F",
                linewidth=0.65,
                label=f"G{int(group_id)}",
                zorder=3,
            )
            if len(subset) >= 3:
                try:
                    hull = ConvexHull(subset[["PCoA1", "PCoA2"]].to_numpy())
                    hull_points = subset[["PCoA1", "PCoA2"]].to_numpy()[hull.vertices]
                    ax.fill(hull_points[:, 0], hull_points[:, 1], color=color, alpha=0.12, zorder=1)
                except QhullError:
                    pass
            point_labels.extend(
                (float(row.PCoA1), float(row.PCoA2), str(row.ponto)) for row in subset.itertuples(index=False)
            )

        x_span = max(float(np.ptp(scores["PCoA1"])), 0.1)
        y_span = max(float(np.ptp(scores["PCoA2"])), 0.1)
        vector_scale = 0.35 * min(x_span, y_span)
        for row in vectors.itertuples(index=False):
            dx = row.corr_pcoa1 * vector_scale
            dy = row.corr_pcoa2 * vector_scale
            ax.annotate(
                "",
                xy=(dx, dy),
                xytext=(0, 0),
                arrowprops={"arrowstyle": "-|>", "color": "#333333", "lw": 1.2, "alpha": 0.88},
                zorder=4,
            )
            vector_labels.append((float(dx), float(dy), str(row.nome_cientifico)))

        explained_1 = float(scores["PCoA1_variancia_pct"].iloc[0])
        explained_2 = float(scores["PCoA2_variancia_pct"].iloc[0])
        ax.axhline(0, color="#D0D0D0", linewidth=0.8, zorder=0)
        ax.axvline(0, color="#D0D0D0", linewidth=0.8, zorder=0)
        ax.set_title(
            f"{product['label']}\nC{start:03d}-C{end:03d}",
            fontsize=12.5,
            fontweight="bold",
            color=COLORS["primary"],
            pad=8,
        )
        ax.set_xlabel(f"PCoA1 ({explained_1:.1f}%)")
        ax.set_ylabel(f"PCoA2 ({explained_2:.1f}%)")
        style_axes(ax, grid_axis="both")
        legend_location = "upper left" if key == "rebaixamento" else "upper right"
        legend = ax.legend(
            loc=legend_location,
            fontsize=8.2,
            frameon=False,
            ncol=2,
            handletextpad=0.3,
            columnspacing=0.8,
        )
        ax.margins(x=0.16, y=0.18)

        fig.canvas.draw()
        occupied = [legend.get_window_extent(renderer=fig.canvas.get_renderer()).expanded(1.05, 1.08)]
        marker_radius = 7.5 * fig.dpi / 72.0
        for x_value, y_value, _label in point_labels:
            display_x, display_y = ax.transData.transform((x_value, y_value))
            occupied.append(
                Bbox.from_extents(
                    display_x - marker_radius,
                    display_y - marker_radius,
                    display_x + marker_radius,
                    display_y + marker_radius,
                )
            )
        for x_value, y_value, label in sorted(vector_labels, key=lambda item: len(item[2]), reverse=True):
            place_annotation(
                ax,
                (x_value, y_value),
                label,
                occupied,
                fontsize=8.0,
                fontstyle="italic",
                leader=True,
            )

        coordinates = np.array([(x_value, y_value) for x_value, y_value, _label in point_labels])
        if len(coordinates) > 1:
            standardized = coordinates / np.maximum(np.ptp(coordinates, axis=0), 1e-9)
            pairwise = np.sqrt(((standardized[:, None, :] - standardized[None, :, :]) ** 2).sum(axis=2))
            pairwise[pairwise == 0] = np.inf
            priority = np.argsort(pairwise.min(axis=1))
        else:
            priority = np.arange(len(point_labels))
        for point_idx in priority:
            x_value, y_value, label = point_labels[int(point_idx)]
            place_annotation(ax, (x_value, y_value), label, occupied, fontsize=8.0, leader=True)

    fig.savefig(PCOA_FIGURE, dpi=DPI, facecolor="white")
    plt.close(fig)


def build_campaign_table() -> tuple[pd.DataFrame, pd.DataFrame]:
    efforts, catches, _units, _species = load_data()
    efforts = efforts.copy()
    efforts["periodo"] = efforts["ordem_campanha"].map(period_key)
    efforts = efforts[efforts["periodo"].notna()].copy()
    catches = catches[catches["id_esforco"].isin(efforts["id_esforco"])].copy()
    names = catches[["id_especie", "nome_cientifico"]].drop_duplicates("id_especie").sort_values("nome_cientifico")

    grid = pd.MultiIndex.from_product(
        [efforts["id_esforco"].tolist(), names["id_especie"].tolist()],
        names=["id_esforco", "id_especie"],
    ).to_frame(index=False)
    grid = grid.merge(
        efforts[["id_esforco", "campanha", "ordem_campanha", "ponto", "esforco_m2", "periodo"]],
        on="id_esforco",
        how="left",
    )
    grid = grid.merge(catches[["id_esforco", "id_especie", "abundancia"]], on=["id_esforco", "id_especie"], how="left")
    grid["abundancia"] = pd.to_numeric(grid["abundancia"], errors="coerce").fillna(0)
    grid["CPUEn"] = grid["abundancia"] / grid["esforco_m2"] * 100
    grid = grid.merge(names, on="id_especie", how="left")

    campaign_species = (
        grid.groupby(["campanha", "ordem_campanha", "periodo", "nome_cientifico"], as_index=False)["CPUEn"]
        .mean()
        .rename(columns={"CPUEn": "CPUEn_media"})
    )
    metadata = (
        efforts.groupby(["campanha", "ordem_campanha", "periodo"], as_index=False)
        .agg(pontos_amostrados=("ponto", "nunique"), unidades_esforco=("id_esforco", "size"))
        .sort_values("ordem_campanha")
    )
    table = (
        campaign_species.pivot_table(
            index="campanha", columns="nome_cientifico", values="CPUEn_media", aggfunc="mean", fill_value=0
        )
        .reindex(index=metadata["campanha"], columns=names["nome_cientifico"], fill_value=0)
    )
    return table, metadata


def silhouette_mean(distance: np.ndarray, labels: np.ndarray) -> float:
    labels = np.asarray(labels)
    unique = np.unique(labels)
    if len(unique) < 2 or len(unique) >= len(labels):
        return float("nan")
    values: list[float] = []
    for idx, label in enumerate(labels):
        own = np.where(labels == label)[0]
        own = own[own != idx]
        if len(own) == 0:
            values.append(0.0)
            continue
        within = float(distance[idx, own].mean())
        between = min(float(distance[idx, np.where(labels == other)[0]].mean()) for other in unique if other != label)
        values.append((between - within) / max(within, between) if max(within, between) > 0 else 0.0)
    return float(np.mean(values))


def build_temporal_pcoa(
    raw: pd.DataFrame,
    metadata: pd.DataFrame,
) -> tuple[dict[str, object], pd.DataFrame, pd.DataFrame]:
    transformed = np.sqrt(raw)
    distance = bray_curtis_matrix(transformed.to_numpy(dtype=float))
    coordinates, explained, eigenvalues = pcoa(distance)
    hierarchy = linkage(squareform(distance, checks=False), method="average")
    groups = fcluster(hierarchy, t=CUT_HEIGHT, criterion="distance")
    vectors = species_vectors(transformed, raw, coordinates, min_occurrence=10)
    scores = metadata.copy()
    scores["periodo_label"] = scores["periodo"].map(period_label)
    scores["grupo_bray_055"] = groups.astype(int)
    scores["PCoA1"] = coordinates[:, 0]
    scores["PCoA2"] = coordinates[:, 1]
    scores["PCoA1_variancia_pct"] = explained[0] * 100
    scores["PCoA2_variancia_pct"] = explained[1] * 100

    contingency = pd.crosstab(scores["periodo_label"], scores["grupo_bray_055"]).reset_index()
    negative_sum = float(np.abs(eigenvalues[eigenvalues < 0]).sum())
    positive_sum = float(eigenvalues[eigenvalues > 0].sum())
    audit = pd.DataFrame(
        [
            {
                "campanhas": len(scores),
                "especies_com_cpuen": int((raw.sum(axis=0) > 0).sum()),
                "grupos_bray_corte_0_55": int(len(np.unique(groups))),
                "PCoA1_variancia_pct": explained[0] * 100,
                "PCoA2_variancia_pct": explained[1] * 100,
                "silhouette_fases": silhouette_mean(distance, scores["periodo"].to_numpy()),
                "silhouette_grupos_055": silhouette_mean(distance, groups),
                "razao_modulo_autovalores_negativos": negative_sum / positive_sum if positive_sum else np.nan,
                "vetores_exibidos": len(vectors),
            }
        ]
    )
    return {
        "raw": raw,
        "transformed": transformed,
        "distance": distance,
        "scores": scores,
        "vectors": vectors,
    }, contingency, audit


def plot_temporal_pcoa(product: dict[str, object]) -> None:
    scores = product["scores"]
    vectors = product["vectors"]
    assert isinstance(scores, pd.DataFrame)
    assert isinstance(vectors, pd.DataFrame)

    fig, ax = plt.subplots(figsize=FIGSIZE, dpi=100)
    fig.subplots_adjust(left=0.075, right=0.985, bottom=0.09, top=0.85)
    ax.plot(scores["PCoA1"], scores["PCoA2"], color="#BFBFBF", linewidth=1.0, alpha=0.65, zorder=1)

    for key, (label, color, marker) in PHASE_STYLES.items():
        subset = scores[scores["periodo"].eq(key)]
        ax.scatter(
            subset["PCoA1"],
            subset["PCoA2"],
            s=72,
            color=color,
            marker=marker,
            edgecolor="#1F1F1F",
            linewidth=0.55,
            alpha=0.90,
            label=label,
            zorder=3,
        )
        if len(subset) >= 3:
            try:
                hull = ConvexHull(subset[["PCoA1", "PCoA2"]].to_numpy())
                hull_points = subset[["PCoA1", "PCoA2"]].to_numpy()[hull.vertices]
                ax.fill(hull_points[:, 0], hull_points[:, 1], color=color, alpha=0.09, zorder=0)
            except QhullError:
                pass

    x_span = max(float(np.ptp(scores["PCoA1"])), 0.1)
    y_span = max(float(np.ptp(scores["PCoA2"])), 0.1)
    vector_scale = 0.33 * min(x_span, y_span)
    vector_labels: list[tuple[float, float, str]] = []
    for row in vectors.itertuples(index=False):
        dx = row.corr_pcoa1 * vector_scale
        dy = row.corr_pcoa2 * vector_scale
        ax.annotate(
            "",
            xy=(dx, dy),
            xytext=(0, 0),
            arrowprops={"arrowstyle": "-|>", "color": "#333333", "lw": 1.25, "alpha": 0.88},
            zorder=4,
        )
        vector_labels.append((float(dx), float(dy), str(row.nome_cientifico)))

    ax.axhline(0, color="#D0D0D0", linewidth=0.8, zorder=0)
    ax.axvline(0, color="#D0D0D0", linewidth=0.8, zorder=0)
    ax.set_xlabel(f"PCoA1 ({float(scores['PCoA1_variancia_pct'].iloc[0]):.1f}%)")
    ax.set_ylabel(f"PCoA2 ({float(scores['PCoA2_variancia_pct'].iloc[0]):.1f}%)")
    style_axes(ax, grid_axis="both")
    legend = ax.legend(
        loc="upper center",
        bbox_to_anchor=(0.5, 1.16),
        ncol=4,
        frameon=False,
        fontsize=12,
        columnspacing=1.4,
    )
    ax.margins(x=0.13, y=0.16)

    fig.canvas.draw()
    occupied = [legend.get_window_extent(renderer=fig.canvas.get_renderer()).expanded(1.03, 1.06)]
    marker_radius = 6.5 * fig.dpi / 72.0
    for row in scores.itertuples(index=False):
        display_x, display_y = ax.transData.transform((row.PCoA1, row.PCoA2))
        occupied.append(
            Bbox.from_extents(
                display_x - marker_radius,
                display_y - marker_radius,
                display_x + marker_radius,
                display_y + marker_radius,
            )
        )
    for x_value, y_value, label in sorted(vector_labels, key=lambda item: len(item[2]), reverse=True):
        place_annotation(
            ax,
            (x_value, y_value),
            label,
            occupied,
            fontsize=9.2,
            fontstyle="italic",
            leader=True,
        )
    key_campaigns = {1, 20, 21, 62, 67, 68, 69}
    for row in scores[scores["ordem_campanha"].isin(key_campaigns)].itertuples(index=False):
        place_annotation(ax, (row.PCoA1, row.PCoA2), row.campanha.split("-")[0], occupied, fontsize=9.2, leader=True)

    fig.savefig(TEMPORAL_PCOA_FIGURE, dpi=DPI, facecolor="white")
    plt.close(fig)


def write_temporal_pcoa_workbook(
    product: dict[str, object],
    contingency: pd.DataFrame,
    audit: pd.DataFrame,
) -> None:
    raw = product["raw"]
    transformed = product["transformed"]
    distance = product["distance"]
    scores = product["scores"]
    vectors = product["vectors"]
    assert isinstance(raw, pd.DataFrame)
    assert isinstance(transformed, pd.DataFrame)
    assert isinstance(distance, np.ndarray)
    assert isinstance(scores, pd.DataFrame)
    assert isinstance(vectors, pd.DataFrame)
    premises = pd.DataFrame(
        [
            ["Pergunta", "As campanhas se organizam segundo as fases e eventos do reservatorio?"],
            ["Unidade", "Campanha; CPUEn media por especie entre os pontos com esforco validado"],
            ["Universo", "C001-C069; 16 pontos regulares; rede de emalhar quantitativa"],
            ["Transformacao", "Raiz quadrada da CPUEn antes de Bray-Curtis"],
            ["Ordenacao", "Uma unica PCoA para todas as 69 campanhas"],
            ["Grupos", "UPGMA/ligacao media, corte Bray-Curtis 0,55"],
            ["Vetores", "Ate 4 especies com forca >= 0,55 e presenca em pelo menos 10 campanhas"],
            ["Cautela", "Agrupamento ou separacao nao demonstra causalidade; controlar hidrologia em etapa posterior"],
        ],
        columns=["Premissa", "Regra"],
    )
    with pd.ExcelWriter(TEMPORAL_PCOA_WORKBOOK, engine="openpyxl") as writer:
        premises.to_excel(writer, sheet_name="Premissas", index=False)
        audit.to_excel(writer, sheet_name="Auditoria", index=False)
        contingency.to_excel(writer, sheet_name="Fase_x_grupo_055", index=False)
        raw.reset_index(names="campanha").to_excel(writer, sheet_name="CPUEn_campanha", index=False)
        transformed.reset_index(names="campanha").to_excel(writer, sheet_name="Raiz_CPUEn", index=False)
        pd.DataFrame(distance, index=raw.index, columns=raw.index).reset_index(names="campanha").to_excel(
            writer, sheet_name="Bray_campanhas", index=False
        )
        scores.to_excel(writer, sheet_name="PCoA_campanhas", index=False)
        vectors.to_excel(writer, sheet_name="Vetores_especies", index=False)
    style_workbook(TEMPORAL_PCOA_WORKBOOK)


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="002060")
    header_font = Font(color="FFFFFF", bold=True)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column_idx, column in enumerate(worksheet.iter_cols(), start=1):
            sample = list(column[: min(len(column), 250)])
            width = max((len(str(cell.value)) if cell.value is not None else 0 for cell in sample), default=0)
            worksheet.column_dimensions[get_column_letter(column_idx)].width = min(max(width + 2, 11), 42)
    workbook.save(path)


def write_beta_workbook(beta: pd.DataFrame, matrix: pd.DataFrame, audit: pd.DataFrame) -> None:
    premises = pd.DataFrame(
        [
            ["Universo", "16 pontos regulares; ICTIO13C/D excluidos; marcacao de ICTIO13A excluida pela base analitica"],
            ["Incidencia", "Presenca/ausencia por campanha, reunindo os metodos de captura validados"],
            ["Comparacao", "Campanhas consecutivas C001-C069, sem separacao por trecho"],
            ["Beta total", "Dissimilaridade de Sorensen (beta-sor)"],
            ["Turnover", "Componente de substituicao de especies de Baselga (beta-sim)"],
            ["Nestedness", "Componente de aninhamento (beta-nes = beta-sor - beta-sim)"],
            ["Interpretacao", "Analise exploratoria; dissimilaridade nao implica impacto ou causalidade"],
        ],
        columns=["Premissa", "Regra"],
    )
    with pd.ExcelWriter(BETA_WORKBOOK, engine="openpyxl") as writer:
        premises.to_excel(writer, sheet_name="Premissas", index=False)
        audit.to_excel(writer, sheet_name="Auditoria", index=False)
        beta.to_excel(writer, sheet_name="Beta_temporal", index=False)
        matrix.to_excel(writer, sheet_name="Matriz_PA", index=False)
    style_workbook(BETA_WORKBOOK)


def write_pcoa_workbook(
    products: dict[str, dict[str, object]],
    effort_audit: pd.DataFrame,
    pcoa_audit: pd.DataFrame,
) -> None:
    premises = pd.DataFrame(
        [
            ["Universo", "16 pontos regulares; somente rede de emalhar quantitativa com esforco validado"],
            ["Metrica", "CPUEn media por especie e ponto em cada periodo, incluindo zeros quando houve esforco"],
            ["Transformacao", "Raiz quadrada da CPUEn antes do calculo de Bray-Curtis"],
            ["Ordenacao", "PCoA classica sobre dissimilaridade de Bray-Curtis"],
            ["Grupos", "UPGMA/ligacao media, corte de Bray-Curtis em 0,55, aprovado apos teste de sensibilidade"],
            ["Vetores", "Correlacao das especies com PCoA1/PCoA2; forca >= 0,55; ate 4 especies"],
            ["Robustez vetores", "Somente especies presentes em pelo menos 3 pontos no periodo"],
            ["Periodos", "Pre C001-C020; pos C021-C061; rebaixamento C062-C067; reenchimento C068-C069"],
            ["Interpretacao", "Grupos sao internos a cada painel e nao equivalentes entre periodos; analise exploratoria"],
        ],
        columns=["Premissa", "Regra"],
    )
    with pd.ExcelWriter(PCOA_WORKBOOK, engine="openpyxl") as writer:
        premises.to_excel(writer, sheet_name="Premissas", index=False)
        effort_audit.to_excel(writer, sheet_name="Auditoria_esforco", index=False)
        pcoa_audit.to_excel(writer, sheet_name="Auditoria_PCoA", index=False)
        for key, _label, _start, _end in PERIODS:
            product = products[key]
            raw = product["raw"]
            transformed = product["transformed"]
            distance = product["distance"]
            scores = product["scores"]
            vectors = product["vectors"]
            assert isinstance(raw, pd.DataFrame)
            assert isinstance(transformed, pd.DataFrame)
            assert isinstance(distance, np.ndarray)
            assert isinstance(scores, pd.DataFrame)
            assert isinstance(vectors, pd.DataFrame)
            raw.reset_index(names="ponto").to_excel(writer, sheet_name=f"CPUEn_{key}", index=False)
            transformed.reset_index(names="ponto").to_excel(writer, sheet_name=f"Raiz_CPUEn_{key}", index=False)
            pd.DataFrame(distance, index=raw.index, columns=raw.index).reset_index(names="ponto").to_excel(
                writer, sheet_name=f"Bray_{key}", index=False
            )
            scores.to_excel(writer, sheet_name=f"PCoA_{key}", index=False)
            vectors.to_excel(writer, sheet_name=f"Vetores_{key}", index=False)
    style_workbook(PCOA_WORKBOOK)


def main() -> None:
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update(
        {
            "font.family": "Arial",
            "font.size": 13,
            "axes.labelsize": 14,
            "xtick.labelsize": 12,
            "ytick.labelsize": 12,
            "legend.fontsize": 12,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "savefig.facecolor": "white",
        }
    )

    beta, beta_matrix, beta_audit = build_beta_temporal()
    plot_beta_temporal(beta)
    write_beta_workbook(beta, beta_matrix, beta_audit)

    period_tables, effort_audit, _species_names = build_period_tables()
    pcoa_products, pcoa_audit = build_pcoa_products(period_tables)
    plot_pcoa(pcoa_products)
    write_pcoa_workbook(pcoa_products, effort_audit, pcoa_audit)

    campaign_table, campaign_metadata = build_campaign_table()
    temporal_product, temporal_contingency, temporal_audit = build_temporal_pcoa(campaign_table, campaign_metadata)
    plot_temporal_pcoa(temporal_product)
    write_temporal_pcoa_workbook(temporal_product, temporal_contingency, temporal_audit)

    print(beta_audit.to_dict("records")[0])
    print(pcoa_audit.to_string(index=False))
    print(temporal_audit.to_string(index=False))
    for path in (
        BETA_FIGURE,
        BETA_WORKBOOK,
        PCOA_FIGURE,
        PCOA_WORKBOOK,
        TEMPORAL_PCOA_FIGURE,
        TEMPORAL_PCOA_WORKBOOK,
    ):
        print(path)


if __name__ == "__main__":
    main()
