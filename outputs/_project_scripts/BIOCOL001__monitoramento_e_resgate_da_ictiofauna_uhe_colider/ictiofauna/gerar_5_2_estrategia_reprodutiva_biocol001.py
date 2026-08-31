from __future__ import annotations

import os
import re
from collections import Counter
from pathlib import Path
from statistics import median

import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(
    "G:/Meu Drive/Opyta/Clientes/Clientes/Clientes/Bios/Col\u00edder/Resultados/2026/Junho-2026/"
    "BIOCOL001_RESULTADOS_ICTIOFAUNA_FINAL_R02"
)
WORKBOOK_PATH = ROOT / "05_02_tabela_caracteristicas_biologicas.xlsx"
FIGURE_PATH = ROOT / "05_02_2_figura_estrategias_reprodutivas.png"
SIZE_FIGURE_PATH = ROOT / "05_02_3_figura_porte_especies.png"
INVENTORY_DIR = Path(
    "G:/Meu Drive/Opyta/Opyta_Data_Analysis/outputs/_project_scripts/"
    "BIOCOL001__monitoramento_e_resgate_da_ictiofauna_uhe_colider/inventory"
)
FISHBASE_SPECIES_PATH = INVENTORY_DIR / "fishbase_v25_04_species.parquet"

CATEGORY_ORDER = (
    "Não migradora",
    "Migradora de curta distância",
    "Migradora de longa distância",
)
COLORS = ("#2E6EA6", "#D4672A", "#6BA547")
SIZE_CLASSES = ("Pequeno", "Médio", "Grande", "Muito grande", "Não determinado")
MANUAL_SIZE_OVERRIDES = {
    "Pamphorichthys cf. scalpridens": {
        "Porte": "Pequeno",
        "Fonte": "Classificação validada pelo responsável técnico do projeto",
        "Observacao": "Nome aberto (cf.); porte validado para o gênero Pamphorichthys.",
    }
}


def style_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30
    for column in range(1, sheet.max_column + 1):
        width = max(len(str(sheet.cell(row, column).value or "")) for row in range(1, sheet.max_row + 1)) + 2
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width, 12), 42)


def save_atomic(workbook) -> None:
    temporary = WORKBOOK_PATH.with_name(f"{WORKBOOK_PATH.stem}.tmp.xlsx")
    workbook.save(temporary)
    os.replace(temporary, WORKBOOK_PATH)


def classify_size(max_length_cm: float) -> str:
    if max_length_cm <= 15:
        return "Pequeno"
    if max_length_cm <= 30:
        return "Médio"
    if max_length_cm <= 60:
        return "Grande"
    return "Muito grande"


def infer_genus(scientific_name: str) -> str:
    first = re.split(r"\s+", str(scientific_name).strip())[0].strip('"\'`“”‘’')
    return first[:1].upper() + first[1:] if first else ""


def load_fishbase_sizes() -> tuple[dict[str, dict[str, object]], dict[str, dict[str, object]]]:
    columns = ["SpecCode", "Genus", "Species", "Length", "LTypeMaxM", "LengthFemale", "LTypeMaxF"]
    fishbase = pd.read_parquet(FISHBASE_SPECIES_PATH, columns=columns)
    result = {}
    genus_lengths: dict[str, list[float]] = {}
    genus_types: dict[str, set[str]] = {}
    for row in fishbase.itertuples(index=False):
        name = f"{row.Genus or ''} {row.Species or ''}".strip()
        candidates = []
        if pd.notna(row.Length) and float(row.Length) > 0:
            candidates.append((float(row.Length), row.LTypeMaxM or "NG", "macho/sexo não informado"))
        if pd.notna(row.LengthFemale) and float(row.LengthFemale) > 0:
            candidates.append((float(row.LengthFemale), row.LTypeMaxF or "NG", "fêmea"))
        if not candidates:
            continue
        length, length_type, sex = max(candidates, key=lambda item: item[0])
        result[name] = {
            "SpecCode": int(row.SpecCode),
            "Comprimento_max_cm": length,
            "Tipo_comprimento": str(length_type),
            "Sexo_referencia": sex,
            "Fonte": f"https://www.fishbase.se/summary/{int(row.SpecCode)}",
        }
        genus = str(row.Genus or "").strip()
        genus_lengths.setdefault(genus, []).append(length)
        genus_types.setdefault(genus, set()).add(str(length_type))
    genus_result = {
        genus: {
            "Comprimento_max_cm": float(median(lengths)),
            "Tipo_comprimento": "/".join(sorted(genus_types[genus])),
            "N_especies_referencia_genero": len(lengths),
            "Fonte": "https://source.coop/cboettig/fishbase/fb/v25.04",
        }
        for genus, lengths in genus_lengths.items()
        if genus
    }
    return result, genus_result


def populate_size_sheet(workbook, species: list[str]) -> Counter:
    previous = {}
    if "Porte_especies" in workbook.sheetnames:
        old_sheet = workbook["Porte_especies"]
        old_headers = {cell.value: cell.column for cell in old_sheet[1]}
        if "nome_cientifico" in old_headers and "Porte" in old_headers:
            for row in range(2, old_sheet.max_row + 1):
                name = old_sheet.cell(row, old_headers["nome_cientifico"]).value
                size = old_sheet.cell(row, old_headers["Porte"]).value
                status = old_sheet.cell(row, old_headers.get("Status_validacao", 0)).value if old_headers.get("Status_validacao") else None
                observation = old_sheet.cell(row, old_headers.get("Observacao", 0)).value if old_headers.get("Observacao") else None
                if name and size and status == "Manual":
                    previous[str(name)] = (str(size), observation)
        del workbook["Porte_especies"]

    fishbase_sizes, genus_sizes = load_fishbase_sizes()
    sheet = workbook.create_sheet("Porte_especies")
    sheet.append(
        [
            "nome_cientifico",
            "Comprimento_max_cm",
            "Tipo_comprimento",
            "Sexo_referencia",
            "N_especies_referencia_genero",
            "Porte",
            "Fonte",
            "Status_validacao",
            "Observacao",
        ]
    )
    counts = Counter()
    for name in sorted(set(species)):
        if name in previous:
            size, observation = previous[name]
            row = [name, None, None, None, None, size, "Classificação fornecida pelo responsável técnico", "Manual", observation]
        elif name in MANUAL_SIZE_OVERRIDES:
            record = MANUAL_SIZE_OVERRIDES[name]
            size = record["Porte"]
            row = [
                name,
                None,
                None,
                None,
                None,
                size,
                record["Fonte"],
                "Manual",
                record["Observacao"],
            ]
        elif name in fishbase_sizes:
            record = fishbase_sizes[name]
            size = classify_size(float(record["Comprimento_max_cm"]))
            row = [
                name,
                record["Comprimento_max_cm"],
                record["Tipo_comprimento"],
                record["Sexo_referencia"],
                None,
                size,
                record["Fonte"],
                "Automático FishBase v25.04 - revisar",
                None,
            ]
        elif infer_genus(name) in genus_sizes:
            genus = infer_genus(name)
            record = genus_sizes[genus]
            size = classify_size(float(record["Comprimento_max_cm"]))
            row = [
                name,
                record["Comprimento_max_cm"],
                record["Tipo_comprimento"],
                "mediana das espécies do gênero",
                record["N_especies_referencia_genero"],
                size,
                record["Fonte"],
                "Estimativa por gênero - mediana FishBase v25.04",
                f"Gênero de referência: {genus}",
            ]
        else:
            size = "Não determinado"
            row = [
                name,
                None,
                None,
                None,
                None,
                size,
                None,
                "Pendente: sem correspondência binomial exata no FishBase",
                None,
            ]
        sheet.append(row)
        counts[size] += 1

    validation = DataValidation(
        type="list",
        formula1='"Pequeno,Médio,Grande,Muito grande,Não determinado"',
        allow_blank=False,
    )
    sheet.add_data_validation(validation)
    validation.add(f"F2:F{sheet.max_row}")
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 2).number_format = "0.0"
    style_sheet(sheet)

    if "Criterio_porte" in workbook.sheetnames:
        del workbook["Criterio_porte"]
    criteria = workbook.create_sheet("Criterio_porte")
    criteria.append(["Porte", "Regra_comprimento_maximo", "Fonte_criterio"])
    criteria.append(["Pequeno", "até 15 cm", "IBAMA, relatório técnico de peixes ornamentais, 2007"])
    criteria.append(["Médio", ">15 a 30 cm", "IBAMA, relatório técnico de peixes ornamentais, 2007"])
    criteria.append(["Grande", ">30 a 60 cm", "IBAMA, relatório técnico de peixes ornamentais, 2007"])
    criteria.append(["Muito grande", ">60 cm", "IBAMA, relatório técnico de peixes ornamentais, 2007"])
    criteria.append(["Não determinado", "Sem correspondência binomial exata ou sem Lmax", "Critério conservador"])
    criteria.append(["Fallback por gênero", "Mediana do Lmax das espécies do gênero no FishBase v25.04", "Premissa aprovada para cf., aff., gr. e sp."])
    criteria.append(
        [
            "URL",
            None,
            "https://www.ibama.gov.br/phocadownload/peixesornamentais/2007/relatorio_peixes_ornamentais_tamandare_out_2007.pdf",
        ]
    )
    style_sheet(criteria)

    if "Resumo_porte" in workbook.sheetnames:
        del workbook["Resumo_porte"]
    summary = workbook.create_sheet("Resumo_porte")
    summary.append(["Porte", "Riqueza", "Percentual_total"])
    for size_class in SIZE_CLASSES:
        summary.append([size_class, counts[size_class], counts[size_class] / len(set(species))])
    summary.append(["Cobertura_classificada", len(set(species)) - counts["Não determinado"], (len(set(species)) - counts["Não determinado"]) / len(set(species))])
    for row in range(2, summary.max_row + 1):
        summary.cell(row, 3).number_format = "0.0%"
    style_sheet(summary)
    return counts


def plot_size_donut(counts: Counter, total: int) -> None:
    categories = ("Pequeno", "Médio", "Grande", "Muito grande")
    values = [counts[category] for category in categories]
    if counts["Não determinado"]:
        return
    colors = ("#2E6EA6", "#D4672A", "#6BA547", "#7B4EA3")
    labels = [
        f"{category} — {value} ({value / total:.1%})".replace(".", ",")
        for category, value in zip(categories, values)
    ]
    plt.rcParams.update({"font.family": "Arial", "font.size": 15})
    fig, ax = plt.subplots(figsize=(18, 10.2), dpi=300, facecolor="white")
    wedges, _ = ax.pie(
        values,
        startangle=90,
        colors=colors,
        wedgeprops={"width": 0.42, "edgecolor": "white", "linewidth": 1.4},
    )
    ax.text(0, 0.04, f"{total}", ha="center", va="center", fontsize=27, fontweight="bold", color="#1F1F1F")
    ax.text(0, -0.11, "espécies", ha="center", va="center", fontsize=15, color="#595959")
    ax.legend(wedges, labels, loc="center left", bbox_to_anchor=(1.02, 0.5), frameon=False, fontsize=15, labelspacing=1.25)
    ax.set_aspect("equal")
    fig.tight_layout()
    fig.savefig(SIZE_FIGURE_PATH, dpi=300, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def main() -> None:
    workbook = load_workbook(WORKBOOK_PATH)
    source = workbook["Tabela_06"]
    headers = {cell.value: cell.column for cell in source[1]}
    species = []
    counts = Counter()
    for row in range(2, source.max_row + 1):
        scientific_name = source.cell(row, headers["nome_cientifico"]).value
        strategy = source.cell(row, headers["estrategia_reprodutiva"]).value
        if scientific_name:
            species.append(str(scientific_name))
        counts[strategy] += 1

    if None in counts or "" in counts:
        raise RuntimeError("Existem especies sem estrategia reprodutiva.")

    if "Estrategias_reprodutivas" in workbook.sheetnames:
        del workbook["Estrategias_reprodutivas"]
    data_sheet = workbook.create_sheet("Estrategias_reprodutivas")
    data_sheet.append(["Estrategia_reprodutiva", "Riqueza", "Percentual"])
    total = sum(counts.values())
    for category in CATEGORY_ORDER:
        data_sheet.append([category, counts[category], counts[category] / total])
    for row in range(2, data_sheet.max_row + 1):
        data_sheet.cell(row, 3).number_format = "0.0%"
    style_sheet(data_sheet)

    size_counts = populate_size_sheet(workbook, species)

    save_atomic(workbook)
    plot_size_donut(size_counts, len(set(species)))

    values = [counts[category] for category in CATEGORY_ORDER]
    labels = [
        f"{category} — {value} ({value / total:.1%})".replace(".", ",")
        for category, value in zip(CATEGORY_ORDER, values)
    ]
    plt.rcParams.update({"font.family": "Arial", "font.size": 15})
    fig, ax = plt.subplots(figsize=(18, 10.2), dpi=300, facecolor="white")
    wedges, _ = ax.pie(
        values,
        startangle=90,
        colors=COLORS,
        wedgeprops={"width": 0.42, "edgecolor": "white", "linewidth": 1.4},
    )
    ax.text(0, 0.04, f"{total}", ha="center", va="center", fontsize=27, fontweight="bold", color="#1F1F1F")
    ax.text(0, -0.11, "espécies", ha="center", va="center", fontsize=15, color="#595959")
    ax.legend(
        wedges,
        labels,
        loc="center left",
        bbox_to_anchor=(1.02, 0.5),
        frameon=False,
        fontsize=15,
        labelspacing=1.25,
    )
    ax.set_aspect("equal")
    fig.tight_layout()
    fig.savefig(FIGURE_PATH, dpi=300, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    print(f"total={total} counts={dict(counts)}")
    print(f"porte={dict(size_counts)}")
    print(FIGURE_PATH)


if __name__ == "__main__":
    main()
