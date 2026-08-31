from __future__ import annotations

import os
from collections import defaultdict
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(
    "G:/Meu Drive/Opyta/Clientes/Clientes/Clientes/Bios/Col\u00edder/Resultados/2026/Junho-2026/"
    "BIOCOL001_RESULTADOS_ICTIOFAUNA_FINAL_R02"
)
BASE_PATH = ROOT / "base_analitica_geral_biocol001.xlsx"
PRODUCT_PATH = ROOT / "05_03_tabela_especies_ameacadas.xlsx"

THREATENED = (
    "Colossoma macropomum",
    "Knodus dorsomaculatus",
    "Harttia dissidens",
    "Scobinancistrus pariolispos",
)

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")


def read_aggregates():
    workbook = load_workbook(BASE_PATH, read_only=True, data_only=True)
    sheet = workbook["base"]
    headers = {cell.value: cell.column - 1 for cell in sheet[1]}

    spatial_n = defaultdict(float)
    spatial_b = defaultdict(float)
    temporal_n = defaultdict(float)
    temporal_b = defaultdict(float)
    points = set()
    campaigns = {}

    for row in sheet.iter_rows(min_row=2, values_only=True):
        point = str(row[headers["ponto"]])
        campaign = str(row[headers["campanha"]])
        points.add(point)
        campaigns[campaign] = (
            row[headers["fase_reservatorio"]] or "",
            row[headers["evento_reservatorio"]] or "",
            row[headers["periodo_rebaixamento_parcial"]] or "",
            row[headers["marco_reenchimento"]] or "",
        )
        species = row[headers["nome_cientifico"]]
        if species not in THREATENED:
            continue
        abundance = float(row[headers["numero_de_individuos"]] or 0)
        biomass = float(row[headers["biomassa_g_linha"]] or 0)
        spatial_n[(point, species)] += abundance
        spatial_b[(point, species)] += biomass
        temporal_n[(campaign, species)] += abundance
        temporal_b[(campaign, species)] += biomass

    # The general analytical base already contains only the 16 operational points.
    all_points = sorted(points)
    all_campaigns = sorted(campaigns)
    workbook.close()
    return spatial_n, spatial_b, temporal_n, temporal_b, all_points, all_campaigns, campaigns


def style_sheet(sheet, abundance_start: int, abundance_end: int) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 42
    for column in range(1, sheet.max_column + 1):
        max_length = max(len(str(sheet.cell(row, column).value or "")) for row in range(1, sheet.max_row + 1))
        sheet.column_dimensions[get_column_letter(column)].width = min(max(max_length + 2, 12), 38)
    if sheet.max_row > 1:
        data_range = (
            f"{get_column_letter(abundance_start)}2:"
            f"{get_column_letter(abundance_end)}{sheet.max_row}"
        )
        sheet.conditional_formatting.add(
            data_range,
            ColorScaleRule(
                start_type="min",
                start_color="FFFFFF",
                mid_type="percentile",
                mid_value=50,
                mid_color="9DC3E6",
                end_type="max",
                end_color="1F4E78",
            ),
        )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00" if "Biomassa" in str(sheet.cell(1, cell.column).value) else "#,##0"


def replace_sheet(workbook, title: str):
    if title in workbook.sheetnames:
        del workbook[title]
    return workbook.create_sheet(title)


def populate_spatial(workbook, title: str, spatial_n, spatial_b, points) -> None:
    sheet = replace_sheet(workbook, title)
    headers = ["Ponto", "Riqueza_ameacadas", "Abundancia_total", "Biomassa_total_g"]
    headers += [f"N_{species}" for species in THREATENED]
    headers += [f"Biomassa_g_{species}" for species in THREATENED]
    sheet.append(headers)
    for point in points:
        abundances = [spatial_n[(point, species)] for species in THREATENED]
        biomasses = [spatial_b[(point, species)] for species in THREATENED]
        sheet.append(
            [point, sum(value > 0 for value in abundances), sum(abundances), sum(biomasses)]
            + abundances
            + biomasses
        )
    style_sheet(sheet, 5, 8)


def populate_temporal(workbook, title: str, temporal_n, temporal_b, campaigns, metadata) -> None:
    sheet = replace_sheet(workbook, title)
    headers = [
        "Campanha",
        "Fase_reservatorio",
        "Evento_reservatorio",
        "Periodo_rebaixamento_parcial",
        "Marco_reenchimento",
        "Riqueza_ameacadas",
        "Abundancia_total",
        "Biomassa_total_g",
    ]
    headers += [f"N_{species}" for species in THREATENED]
    headers += [f"Biomassa_g_{species}" for species in THREATENED]
    sheet.append(headers)
    for campaign in campaigns:
        abundances = [temporal_n[(campaign, species)] for species in THREATENED]
        biomasses = [temporal_b[(campaign, species)] for species in THREATENED]
        phase, event, lowering, refill = metadata.get(campaign, ("", "", "", ""))
        sheet.append(
            [campaign, phase, event, lowering, refill, sum(value > 0 for value in abundances), sum(abundances), sum(biomasses)]
            + abundances
            + biomasses
        )
    style_sheet(sheet, 9, 12)


def save_atomic(workbook, path: Path) -> None:
    temporary = path.with_name(f"{path.stem}.tmp.xlsx")
    workbook.save(temporary)
    os.replace(temporary, path)


def main() -> None:
    spatial_n, spatial_b, temporal_n, temporal_b, points, campaigns, metadata = read_aggregates()

    product = load_workbook(PRODUCT_PATH)
    populate_spatial(product, "Variacao_espacial", spatial_n, spatial_b, points)
    populate_temporal(product, "Variacao_temporal", temporal_n, temporal_b, campaigns, metadata)
    save_atomic(product, PRODUCT_PATH)

    print(f"pontos={len(points)} campanhas={len(campaigns)}")
    print(f"abundancia={sum(spatial_n.values()):.0f} biomassa_g={sum(spatial_b.values()):.4f}")


if __name__ == "__main__":
    main()
