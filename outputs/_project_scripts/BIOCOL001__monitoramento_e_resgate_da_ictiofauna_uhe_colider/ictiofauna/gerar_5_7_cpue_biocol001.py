from __future__ import annotations

import os
import re
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.patches import Patch
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy.stats import t
from sqlalchemy import create_engine, text


PROJECT_ID = 206
ROOT = Path(
    "G:/Meu Drive/Opyta/Clientes/Clientes/Clientes/Bios/Colíder/Resultados/2026/Junho-2026/"
    "BIOCOL001_RESULTADOS_ICTIOFAUNA_FINAL_R02"
)
ANALYSIS_ROOT = Path("G:/Meu Drive/Opyta/Opyta_Data_Analysis")
INVENTORY = ANALYSIS_ROOT / "outputs/_project_scripts/BIOCOL001__monitoramento_e_resgate_da_ictiofauna_uhe_colider/inventory"
WORKBOOK = ROOT / "05_07_captura_unidade_esforco.xlsx"
FIGURES = {
    "cpuen_trecho": ROOT / "05_07_1_cpuen_trecho.png",
    "cpueb_trecho": ROOT / "05_07_2_cpueb_trecho.png",
    "cpuen_temporal": ROOT / "05_07_3_cpuen_temporal.png",
    "cpueb_temporal": ROOT / "05_07_4_cpueb_temporal.png",
    "top20_cpuen": ROOT / "05_07_5_top20_cpuen_especies.png",
    "top20_cpueb": ROOT / "05_07_6_top20_cpueb_especies.png",
}
POINTS = [f"ICTIO{i:02d}" for i in range(1, 13)] + ["ICTIO13A", "ICTIO13B", "ICTIO14", "ICTIO15"]
COLORS = {
    "primary": "#002060",
    "secondary": "#5B9BD5",
    "highlight": "#1F4E79",
    "pre": "#F0F0F0",
    "lowering": "#DBE5F1",
    "refill": "#D4672A",
    "grid": "#D9D9D9",
}
FIGSIZE = (18, 10.2)


def style_axes(ax, grid_axis: str = "y"):
    ax.grid(False)
    ax.grid(axis=grid_axis, color=COLORS["grid"], alpha=0.25, linewidth=1)
    for spine in ax.spines.values():
        spine.set_visible(True)
        spine.set_color("black")
        spine.set_linewidth(1.2)
    ax.tick_params(colors="black", labelsize=13)


def engine():
    for path in (ANALYSIS_ROOT / ".env", Path("G:/Meu Drive/Opyta/Opyta_Data/.env")):
        if path.exists():
            for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
                if line.strip() and not line.lstrip().startswith("#") and "=" in line:
                    key, value = line.split("=", 1)
                    os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))
    url = os.getenv("FISICO_DB_URL") or os.getenv("DATABASE_URL") or os.getenv("SUPABASE_DB_URL")
    if not url:
        raise RuntimeError("URL do banco nao encontrada.")
    return create_engine(url + ("&" if "?" in url else "?") + "connect_timeout=20", pool_pre_ping=True)


def campaign_number(value: str) -> int:
    match = re.match(r"C(\d{3})", str(value))
    return int(match.group(1)) if match else 999


def mean_ci(data: pd.DataFrame, groups: list[str]) -> pd.DataFrame:
    rows = []
    for keys, group in data.groupby(groups, sort=False, dropna=False):
        keys = keys if isinstance(keys, tuple) else (keys,)
        row = dict(zip(groups, keys))
        for metric in ("CPUEn", "CPUEb_kg"):
            values = group[metric].astype(float).to_numpy()
            n = len(values)
            mean = float(np.mean(values))
            sd = float(np.std(values, ddof=1)) if n > 1 else 0.0
            ci = float(t.ppf(0.975, n - 1) * sd / np.sqrt(n)) if n > 1 else 0.0
            row.update({f"{metric}_media": mean, f"{metric}_dp": sd, f"{metric}_ic95": ci})
        row["n_unidades"] = len(group)
        row["esforco_total_m2"] = group["esforco_m2"].sum()
        row["abundancia_total"] = group["N"].sum()
        row["biomassa_total_kg"] = group["biomassa_kg"].sum()
        rows.append(row)
    return pd.DataFrame(rows)


def load_data():
    effort_sql = text(
        """
        SELECT ea.id_esforco, c.nome_campanha AS campanha, pc.nome_ponto AS ponto,
               ea.metodo_de_captura, ea.tipo_de_amostragem, ea.esforco AS esforco_m2,
               ea.unidade_esforco
        FROM esforcos_amostragem ea
        JOIN pontos_coleta pc ON pc.id_ponto_coleta = ea.id_ponto_coleta
        JOIN campanhas c ON c.id_campanha = pc.id_campanha
        WHERE pc.id_projeto = :pid
          AND ea.grupo_biologico = 'Ictiofauna'
          AND lower(ea.metodo_de_captura) = lower('Rede de Emalhar')
          AND lower(coalesce(ea.tipo_de_amostragem, ea.tipo_amostragem, '')) LIKE 'quant%'
        """
    )
    detail_sql = text(
        """
        SELECT d.id_esforco, d.id_especie, e.nome_cientifico, e.nome_popular,
               sum(d.numero_de_individuos) AS abundancia,
               sum(d.numero_de_individuos * d.pc_g) / 1000.0 AS biomassa_kg,
               sum(CASE WHEN d.pc_g IS NULL THEN d.numero_de_individuos ELSE 0 END) AS individuos_sem_pc
        FROM resultados_ictiofauna_detalhe d
        JOIN especies e ON e.id_especie = d.id_especie
        WHERE d.codigo_opyta = 'BIOCOL001'
          AND lower(d.metodo_de_captura) = lower('Rede de Emalhar')
        GROUP BY d.id_esforco, d.id_especie, e.nome_cientifico, e.nome_popular
        """
    )
    with engine().connect() as connection:
        efforts = pd.read_sql(effort_sql, connection, params={"pid": PROJECT_ID})
        catches = pd.read_sql(detail_sql, connection)
    efforts["ordem_campanha"] = efforts["campanha"].map(campaign_number)
    efforts = efforts[
        efforts["ponto"].isin(POINTS)
        & efforts["ordem_campanha"].between(1, 69)
        & pd.to_numeric(efforts["esforco_m2"], errors="coerce").gt(0)
    ].copy()
    efforts["esforco_m2"] = pd.to_numeric(efforts["esforco_m2"], errors="coerce")
    efforts = efforts.sort_values(["ordem_campanha", "ponto"]).drop_duplicates("id_esforco")

    totals = catches.groupby("id_esforco", as_index=False).agg(
        N=("abundancia", "sum"), biomassa_kg=("biomassa_kg", "sum"), individuos_sem_pc=("individuos_sem_pc", "sum")
    )
    units = efforts.merge(totals, on="id_esforco", how="left")
    for column in ("N", "biomassa_kg", "individuos_sem_pc"):
        units[column] = pd.to_numeric(units[column], errors="coerce").fillna(0)
    units["captura_zero"] = units["N"].eq(0)
    units["CPUEn"] = units["N"] / units["esforco_m2"] * 100
    units["CPUEb_kg"] = units["biomassa_kg"] / units["esforco_m2"] * 100

    species = efforts[["id_esforco", "campanha", "ordem_campanha", "ponto", "esforco_m2"]].merge(catches, on="id_esforco", how="left")
    species = species[species["id_especie"].notna()].copy()
    return efforts, catches, units, species


def species_cpue(efforts: pd.DataFrame, catches: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    names = catches[["id_especie", "nome_cientifico", "nome_popular"]].drop_duplicates("id_especie")
    ids = names["id_especie"].tolist()
    effort_ids = efforts["id_esforco"].tolist()
    grid = pd.MultiIndex.from_product([effort_ids, ids], names=["id_esforco", "id_especie"]).to_frame(index=False)
    grid = grid.merge(efforts[["id_esforco", "esforco_m2"]], on="id_esforco", how="left")
    grid = grid.merge(catches[["id_esforco", "id_especie", "abundancia", "biomassa_kg"]], on=["id_esforco", "id_especie"], how="left")
    grid[["abundancia", "biomassa_kg"]] = grid[["abundancia", "biomassa_kg"]].fillna(0)
    grid = grid.rename(columns={"abundancia": "N"})
    grid["CPUEn"] = grid["N"] / grid["esforco_m2"] * 100
    grid["CPUEb_kg"] = grid["biomassa_kg"] / grid["esforco_m2"] * 100
    summary = mean_ci(grid, ["id_especie"]).merge(names, on="id_especie", how="left")
    summary["nome_popular"] = summary["nome_popular"].fillna("N.I.")
    leading = ["nome_cientifico", "nome_popular", "id_especie"]
    summary = summary[leading + [column for column in summary.columns if column not in leading]]
    top_n = summary.nlargest(20, "CPUEn_media").reset_index(drop=True)
    top_b = summary.nlargest(20, "CPUEb_kg_media").reset_index(drop=True)
    return top_n, top_b


def reservoir_layer():
    layer = pd.read_csv(INVENTORY / "camada_temporal_reservatorio_biocol001.csv")
    return layer[layer["ordem_campanha"].between(1, 69)].copy()


def shade_events(ax, temporal: pd.DataFrame):
    pre = temporal["ordem_campanha"].le(20).to_numpy()
    if pre.any():
        ax.axvspan(-0.5, np.where(pre)[0][-1] + 0.5, color=COLORS["pre"], alpha=0.75, linewidth=0, zorder=0)
    lowering = temporal["periodo_rebaixamento_parcial"].astype(str).str.lower().eq("sim").to_numpy()
    if lowering.any():
        idx = np.where(lowering)[0]
        ax.axvspan(idx.min() - 0.5, idx.max() + 0.5, color=COLORS["lowering"], alpha=0.72, linewidth=0, zorder=0)
    refill = temporal["marco_reenchimento"].astype(str).str.lower().eq("sim").to_numpy()
    if refill.any():
        idx = np.where(refill)[0][0]
        ax.axvline(idx, color=COLORS["refill"], linewidth=1.6, linestyle="--", zorder=1)


def plot_points(table: pd.DataFrame, units: pd.DataFrame, metric: str, ylabel: str, path: Path, color: str):
    x = np.arange(len(table))
    fig, ax = plt.subplots(figsize=FIGSIZE, dpi=300)
    point_positions = {point: index for index, point in enumerate(table["ponto"])}
    raw_x = units["ponto"].map(point_positions).to_numpy(float)
    jitter = ((units["ordem_campanha"].to_numpy(int) % 11) - 5) / 32
    ax.scatter(raw_x + jitter, units[metric], s=16, color=color, alpha=0.13, edgecolors="none", zorder=1)
    mean = table[f"{metric}_media"].to_numpy(float)
    ci = table[f"{metric}_ic95"].to_numpy(float)
    yerr = np.vstack([np.minimum(ci, mean), ci])
    ax.errorbar(x, mean, yerr=yerr, fmt="o", markersize=9,
                color=color, ecolor=color, elinewidth=2, capsize=5, label="Média ± IC 95%", zorder=3)
    ax.set_xticks(x, table["ponto"], rotation=45, ha="right")
    ax.set_ylabel(ylabel)
    ax.set_xlabel("Trecho amostral")
    style_axes(ax)
    ax.legend(frameon=False, loc="upper center", bbox_to_anchor=(0.5, 1.08))
    fig.tight_layout()
    fig.savefig(path, dpi=300, facecolor="white")
    plt.close(fig)


def plot_temporal(table: pd.DataFrame, metric: str, ylabel: str, path: Path, color: str):
    x = np.arange(len(table))
    mean = table[f"{metric}_media"].to_numpy(float)
    ci = table[f"{metric}_ic95"].to_numpy(float)
    yerr = np.vstack([np.minimum(ci, mean), ci])
    fig, ax = plt.subplots(figsize=FIGSIZE, dpi=300)
    shade_events(ax, table)
    ax.errorbar(
        x,
        mean,
        yerr=yerr,
        fmt="none",
        ecolor=color,
        elinewidth=1.1,
        capsize=2.5,
        capthick=1.1,
        alpha=0.38,
        zorder=2,
    )
    ax.plot(x, mean, color=color, linewidth=2.8, zorder=3)
    ax.fill_between(x, mean, color=color, alpha=0.12, zorder=1)
    step = max(1, len(table) // 12)
    ax.set_xticks(x[::step], table["campanha"].iloc[::step], rotation=45, ha="right")
    ax.set_ylabel(ylabel)
    ax.set_xlabel("Campanha")
    style_axes(ax)
    handles = [
        plt.Line2D([0], [0], color=color, linewidth=2.8, label="Média com IC 95%"),
        Patch(facecolor=COLORS["pre"], alpha=0.75, label="Pré-enchimento"),
        Patch(facecolor=COLORS["lowering"], alpha=0.72, label="Rebaixamento parcial"),
    ]
    if table["marco_reenchimento"].astype(str).str.lower().eq("sim").any():
        handles.append(plt.Line2D([0], [0], color=COLORS["refill"], linewidth=1.6, linestyle="--", label="Reenchimento"))
    ax.legend(
        handles=handles,
        frameon=False,
        loc="lower center",
        bbox_to_anchor=(0.5, 1.01),
        ncol=len(handles),
        columnspacing=1.8,
        handlelength=2.4,
    )
    fig.tight_layout()
    fig.savefig(path, dpi=300, facecolor="white", bbox_inches="tight")
    plt.close(fig)


def plot_top(table: pd.DataFrame, metric: str, xlabel: str, path: Path, color: str):
    plot = table.sort_values(f"{metric}_media", ascending=True).copy()
    labels = [f"{row.nome_popular}  |  $\\it{{{row.nome_cientifico.replace(' ', r'\ ')}}}$" for row in plot.itertuples()]
    fig, ax = plt.subplots(figsize=FIGSIZE, dpi=300)
    ax.barh(np.arange(len(plot)), plot[f"{metric}_media"], xerr=plot[f"{metric}_ic95"],
            color=color, ecolor="#404040", capsize=3)
    ax.set_yticks(np.arange(len(plot)), labels, fontsize=13)
    ax.set_xlabel(xlabel)
    style_axes(ax, grid_axis="x")
    fig.tight_layout()
    fig.savefig(path, dpi=300, facecolor="white", bbox_inches="tight")
    plt.close(fig)


def style_excel(path: Path):
    workbook = load_workbook(path)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column in range(1, sheet.max_column + 1):
            values = [sheet.cell(row, column).value for row in range(1, min(sheet.max_row, 200) + 1)]
            width = max(len(str(value or "")) for value in values) + 2
            sheet.column_dimensions[get_column_letter(column)].width = min(max(width, 12), 42)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
    temporary = path.with_name(f"{path.stem}.tmp.xlsx")
    workbook.save(temporary)
    os.replace(temporary, path)


def main():
    plt.rcParams.update({
        "font.family": "Arial",
        "font.size": 17,
        "axes.labelsize": 17,
        "axes.titlesize": 16,
        "xtick.labelsize": 13,
        "ytick.labelsize": 13,
        "legend.fontsize": 15,
        "figure.facecolor": "white",
        "axes.facecolor": "white",
        "savefig.facecolor": "white",
    })
    efforts, catches, units, _ = load_data()
    spatial = mean_ci(units, ["ponto"]).set_index("ponto").reindex(POINTS).reset_index()
    temporal = mean_ci(units, ["campanha", "ordem_campanha"]).sort_values("ordem_campanha")
    temporal = temporal.merge(
        reservoir_layer()[["campanha", "fase_reservatorio", "evento_reservatorio", "periodo_rebaixamento_parcial", "marco_reenchimento"]],
        on="campanha", how="left",
    )
    top_n, top_b = species_cpue(efforts, catches[catches["id_esforco"].isin(efforts["id_esforco"])].copy())

    premises = pd.DataFrame(
        [
            ["Período", "Campanhas C001 a C069 (dezembro/2011 a junho/2026)"],
            ["Método", "Somente Rede de Emalhar, amostragem quantitativa"],
            ["Universo espacial", "16 pontos regulares; excluídos ICTIO13A-Marcação e ICTIO13C/D"],
            ["CPUEn", "N / esforço (m²) x 100; unidade: ind./100 m²"],
            ["CPUEb", "Biomassa (kg) / esforço (m²) x 100; unidade: kg/100 m²"],
            ["Biomassa", "Soma linha a linha de N x PC_g, convertida de g para kg"],
            ["Média", "Média aritmética das unidades campanha x ponto com esforço"],
            ["Incerteza", "Intervalo de confiança de 95% pela distribuição t"],
            ["Captura zero", "Incluída quando existe esforço cadastrado sem captura"],
            ["Top 20", "Média por espécie calculada em todas as unidades com esforço, incluindo zeros"],
        ], columns=["Premissa", "Regra"]
    )
    audit = pd.DataFrame(
        [{
            "unidades_esforco": len(units), "pontos": units["ponto"].nunique(), "campanhas": units["campanha"].nunique(),
            "capturas_zero": int(units["captura_zero"].sum()), "esforco_total_m2": units["esforco_m2"].sum(),
            "abundancia_total": units["N"].sum(), "biomassa_total_kg": units["biomassa_kg"].sum(),
            "individuos_sem_pc": units["individuos_sem_pc"].sum(),
        }]
    )
    with pd.ExcelWriter(WORKBOOK, engine="openpyxl") as writer:
        premises.to_excel(writer, sheet_name="Premissas", index=False)
        audit.to_excel(writer, sheet_name="Auditoria", index=False)
        efforts.to_excel(writer, sheet_name="Esforco_validado", index=False)
        units.to_excel(writer, sheet_name="Base_CPUE", index=False)
        spatial.to_excel(writer, sheet_name="CPUEn_CPUEb_trecho", index=False)
        temporal.to_excel(writer, sheet_name="CPUEn_CPUEb_temporal", index=False)
        top_n.to_excel(writer, sheet_name="Top20_CPUEn", index=False)
        top_b.to_excel(writer, sheet_name="Top20_CPUEb", index=False)
    style_excel(WORKBOOK)

    plot_points(spatial, units, "CPUEn", "CPUEn (ind./100 m²)", FIGURES["cpuen_trecho"], COLORS["primary"])
    plot_points(spatial, units, "CPUEb_kg", "CPUEb (kg/100 m²)", FIGURES["cpueb_trecho"], COLORS["secondary"])
    plot_temporal(temporal, "CPUEn", "CPUEn (ind./100 m²)", FIGURES["cpuen_temporal"], COLORS["primary"])
    plot_temporal(temporal, "CPUEb_kg", "CPUEb (kg/100 m²)", FIGURES["cpueb_temporal"], COLORS["secondary"])
    plot_top(top_n, "CPUEn", "CPUEn média (ind./100 m²)", FIGURES["top20_cpuen"], COLORS["primary"])
    plot_top(top_b, "CPUEb_kg", "CPUEb média (kg/100 m²)", FIGURES["top20_cpueb"], COLORS["secondary"])
    print(audit.to_dict("records")[0])
    print(WORKBOOK)
    for path in FIGURES.values():
        print(path)


if __name__ == "__main__":
    main()
