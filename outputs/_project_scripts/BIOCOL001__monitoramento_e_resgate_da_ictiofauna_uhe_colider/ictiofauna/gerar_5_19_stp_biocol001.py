from __future__ import annotations

import os
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows


ROOT = Path(
    "G:/Meu Drive/Opyta/Clientes/Clientes/Clientes/Bios/Col\u00edder/Resultados/2026/Junho-2026/"
    "BIOCOL001_RESULTADOS_ICTIOFAUNA_FINAL_R02"
)
SOURCE = ROOT / "base_analitica_stp_pt13c_pt13d_biocol001.xlsx"
OUTPUT = ROOT / "05_19_sistema_transposicao_peixes.xlsx"
POINTS = ["ICTIO13C", "ICTIO13D"]
CATEGORY_COLUMNS = [
    "nome_popular",
    "ordem",
    "familia",
    "origem",
    "classe_migratoria",
    "estrategia_reprodutiva",
    "status_ameaca_estadual",
    "status_ameaca_nacional",
    "status_ameaca_global",
]


def campaign_order(value: str) -> int:
    match = re.match(r"C(\d{3})", str(value))
    return int(match.group(1)) if match else 999


def joined_unique(values: pd.Series) -> str:
    unique = sorted({str(value).strip() for value in values.dropna() if str(value).strip()})
    return "; ".join(unique)


def species_matrix(data: pd.DataFrame, dimension: str, value: str, columns: list[str], total_name: str):
    matrix = data.pivot_table(
        index=["nome_cientifico", "nome_popular"],
        columns=dimension,
        values=value,
        aggfunc="sum",
        fill_value=0,
    ).reindex(columns=columns, fill_value=0)
    matrix[total_name] = matrix.sum(axis=1)
    return matrix.reset_index().sort_values([total_name, "nome_cientifico"], ascending=[False, True]).reset_index(drop=True)


def occurrence_matrix(data: pd.DataFrame, dimension: str, unit: str, columns: list[str], total_name: str):
    local = data[["nome_cientifico", "nome_popular", dimension, unit]].drop_duplicates()
    matrix = local.pivot_table(
        index=["nome_cientifico", "nome_popular"],
        columns=dimension,
        values=unit,
        aggfunc="nunique",
        fill_value=0,
    ).reindex(columns=columns, fill_value=0)
    matrix[total_name] = matrix.sum(axis=1)
    return matrix.reset_index().sort_values([total_name, "nome_cientifico"], ascending=[False, True]).reset_index(drop=True)


def build_tables(data: pd.DataFrame):
    campaigns = sorted(data["campanha"].dropna().unique(), key=campaign_order)
    categories = (
        data.groupby("nome_cientifico", as_index=False)
        .agg(
            nome_popular=("nome_popular", "first"),
            ordem=("ordem", "first"),
            familia=("familia", "first"),
            origem=("origem", "first"),
            classe_migratoria=("classe_migratoria", "first"),
            estrategia_reprodutiva=("estrategia_reprodutiva", "first"),
            status_ameaca_estadual=("status_ameaca_estadual", "first"),
            status_ameaca_nacional=("status_ameaca_nacional", "first"),
            status_ameaca_global=("status_ameaca_global", "first"),
            abundancia_total=("numero_de_individuos", "sum"),
            biomassa_total_g=("biomassa_g_linha", "sum"),
            campanhas_ocorrencia=("campanha", "nunique"),
            pontos_ocorrencia=("ponto", "nunique"),
            metodos_captura=("metodo_de_captura", joined_unique),
        )
        .sort_values(["ordem", "familia", "nome_cientifico"])
        .reset_index(drop=True)
    )

    spatial_summary = (
        data.groupby("ponto", as_index=False)
        .agg(
            campanhas_amostradas=("campanha", "nunique"),
            riqueza=("nome_cientifico", "nunique"),
            abundancia=("numero_de_individuos", "sum"),
            biomassa_g=("biomassa_g_linha", "sum"),
            metodos_captura=("metodo_de_captura", joined_unique),
        )
        .set_index("ponto")
        .reindex(POINTS)
        .reset_index()
    )
    spatial_abundance = species_matrix(data, "ponto", "numero_de_individuos", POINTS, "Total")
    spatial_biomass = species_matrix(data, "ponto", "biomassa_g_linha", POINTS, "Total_g")
    spatial_occurrence = occurrence_matrix(data, "ponto", "campanha", POINTS, "Total_ponto_campanha")

    temporal_summary = (
        data.groupby("campanha", as_index=False)
        .agg(
            pontos_amostrados=("ponto", "nunique"),
            riqueza=("nome_cientifico", "nunique"),
            abundancia=("numero_de_individuos", "sum"),
            biomassa_g=("biomassa_g_linha", "sum"),
            metodos_captura=("metodo_de_captura", joined_unique),
            fase_reservatorio=("fase_reservatorio", "first"),
            evento_reservatorio=("evento_reservatorio", "first"),
            periodo_rebaixamento_parcial=("periodo_rebaixamento_parcial", "first"),
            marco_reenchimento=("marco_reenchimento", "first"),
        )
    )
    temporal_summary["ordem_campanha"] = temporal_summary["campanha"].map(campaign_order)
    temporal_summary = temporal_summary.sort_values("ordem_campanha").drop(columns="ordem_campanha").reset_index(drop=True)
    temporal_abundance = species_matrix(data, "campanha", "numero_de_individuos", campaigns, "Total")
    temporal_biomass = species_matrix(data, "campanha", "biomassa_g_linha", campaigns, "Total_g")
    temporal_occurrence = occurrence_matrix(data, "campanha", "ponto", campaigns, "Total_ponto_campanha")

    audit_rows = [
        {"verificacao": "Pontos exclusivos da camada", "resultado": ", ".join(sorted(data["ponto"].unique()))},
        {"verificacao": "Camada operacional", "resultado": joined_unique(data["camada_operacional"])},
        {"verificacao": "Linhas", "resultado": len(data)},
        {"verificacao": "Individuos", "resultado": data["numero_de_individuos"].sum()},
        {"verificacao": "Biomassa total g", "resultado": data["biomassa_g_linha"].sum()},
        {"verificacao": "Especies", "resultado": data["nome_cientifico"].nunique()},
        {"verificacao": "Campanhas com registros STP", "resultado": data["campanha"].nunique()},
    ]
    for column in CATEGORY_COLUMNS:
        missing = (
            data.groupby("nome_cientifico")[column]
            .agg(lambda values: values.dropna().astype(str).str.strip().replace("", pd.NA).dropna().nunique())
            .eq(0)
            .sum()
        )
        divergent = data.groupby("nome_cientifico")[column].nunique(dropna=True).gt(1).sum()
        audit_rows.append({"verificacao": f"{column}: especies sem categoria", "resultado": int(missing)})
        audit_rows.append({"verificacao": f"{column}: especies divergentes", "resultado": int(divergent)})
    audit = pd.DataFrame(audit_rows)

    premises = pd.DataFrame(
        {
            "premissa": [
                "Universo",
                "Camada",
                "Métodos",
                "Período temporal",
                "Abundância",
                "Biomassa",
                "Ocorrência espacial",
                "Ocorrência temporal",
            ],
            "regra": [
                "Somente ICTIO13C e ICTIO13D, equivalentes a PT-13C/PT-13D",
                "sistema_transposicao_condicional; não misturar ao universo geral",
                "Todos os métodos quantitativos e qualitativos registrados no STP",
                "Somente as 37 campanhas com registros STP; campanhas não amostradas não recebem zero",
                "Soma do número de indivíduos por espécie",
                "Soma da biomassa linha a linha, em gramas",
                "Número de campanhas com ocorrência da espécie em cada ponto",
                "Número de pontos com ocorrência da espécie em cada campanha",
            ],
        }
    )
    return {
        "Classificacao_geral": categories,
        "Espacial_resumo": spatial_summary,
        "Espacial_abundancia": spatial_abundance,
        "Espacial_biomassa": spatial_biomass,
        "Espacial_ocorrencia": spatial_occurrence,
        "Temporal_resumo": temporal_summary,
        "Temporal_abundancia": temporal_abundance,
        "Temporal_biomassa": temporal_biomass,
        "Temporal_ocorrencia": temporal_occurrence,
        "Auditoria": audit,
        "Premissas": premises,
    }


def write_workbook(path: Path, sheets: dict[str, pd.DataFrame]):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, data in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in dataframe_to_rows(data, index=False, header=True):
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column in range(1, sheet.max_column + 1):
            values = [sheet.cell(row, column).value for row in range(1, min(sheet.max_row, 300) + 1)]
            width = max(len(str(value or "")) for value in values) + 2
            sheet.column_dimensions[get_column_letter(column)].width = min(max(width, 12), 42)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
    temporary = path.with_name(f"{path.stem}.tmp.xlsx")
    workbook.save(temporary)
    os.replace(temporary, path)


def main():
    data = pd.read_excel(SOURCE, sheet_name="base")
    data = data[
        data["ponto"].isin(POINTS)
        & data["camada_operacional"].eq("sistema_transposicao_condicional")
    ].copy()
    data["numero_de_individuos"] = pd.to_numeric(data["numero_de_individuos"], errors="coerce").fillna(0)
    data["biomassa_g_linha"] = pd.to_numeric(data["biomassa_g_linha"], errors="coerce").fillna(0)
    tables = build_tables(data)
    write_workbook(OUTPUT, tables)
    print(
        {
            "linhas": len(data),
            "individuos": float(data["numero_de_individuos"].sum()),
            "biomassa_g": float(data["biomassa_g_linha"].sum()),
            "especies": int(data["nome_cientifico"].nunique()),
            "campanhas": int(data["campanha"].nunique()),
            "pontos": sorted(data["ponto"].unique()),
        }
    )
    print(OUTPUT)


if __name__ == "__main__":
    main()
