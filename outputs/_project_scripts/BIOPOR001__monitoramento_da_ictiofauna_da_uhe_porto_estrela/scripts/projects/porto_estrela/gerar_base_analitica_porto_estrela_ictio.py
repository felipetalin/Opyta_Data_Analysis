from __future__ import annotations

import re
import unicodedata
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd


PROJECT_CODE = "BIOPOR001"
PROJECT_NAME = "Porto Estrela"
DATE_TAG = "20260602"
CUT_YYYYMM = 202512

PLANILHA_ROOT = Path(
    r"G:\Meu Drive\Opyta\Clientes\Clientes\Clientes\Bios\Porto Estrela\Planilha"
)
RESULTADOS_DIR = PLANILHA_ROOT / "Resultados"


def _find_migration_dir() -> Path:
    for child in PLANILHA_ROOT.iterdir():
        if child.is_dir() and child.name.lower().startswith("migra"):
            return child
    raise FileNotFoundError("Pasta de migracao nao encontrada em Porto Estrela.")


def _strip_accents(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", str(value))
    return "".join(ch for ch in normalized if not unicodedata.combining(ch))


def _norm_text(value) -> str:
    if pd.isna(value):
        return ""
    return re.sub(r"\s+", " ", str(value).strip())


def _norm_key(value) -> str:
    return _strip_accents(_norm_text(value)).lower()


def _is_missing_token(value) -> bool:
    key = _norm_key(value).replace(".", "").replace("-", "")
    return key in {"", "na", "nan", "none", "null", "ni", "n/i", "nao informado"}


def _to_number(series: pd.Series) -> pd.Series:
    return pd.to_numeric(
        series.astype(str).str.replace(",", ".", regex=False).str.strip(),
        errors="coerce",
    )


def _tipo_base(value) -> str:
    txt = _norm_key(value)
    if "quanti" in txt:
        return "Quantitativa"
    if "quali" in txt:
        return "Qualitativa"
    return _norm_text(value) or "Nao informado"


def _sexo_padronizado(value) -> str:
    if _is_missing_token(value):
        return "Nao informado"
    txt = _norm_key(value)
    if txt in {"f", "femea", "fem", "female"}:
        return "Femea"
    if txt in {"m", "macho", "male"}:
        return "Macho"
    return _norm_text(value)


def _emg_codigo(value, sexo_value=None) -> str:
    if _is_missing_token(value):
        return ""
    txt = _strip_accents(_norm_text(value)).upper().replace(" ", "")
    if not txt:
        return ""
    match = re.search(r"([FM])\s*([1-4])", txt)
    if match:
        return f"{match.group(1)}{match.group(2)}"
    match = re.search(r"(^|[^0-9])([1-4])([^0-9]|$)", txt)
    if match:
        sexo_txt = _sexo_padronizado(sexo_value)
        prefix = "F" if sexo_txt == "Femea" else "M" if sexo_txt == "Macho" else ""
        return f"{prefix}{match.group(2)}" if prefix else match.group(2)
    return txt


def _emg_estadio(code: str) -> str:
    if not code:
        return "Nao informado"
    stage = str(code)[-1]
    return {
        "1": "Repouso",
        "2": "Maturacao inicial",
        "3": "Maturacao avancada/maduro",
        "4": "Desovado/esgotado",
    }.get(stage, "Nao classificado")


def _emg_ordem(code: str) -> float:
    if not code:
        return np.nan
    stage = str(code)[-1]
    return float(stage) if stage in {"1", "2", "3", "4"} else np.nan


def _parse_campanha(campanha: str) -> dict[str, object]:
    txt = _norm_text(campanha)
    match = re.match(r"PE(\d{3})_AH(\d{4})_(\d{6})(?:_.+)?$", txt)
    if not match:
        return {
            "campanha_ordem": np.nan,
            "ano_hidrologico": "",
            "ano_hidrologico_rotulo": "",
            "campanha_aaaamm": np.nan,
            "data_campanha": pd.NaT,
        }

    ordem = int(match.group(1))
    ah = match.group(2)
    aaaamm = int(match.group(3))
    ano = int(match.group(3)[:4])
    mes = int(match.group(3)[4:6])
    return {
        "campanha_ordem": ordem,
        "ano_hidrologico": f"AH{ah}",
        "ano_hidrologico_rotulo": f"ano {ah[:2]}-{ah[2:]}",
        "campanha_aaaamm": aaaamm,
        "data_campanha": pd.Timestamp(year=ano, month=mes, day=1),
    }


def _de_para_trechos() -> pd.DataFrame:
    rows = []
    global_order = 1
    groups = [
        ("Montante", ["P4", "P5", "P2", "P1"]),
        ("Jusante", ["P3", "P6", "P7", "P8", "P9"]),
    ]
    for trecho, points in groups:
        for ordem_trecho, ponto in enumerate(points, start=1):
            rows.append(
                {
                    "Ponto": ponto,
                    "Trecho": trecho,
                    "Ordem_Trecho_Montante_Jusante": ordem_trecho,
                    "Ordem_Global_Montante_Jusante": global_order,
                    "Criterio": "Definido pelo usuario em 2026-06-02",
                }
            )
            global_order += 1
    return pd.DataFrame(rows)


def _sum_min_count(series: pd.Series) -> float:
    valid = pd.to_numeric(series, errors="coerce").dropna()
    if valid.empty:
        return np.nan
    return float(valid.sum())


def _read_inputs() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, Path]:
    migration_dir = _find_migration_dir()
    workbook = migration_dir / "Opyta-Bios-Porto_Estrela-Ictio-2026_MIGRACAO_VALIDADA_260602.xlsx"
    caracterizacao = RESULTADOS_DIR / f"caracterizacao_especies_porto_estrela_{DATE_TAG}.xlsx"

    if not workbook.exists():
        raise FileNotFoundError(workbook)
    if not caracterizacao.exists():
        raise FileNotFoundError(caracterizacao)

    resultados = pd.read_excel(workbook, sheet_name="Resultados_Ictiofauna")
    esforcos = pd.read_excel(workbook, sheet_name="Metadados_Esforco")
    especies = pd.read_excel(caracterizacao, sheet_name="Caracterizacao_Especies")
    return resultados, esforcos, especies, workbook


def _prepare_efforts(esforcos: pd.DataFrame, de_para: pd.DataFrame) -> pd.DataFrame:
    df = esforcos.copy()
    for col in ["Campanha", "Ponto", "Metodo_de_Captura", "Tipo_de_Amostragem"]:
        df[col] = df[col].map(_norm_text)
    df["Tipo_Amostragem_Base"] = df["Tipo_de_Amostragem"].map(_tipo_base)
    df["Esforco_Metadata"] = _to_number(df["Esforco"])

    campanha_info = pd.DataFrame([_parse_campanha(v) for v in df["Campanha"]])
    df = pd.concat([df.reset_index(drop=True), campanha_info], axis=1)
    df = df[df["campanha_aaaamm"].le(CUT_YYYYMM)].copy()

    df = df.merge(de_para, on="Ponto", how="left")
    return df


def _prepare_results(
    resultados: pd.DataFrame,
    esforcos: pd.DataFrame,
    especies: pd.DataFrame,
    de_para: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    df = resultados.copy()
    for col in ["Campanha", "Ponto", "Metodo_de_Captura", "Tipo_de_Amostragem", "Nome_Cientifico"]:
        df[col] = df[col].map(_norm_text)

    df["Tipo_Amostragem_Base"] = df["Tipo_de_Amostragem"].map(_tipo_base)
    df["Numero_de_Individuos"] = _to_number(df["Numero_de_Individuos"]).fillna(0)
    df["PC_g_individual"] = _to_number(df["PC_g"])
    df["CT_cm"] = _to_number(df["CT_cm"])
    df["CP_cm"] = _to_number(df["CP_cm"])
    df["Esforco_Resultado"] = _to_number(df["Esforco_Amostral"])
    df["Sexo_Padronizado"] = df["Sexo"].map(_sexo_padronizado)
    df["EMG_Codigo"] = [
        _emg_codigo(emg, sexo) for emg, sexo in zip(df["EMG"], df["Sexo"], strict=False)
    ]
    df["EMG_Estadio"] = df["EMG_Codigo"].map(_emg_estadio)
    df["EMG_Ordem"] = df["EMG_Codigo"].map(_emg_ordem)
    df["Evidencia_Reprodutiva_Forte"] = df["EMG_Codigo"].isin({"F3", "M3", "F4", "M4"})
    df["Tem_EMG_Informado"] = df["EMG_Codigo"].astype(str).str.len().gt(0)

    campanha_info = pd.DataFrame([_parse_campanha(v) for v in df["Campanha"]])
    df = pd.concat([df.reset_index(drop=True), campanha_info], axis=1)
    df = df[df["campanha_aaaamm"].le(CUT_YYYYMM)].copy()

    effort_key = ["Campanha", "Ponto", "Metodo_de_Captura", "Tipo_de_Amostragem"]
    effort_lookup = esforcos[effort_key + ["Esforco_Metadata", "Unidade_Esforco"]].drop_duplicates(effort_key)
    df = df.merge(effort_lookup, on=effort_key, how="left")
    df["Esforco_Linha"] = df["Esforco_Resultado"].combine_first(df["Esforco_Metadata"])
    df["Unidade_Esforco_Linha"] = df["Unidade_Esforco_x"].combine_first(df["Unidade_Esforco_y"])
    df = df.drop(columns=[c for c in ["Unidade_Esforco_x", "Unidade_Esforco_y"] if c in df.columns])

    df["Biomassa_g_linha"] = df["Numero_de_Individuos"] * df["PC_g_individual"]
    quant_mask = df["Tipo_Amostragem_Base"].eq("Quantitativa") & df["Esforco_Linha"].gt(0)
    df["CPUEn_linha"] = np.where(
        quant_mask,
        df["Numero_de_Individuos"] / df["Esforco_Linha"] * 100.0,
        np.nan,
    )
    df["CPUEb_linha"] = np.where(
        quant_mask & df["Biomassa_g_linha"].notna(),
        df["Biomassa_g_linha"] / df["Esforco_Linha"] * 100.0,
        np.nan,
    )

    df = df.merge(de_para, on="Ponto", how="left")

    species_cols = [
        "Nome_Cientifico",
        "Nome_Popular",
        "Migradora_Nao_Migradora",
        "Nativa_Nao_Nativa",
        "Ameacada_Extincao",
        "Origem_Distribuicao",
        "Status_Ameaca_Estadual",
        "Status_Ameaca_Nacional",
        "Status_Ameaca_Global",
    ]
    species_lookup = especies[species_cols].drop_duplicates("Nome_Cientifico")
    df = df.merge(species_lookup, on="Nome_Cientifico", how="left")

    df["Projeto"] = PROJECT_NAME
    df["Codigo_Projeto"] = PROJECT_CODE

    key_cpue = ["Campanha", "Ponto", "Metodo_de_Captura", "Tipo_de_Amostragem"]
    quant_results = df[df["Tipo_Amostragem_Base"].eq("Quantitativa")].copy()
    result_totals = (
        quant_results.groupby(key_cpue, dropna=False)
        .agg(
            Linhas_Resultados=("Nome_Cientifico", "size"),
            Riqueza=("Nome_Cientifico", "nunique"),
            Abundancia_Total=("Numero_de_Individuos", "sum"),
            Biomassa_Total_g=("Biomassa_g_linha", _sum_min_count),
            CPUEn_Total=("CPUEn_linha", "sum"),
            CPUEb_Total=("CPUEb_linha", _sum_min_count),
        )
        .reset_index()
    )

    effort_quant = esforcos[esforcos["Tipo_Amostragem_Base"].eq("Quantitativa")].copy()
    cpue_effort = effort_quant.merge(result_totals, on=key_cpue, how="left")
    for col in ["Linhas_Resultados", "Riqueza", "Abundancia_Total", "CPUEn_Total"]:
        cpue_effort[col] = pd.to_numeric(cpue_effort[col], errors="coerce").fillna(0)
    cpue_effort["Biomassa_Total_g"] = pd.to_numeric(cpue_effort["Biomassa_Total_g"], errors="coerce").fillna(0)
    cpue_effort["CPUEb_Total"] = pd.to_numeric(cpue_effort["CPUEb_Total"], errors="coerce").fillna(0)
    cpue_effort["Captura_Zero"] = cpue_effort["Linhas_Resultados"].eq(0)

    return df, cpue_effort


def _summaries(base: pd.DataFrame, cpue_effort: pd.DataFrame) -> dict[str, pd.DataFrame]:
    summaries: dict[str, pd.DataFrame] = {}

    summaries["Resumo_Amostragem"] = (
        base.groupby("Tipo_Amostragem_Base", dropna=False)
        .agg(
            Linhas=("Nome_Cientifico", "size"),
            Campanhas=("Campanha", "nunique"),
            Pontos=("Ponto", "nunique"),
            Especies=("Nome_Cientifico", "nunique"),
            Abundancia_Total=("Numero_de_Individuos", "sum"),
            Biomassa_Total_g=("Biomassa_g_linha", _sum_min_count),
        )
        .reset_index()
    )

    summaries["PA_Campanha_Ponto_Especie"] = (
        base.groupby(
            [
                "Campanha",
                "campanha_ordem",
                "ano_hidrologico",
                "data_campanha",
                "Ponto",
                "Trecho",
                "Nome_Cientifico",
            ],
            dropna=False,
        )
        .agg(
            Presenca=("Nome_Cientifico", lambda s: 1),
            Abundancia_Total=("Numero_de_Individuos", "sum"),
            Tipos_Amostragem=("Tipo_Amostragem_Base", lambda s: "; ".join(sorted(set(map(str, s))))),
        )
        .reset_index()
    )

    summaries["CPUE_Esforco_Quanti"] = cpue_effort.copy()

    summaries["CPUE_Campanha"] = (
        cpue_effort.groupby(["Campanha", "campanha_ordem", "ano_hidrologico", "data_campanha"], dropna=False)
        .agg(
            Pontos=("Ponto", "nunique"),
            Esforcos_Quantitativos=("Ponto", "size"),
            Abundancia_Total=("Abundancia_Total", "sum"),
            Biomassa_Total_g=("Biomassa_Total_g", "sum"),
            CPUEn_Total=("CPUEn_Total", "sum"),
            CPUEb_Total=("CPUEb_Total", "sum"),
            Esforcos_Captura_Zero=("Captura_Zero", "sum"),
        )
        .reset_index()
    )

    summaries["CPUE_Ano_Hidrologico"] = (
        cpue_effort.groupby(["ano_hidrologico", "ano_hidrologico_rotulo"], dropna=False)
        .agg(
            Campanhas=("Campanha", "nunique"),
            Pontos=("Ponto", "nunique"),
            Esforcos_Quantitativos=("Ponto", "size"),
            Abundancia_Total=("Abundancia_Total", "sum"),
            Biomassa_Total_g=("Biomassa_Total_g", "sum"),
            CPUEn_Total=("CPUEn_Total", "sum"),
            CPUEb_Total=("CPUEb_Total", "sum"),
        )
        .reset_index()
    )

    summaries["CPUE_Trecho"] = (
        cpue_effort.groupby(["Trecho"], dropna=False)
        .agg(
            Campanhas=("Campanha", "nunique"),
            Pontos=("Ponto", "nunique"),
            Esforcos_Quantitativos=("Ponto", "size"),
            Abundancia_Total=("Abundancia_Total", "sum"),
            Biomassa_Total_g=("Biomassa_Total_g", "sum"),
            CPUEn_Total=("CPUEn_Total", "sum"),
            CPUEb_Total=("CPUEb_Total", "sum"),
        )
        .reset_index()
    )

    quant = base[base["Tipo_Amostragem_Base"].eq("Quantitativa")].copy()
    group_cols = [
        "Campanha",
        "campanha_ordem",
        "ano_hidrologico",
        "data_campanha",
        "Migradora_Nao_Migradora",
        "Nativa_Nao_Nativa",
        "Ameacada_Extincao",
    ]
    summaries["CPUE_Grupos_Quanti"] = (
        quant.groupby(group_cols, dropna=False)
        .agg(
            Especies=("Nome_Cientifico", "nunique"),
            Abundancia_Total=("Numero_de_Individuos", "sum"),
            Biomassa_Total_g=("Biomassa_g_linha", _sum_min_count),
            CPUEn_Total=("CPUEn_linha", "sum"),
            CPUEb_Total=("CPUEb_linha", _sum_min_count),
        )
        .reset_index()
    )

    summaries["CPUE_Especie_Campanha"] = (
        quant.groupby(
            [
                "Campanha",
                "campanha_ordem",
                "ano_hidrologico",
                "data_campanha",
                "Nome_Cientifico",
                "Migradora_Nao_Migradora",
                "Nativa_Nao_Nativa",
                "Ameacada_Extincao",
            ],
            dropna=False,
        )
        .agg(
            Abundancia_Total=("Numero_de_Individuos", "sum"),
            Biomassa_Total_g=("Biomassa_g_linha", _sum_min_count),
            CPUEn_Total=("CPUEn_linha", "sum"),
            CPUEb_Total=("CPUEb_linha", _sum_min_count),
        )
        .reset_index()
    )

    repro = base[base["Tem_EMG_Informado"]].copy()
    repro["Grupo_Reprodutivo_Principal"] = (
        repro["Migradora_Nao_Migradora"].eq("Migradora")
        | repro["Ameacada_Extincao"].eq("Sim")
    )

    if repro.empty:
        empty_cols = [
            "Nome_Cientifico",
            "Sexo_Padronizado",
            "EMG_Codigo",
            "EMG_Estadio",
            "Abundancia_Total",
            "Abundancia_Evidencia_Forte",
        ]
        summaries["Repro_EMG_Especie"] = pd.DataFrame(columns=empty_cols)
        summaries["Repro_Especie"] = pd.DataFrame(columns=empty_cols)
        summaries["Repro_Ponto"] = pd.DataFrame(columns=empty_cols)
        summaries["Repro_Campanha"] = pd.DataFrame(columns=empty_cols)
        summaries["Repro_Ano_Hidrologico"] = pd.DataFrame(columns=empty_cols)
        summaries["Repro_Trecho"] = pd.DataFrame(columns=empty_cols)
    else:
        repro["Abundancia_Evidencia_Forte_Linha"] = np.where(
            repro["Evidencia_Reprodutiva_Forte"],
            repro["Numero_de_Individuos"],
            0,
        )

        summaries["Repro_EMG_Especie"] = (
            repro.groupby(
                [
                    "Nome_Cientifico",
                    "Migradora_Nao_Migradora",
                    "Ameacada_Extincao",
                    "Grupo_Reprodutivo_Principal",
                    "Sexo_Padronizado",
                    "EMG_Codigo",
                    "EMG_Estadio",
                    "EMG_Ordem",
                ],
                dropna=False,
            )
            .agg(
                Linhas=("Nome_Cientifico", "size"),
                Abundancia_Total=("Numero_de_Individuos", "sum"),
                Abundancia_Evidencia_Forte=("Abundancia_Evidencia_Forte_Linha", "sum"),
            )
            .reset_index()
            .sort_values(["Grupo_Reprodutivo_Principal", "Nome_Cientifico", "EMG_Ordem"], ascending=[False, True, True])
        )

        summaries["Repro_Especie"] = (
            repro.groupby(
                [
                    "Nome_Cientifico",
                    "Migradora_Nao_Migradora",
                    "Ameacada_Extincao",
                    "Grupo_Reprodutivo_Principal",
                ],
                dropna=False,
            )
            .agg(
                Linhas_EMG=("Nome_Cientifico", "size"),
                Abundancia_Total_EMG=("Numero_de_Individuos", "sum"),
                Abundancia_Evidencia_Forte=("Abundancia_Evidencia_Forte_Linha", "sum"),
                Campanhas=("Campanha", "nunique"),
                Pontos=("Ponto", "nunique"),
            )
            .reset_index()
        )
        summaries["Repro_Especie"]["Perc_Evidencia_Forte"] = np.where(
            summaries["Repro_Especie"]["Abundancia_Total_EMG"].gt(0),
            summaries["Repro_Especie"]["Abundancia_Evidencia_Forte"]
            / summaries["Repro_Especie"]["Abundancia_Total_EMG"]
            * 100,
            np.nan,
        )
        summaries["Repro_Especie"] = summaries["Repro_Especie"].sort_values(
            ["Grupo_Reprodutivo_Principal", "Abundancia_Evidencia_Forte", "Abundancia_Total_EMG"],
            ascending=[False, False, False],
        )

        for sheet, group_cols in {
            "Repro_Ponto": ["Ponto", "Trecho", "Ordem_Global_Montante_Jusante"],
            "Repro_Campanha": ["Campanha", "campanha_ordem", "ano_hidrologico", "data_campanha"],
            "Repro_Ano_Hidrologico": ["ano_hidrologico", "ano_hidrologico_rotulo"],
            "Repro_Trecho": ["Trecho"],
        }.items():
            out = (
                repro.groupby(group_cols, dropna=False)
                .agg(
                    Linhas_EMG=("Nome_Cientifico", "size"),
                    Especies_EMG=("Nome_Cientifico", "nunique"),
                    Especies_Principais=(
                        "Nome_Cientifico",
                        lambda s: repro.loc[s.index, "Nome_Cientifico"][
                            repro.loc[s.index, "Grupo_Reprodutivo_Principal"]
                        ].nunique(),
                    ),
                    Abundancia_Total_EMG=("Numero_de_Individuos", "sum"),
                    Abundancia_Evidencia_Forte=("Abundancia_Evidencia_Forte_Linha", "sum"),
                )
                .reset_index()
            )
            out["Perc_Evidencia_Forte"] = np.where(
                out["Abundancia_Total_EMG"].gt(0),
                out["Abundancia_Evidencia_Forte"] / out["Abundancia_Total_EMG"] * 100,
                np.nan,
            )
            summaries[sheet] = out

    return summaries


def _validations(base: pd.DataFrame, efforts: pd.DataFrame, cpue_effort: pd.DataFrame) -> pd.DataFrame:
    quant = base["Tipo_Amostragem_Base"].eq("Quantitativa")
    emg = base["Tem_EMG_Informado"]
    strong = base["Evidencia_Reprodutiva_Forte"]
    principal = base["Migradora_Nao_Migradora"].eq("Migradora") | base["Ameacada_Extincao"].eq("Sim")
    effort_mismatch = (
        base["Esforco_Resultado"].notna()
        & base["Esforco_Metadata"].notna()
        & (base["Esforco_Resultado"] - base["Esforco_Metadata"]).abs().gt(1e-9)
    )
    rows = [
        ("Projeto", PROJECT_NAME, PROJECT_CODE),
        ("Corte temporal", f"<= {CUT_YYYYMM}", "Banco confirmado ate PE090_AH2526_202512"),
        ("Linhas de resultado apos corte", len(base), ""),
        ("Campanhas apos corte", base["Campanha"].nunique(), "Esperado: 90"),
        ("Pontos apos corte", base["Ponto"].nunique(), "Esperado: 9"),
        ("Especies apos corte", base["Nome_Cientifico"].nunique(), "Esperado: 61"),
        ("Linhas quantitativas", int(quant.sum()), ""),
        ("Linhas qualitativas", int((~quant).sum()), ""),
        ("Esforcos quantitativos", int(efforts["Tipo_Amostragem_Base"].eq("Quantitativa").sum()), ""),
        ("Esforcos qualitativos", int(efforts["Tipo_Amostragem_Base"].eq("Qualitativa").sum()), ""),
        ("Resultados sem trecho espacial", int(base["Trecho"].isna().sum()), ""),
        ("Resultados sem classe de especie", int(base["Migradora_Nao_Migradora"].isna().sum()), ""),
        ("Quantitativas com esforco ausente/zero", int((quant & base["Esforco_Linha"].fillna(0).le(0)).sum()), ""),
        ("Quantitativas sem PC_g para CPUEb", int((quant & base["PC_g_individual"].isna()).sum()), ""),
        ("Linhas com mais de um individuo", int(base["Numero_de_Individuos"].gt(1).sum()), "PC_g tratado como peso individual"),
        ("Divergencia esforco resultado vs metadata", int(effort_mismatch.sum()), ""),
        ("Esforcos quantitativos com captura zero", int(cpue_effort["Captura_Zero"].sum()), ""),
        ("Linhas com EMG informado", int(emg.sum()), ""),
        ("Abundancia com EMG informado", int(base.loc[emg, "Numero_de_Individuos"].sum()), ""),
        ("Linhas com evidencia reprodutiva forte", int(strong.sum()), "F3/M3 + F4/M4"),
        ("Abundancia evidencia reprodutiva forte", int(base.loc[strong, "Numero_de_Individuos"].sum()), "F3/M3 + F4/M4"),
        (
            "Abundancia evidencia forte migradoras/ameacadas",
            int(base.loc[strong & principal, "Numero_de_Individuos"].sum()),
            "Recorte principal da analise reprodutiva",
        ),
    ]
    return pd.DataFrame(rows, columns=["Item", "Valor", "Observacao"])


def _format_excel(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for col in ws.columns:
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 10), 55)
    wb.save(path)


def _write_outputs(
    base: pd.DataFrame,
    efforts: pd.DataFrame,
    de_para: pd.DataFrame,
    summaries: dict[str, pd.DataFrame],
    validations: pd.DataFrame,
    source_workbook: Path,
) -> tuple[Path, Path, Path]:
    RESULTADOS_DIR.mkdir(parents=True, exist_ok=True)
    out_xlsx = RESULTADOS_DIR / f"base_analitica_ictiofauna_porto_estrela_{DATE_TAG}.xlsx"
    out_depara = RESULTADOS_DIR / f"de_para_pontos_trechos_porto_estrela_{DATE_TAG}.xlsx"
    out_md = RESULTADOS_DIR / f"base_analitica_ictiofauna_porto_estrela_{DATE_TAG}.md"

    de_para.to_excel(out_depara, index=False, engine="openpyxl")
    _format_excel(out_depara)

    with pd.ExcelWriter(out_xlsx, engine="openpyxl") as writer:
        validations.to_excel(writer, sheet_name="Validacoes", index=False)
        de_para.to_excel(writer, sheet_name="DePara_Trechos", index=False)
        efforts.to_excel(writer, sheet_name="Metadados_Esforco", index=False)
        base.to_excel(writer, sheet_name="Base_Linhas", index=False)
        for sheet_name, df in summaries.items():
            df.to_excel(writer, sheet_name=sheet_name[:31], index=False)
    _format_excel(out_xlsx)

    resumo_amostragem = summaries["Resumo_Amostragem"]
    cpue_campanha = summaries["CPUE_Campanha"]
    cpue_trecho = summaries["CPUE_Trecho"]
    repro_especie = summaries.get("Repro_Especie", pd.DataFrame())
    repro_principal = repro_especie[
        repro_especie.get("Grupo_Reprodutivo_Principal", pd.Series(dtype=bool)).eq(True)
    ].copy() if not repro_especie.empty else pd.DataFrame()

    def _fmt_table(df: pd.DataFrame) -> str:
        return df.to_markdown(index=False)

    md = f"""# Base analitica ictiofauna - Porto Estrela

Gerado em {datetime.now().strftime("%Y-%m-%d %H:%M")}.

## Fontes

- Workbook validado: `{source_workbook}`
- Caracterizacao aprovada: `caracterizacao_especies_porto_estrela_{DATE_TAG}.xlsx`

## Premissas aplicadas

- Corte temporal: ate dezembro de 2025 (`AAAAMM <= {CUT_YYYYMM}`).
- Campanhas no corte: {base["Campanha"].nunique()}.
- Pontos: {base["Ponto"].nunique()}.
- Especies: {base["Nome_Cientifico"].nunique()}.
- `PC_g` tratado como peso individual.
- `Biomassa_g_linha = Numero_de_Individuos * PC_g`.
- `CPUEn_linha = Numero_de_Individuos / Esforco * 100`.
- `CPUEb_linha = Biomassa_g_linha / Esforco * 100`.
- CPUE calculada apenas para amostragem quantitativa.
- `CPUEn` e a metrica central de abundancia padronizada e deve orientar
  estatistica, diversidade quantitativa, similaridade e series temporais de
  abundancia.
- Reproducao: `Sexo` + `EMG` padronizados conforme Bazzoli (2003).
- Evidencia reprodutiva forte: `F3`, `M3`, `F4` e `M4`.

## Amostragem

{_fmt_table(resumo_amostragem)}

## CPUE por trecho

{_fmt_table(cpue_trecho)}

## Reproducao - especies migradoras e/ou ameacadas

{_fmt_table(repro_principal.head(30)) if not repro_principal.empty else "Sem registros reprodutivos no recorte principal."}

## Validacoes

{_fmt_table(validations)}

## Saidas

- `{out_xlsx}`
- `{out_depara}`
"""
    out_md.write_text(md, encoding="utf-8")
    return out_xlsx, out_depara, out_md


def main() -> None:
    resultados, esforcos_raw, especies, source_workbook = _read_inputs()
    de_para = _de_para_trechos()
    efforts = _prepare_efforts(esforcos_raw, de_para)
    base, cpue_effort = _prepare_results(resultados, efforts, especies, de_para)
    summaries = _summaries(base, cpue_effort)
    validations = _validations(base, efforts, cpue_effort)
    out_xlsx, out_depara, out_md = _write_outputs(
        base, efforts, de_para, summaries, validations, source_workbook
    )

    print("OK - base analitica gerada")
    print(out_xlsx)
    print(out_depara)
    print(out_md)
    print(validations.to_string(index=False))


if __name__ == "__main__":
    main()
